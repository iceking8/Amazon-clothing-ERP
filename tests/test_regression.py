import importlib
import json
import os
import re
import secrets
import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from werkzeug.datastructures import FileStorage


ADMIN_PASSWORD = "TestAdmin!234"
EMERGENCY_ADMIN_PASSWORD = "LocalAdmin!234"


class ERPRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.runtime_root = Path.cwd() / "tests_runtime"
        cls.runtime_root.mkdir(exist_ok=True)
        cls.db_name = f"regression-{secrets.token_hex(4)}.db"
        cls.db_file = Path.cwd() / "instance" / cls.db_name
        cls.db_file.parent.mkdir(exist_ok=True)
        if cls.db_file.exists():
            cls.db_file.unlink()
        os.environ["ERP_DATABASE_URL"] = f"sqlite:///{cls.db_name}"
        os.environ["ERP_INIT_ADMIN_PASSWORD"] = ADMIN_PASSWORD
        os.environ["ERP_EMERGENCY_ADMIN_PASSWORD"] = EMERGENCY_ADMIN_PASSWORD
        os.environ["ERP_SECRET_KEY"] = "test-secret-key"

        cls.app_module = importlib.import_module("app")
        cls.app = cls.app_module.app
        cls.db = cls.app_module.db
        cls.User = cls.app_module.User
        cls.Role = cls.app_module.Role
        cls.Supplier = cls.app_module.Supplier
        cls.Customer = cls.app_module.Customer
        cls.Warehouse = cls.app_module.Warehouse
        cls.SKU = cls.app_module.SKU
        cls.PurchaseOrder = cls.app_module.PurchaseOrder
        cls.SalesOrder = cls.app_module.SalesOrder
        cls.InventoryBalance = cls.app_module.InventoryBalance
        cls.InventoryTransaction = cls.app_module.InventoryTransaction
        cls.DefectRepair = cls.app_module.DefectRepair
        cls.LoginAudit = cls.app_module.LoginAudit
        cls.SystemSetting = cls.app_module.SystemSetting
        cls.WarehousePostalCode = cls.app_module.WarehousePostalCode

    @classmethod
    def tearDownClass(cls):
        if getattr(cls, "db", None):
            with cls.app.app_context():
                cls.db.session.remove()
                cls.db.engine.dispose()
        if getattr(cls, "db_file", None) and cls.db_file.exists():
            try:
                cls.db_file.unlink()
            except PermissionError:
                pass
        for env_key in ("ERP_DATABASE_URL", "ERP_INIT_ADMIN_PASSWORD", "ERP_EMERGENCY_ADMIN_PASSWORD", "ERP_SECRET_KEY"):
            os.environ.pop(env_key, None)

    def setUp(self):
        self.client = self.app.test_client()
        self.app_module.LOGIN_ATTEMPTS.clear()

    def extract_csrf(self, html):
        match = re.search(r'name="csrf_token" value="([^"]+)"', html)
        self.assertIsNotNone(match, "CSRF token not found in HTML")
        return match.group(1)

    def login(self, username="admin", password=ADMIN_PASSWORD):
        login_page = self.client.get("/login")
        self.assertEqual(login_page.status_code, 200)
        csrf = self.extract_csrf(login_page.get_data(as_text=True))
        response = self.client.post(
            "/login",
            data={"csrf_token": csrf, "username": username, "password": password},
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        return response

    def current_csrf(self):
        with self.client.session_transaction() as session:
            return session["_csrf_token"]

    def unique_suffix(self):
        return secrets.token_hex(4).upper()

    def current_fba_calc_result(self):
        with self.client.session_transaction() as session:
            legacy_result = session.get("_fba_calc_result")
            if legacy_result:
                return legacy_result
            result_key = session.get("_fba_calc_result_key")
        if not result_key:
            return {}
        result_path = Path(self.app_module.get_fba_calc_result_path(result_key))
        if not result_path.exists():
            return {}
        return json.loads(result_path.read_text(encoding="utf-8"))

    def current_purchase_calc_result(self):
        with self.client.session_transaction() as session:
            legacy_result = session.get("_purchase_calc_result")
            if legacy_result:
                return legacy_result
            result_key = session.get("_purchase_calc_result_key")
        if not result_key:
            return {}
        result_path = Path(self.app_module.get_purchase_calc_result_path(result_key))
        if not result_path.exists():
            return {}
        return json.loads(result_path.read_text(encoding="utf-8"))

    def create_master_data(self, suffix):
        csrf = self.current_csrf()
        warehouse_code = f"T{suffix[:5]}"
        sku_code = f"SKU-{suffix}"
        barcode = f"98{suffix[:6].translate(str.maketrans('ABCDEF', '123456'))}1234"

        supplier_response = self.client.post(
            "/suppliers",
            data={
                "csrf_token": csrf,
                "name": f"Supplier {suffix}",
                "contact_name": "Alice",
                "phone": "13800138000",
                "address": "Shanghai",
                "remark": "Regression test supplier",
            },
            follow_redirects=True,
        )
        self.assertIn("供应商创建成功。", supplier_response.get_data(as_text=True))

        customer_response = self.client.post(
            "/customers",
            data={
                "csrf_token": csrf,
                "name": f"Customer {suffix}",
                "contact_name": "Bob",
                "phone": "13900139000",
                "address": "Suzhou",
                "remark": "Regression test customer",
            },
            follow_redirects=True,
        )
        self.assertIn("客户创建成功。", customer_response.get_data(as_text=True))

        warehouse_response = self.client.post(
            "/warehouses",
            data={
                "csrf_token": csrf,
                "name": f"Warehouse {suffix}",
                "code": warehouse_code,
                "address": "Ningbo",
                "manager": "Manager",
                "remark": "Regression test warehouse",
            },
            follow_redirects=True,
        )
        self.assertIn("仓库创建成功。", warehouse_response.get_data(as_text=True))

        sku_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": sku_code,
                "barcode": barcode,
                "name": f"Product {suffix}",
                "category": "Test",
                "color": "Black",
                "size": "L",
                "unit": "pcs",
                "cost_price": "25.50",
                "sale_price": "55.00",
                "safety_stock": "5",
                "remark": "Regression test sku",
            },
            follow_redirects=True,
        )
        self.assertIn("SKU 创建成功。", sku_response.get_data(as_text=True))

        with self.app.app_context():
            supplier = self.Supplier.query.filter_by(name=f"Supplier {suffix}").first()
            customer = self.Customer.query.filter_by(name=f"Customer {suffix}").first()
            warehouse = self.Warehouse.query.filter_by(code=warehouse_code).first()
            sku = self.SKU.query.filter_by(sku_code=sku_code).first()

        self.assertIsNotNone(supplier)
        self.assertIsNotNone(customer)
        self.assertIsNotNone(warehouse)
        self.assertIsNotNone(sku)
        return supplier, customer, warehouse, sku

    def test_login_page_renders_csrf_and_no_public_default_password(self):
        response = self.client.get("/login")
        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('name="csrf_token"', body)
        self.assertIn("系统不再内置公开默认密码", body)

    def test_import_sku_resolution_normalizes_codes_and_honors_mapping_scope(self):
        suffix = self.unique_suffix()
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            sku = self.SKU.query.first()
            operator = self.User(
                username=f"mapping-owner-{suffix}",
                full_name=f"Mapping Owner {suffix}",
                password_hash="unused",
            )
            other_operator = self.User(
                username=f"mapping-other-{suffix}",
                full_name=f"Mapping Other {suffix}",
                password_hash="unused",
            )
            self.db.session.add_all([operator, other_operator])
            self.db.session.flush()
            external_sku = f"E1081-WinRed-{suffix}"
            self.db.session.add(
                self.app_module.SKUMapping(
                    user_id=operator.id,
                    sku_id=sku.id,
                    external_sku_code=external_sku.upper(),
                )
            )
            self.db.session.commit()

            imported_code = f"\ufeffE1081T-Wine-Red-{suffix.lower()}\u200b"
            operator_map, operator_missing = self.app_module.resolve_import_sku_codes(
                [imported_code],
                user=operator,
            )
            admin_map, admin_missing = self.app_module.resolve_import_sku_codes(
                [imported_code],
                user=admin,
            )
            other_map, other_missing = self.app_module.resolve_import_sku_codes(
                [imported_code],
                user=other_operator,
            )
            ambiguous_sku = self.SKU(
                sku_code=f"SKU-AMB-{suffix}",
                name=f"Ambiguous Product {suffix}",
                category="Test",
            )
            self.db.session.add(ambiguous_sku)
            self.db.session.flush()
            self.db.session.add(
                self.app_module.SKUMapping(
                    user_id=operator.id,
                    sku_id=ambiguous_sku.id,
                    external_sku_code=f"E1081T-WineRed-{suffix}",
                )
            )
            self.db.session.commit()
            ambiguous_map, ambiguous_missing = self.app_module.resolve_import_sku_codes(
                [imported_code],
                user=operator,
            )

        normalized_code = f"E1081T-Wine-Red-{suffix.lower()}"
        self.assertEqual(operator_map[normalized_code].id, sku.id)
        self.assertEqual(operator_missing, [])
        self.assertEqual(admin_map[normalized_code].id, sku.id)
        self.assertEqual(admin_missing, [])
        self.assertEqual(other_map, {})
        self.assertEqual(other_missing, [normalized_code])
        self.assertEqual(ambiguous_map, {})
        self.assertEqual(ambiguous_missing, [normalized_code])

    def test_system_info_page_is_available_for_admin(self):
        self.login()
        response = self.client.get("/system-info")
        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("系统信息 / 版本信息", body)
        self.assertIn("当前版本", body)
        self.assertIn("数据库位置", body)

    def test_dashboard_shows_inbound_and_outbound_sku_rankings(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU-RANK-{suffix}"
        second_sku_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": f"77{suffix[:6].translate(str.maketrans('ABCDEF', '123456'))}8899",
                "name": f"Rank Product {suffix}",
                "category": "Rank",
                "color": "White",
                "size": "M",
                "unit": "pcs",
                "cost_price": "18.00",
                "sale_price": "39.00",
                "safety_stock": "2",
                "remark": "Ranking test sku",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_sku_response.status_code, 200)

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
            self.assertIsNotNone(second_sku)

            self.app_module.add_inventory_transaction(
                sku.id,
                warehouse.id,
                30,
                "manual_inbound",
                reference_type="ranking_test",
                reference_no=f"RANK-IN-A-{suffix}",
                operator_name="admin",
            )
            self.app_module.add_inventory_transaction(
                sku.id,
                warehouse.id,
                -8,
                "manual_outbound",
                reference_type="ranking_test",
                reference_no=f"RANK-OUT-A-{suffix}",
                operator_name="admin",
            )
            self.app_module.add_inventory_transaction(
                sku.id,
                warehouse.id,
                -5,
                "sales_outbound",
                reference_type="ranking_test",
                reference_no=f"RANK-SALE-A-{suffix}",
                operator_name="admin",
            )
            self.app_module.add_inventory_transaction(
                second_sku.id,
                warehouse.id,
                12,
                "purchase_inbound",
                reference_type="ranking_test",
                reference_no=f"RANK-IN-B-{suffix}",
                operator_name="admin",
            )
            self.app_module.add_inventory_transaction(
                second_sku.id,
                warehouse.id,
                6,
                "initial_inbound",
                reference_type="ranking_test",
                reference_no=f"RANK-INIT-B-{suffix}",
                operator_name="admin",
            )
            self.app_module.add_inventory_transaction(
                second_sku.id,
                warehouse.id,
                -3,
                "manual_outbound",
                reference_type="ranking_test",
                reference_no=f"RANK-OUT-B-{suffix}",
                operator_name="admin",
            )
            self.db.session.commit()

        dashboard_response = self.client.get("/dashboard")
        dashboard_body = dashboard_response.get_data(as_text=True)
        self.assertEqual(dashboard_response.status_code, 200)
        self.assertIn("SKU 出库/销售数量排行榜", dashboard_body)
        self.assertIn("SKU 入库数量排行榜", dashboard_body)
        self.assertLess(dashboard_body.index(sku.sku_code), dashboard_body.index(second_sku_code))

        filtered_response = self.client.get(f"/dashboard?sku_keyword={second_sku_code}")
        filtered_body = filtered_response.get_data(as_text=True)
        self.assertEqual(filtered_response.status_code, 200)
        self.assertIn(second_sku_code, filtered_body)
        self.assertNotIn(sku.sku_code, filtered_body)

    def test_roles_can_be_updated_and_deleted(self):
        self.login()
        csrf = self.current_csrf()

        create_response = self.client.post(
            "/roles",
            data={
                "csrf_token": csrf,
                "name": "运营测试角色",
                "description": "初始说明",
            },
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)

        with self.app.app_context():
            role = self.Role.query.filter_by(name="运营测试角色").first()
            self.assertIsNotNone(role)
            dashboard_permission = self.app_module.Permission.query.filter_by(code="dashboard.view").first()
        self.assertIsNotNone(dashboard_permission)

        update_response = self.client.post(
            f"/roles/{role.id}/update",
            data={
                "csrf_token": csrf,
                "name": "运营测试角色-已更新",
                "description": "更新后的说明",
                "permission_ids": str(dashboard_permission.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertIn("角色已更新", update_response.get_data(as_text=True))

        with self.app.app_context():
            refreshed_role = self.Role.query.get(role.id)
            self.assertIsNotNone(refreshed_role)
            self.assertEqual(refreshed_role.name, "运营测试角色-已更新")
            self.assertEqual(refreshed_role.description, "更新后的说明")
            self.assertEqual([permission.code for permission in refreshed_role.permissions], ["dashboard.view"])

        delete_response = self.client.post(
            f"/roles/{role.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(delete_response.status_code, 200)
        self.assertIn("角色已删除", delete_response.get_data(as_text=True))

        with self.app.app_context():
            self.assertIsNone(self.Role.query.get(role.id))

    def test_sku_permissions_support_custom_sales_roles(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        with self.app.app_context():
            sales_manage_permission = self.app_module.Permission.query.filter_by(code="sales.manage").first()
            inventory_view_permission = self.app_module.Permission.query.filter_by(code="inventory.view").first()
        self.assertIsNotNone(sales_manage_permission)
        self.assertIsNotNone(inventory_view_permission)

        role_response = self.client.post(
            "/roles",
            data={
                "csrf_token": csrf,
                "name": f"运营{suffix}",
                "description": "自定义销售角色",
                "permission_ids": [str(sales_manage_permission.id), str(inventory_view_permission.id)],
            },
            follow_redirects=True,
        )
        self.assertEqual(role_response.status_code, 200)

        with self.app.app_context():
            custom_role = self.Role.query.filter_by(name=f"运营{suffix}").first()
        self.assertIsNotNone(custom_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"ops{suffix.lower()}",
                "full_name": f"运营 {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(custom_role.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        with self.app.app_context():
            ops_user = self.User.query.filter_by(username=f"ops{suffix.lower()}").first()
        self.assertIsNotNone(ops_user)

        assign_response = self.client.post(
            f"/skus/{sku.id}/permissions",
            data={
                "csrf_token": csrf,
                "authorized_user_ids": str(ops_user.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(assign_response.status_code, 200)

        sku_page = self.client.get(f"/skus?sku_keyword={suffix}")
        sku_body = sku_page.get_data(as_text=True)
        self.assertEqual(sku_page.status_code, 200)
        self.assertIn(ops_user.full_name, sku_body)

        remove_response = self.client.post(
            f"/skus/{sku.id}/permissions",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(remove_response.status_code, 200)

        with self.app.app_context():
            refreshed_sku = self.SKU.query.get(sku.id)
            self.assertFalse(any(user.id == ops_user.id for user in refreshed_sku.authorized_users))

    def test_custom_operations_role_logs_into_first_accessible_page_without_permission_loop(self):
        self.login()
        suffix = self.unique_suffix()
        csrf = self.current_csrf()

        with self.app.app_context():
            sales_manage_permission = self.app_module.Permission.query.filter_by(code="sales.manage").first()
            inventory_view_permission = self.app_module.Permission.query.filter_by(code="inventory.view").first()
        self.assertIsNotNone(sales_manage_permission)
        self.assertIsNotNone(inventory_view_permission)

        role_response = self.client.post(
            "/roles",
            data={
                "csrf_token": csrf,
                "name": f"运营登录{suffix}",
                "description": "运营登录测试",
                "permission_ids": [str(sales_manage_permission.id), str(inventory_view_permission.id)],
            },
            follow_redirects=True,
        )
        self.assertEqual(role_response.status_code, 200)

        with self.app.app_context():
            operations_role = self.Role.query.filter_by(name=f"运营登录{suffix}").first()
        self.assertIsNotNone(operations_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"operate{suffix.lower()}",
                "full_name": f"运营登录 {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(operations_role.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        self.client.post("/logout", data={"csrf_token": csrf}, follow_redirects=True)

        operations_client = self.app.test_client()
        login_page = operations_client.get("/login")
        login_csrf = self.extract_csrf(login_page.get_data(as_text=True))
        login_response = operations_client.post(
            "/login",
            data={
                "csrf_token": login_csrf,
                "username": f"operate{suffix.lower()}",
                "password": "SalesUser!234",
            },
            follow_redirects=True,
        )
        login_body = login_response.get_data(as_text=True)
        self.assertEqual(login_response.status_code, 200)
        self.assertIn("登录成功", login_body)
        self.assertNotIn("权限不足，无法访问该页面。", login_body)

    def test_session_expires_after_one_hour_of_inactivity(self):
        self.login()
        with self.client.session_transaction() as session:
            session["_last_seen_at"] = 0
        response = self.client.get("/dashboard", follow_redirects=True)
        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("登录已超时，请重新登录。", body)
        self.assertIn("登录", body)

        with self.app.app_context():
            latest_audit = self.LoginAudit.query.order_by(self.LoginAudit.id.desc()).first()
            self.assertIsNotNone(latest_audit)
            self.assertEqual(latest_audit.logout_reason, "超时退出")
            self.assertIsNotNone(latest_audit.logout_at)

    def test_login_audit_page_records_login_and_logout(self):
        self.login()
        page = self.client.get("/login-audits")
        body = page.get_data(as_text=True)
        self.assertEqual(page.status_code, 200)
        self.assertIn("登录日志", body)

        with self.app.app_context():
            latest_audit = self.LoginAudit.query.order_by(self.LoginAudit.id.desc()).first()
            self.assertIsNotNone(latest_audit)
            self.assertEqual(latest_audit.username, "admin")
            self.assertIsNone(latest_audit.logout_at)

        csrf = self.current_csrf()
        logout_response = self.client.post("/logout", data={"csrf_token": csrf}, follow_redirects=True)
        self.assertEqual(logout_response.status_code, 200)

        with self.app.app_context():
            latest_audit = self.LoginAudit.query.order_by(self.LoginAudit.id.desc()).first()
            self.assertIsNotNone(latest_audit.logout_at)
            self.assertEqual(latest_audit.logout_reason, "主动退出")

    def test_system_info_can_update_login_security_config(self):
        self.login()
        csrf = self.current_csrf()
        response = self.client.post(
            "/system-info/login-security",
            data={
                "csrf_token": csrf,
                "login_lock_threshold": "5",
                "login_lock_duration_minutes": "30",
                "emergency_admin_ip_whitelist": "127.0.0.1,::1,192.168.5.69",
            },
            follow_redirects=True,
        )
        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("登录安全配置已更新", body)

        with self.app.app_context():
            threshold = self.SystemSetting.query.get("login_lock_threshold")
            duration = self.SystemSetting.query.get("login_lock_duration_minutes")
            self.assertIsNotNone(threshold)
            self.assertIsNotNone(duration)
            self.assertEqual(threshold.value, "5")
            self.assertEqual(duration.value, "30")
            whitelist = self.SystemSetting.query.get("emergency_admin_ip_whitelist")
            self.assertIsNotNone(whitelist)
            self.assertEqual(whitelist.value, "127.0.0.1,::1,192.168.5.69")

    def test_emergency_admin_respects_ip_whitelist(self):
        with self.app.app_context():
            emergency_user = self.User.query.filter_by(is_emergency_account=True).first()
        self.assertIsNotNone(emergency_user)

        local_client = self.app.test_client()
        local_page = local_client.get("/login")
        local_csrf = self.extract_csrf(local_page.get_data(as_text=True))
        local_success = local_client.post(
            "/login",
            data={"csrf_token": local_csrf, "username": emergency_user.username, "password": EMERGENCY_ADMIN_PASSWORD},
            follow_redirects=True,
        )
        self.assertEqual(local_success.status_code, 200)
        self.assertIn("登录成功", local_success.get_data(as_text=True))

        remote_client = self.app.test_client()
        remote_page = remote_client.get("/login", environ_overrides={"REMOTE_ADDR": "10.10.10.10"})
        remote_csrf = self.extract_csrf(remote_page.get_data(as_text=True))
        remote_response = remote_client.post(
            "/login",
            data={"csrf_token": remote_csrf, "username": emergency_user.username, "password": EMERGENCY_ADMIN_PASSWORD},
            environ_overrides={"REMOTE_ADDR": "10.10.10.10"},
            follow_redirects=True,
        )
        self.assertIn("应急管理员仅允许从白名单 IP 登录", remote_response.get_data(as_text=True))

        self.login()
        csrf = self.current_csrf()
        update_response = self.client.post(
            "/system-info/login-security",
            data={
                "csrf_token": csrf,
                "login_lock_threshold": "3",
                "login_lock_duration_minutes": str(7 * 24 * 60),
                "emergency_admin_ip_whitelist": "127.0.0.1,::1,192.168.5.69",
            },
            follow_redirects=True,
        )
        self.assertEqual(update_response.status_code, 200)

        whitelisted_client = self.app.test_client()
        whitelisted_page = whitelisted_client.get("/login", environ_overrides={"REMOTE_ADDR": "192.168.5.69"})
        whitelisted_csrf = self.extract_csrf(whitelisted_page.get_data(as_text=True))
        whitelisted_response = whitelisted_client.post(
            "/login",
            data={"csrf_token": whitelisted_csrf, "username": emergency_user.username, "password": EMERGENCY_ADMIN_PASSWORD},
            environ_overrides={"REMOTE_ADDR": "192.168.5.69"},
            follow_redirects=True,
        )
        self.assertIn("登录成功", whitelisted_response.get_data(as_text=True))

        with self.app.app_context():
            whitelist_setting = self.SystemSetting.query.get("emergency_admin_ip_whitelist")
            self.assertIsNotNone(whitelist_setting)
            self.assertIn("192.168.5.69", whitelist_setting.value)

    def test_login_lock_triggers_after_three_failed_attempts_and_can_be_unlocked(self):
        client = self.app.test_client()
        last_body = ""
        for _ in range(3):
            page = client.get("/login")
            csrf = self.extract_csrf(page.get_data(as_text=True))
            response = client.post(
                "/login",
                data={"csrf_token": csrf, "username": "admin", "password": "wrong-password"},
                follow_redirects=True,
            )
            last_body = response.get_data(as_text=True)
        self.assertIn("账号已锁定", last_body)

        blocked_page = client.get("/login")
        blocked_csrf = self.extract_csrf(blocked_page.get_data(as_text=True))
        blocked_response = client.post(
            "/login",
            data={"csrf_token": blocked_csrf, "username": "admin", "password": ADMIN_PASSWORD},
            follow_redirects=True,
        )
        self.assertIn("账号已锁定", blocked_response.get_data(as_text=True))

        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            emergency_user = self.User.query.filter_by(is_emergency_account=True).first()
        self.assertIsNotNone(admin)
        self.assertIsNotNone(emergency_user)

        emergency_login = self.login(username=emergency_user.username, password=EMERGENCY_ADMIN_PASSWORD)
        self.assertEqual(emergency_login.status_code, 200)
        csrf = self.current_csrf()
        unlock_response = self.client.post(f"/users/{admin.id}/unlock", data={"csrf_token": csrf}, follow_redirects=True)
        self.assertEqual(unlock_response.status_code, 200)
        self.assertIn("账号已解锁", unlock_response.get_data(as_text=True))

        fresh_client = self.app.test_client()
        page = fresh_client.get("/login")
        csrf = self.extract_csrf(page.get_data(as_text=True))
        success_response = fresh_client.post(
            "/login",
            data={"csrf_token": csrf, "username": "admin", "password": ADMIN_PASSWORD},
            follow_redirects=True,
        )
        self.assertIn("登录成功", success_response.get_data(as_text=True))

        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            self.assertIsNotNone(admin)
            self.assertEqual(admin.failed_login_attempts, 0)
            self.assertIsNone(admin.locked_until)
            latest_failure = self.LoginAudit.query.filter_by(username="admin", is_success=False).order_by(self.LoginAudit.id.desc()).first()
            self.assertIsNotNone(latest_failure)
            self.assertEqual(latest_failure.failure_reason, "账号已锁定")

    def test_sku_barcode_is_optional_and_blank_values_do_not_conflict(self):
        self.login()
        csrf = self.current_csrf()

        for suffix in ("NOBC01", "NOBC02"):
            response = self.client.post(
                "/skus",
                data={
                    "csrf_token": csrf,
                    "sku_code": f"SKU-{suffix}",
                    "barcode": "",
                    "name": f"无条码商品 {suffix}",
                    "category": "Test",
                    "color": "Black",
                    "size": "L",
                    "unit": "pcs",
                    "cost_price": "12.50",
                    "sale_price": "29.90",
                    "safety_stock": "2",
                    "remark": "Optional barcode test",
                },
                follow_redirects=True,
            )
            self.assertIn("SKU 创建成功。", response.get_data(as_text=True))

        with self.app.app_context():
            first_sku = self.SKU.query.filter_by(sku_code="SKU-NOBC01").first()
            second_sku = self.SKU.query.filter_by(sku_code="SKU-NOBC02").first()
            self.assertIsNotNone(first_sku)
            self.assertIsNotNone(second_sku)
            self.assertIsNone(first_sku.barcode)
            self.assertIsNone(second_sku.barcode)

    def test_search_filters_and_report_export_respect_sku_keyword(self):
        self.login()
        csrf = self.current_csrf()

        target_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": "1008-Black-M",
                "barcode": "",
                "name": "测试上衣 1008",
                "category": "Top",
                "color": "Black",
                "size": "M",
                "unit": "pcs",
                "cost_price": "10",
                "sale_price": "20",
                "safety_stock": "1",
                "remark": "",
            },
            follow_redirects=True,
        )
        other_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": "2001-White-L",
                "barcode": "",
                "name": "测试上衣 2001",
                "category": "Top",
                "color": "White",
                "size": "L",
                "unit": "pcs",
                "cost_price": "10",
                "sale_price": "20",
                "safety_stock": "1",
                "remark": "",
            },
            follow_redirects=True,
        )
        self.assertIn("SKU 创建成功。", target_response.get_data(as_text=True))
        self.assertIn("SKU 创建成功。", other_response.get_data(as_text=True))

        sku_search_page = self.client.get("/skus?sku_keyword=1008")
        sku_search_body = sku_search_page.get_data(as_text=True)
        self.assertIn("1008-Black-M", sku_search_body)
        self.assertNotIn("2001-White-L", sku_search_body)

        export_response = self.client.get("/reports/export/inventory?sku_keyword=1008")
        self.assertEqual(export_response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export_response.data))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        flattened = "\n".join("" if cell is None else str(cell) for row in rows for cell in row)
        self.assertIn("1008-Black-M", flattened)
        self.assertNotIn("2001-White-L", flattened)

    def test_inventory_transactions_page_and_manual_outbound_multi_sku(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU2-{suffix}"
        second_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": "",
                "name": f"Second Product {suffix}",
                "category": "Test",
                "color": "Blue",
                "size": "XL",
                "unit": "pcs",
                "cost_price": "30.00",
                "sale_price": "66.00",
                "safety_stock": "3",
                "remark": "second sku",
            },
            follow_redirects=True,
        )
        self.assertIn("SKU 创建成功。", second_response.get_data(as_text=True))

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
        self.assertIsNotNone(second_sku)

        for current_sku, qty, ref in ((sku, "10", f"IN1-{suffix}"), (second_sku, "7", f"IN2-{suffix}")):
            inbound_response = self.client.post(
                "/inventory/in",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(current_sku.id),
                    "warehouse_id": str(warehouse.id),
                    "supplier_id": str(supplier.id),
                    "quantity": qty,
                    "reference_type": "test",
                    "reference_no": ref,
                    "note": "stock for manual outbound",
                },
                follow_redirects=True,
            )
            self.assertIn("入库操作成功。", inbound_response.get_data(as_text=True))

        outbound_response = self.client.post(
            "/inventory/out",
            data={
                "csrf_token": csrf,
                "warehouse_id": str(warehouse.id),
                "reference_type": "manual_outbound_order",
                "reference_no": f"MO-{suffix}",
                "note": "multi sku outbound",
                "manual_outbound_sku_id": [str(sku.id), str(second_sku.id)],
                "manual_outbound_quantity": ["4", "2"],
                "manual_outbound_price": ["0", "0"],
            },
            follow_redirects=True,
        )
        self.assertIn("手工出库单创建成功。", outbound_response.get_data(as_text=True))
        self.assertIn(f"MO-{suffix}", outbound_response.get_data(as_text=True))

        with self.app.app_context():
            first_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            second_balance = self.InventoryBalance.query.filter_by(sku_id=second_sku.id, warehouse_id=warehouse.id).first()
            self.assertEqual(first_balance.quantity, 6)
            self.assertEqual(second_balance.quantity, 5)

        transactions_page = self.client.get(f"/inventory/transactions?sku_keyword={suffix}")
        transactions_body = transactions_page.get_data(as_text=True)
        self.assertEqual(transactions_page.status_code, 200)
        self.assertIn(f"MO-{suffix}", transactions_body)
        self.assertIn("库存流水", transactions_body)

    def test_inventory_history_can_be_filtered_by_month_and_paginated(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)

        with self.app.app_context():
            for index in range(51):
                self.db.session.add(
                    self.InventoryTransaction(
                        sku_id=sku.id,
                        warehouse_id=warehouse.id,
                        supplier_id=supplier.id,
                        quantity=1,
                        transaction_type="manual_inbound",
                        reference_type="manual_inbound_order",
                        reference_no=f"APR-{suffix}-{index:02d}",
                        created_at=datetime(2026, 4, 1) + timedelta(days=index % 29, minutes=index),
                    )
                )
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    supplier_id=supplier.id,
                    quantity=1,
                    transaction_type="manual_inbound",
                    reference_type="manual_inbound_order",
                    reference_no=f"MAY-{suffix}",
                    created_at=datetime(2026, 5, 1),
                )
            )
            self.db.session.commit()

        inbound_page = self.client.get(f"/inventory/in?sku_keyword={suffix}&month=2026-04")
        inbound_body = inbound_page.get_data(as_text=True)
        self.assertEqual(inbound_page.status_code, 200)
        self.assertIn("查询月份", inbound_body)
        self.assertIn("共 51 条，每页 50 条", inbound_body)
        self.assertIn("第 1 / 2 页", inbound_body)
        self.assertNotIn(f"MAY-{suffix}", inbound_body)

        inbound_second_page = self.client.get(f"/inventory/in?sku_keyword={suffix}&month=2026-04&page=2")
        inbound_second_body = inbound_second_page.get_data(as_text=True)
        self.assertIn("第 2 / 2 页", inbound_second_body)
        self.assertIn(sku.sku_code, inbound_second_body)

        transactions_page = self.client.get(f"/inventory/transactions?sku_keyword={suffix}&month=2026-04&page=2")
        transactions_body = transactions_page.get_data(as_text=True)
        self.assertEqual(transactions_page.status_code, 200)
        self.assertIn("第 2 / 2 页", transactions_body)
        self.assertIn(f"APR-{suffix}-00", transactions_body)
        self.assertNotIn(f"MAY-{suffix}", transactions_body)

    def test_inventory_flow_and_export_report(self):
        login_response = self.login()
        self.assertIn("登录成功。", login_response.get_data(as_text=True))

        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "20",
                "reference_type": "test",
                "reference_no": f"IN-{suffix}",
                "note": "Initial inbound for regression",
            },
            follow_redirects=True,
        )
        self.assertIn("入库操作成功。", inbound_response.get_data(as_text=True))

        purchase_response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "order_no": "",
                "expected_date": "2026-04-10",
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "status": "submitted",
                "remark": "Regression PO",
                "purchase_sku_id": str(sku.id),
                "purchase_quantity": "8",
                "purchase_price": "20.00",
            },
            follow_redirects=True,
        )
        self.assertIn("采购订单创建成功。", purchase_response.get_data(as_text=True))

        with self.app.app_context():
            purchase_order = self.PurchaseOrder.query.order_by(self.PurchaseOrder.id.desc()).first()

        receive_response = self.client.post(
            f"/purchase-orders/{purchase_order.id}/receive",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("采购订单已收货，库存已更新。", receive_response.get_data(as_text=True))

        sales_response = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": csrf,
                "order_no": "",
                "customer_id": str(customer.id),
                "customer_name": "",
                "phone": "",
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "shipping_address": "",
                "remark": "Regression SO",
                "sales_sku_id": str(sku.id),
                "sales_quantity": "6",
                "sales_price": "55.00",
            },
            follow_redirects=True,
        )
        self.assertIn("销售单创建成功。", sales_response.get_data(as_text=True))

        with self.app.app_context():
            sales_order = self.SalesOrder.query.order_by(self.SalesOrder.id.desc()).first()

        ship_response = self.client.post(
            f"/sales-orders/{sales_order.id}/ship",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("销售单已发货，库存已更新。", ship_response.get_data(as_text=True))

        with self.app.app_context():
            inventory_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertIsNotNone(inventory_balance)
            self.assertEqual(inventory_balance.quantity, 22)

        inventory_page = self.client.get("/inventory")
        self.assertEqual(inventory_page.status_code, 200)

        scan_page = self.client.get(f"/scan?q={sku.sku_code}")
        self.assertEqual(scan_page.status_code, 200)
        self.assertIn(sku.sku_code, scan_page.get_data(as_text=True))

        export_response = self.client.get("/reports/export/inventory")
        self.assertEqual(export_response.status_code, 200)
        self.assertTrue(
            export_response.headers["Content-Type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )

    def test_manual_inbound_records_supplier_per_order_for_same_sku(self):
        self.login()
        suffix = self.unique_suffix()
        supplier_a, _customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        supplier_b_response = self.client.post(
            "/suppliers",
            data={
                "csrf_token": csrf,
                "name": f"Factory B {suffix}",
                "contact_name": "Factory B Owner",
                "phone": "13700137000",
                "address": "Dongguan",
                "remark": "Alternative factory for same SKU",
            },
            follow_redirects=True,
        )
        self.assertIn("供应商创建成功。", supplier_b_response.get_data(as_text=True))

        with self.app.app_context():
            supplier_b = self.Supplier.query.filter_by(name=f"Factory B {suffix}").first()
        self.assertIsNotNone(supplier_b)

        for supplier, qty, ref in (
            (supplier_a, "6", f"FACTORY-A-{suffix}"),
            (supplier_b, "4", f"FACTORY-B-{suffix}"),
        ):
            inbound_response = self.client.post(
                "/inventory/in",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(sku.id),
                    "warehouse_id": str(warehouse.id),
                    "supplier_id": str(supplier.id),
                    "quantity": qty,
                    "reference_type": "factory-inbound",
                    "reference_no": ref,
                    "note": "same SKU from different factories",
                },
                follow_redirects=True,
            )
            self.assertIn("入库操作成功。", inbound_response.get_data(as_text=True))

        with self.app.app_context():
            factory_a_txn = self.InventoryTransaction.query.filter_by(reference_no=f"FACTORY-A-{suffix}").first()
            factory_b_txn = self.InventoryTransaction.query.filter_by(reference_no=f"FACTORY-B-{suffix}").first()
            balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertIsNotNone(factory_a_txn)
            self.assertIsNotNone(factory_b_txn)
            self.assertEqual(factory_a_txn.supplier_id, supplier_a.id)
            self.assertEqual(factory_b_txn.supplier_id, supplier_b.id)
            self.assertEqual(balance.quantity, 10)

        inbound_page = self.client.get(f"/inventory/in?sku_keyword={suffix}")
        inbound_body = inbound_page.get_data(as_text=True)
        self.assertIn("供应商", inbound_body)
        self.assertIn(supplier_a.name, inbound_body)
        self.assertIn(supplier_b.name, inbound_body)

    def test_operator_name_is_recorded_for_business_documents(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        purchase_response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "status": "submitted",
                "remark": "Operator tracking PO",
                "purchase_sku_id": str(sku.id),
                "purchase_quantity": "3",
                "purchase_price": "10.00",
            },
            follow_redirects=True,
        )
        self.assertEqual(purchase_response.status_code, 200)

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "9",
                "reference_type": "operator-test",
                "reference_no": f"OP-IN-{suffix}",
                "note": "operator inbound",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        sales_response = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": csrf,
                "customer_id": str(customer.id),
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "remark": "Operator tracking SO",
                "sales_sku_id": str(sku.id),
                "sales_quantity": "2",
                "sales_price": "55.00",
            },
            follow_redirects=True,
        )
        self.assertEqual(sales_response.status_code, 200)

        defect_response = self.client.post(
            "/defects",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "1",
                "status": "pending",
                "defect_reason": "operator tracking defect",
                "repair_result": "",
            },
            follow_redirects=True,
        )
        self.assertEqual(defect_response.status_code, 200)

        purchase_page = self.client.get("/purchase-orders")
        sales_page = self.client.get("/sales-orders")
        in_page = self.client.get("/inventory/in")
        defects_page = self.client.get("/defects")
        txns_page = self.client.get("/inventory/transactions")

        with self.app.app_context():
            purchase_order = self.PurchaseOrder.query.order_by(self.PurchaseOrder.id.desc()).first()
            sales_order = self.SalesOrder.query.order_by(self.SalesOrder.id.desc()).first()
            defect_repair = self.DefectRepair.query.order_by(self.DefectRepair.id.desc()).first()
            inbound_txn = self.InventoryTransaction.query.filter_by(reference_no=f"OP-IN-{suffix}").first()
            self.assertEqual(purchase_order.operator_name, "系统管理员")
            self.assertEqual(sales_order.operator_name, "系统管理员")
            self.assertEqual(defect_repair.operator_name, "系统管理员")
            self.assertEqual(inbound_txn.operator_name, "系统管理员")

        for response in (purchase_page, sales_page, in_page, defects_page, txns_page):
            self.assertEqual(response.status_code, 200)
            self.assertIn("系统管理员", response.get_data(as_text=True))

    def test_salesperson_can_only_access_authorized_skus(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU-LIMIT-{suffix}"
        second_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": "",
                "name": f"Limited Product {suffix}",
                "category": "Test",
                "color": "Blue",
                "size": "XL",
                "unit": "pcs",
                "cost_price": "30.00",
                "sale_price": "60.00",
                "safety_stock": "2",
                "remark": "second limited sku",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_response.status_code, 200)

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
            sales_role = self.Role.query.filter_by(name="销售员").first()
        self.assertIsNotNone(second_sku)
        self.assertIsNotNone(sales_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"sales{suffix.lower()}",
                "full_name": f"Sales {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(sales_role.id),
                "authorized_sku_ids": str(sku.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        for current_sku, ref in ((sku, f"AUTH-IN-{suffix}"), (second_sku, f"UNAUTH-IN-{suffix}")):
            inbound_response = self.client.post(
                "/inventory/in",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(current_sku.id),
                    "warehouse_id": str(warehouse.id),
                    "supplier_id": str(supplier.id),
                    "quantity": "10",
                    "reference_type": "auth-test",
                    "reference_no": ref,
                    "note": "inventory for sku auth test",
                },
                follow_redirects=True,
            )
            self.assertEqual(inbound_response.status_code, 200)

        self.client.post("/logout", data={"csrf_token": csrf}, follow_redirects=True)
        sales_login = self.login(username=f"sales{suffix.lower()}", password="SalesUser!234")
        self.assertEqual(sales_login.status_code, 200)
        sales_csrf = self.current_csrf()

        inventory_page = self.client.get("/inventory")
        inventory_body = inventory_page.get_data(as_text=True)
        self.assertEqual(inventory_page.status_code, 200)
        self.assertIn(sku.sku_code, inventory_body)
        self.assertNotIn(second_sku.sku_code, inventory_body)

        sales_page = self.client.get("/sales-orders")
        self.assertIn(sku.sku_code, sales_page.get_data(as_text=True))
        self.assertNotIn(second_sku.sku_code, sales_page.get_data(as_text=True))

        unauthorized_sales = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": sales_csrf,
                "customer_id": str(customer.id),
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "remark": "unauthorized sales order",
                "sales_sku_id": str(second_sku.id),
                "sales_quantity": "1",
                "sales_price": "60.00",
            },
            follow_redirects=True,
        )
        self.assertIn("未授权", unauthorized_sales.get_data(as_text=True))

        unauthorized_outbound = self.client.post(
            "/inventory/out",
            data={
                "csrf_token": sales_csrf,
                "warehouse_id": str(warehouse.id),
                "reference_type": "manual_outbound_order",
                "reference_no": f"NOAUTH-{suffix}",
                "note": "unauthorized outbound",
                "manual_outbound_sku_id": str(second_sku.id),
                "manual_outbound_quantity": "1",
                "manual_outbound_price": "0",
            },
            follow_redirects=True,
        )
        self.assertIn("未授权", unauthorized_outbound.get_data(as_text=True))

    def test_sku_permissions_can_be_assigned_from_sku_list(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        with self.app.app_context():
            sales_role = self.Role.query.filter_by(name="销售员").first()
        self.assertIsNotNone(sales_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"skuassign{suffix.lower()}",
                "full_name": f"SKU Assign {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(sales_role.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        with self.app.app_context():
            sales_user = self.User.query.filter_by(username=f"skuassign{suffix.lower()}").first()
        self.assertIsNotNone(sales_user)

        assign_response = self.client.post(
            f"/skus/{sku.id}/permissions",
            data={
                "csrf_token": csrf,
                "authorized_user_ids": str(sales_user.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(assign_response.status_code, 200)

        with self.app.app_context():
            refreshed_sku = self.SKU.query.get(sku.id)
            refreshed_user = self.User.query.get(sales_user.id)
            self.assertTrue(any(user.id == sales_user.id for user in refreshed_sku.authorized_users))
            self.assertTrue(any(item.id == sku.id for item in refreshed_user.authorized_skus))

    def test_sku_list_can_filter_by_sales_user_permission(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU-FILTER-{suffix}"
        second_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": "",
                "name": f"Filter Product {suffix}",
                "category": "Test",
                "color": "Gray",
                "size": "L",
                "unit": "pcs",
                "cost_price": "25.00",
                "sale_price": "55.00",
                "safety_stock": "1",
                "remark": "sku filter test",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_response.status_code, 200)

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
            sales_role = self.Role.query.filter_by(name="销售员").first()
        self.assertIsNotNone(second_sku)
        self.assertIsNotNone(sales_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"skufilter{suffix.lower()}",
                "full_name": f"SKU Filter {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(sales_role.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        with self.app.app_context():
            sales_user = self.User.query.filter_by(username=f"skufilter{suffix.lower()}").first()
        self.assertIsNotNone(sales_user)

        assign_response = self.client.post(
            f"/skus/{sku.id}/permissions?sales_user_id={sales_user.id}",
            data={
                "csrf_token": csrf,
                "authorized_user_ids": str(sales_user.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(assign_response.status_code, 200)

        filtered_page = self.client.get(f"/skus?sales_user_id={sales_user.id}")
        filtered_body = filtered_page.get_data(as_text=True)
        self.assertEqual(filtered_page.status_code, 200)
        self.assertIn(sku.sku_code, filtered_body)
        self.assertNotIn(second_sku.sku_code, filtered_body)
        self.assertIn("全部销售员", filtered_body)
        self.assertIn(sales_user.full_name, filtered_body)

    def test_sku_list_can_filter_unassigned_skus(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU-UNASSIGNED-{suffix}"
        second_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": "",
                "name": f"Unassigned Product {suffix}",
                "category": "Test",
                "color": "Black",
                "size": "M",
                "unit": "pcs",
                "cost_price": "21.00",
                "sale_price": "45.00",
                "safety_stock": "1",
                "remark": "unassigned sku test",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_response.status_code, 200)

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
            sales_role = self.Role.query.filter_by(name="销售员").first()
        self.assertIsNotNone(second_sku)
        self.assertIsNotNone(sales_role)

        user_response = self.client.post(
            "/users",
            data={
                "csrf_token": csrf,
                "username": f"unassigned{suffix.lower()}",
                "full_name": f"Unassigned Filter {suffix}",
                "password": "SalesUser!234",
                "is_active": "on",
                "role_ids": str(sales_role.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(user_response.status_code, 200)

        with self.app.app_context():
            sales_user = self.User.query.filter_by(username=f"unassigned{suffix.lower()}").first()
        self.assertIsNotNone(sales_user)

        assign_response = self.client.post(
            f"/skus/{sku.id}/permissions",
            data={
                "csrf_token": csrf,
                "authorized_user_ids": str(sales_user.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(assign_response.status_code, 200)

        filtered_page = self.client.get("/skus?authorization_state=unassigned")
        filtered_body = filtered_page.get_data(as_text=True)
        self.assertEqual(filtered_page.status_code, 200)
        self.assertIn(second_sku.sku_code, filtered_body)
        self.assertNotIn(sku.sku_code, filtered_body)
        self.assertIn("仅看未分配", filtered_body)

    def test_sku_batch_permissions_append_without_overwriting_existing_users(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_sku_code = f"SKU-BATCH-{suffix}"
        second_response = self.client.post(
            "/skus",
            data={
                "csrf_token": csrf,
                "sku_code": second_sku_code,
                "barcode": "",
                "name": f"Batch Product {suffix}",
                "category": "Test",
                "color": "White",
                "size": "S",
                "unit": "pcs",
                "cost_price": "20.00",
                "sale_price": "50.00",
                "safety_stock": "1",
                "remark": "batch permission test",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_response.status_code, 200)

        with self.app.app_context():
            second_sku = self.SKU.query.filter_by(sku_code=second_sku_code).first()
            sales_role = self.Role.query.filter_by(name="销售员").first()
        self.assertIsNotNone(second_sku)
        self.assertIsNotNone(sales_role)

        for username_prefix, full_name in (
            ("batcha", f"Batch A {suffix}"),
            ("batchb", f"Batch B {suffix}"),
        ):
            response = self.client.post(
                "/users",
                data={
                    "csrf_token": csrf,
                    "username": f"{username_prefix}{suffix.lower()}",
                    "full_name": full_name,
                    "password": "SalesUser!234",
                    "is_active": "on",
                    "role_ids": str(sales_role.id),
                },
                follow_redirects=True,
            )
            self.assertEqual(response.status_code, 200)

        with self.app.app_context():
            sales_user_a = self.User.query.filter_by(username=f"batcha{suffix.lower()}").first()
            sales_user_b = self.User.query.filter_by(username=f"batchb{suffix.lower()}").first()
        self.assertIsNotNone(sales_user_a)
        self.assertIsNotNone(sales_user_b)

        initial_assign = self.client.post(
            f"/skus/{sku.id}/permissions",
            data={
                "csrf_token": csrf,
                "authorized_user_ids": str(sales_user_a.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(initial_assign.status_code, 200)

        batch_assign = self.client.post(
            "/skus/permissions/batch",
            data={
                "csrf_token": csrf,
                "selected_sku_ids": [str(sku.id), str(second_sku.id)],
                "authorized_user_ids": str(sales_user_b.id),
            },
            follow_redirects=True,
        )
        self.assertEqual(batch_assign.status_code, 200)

        with self.app.app_context():
            refreshed_first_sku = self.SKU.query.get(sku.id)
            refreshed_second_sku = self.SKU.query.get(second_sku.id)
            first_usernames = sorted(user.username for user in refreshed_first_sku.authorized_users)
            second_usernames = sorted(user.username for user in refreshed_second_sku.authorized_users)
            self.assertEqual(first_usernames, sorted([sales_user_a.username, sales_user_b.username]))
            self.assertEqual(second_usernames, [sales_user_b.username])

    def test_csrf_duplicate_actions_and_insufficient_inventory(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        no_csrf_response = self.client.post(
            "/inventory/in",
            data={"sku_id": str(sku.id), "warehouse_id": str(warehouse.id), "quantity": "1"},
            follow_redirects=True,
        )
        self.assertIn("请求已过期或无效", no_csrf_response.get_data(as_text=True))

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "10",
                "reference_type": "test",
                "reference_no": f"SAFE-{suffix}",
                "note": "stock for negative tests",
            },
            follow_redirects=True,
        )
        self.assertIn("入库操作成功。", inbound_response.get_data(as_text=True))

        purchase_response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "status": "submitted",
                "remark": "Duplicate receive test",
                "purchase_sku_id": str(sku.id),
                "purchase_quantity": "2",
                "purchase_price": "10.00",
            },
            follow_redirects=True,
        )
        self.assertIn("采购订单创建成功。", purchase_response.get_data(as_text=True))

        with self.app.app_context():
            purchase_order = self.PurchaseOrder.query.order_by(self.PurchaseOrder.id.desc()).first()

        first_receive = self.client.post(
            f"/purchase-orders/{purchase_order.id}/receive",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("采购订单已收货，库存已更新。", first_receive.get_data(as_text=True))

        duplicate_receive = self.client.post(
            f"/purchase-orders/{purchase_order.id}/receive",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("已经收货", duplicate_receive.get_data(as_text=True))

        sales_response = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": csrf,
                "customer_id": str(customer.id),
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "remark": "Duplicate ship test",
                "sales_sku_id": str(sku.id),
                "sales_quantity": "4",
                "sales_price": "55.00",
            },
            follow_redirects=True,
        )
        self.assertIn("销售单创建成功。", sales_response.get_data(as_text=True))

        with self.app.app_context():
            sales_order = self.SalesOrder.query.order_by(self.SalesOrder.id.desc()).first()

        first_ship = self.client.post(
            f"/sales-orders/{sales_order.id}/ship",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("销售单已发货，库存已更新。", first_ship.get_data(as_text=True))

        duplicate_ship = self.client.post(
            f"/sales-orders/{sales_order.id}/ship",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("已经发货", duplicate_ship.get_data(as_text=True))

        overdraw_response = self.client.post(
            "/inventory/out",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "9999",
                "reference_type": "stress",
                "reference_no": "OVERDRAW",
                "note": "should fail",
            },
            follow_redirects=True,
        )
        self.assertIn("库存不足", overdraw_response.get_data(as_text=True))

        with self.app.app_context():
            before_qty = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first().quantity

        huge_sales_response = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": csrf,
                "customer_id": "",
                "customer_name": "Stress Customer",
                "phone": "13000000000",
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "shipping_address": "Test Address",
                "remark": "Should fail on ship",
                "sales_sku_id": str(sku.id),
                "sales_quantity": "9999",
                "sales_price": "55.00",
            },
            follow_redirects=True,
        )
        self.assertIn("销售单创建成功。", huge_sales_response.get_data(as_text=True))

        with self.app.app_context():
            huge_sales_order = self.SalesOrder.query.order_by(self.SalesOrder.id.desc()).first()

        failed_ship = self.client.post(
            f"/sales-orders/{huge_sales_order.id}/ship",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertIn("库存不足", failed_ship.get_data(as_text=True))

        with self.app.app_context():
            after_qty = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first().quantity
        self.assertEqual(before_qty, after_qty)

    def test_delete_routes_cleanup_inventory_and_master_data(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "5",
                "reference_type": "cleanup",
                "reference_no": f"IN-{suffix}",
                "note": "cleanup inbound",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        with self.app.app_context():
            manual_txn = self.app_module.InventoryTransaction.query.filter_by(reference_no=f"IN-{suffix}").first()
        self.assertIsNotNone(manual_txn)

        purchase_response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "status": "submitted",
                "remark": "Delete cleanup PO",
                "purchase_sku_id": str(sku.id),
                "purchase_quantity": "4",
                "purchase_price": "12.00",
            },
            follow_redirects=True,
        )
        self.assertEqual(purchase_response.status_code, 200)

        with self.app.app_context():
            purchase_order = self.PurchaseOrder.query.order_by(self.PurchaseOrder.id.desc()).first()

        self.client.post(
            f"/purchase-orders/{purchase_order.id}/receive",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )

        sales_response = self.client.post(
            "/sales-orders",
            data={
                "csrf_token": csrf,
                "customer_id": str(customer.id),
                "warehouse_id": str(warehouse.id),
                "status": "confirmed",
                "remark": "Delete cleanup SO",
                "sales_sku_id": str(sku.id),
                "sales_quantity": "3",
                "sales_price": "55.00",
            },
            follow_redirects=True,
        )
        self.assertEqual(sales_response.status_code, 200)

        with self.app.app_context():
            sales_order = self.SalesOrder.query.order_by(self.SalesOrder.id.desc()).first()

        self.client.post(
            f"/sales-orders/{sales_order.id}/ship",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )

        blocked_supplier_delete = self.client.post(
            f"/suppliers/{supplier.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(blocked_supplier_delete.status_code, 200)
        with self.app.app_context():
            self.assertIsNotNone(self.db.session.get(self.Supplier, supplier.id))

        self.client.post(
            f"/sales-orders/{sales_order.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        with self.app.app_context():
            balance_after_sales_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertIsNotNone(balance_after_sales_delete)
            self.assertEqual(balance_after_sales_delete.quantity, 9)

        self.client.post(
            f"/purchase-orders/{purchase_order.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        with self.app.app_context():
            balance_after_purchase_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertIsNotNone(balance_after_purchase_delete)
            self.assertEqual(balance_after_purchase_delete.quantity, 5)

        self.client.post(
            f"/inventory/transactions/{manual_txn.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        with self.app.app_context():
            remaining_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertTrue(remaining_balance is None or remaining_balance.quantity == 0)

        self.client.post(
            f"/customers/{customer.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.client.post(
            f"/suppliers/{supplier.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.client.post(
            f"/warehouses/{warehouse.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.client.post(
            f"/skus/{sku.id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )

        with self.app.app_context():
            self.assertIsNone(self.db.session.get(self.Customer, customer.id))
            self.assertIsNone(self.db.session.get(self.Supplier, supplier.id))
            self.assertIsNone(self.db.session.get(self.Warehouse, warehouse.id))
            self.assertIsNone(self.db.session.get(self.SKU, sku.id))

    def test_inventory_summary_shows_defect_status_quantities_and_defect_list_sku(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "20",
                "reference_type": "defect-test",
                "reference_no": f"DEFECT-IN-{suffix}",
                "note": "inventory for defect summary",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        for qty, status, reason in (
            ("2", "pending", "pending defect"),
            ("3", "repairing", "repairing defect"),
            ("4", "done", "done defect"),
            ("5", "scrapped", "scrapped defect"),
        ):
            defect_response = self.client.post(
                "/defects",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(sku.id),
                    "warehouse_id": str(warehouse.id),
                    "quantity": qty,
                    "status": status,
                    "defect_reason": reason,
                    "repair_result": "",
                },
                follow_redirects=True,
            )
            self.assertEqual(defect_response.status_code, 200)

        defects_page = self.client.get(f"/defects?sku_keyword={suffix}")
        defects_body = defects_page.get_data(as_text=True)
        self.assertEqual(defects_page.status_code, 200)
        self.assertIn(sku.sku_code, defects_body)

        inventory_page = self.client.get(f"/inventory?sku_keyword={suffix}")
        inventory_body = inventory_page.get_data(as_text=True)
        self.assertEqual(inventory_page.status_code, 200)
        self.assertIn(">10</td>", inventory_body)
        self.assertIn(">2</td>", inventory_body)
        self.assertIn(">3</td>", inventory_body)
        self.assertIn(">4</td>", inventory_body)
        self.assertIn(">5</td>", inventory_body)

        export_response = self.client.get(f"/reports/export/inventory?sku_keyword={suffix}")
        self.assertEqual(export_response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export_response.data))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0]
        flattened = "\n".join("" if cell is None else str(cell) for row in rows for cell in row)
        self.assertIn("待处理", header)
        self.assertIn("返修中", header)
        self.assertIn("已完成", header)
        self.assertIn("已报废", header)
        self.assertIn("当前库存", header)
        self.assertIn("安全库存", header)
        self.assertIn("库存价值", header)
        self.assertIn("10", flattened)
        self.assertIn("2", flattened)
        self.assertIn("3", flattened)
        self.assertIn("4", flattened)
        self.assertIn("5", flattened)

    def test_inventory_summary_filters_defect_status_by_warehouse(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        second_warehouse_response = self.client.post(
            "/warehouses",
            data={
                "csrf_token": csrf,
                "name": f"Warehouse 2 {suffix}",
                "code": f"W2{suffix[:4]}",
                "address": "Hangzhou",
                "manager": "Manager 2",
                "remark": "Second regression warehouse",
            },
            follow_redirects=True,
        )
        self.assertEqual(second_warehouse_response.status_code, 200)

        with self.app.app_context():
            second_warehouse = self.Warehouse.query.filter_by(name=f"Warehouse 2 {suffix}").first()
        self.assertIsNotNone(second_warehouse)

        for current_warehouse, qty, ref in (
            (warehouse, "20", f"DEFECT-IN-A-{suffix}"),
            (second_warehouse, "8", f"DEFECT-IN-B-{suffix}"),
        ):
            inbound_response = self.client.post(
                "/inventory/in",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(sku.id),
                    "warehouse_id": str(current_warehouse.id),
                    "supplier_id": str(supplier.id),
                    "quantity": qty,
                    "reference_type": "defect-test",
                    "reference_no": ref,
                    "note": "inventory for defect warehouse summary",
                },
                follow_redirects=True,
            )
            self.assertEqual(inbound_response.status_code, 200)

        for current_warehouse, qty, status, reason in (
            (warehouse, "2", "pending", "warehouse a pending"),
            (warehouse, "3", "repairing", "warehouse a repairing"),
            (second_warehouse, "7", "scrapped", "warehouse b scrapped"),
        ):
            defect_response = self.client.post(
                "/defects",
                data={
                    "csrf_token": csrf,
                    "sku_id": str(sku.id),
                    "warehouse_id": str(current_warehouse.id),
                    "quantity": qty,
                    "status": status,
                    "defect_reason": reason,
                    "repair_result": "",
                },
                follow_redirects=True,
            )
            self.assertEqual(defect_response.status_code, 200)

        warehouse_a_page = self.client.get(f"/inventory?warehouse_id={warehouse.id}&sku_keyword={suffix}")
        warehouse_a_body = warehouse_a_page.get_data(as_text=True)
        self.assertEqual(warehouse_a_page.status_code, 200)
        self.assertIn(">15</td>", warehouse_a_body)
        self.assertIn(">2</td>", warehouse_a_body)
        self.assertIn(">3</td>", warehouse_a_body)
        self.assertNotIn(">7</td>", warehouse_a_body)

        warehouse_b_page = self.client.get(f"/inventory?warehouse_id={second_warehouse.id}&sku_keyword={suffix}")
        warehouse_b_body = warehouse_b_page.get_data(as_text=True)
        self.assertEqual(warehouse_b_page.status_code, 200)
        self.assertIn(">1</td>", warehouse_b_body)
        self.assertIn(">7</td>", warehouse_b_body)

        export_response = self.client.get(f"/reports/export/inventory?warehouse_id={warehouse.id}&sku_keyword={suffix}")
        self.assertEqual(export_response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export_response.data))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        flattened = "\n".join("" if cell is None else str(cell) for row in rows for cell in row)
        self.assertIn(warehouse.name, flattened)
        self.assertNotIn(second_warehouse.name, flattened)

    def test_defect_status_update_rebuilds_inventory_transactions(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "12",
                "reference_type": "defect-update-test",
                "reference_no": f"DEFECT-UP-{suffix}",
                "note": "inventory for defect update",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        create_response = self.client.post(
            "/defects",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "5",
                "status": "pending",
                "defect_reason": "update flow defect",
                "repair_result": "",
            },
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)

        with self.app.app_context():
            repair = self.app_module.DefectRepair.query.order_by(self.app_module.DefectRepair.id.desc()).first()
            repair_id = repair.id
            repair_no = repair.repair_no
            pending_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            pending_txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertEqual(pending_balance.quantity, 7)
            self.assertEqual(len(pending_txns), 1)
            self.assertEqual(pending_txns[0].transaction_type, "defect_hold")

        update_response = self.client.post(
            f"/defects/{repair_id}/update",
            data={
                "csrf_token": csrf,
                "status": "done",
                "completed_quantity": "0",
                "scrapped_quantity": "0",
                "repair_result": "repaired successfully",
            },
            follow_redirects=True,
        )
        self.assertEqual(update_response.status_code, 200)

        with self.app.app_context():
            done_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            done_txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).order_by(self.app_module.InventoryTransaction.id.asc()).all()
            self.assertEqual(done_balance.quantity, 12)
            self.assertEqual([txn.transaction_type for txn in done_txns], ["defect_hold", "defect_return"])

        scrap_response = self.client.post(
            f"/defects/{repair_id}/update",
            data={
                "csrf_token": csrf,
                "status": "scrapped",
                "repair_result": "cannot repair",
            },
            follow_redirects=True,
        )
        self.assertEqual(scrap_response.status_code, 200)

        with self.app.app_context():
            scrapped_balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            scrapped_txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertEqual(scrapped_balance.quantity, 7)
            self.assertEqual(len(scrapped_txns), 1)
            self.assertEqual(scrapped_txns[0].transaction_type, "defect_scrap")

    def test_defect_repair_allows_partial_completion_and_scrap(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "34",
                "reference_type": "defect-partial-test",
                "reference_no": f"DEFECT-PARTIAL-IN-{suffix}",
                "note": "inventory for partial defect repair",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        create_response = self.client.post(
            "/defects",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "34",
                "status": "repairing",
                "defect_reason": "partial repair flow defect",
                "repair_result": "",
            },
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)

        with self.app.app_context():
            repair = self.app_module.DefectRepair.query.order_by(self.app_module.DefectRepair.id.desc()).first()
            repair_id = repair.id
            repair_no = repair.repair_no
            balance_after_create = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertEqual(balance_after_create.quantity, 0)

        update_response = self.client.post(
            f"/defects/{repair_id}/update",
            data={
                "csrf_token": csrf,
                "status": "done",
                "completed_quantity": "32",
                "scrapped_quantity": "2",
                "repair_result": "32 repaired, 2 scrapped",
            },
            follow_redirects=True,
        )
        self.assertEqual(update_response.status_code, 200)
        update_body = update_response.get_data(as_text=True)
        self.assertIn("部分完成/部分报废", update_body)
        self.assertIn("成功 32", update_body)
        self.assertIn("报废 2", update_body)

        with self.app.app_context():
            updated_repair = self.app_module.DefectRepair.query.get(repair_id)
            balance_after_update = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).order_by(self.app_module.InventoryTransaction.id.asc()).all()
            self.assertEqual(updated_repair.completed_quantity, 32)
            self.assertEqual(updated_repair.scrapped_quantity, 2)
            self.assertEqual(updated_repair.remaining_quantity, 0)
            self.assertEqual(balance_after_update.quantity, 32)
            self.assertEqual([txn.transaction_type for txn in txns], ["defect_hold", "defect_return"])
            self.assertEqual([txn.quantity for txn in txns], [-34, 32])

        inventory_page = self.client.get(f"/inventory?sku_keyword={suffix}")
        inventory_body = inventory_page.get_data(as_text=True)
        self.assertEqual(inventory_page.status_code, 200)
        self.assertIn("32", inventory_body)
        self.assertIn("2", inventory_body)

    def test_defect_repair_can_be_deleted_and_rolls_back_inventory(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "10",
                "reference_type": "defect-delete-test",
                "reference_no": f"DEFECT-DEL-IN-{suffix}",
                "note": "inventory for defect delete",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        create_response = self.client.post(
            "/defects",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "4",
                "status": "pending",
                "defect_reason": "delete defect repair",
                "repair_result": "",
            },
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)

        with self.app.app_context():
            repair = self.app_module.DefectRepair.query.order_by(self.app_module.DefectRepair.id.desc()).first()
            self.assertIsNotNone(repair)
            repair_id = repair.id
            repair_no = repair.repair_no
            balance_after_create = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            repair_txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertEqual(balance_after_create.quantity, 6)
            self.assertEqual(len(repair_txns), 1)

        delete_response = self.client.post(
            f"/defects/{repair_id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(delete_response.status_code, 200)
        self.assertIn("返修单已删除", delete_response.get_data(as_text=True))

        with self.app.app_context():
            deleted_repair = self.app_module.DefectRepair.query.get(repair_id)
            balance_after_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            repair_txns = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertIsNone(deleted_repair)
            self.assertEqual(balance_after_delete.quantity, 10)
            self.assertEqual(repair_txns, [])

    def test_parse_reserved_inventory_report_supports_header_based_columns(self):
        report = (
            "站点,店铺SKU,商品名称,可售,处理中,预留库存,备注\n"
            "US,E1081T-Navy-M,Test Product,0,0,12,\n"
            "US,E1081T-Navy-L,Test Product,0,0,7,\n"
        ).encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="reserved.csv")

        parsed = self.app_module.parse_reserved_inventory_report(file_storage)

        self.assertEqual(parsed["E1081T-Navy-M"], 12)
        self.assertEqual(parsed["E1081T-Navy-L"], 7)

    def test_parse_reserved_inventory_report_prefers_reserved_fc_transfers_column(self):
        report = (
            "merchant-sku,reserved-customerorders,reserved_fc-transfers,reserved_fc-processing\n"
            "E1081T-Navy-M,99,12,88\n"
            "E1081T-Navy-L,77,7,66\n"
        ).encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="reserved.csv")

        parsed = self.app_module.parse_reserved_inventory_report(file_storage)

        self.assertEqual(parsed["E1081T-Navy-M"], 12)
        self.assertEqual(parsed["E1081T-Navy-L"], 7)

    def test_parse_reserved_inventory_report_handles_decimal_and_thousand_formats(self):
        report = (
            "merchant-sku,reserved_fc-transfers\n"
            "E1081T-Black-S,12.0\n"
            "E1081T-Black-M,\"1,234\"\n"
        ).encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="reserved.csv")

        parsed = self.app_module.parse_reserved_inventory_report(file_storage)

        self.assertEqual(parsed["E1081T-Black-S"], 12)
        self.assertEqual(parsed["E1081T-Black-M"], 1234)

    def test_parse_reserved_inventory_report_supports_cp1252_csv_exports(self):
        report = (
            '"sku","fnsku","asin","product-name","reserved_qty","reserved_customerorders","reserved_fc-transfers","reserved_fc-processing","program"\n'
            '"E1081T-White-S","X004ZKRP3R","B0GJKPFT66","Women\'s Dress","2","2","1","0",""\n'
            '"E1081T-Red-S","X004ZU0PEX","B0GKNWRFNY","Holiday Business’ Dress","1","1","2","0",""\n'
        ).encode("cp1252")
        file_storage = FileStorage(stream=BytesIO(report), filename="reserved.csv")

        parsed = self.app_module.parse_reserved_inventory_report(file_storage)

        self.assertEqual(parsed["E1081T-White-S"], 1)
        self.assertEqual(parsed["E1081T-Red-S"], 2)

    def test_parse_uploaded_table_prefers_utf8_tsv_over_latin1_mojibake(self):
        shipment_no_label = "\u8d27\u4ef6\u7f16\u53f7"
        delivery_label = "\u914d\u9001\u5730\u5740"
        seller_sku_label = "\u5356\u5bb6 SKU"
        product_name_label = "\u5546\u54c1\u540d\u79f0"
        shipped_label = "\u5df2\u53d1\u8d27"
        report = (
            f"{shipment_no_label}\tFBA19D64Q6VT\n"
            f"{delivery_label}\tABQ2\n"
            "\n"
            f"{seller_sku_label}\t{product_name_label}\t{shipped_label}\n"
            "STY1006-LightBlue-XL\tTest Product\t4\n"
        ).encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="FBA19D64Q6VT.tsv")

        rows = self.app_module.parse_uploaded_table(file_storage)

        self.assertEqual(rows[0], [shipment_no_label, "FBA19D64Q6VT"])
        self.assertIn([seller_sku_label, product_name_label, shipped_label], rows)

    def test_parse_reserved_inventory_report_falls_back_to_legacy_column_positions(self):
        report = (
            "店铺SKU,a,b,c,d,e,预留库存\n"
            "E1081T-Navy-XL,1,2,3,4,5,9\n"
            "E1081T-Navy-XXL,1,2,3,4,5,6\n"
        ).encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="reserved.csv")

        parsed = self.app_module.parse_reserved_inventory_report(file_storage)

        self.assertEqual(parsed["E1081T-Navy-XL"], 9)
        self.assertEqual(parsed["E1081T-Navy-XXL"], 6)

    def test_parse_sales_velocity_windows_derives_7_14_30_from_one_dated_file(self):
        rows = ["purchase-date,merchant-sku,quantity-purchased,marketplace"]
        rows.append("2026-03-31,WINDOW-SKU,1,US")
        for day in range(1, 30):
            rows.append(f"2026-04-{day:02d},WINDOW-SKU,1,US")
        rows.append("2026-04-30,WINDOW-SKU,99,US")
        rows.append("2026-04-30,WINDOW-SKU,5,CA")
        report = ("\n".join(rows) + "\n").encode("utf-8")
        file_storage = FileStorage(stream=BytesIO(report), filename="sales.csv")

        parsed = self.app_module.parse_sales_velocity_windows(
            file_storage,
            marketplace="US",
            as_of_datetime=datetime(2026, 4, 30, 15, tzinfo=timezone.utc),
        )

        self.assertEqual(parsed["latest_sale_date"], "2026-04-30")
        self.assertEqual(parsed["window_end_date"], "2026-04-29")
        self.assertEqual(parsed["sales_7"]["WINDOW-SKU"], 7)
        self.assertEqual(parsed["sales_14"]["WINDOW-SKU"], 14)
        self.assertEqual(parsed["sales_30"]["WINDOW-SKU"], 30)

    def test_parse_sales_velocity_windows_matches_canada_order_report(self):
        report = (
            "amazon-order-id\tmerchant-order-id\tpurchase-date\tlast-updated-date\torder-status\tfulfillment-channel\tsales-channel\tsku\tquantity\tship-country\n"
            "1\t1\t2026-04-26T12:39:37+00:00\t2026-04-26T12:41:20+00:00\tShipped\tAmazon\tAmazon.ca\tCA-SKU\t2\tUS\n"
            "2\t2\t2026-04-27T12:39:37+00:00\t2026-04-27T12:41:20+00:00\tShipped\tAmazon\tAmazon.ca\tCA-SKU\t99\tUS\n"
            "2\t2\t2026-04-27T12:39:37+00:00\t2026-04-27T12:41:20+00:00\tShipped\tAmazon\tAmazon.com\tUS-SKU\t3\tUS\n"
        ).encode("utf-8")

        parsed = self.app_module.parse_sales_velocity_windows(
            FileStorage(stream=BytesIO(report), filename="sales.txt"),
            marketplace="CA",
            as_of_datetime=datetime(2026, 4, 27, 15, tzinfo=timezone.utc),
        )

        self.assertEqual(parsed["window_end_date"], "2026-04-26")
        self.assertEqual(parsed["sales_30"], {"CA-SKU": 2})

    def test_calculate_fba_replenishment_uses_coverage_days(self):
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            sku = self.SKU.query.first()
            warehouse = self.app_module.get_default_warehouse()
            external_sku = f"COVERAGE-{self.unique_suffix()}"
            mapping = self.app_module.SKUMapping(user_id=admin.id, sku_id=sku.id, external_sku_code=external_sku)
            self.db.session.add(mapping)
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=15,
                    transaction_type="initial_inbound",
                    reference_type="test",
                    reference_no=external_sku,
                    note="coverage-days regression",
                )
            )
            self.db.session.commit()
            self.app_module.sync_inventory_balances()

            result_30 = self.app_module.calculate_fba_replenishment(
                inventory_data={external_sku: {"sellable_qty": 0, "received_qty": 0}},
                reserved_data={},
                sales_7_data={external_sku: 7},
                sales_14_data={},
                sales_30_data={external_sku: 30},
                inbound_data={},
                multiplier=1,
                new_product_multiplier=1.5,
                coverage_days=30,
                user=admin,
            )
            result_10 = self.app_module.calculate_fba_replenishment(
                inventory_data={external_sku: {"sellable_qty": 0, "received_qty": 0}},
                reserved_data={},
                sales_7_data={external_sku: 7},
                sales_14_data={},
                sales_30_data={external_sku: 30},
                inbound_data={},
                multiplier=1,
                new_product_multiplier=1.5,
                coverage_days=10,
                user=admin,
            )

        self.assertEqual(result_30["rows"][0]["suggested_qty"], 30)
        self.assertEqual(result_10["rows"][0]["suggested_qty"], 10)

    def test_calculate_fba_replenishment_marks_stockout_drop_without_suppressing_forecast(self):
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            sku = self.SKU.query.first()
            external_sku = f"STOCKOUT-{self.unique_suffix()}"
            mapping = self.app_module.SKUMapping(user_id=admin.id, sku_id=sku.id, external_sku_code=external_sku)
            self.db.session.add(mapping)
            self.db.session.commit()

            result = self.app_module.calculate_fba_replenishment(
                inventory_data={external_sku: {"sellable_qty": 0, "received_qty": 0}},
                reserved_data={},
                sales_7_data={external_sku: 1},
                sales_14_data={external_sku: 14},
                sales_30_data={external_sku: 30},
                inbound_data={},
                multiplier=1,
                new_product_multiplier=1.5,
                coverage_days=30,
                user=admin,
            )

        row = result["rows"][0]
        self.assertEqual(row["sales_health_status"], "疑似断货")
        self.assertTrue(row["stockout_suspected"])
        self.assertEqual(row["trend_factor"], 1)
        self.assertEqual(row["suggested_qty"], 30)

    def test_calculate_fba_replenishment_includes_14_day_sales(self):
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            sku = self.SKU.query.first()
            external_sku = f"TREND-{self.unique_suffix()}"
            mapping = self.app_module.SKUMapping(user_id=admin.id, sku_id=sku.id, external_sku_code=external_sku)
            self.db.session.add(mapping)
            self.db.session.commit()

            result = self.app_module.calculate_fba_replenishment(
                inventory_data={external_sku: {"sellable_qty": 0, "received_qty": 0}},
                reserved_data={},
                sales_7_data={external_sku: 7},
                sales_14_data={external_sku: 28},
                sales_30_data={external_sku: 30},
                inbound_data={},
                multiplier=1,
                new_product_multiplier=1.5,
                coverage_days=30,
                user=admin,
            )

        row = result["rows"][0]
        self.assertEqual(row["sales_14_qty"], 28)
        self.assertEqual(row["daily_14"], 2)
        self.assertEqual(row["weighted_daily"], 1.6)
        self.assertEqual(row["trend_factor"], 0.7)
        self.assertEqual(row["suggested_qty"], 34)

    def test_purchase_calculator_uses_stock_open_purchase_and_outbound_consumption(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        with self.app.app_context():
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=10,
                    transaction_type="initial_inbound",
                    reference_type="test",
                    reference_no=f"INIT-{suffix}",
                    note="purchase calculator stock",
                )
            )
            open_order = self.PurchaseOrder(
                order_no=f"PO-OPEN-{suffix}",
                supplier_id=supplier.id,
                warehouse_id=warehouse.id,
                operator_name="Tester",
                status="submitted",
            )
            open_order.items.append(self.app_module.PurchaseOrderItem(sku_id=sku.id, quantity=5, unit_price=sku.cost_price))
            self.db.session.add(open_order)
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=-14,
                    transaction_type="manual_outbound",
                    reference_type="test",
                    reference_no=f"OUT-{suffix}",
                    note="purchase calculator consumption",
                )
            )
            self.db.session.commit()
            self.app_module.sync_inventory_balances()

            result = self.app_module.build_purchase_calculator_result(
                warehouse_id=warehouse.id,
                coverage_days=10,
                lead_days=5,
                sku_keyword=sku.sku_code,
                suggested_min=0,
                positive_only=True,
                user=None,
            )

        row = result["rows"][0]
        self.assertEqual(row["current_stock"], -4)
        self.assertEqual(row["open_purchase_qty"], 5)
        self.assertEqual(row["outbound_7"], 14)
        self.assertEqual(row["daily_consumption"], 2)
        self.assertEqual(row["suggested_qty"], 34)

    def test_purchase_calculator_uses_uploaded_sales_trend_for_demand(self):
        self.login()
        suffix = self.unique_suffix()
        _supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        with self.app.app_context():
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=50,
                    transaction_type="initial_inbound",
                    reference_type="test",
                    reference_no=f"STOCK-{suffix}",
                    note="purchase calculator sales trend stock",
                )
            )
            self.db.session.commit()
            self.app_module.sync_inventory_balances()

            result = self.app_module.build_purchase_calculator_result(
                warehouse_id=warehouse.id,
                coverage_days=10,
                lead_days=5,
                sku_keyword=sku.sku_code,
                suggested_min=0,
                positive_only=True,
                user=None,
                sales_windows={
                    "sales_7": {sku.id: 70},
                    "sales_14": {sku.id: 84},
                    "sales_30": {sku.id: 100},
                },
            )

        row = result["rows"][0]
        self.assertEqual(row["current_stock"], 50)
        self.assertEqual(row["sales_30"], 100)
        self.assertEqual(row["sales_weighted_daily"], 4.93)
        self.assertEqual(row["sales_trend_factor"], 1.3)
        self.assertEqual(row["daily_consumption"], 6.41)
        self.assertEqual(row["demand_source"], "销量趋势")
        self.assertEqual(row["suggested_qty"], 53)

    def test_purchase_calculator_limits_rows_to_uploaded_sales_skus(self):
        self.login()
        suffix = self.unique_suffix()
        _supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        with self.app.app_context():
            unrelated_sku = self.SKU(
                sku_code=f"UNRELATED-{suffix}",
                barcode=f"UNRELATED{suffix}",
                name="Unrelated Product",
                category="Test",
                color="Blue",
                size="S",
                unit="pcs",
                cost_price=Decimal("10.00"),
                sale_price=Decimal("20.00"),
                safety_stock=5,
                remark="purchase calculator unrelated sku",
            )
            self.db.session.add(unrelated_sku)
            self.db.session.commit()

            result = self.app_module.build_purchase_calculator_result(
                warehouse_id=warehouse.id,
                coverage_days=10,
                lead_days=5,
                sku_keyword="",
                suggested_min=0,
                positive_only=False,
                user=None,
                sales_windows={
                    "sales_7": {sku.id: 7},
                    "sales_14": {sku.id: 14},
                    "sales_30": {sku.id: 30},
                },
            )

        self.assertEqual([row["sku_id"] for row in result["rows"]], [sku.id])
        self.assertEqual(result["summary"]["sku_count"], 1)
        self.assertEqual(result["summary"]["sales_30"], 30)

    def test_purchase_calculator_page_calculates_and_renders(self):
        self.login()
        suffix = self.unique_suffix()
        _supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()
        with self.app.app_context():
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=3,
                    transaction_type="initial_inbound",
                    reference_type="test",
                    reference_no=f"PAGE-{suffix}",
                    note="purchase calculator render",
                )
            )
            self.db.session.commit()
            self.app_module.sync_inventory_balances()

        response = self.client.post(
            "/purchase-calculator",
            data={
                "csrf_token": csrf,
                "warehouse_id": str(warehouse.id),
                "coverage_days": "30",
                "lead_days": "15",
                "sales_30_file": (
                    BytesIO(
                        (
                            "purchase-date,merchant-sku,quantity-purchased\n"
                            f"{(datetime.now(timezone.utc).date() - timedelta(days=1)).isoformat()},{sku.sku_code},1\n"
                        ).encode("utf-8")
                    ),
                    "purchase-sales.csv",
                ),
                "sku_keyword": sku.sku_code,
                "suggested_min": "",
                "positive_only": "0",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn(sku.sku_code, body)
        with self.client.session_transaction() as session:
            self.assertNotIn("_purchase_calc_result", session)
            self.assertIn("_purchase_calc_result_key", session)
        self.assertEqual(self.current_purchase_calc_result()["rows"][0]["sku_id"], sku.id)

    def test_purchase_calculator_generates_purchase_order_draft(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()
        with self.client.session_transaction() as session:
            session["_purchase_calc_result"] = {
                "rows": [
                    {
                        "sku_id": sku.id,
                        "sku_code": sku.sku_code,
                        "name": sku.name,
                        "suggested_qty": 7,
                    }
                ],
                "summary": {"sku_count": 1, "suggested_sku_count": 1, "suggested_qty": 7, "current_stock": 0, "open_purchase_qty": 0},
                "params": {"warehouse_id": warehouse.id},
            }

        export_response = self.client.get("/purchase-calculator/export")
        self.assertEqual(export_response.status_code, 200)
        with self.client.session_transaction() as session:
            self.assertNotIn("_purchase_calc_result", session)
            self.assertIn("_purchase_calc_result_key", session)

        response = self.client.post(
            "/purchase-calculator/generate-order",
            data={
                "csrf_token": csrf,
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "purchase_calc_sku_id": str(sku.id),
                f"purchase_calc_quantity__{sku.id}": "7",
                "remark": "calculator draft",
            },
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            order = self.PurchaseOrder.query.order_by(self.PurchaseOrder.id.desc()).first()
            self.assertEqual(order.status, "draft")
            self.assertEqual(order.supplier_id, supplier.id)
            self.assertEqual(order.warehouse_id, warehouse.id)
            self.assertEqual(order.items[0].sku_id, sku.id)
            self.assertEqual(order.items[0].quantity, 7)

        export_order_response = self.client.get(f"/purchase-orders/{order.id}/export")
        self.assertEqual(export_order_response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", export_order_response.content_type)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export_order_response.data))
        sheet = workbook.active
        self.assertEqual(sheet["A1"].value, f"采购订单 {order.order_no}")
        exported_values = [cell.value for row in sheet.iter_rows(values_only=False) for cell in row]
        self.assertIn(sku.sku_code, exported_values)
        self.assertIn(7, exported_values)

    def test_purchase_order_import_template_and_preview(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        template_response = self.client.get("/purchase-orders/import-template")
        self.assertEqual(template_response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", template_response.content_type)
        from openpyxl import Workbook, load_workbook

        template_workbook = load_workbook(BytesIO(template_response.data))
        self.assertEqual(template_workbook.sheetnames, ["采购明细", "填写说明", "SKU参考"])
        self.assertEqual(template_workbook["采购明细"]["A1"].value, "SKU编码（必填）")
        reference_values = [cell.value for cell in template_workbook["SKU参考"]["A"]]
        self.assertIn(sku.sku_code, reference_values)

        upload_workbook = Workbook()
        upload_sheet = upload_workbook.active
        upload_sheet.append(["SKU编码", "采购数量", "采购单价"])
        upload_sheet.append([sku.sku_code, 12, 24.5])
        upload_stream = BytesIO()
        upload_workbook.save(upload_stream)
        upload_stream.seek(0)

        response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "action": "import",
                "order_no": f"PO-IMPORT-{suffix}",
                "expected_date": "2026-06-30",
                "supplier_id": str(supplier.id),
                "warehouse_id": str(warehouse.id),
                "status": "draft",
                "remark": "import preview",
                "purchase_import_file": (upload_stream, "purchase-import.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("已读取 1 条采购明细", body)
        self.assertIn(sku.sku_code, body)
        self.assertIn('value="12"', body)
        self.assertIn('value="24.50"', body)
        self.assertIn(f'option value="{supplier.id}" selected', body)
        with self.app.app_context():
            self.assertIsNone(self.PurchaseOrder.query.filter_by(order_no=f"PO-IMPORT-{suffix}").first())

    def test_purchase_order_import_rejects_unknown_sku(self):
        self.login()
        csrf = self.current_csrf()
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["SKU编码", "采购数量"])
        sheet.append(["UNKNOWN-SKU", 3])
        upload_stream = BytesIO()
        workbook.save(upload_stream)
        upload_stream.seek(0)

        response = self.client.post(
            "/purchase-orders",
            data={
                "csrf_token": csrf,
                "action": "import",
                "purchase_import_file": (upload_stream, "purchase-import.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("不存在或当前账号无权使用", response.get_data(as_text=True))

    def test_fba_rule_batch_mark_new_does_not_error(self):
        self.login()
        csrf = self.current_csrf()
        external_sku = f"BATCH-{self.unique_suffix()}"
        with self.client.session_transaction() as session:
            session["_fba_calc_form"] = {
                "multiplier": "1",
                "new_product_multiplier": "1.5",
                "coverage_days": "30",
                "sea_days": "30",
                "marketplace": "US",
            }
            session["_fba_calc_result"] = {
                "rows": [
                    {
                        "external_sku_code": external_sku,
                        "sellable_qty": 0,
                        "inbound_qty": 0,
                        "daily_7": 1,
                        "daily_30": 1,
                        "trend_factor": 1,
                    }
                ],
                "summary": {"sku_count": 1, "suggested_qty": 0, "sellable_qty": 0, "inbound_qty": 0},
            }

        response = self.client.post(
            "/fba-calculator/rules/batch",
            data={
                "csrf_token": csrf,
                "external_sku_codes": external_sku,
                "batch_action": "mark_new",
                "shipment_store_name": "",
                "shipment_box_count": "1",
                "shipment_remark": "",
            },
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            rule = self.app_module.FBAReplenishmentRule.query.filter_by(external_sku_code=external_sku).first()
            self.assertIsNotNone(rule)
            self.assertTrue(rule.is_new_product)

    def test_fba_rule_batch_refresh_calc_updates_session_parameters(self):
        self.login()
        csrf = self.current_csrf()
        external_sku = f"REFRESH-{self.unique_suffix()}"
        with self.client.session_transaction() as session:
            session["_fba_calc_form"] = {
                "multiplier": "1",
                "new_product_multiplier": "1.5",
                "coverage_days": "30",
                "sea_days": "30",
                "marketplace": "US",
            }
            session["_fba_calc_result"] = {
                "rows": [
                    {
                        "external_sku_code": external_sku,
                        "sellable_qty": 0,
                        "sales_7_qty": 7,
                        "sales_14_qty": 28,
                        "sales_30_qty": 30,
                        "inbound_qty": 0,
                        "daily_7": 1,
                        "daily_14": 2,
                        "daily_30": 1,
                        "trend_factor": 1,
                        "suggested_qty": 30,
                    }
                ],
                "summary": {"sku_count": 1, "suggested_qty": 30, "sellable_qty": 0, "inbound_qty": 0},
            }

        response = self.client.post(
            "/fba-calculator/rules/batch",
            data={
                "csrf_token": csrf,
                "batch_action": "refresh_calc",
                "multiplier": "2",
                "new_product_multiplier": "3",
                "coverage_days": "10",
                "production_days": "0",
                "sea_days": "10",
                "inbound_processing_days": "0",
                "safety_stock_days": "0",
                "shipment_store_name": "",
                "shipment_box_count": "1",
                "shipment_remark": "",
            },
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 302)
        with self.client.session_transaction() as session:
            self.assertEqual(session["_fba_calc_form"]["multiplier"], "2")
            self.assertEqual(session["_fba_calc_form"]["new_product_multiplier"], "3")
            self.assertEqual(session["_fba_calc_form"]["coverage_days"], "10")
        refreshed_result = self.current_fba_calc_result()
        self.assertEqual(refreshed_result["rows"][0]["daily_14"], 2)
        self.assertEqual(refreshed_result["rows"][0]["weighted_daily"], 1.6)
        self.assertEqual(refreshed_result["rows"][0]["target_days"], 10)
        self.assertEqual(refreshed_result["rows"][0]["trend_factor"], 0.7)
        self.assertEqual(refreshed_result["rows"][0]["suggested_qty"], 23)

    def test_fba_calculator_page_refreshes_inbound_qty_from_active_records(self):
        self.login()
        suffix = self.unique_suffix()
        external_sku = f"INBOUND-{suffix}"
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            self.db.session.add(
                self.app_module.FBAInboundRecord(
                    shipment_ref=f"SHIP-{suffix}",
                    external_sku_code=external_sku,
                    quantity=12,
                    status="active",
                    uploader_user_id=admin.id,
                )
            )
            self.db.session.commit()

        with self.client.session_transaction() as session:
            session["_fba_calc_form"] = {
                "multiplier": "1",
                "new_product_multiplier": "1.5",
                "coverage_days": "30",
                "sea_days": "30",
                "marketplace": "US",
            }
            session["_fba_calc_result"] = {
                "rows": [
                    {
                        "external_sku_code": external_sku,
                        "sellable_qty": 0,
                        "sales_7_qty": 7,
                        "sales_14_qty": 28,
                        "sales_30_qty": 30,
                        "inbound_qty": 0,
                        "daily_7": 1,
                        "daily_14": 2,
                        "daily_30": 1,
                        "trend_factor": 0.8,
                        "suggested_qty": 48,
                    }
                ],
                "summary": {"sku_count": 1, "suggested_qty": 48, "sellable_qty": 0, "inbound_qty": 0},
            }

        response = self.client.get("/fba-calculator")
        self.assertEqual(response.status_code, 200)

        refreshed_result = self.current_fba_calc_result()
        refreshed_row = refreshed_result["rows"][0]
        self.assertEqual(refreshed_row["inbound_qty"], 12)
        self.assertEqual(refreshed_row["suggested_qty"], 22)
        self.assertEqual(refreshed_result["summary"]["inbound_qty"], 12)

    def test_fba_filter_and_refresh_use_server_side_cached_result(self):
        self.login()
        csrf = self.current_csrf()
        keep_sku = f"KEEP-{self.unique_suffix()}"
        drop_sku = f"DROP-{self.unique_suffix()}"
        with self.client.session_transaction() as session:
            session["_fba_calc_form"] = {
                "multiplier": "1",
                "new_product_multiplier": "1.5",
                "coverage_days": "30",
                "sea_days": "30",
                "marketplace": "US",
            }
            session["_fba_calc_result"] = {
                "rows": [
                    {
                        "external_sku_code": keep_sku,
                        "warehouse_sku_code": keep_sku,
                        "warehouse_sku_name": "Keep product",
                        "sellable_qty": 0,
                        "sales_7_qty": 7,
                        "sales_14_qty": 28,
                        "sales_30_qty": 30,
                        "inbound_qty": 0,
                        "daily_7": 1,
                        "daily_14": 2,
                        "daily_30": 1,
                        "trend_factor": 0.8,
                        "suggested_qty": 48,
                    },
                    {
                        "external_sku_code": drop_sku,
                        "warehouse_sku_code": drop_sku,
                        "warehouse_sku_name": "Drop product",
                        "sellable_qty": 0,
                        "sales_7_qty": 7,
                        "sales_14_qty": 14,
                        "sales_30_qty": 30,
                        "inbound_qty": 0,
                        "daily_7": 1,
                        "daily_14": 1,
                        "daily_30": 1,
                        "trend_factor": 1,
                        "suggested_qty": 30,
                    },
                ],
                "summary": {"sku_count": 2, "suggested_qty": 78, "sellable_qty": 0, "inbound_qty": 0},
            }

        filter_response = self.client.get(f"/fba-calculator?sku_keyword={keep_sku}")
        self.assertEqual(filter_response.status_code, 200)
        body = filter_response.get_data(as_text=True)
        self.assertIn(keep_sku, body)
        self.assertNotIn(drop_sku, body)
        with self.client.session_transaction() as session:
            self.assertNotIn("_fba_calc_result", session)
            self.assertIn("_fba_calc_result_key", session)

        refresh_response = self.client.post(
            "/fba-calculator/rules/batch",
            data={
                "csrf_token": csrf,
                "batch_action": "refresh_calc",
                "multiplier": "2",
                "new_product_multiplier": "1.5",
                "coverage_days": "10",
                "production_days": "0",
                "sea_days": "10",
                "inbound_processing_days": "0",
                "safety_stock_days": "0",
                "shipment_store_name": "",
                "shipment_box_count": "1",
                "shipment_remark": "",
            },
            follow_redirects=False,
        )
        self.assertEqual(refresh_response.status_code, 302)
        refreshed_result = self.current_fba_calc_result()
        self.assertEqual(len(refreshed_result["rows"]), 2)
        self.assertGreater(refreshed_result["summary"]["suggested_qty"], 0)

    def test_fba_generate_shipment_allows_zero_or_blank_quantities(self):
        self.login()
        csrf = self.current_csrf()
        suffix = self.unique_suffix()
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            sku = self.SKU.query.first()
            sku_id = sku.id
            sku_code = sku.sku_code
            mapping = self.app_module.SKUMapping(user_id=admin.id, sku_id=sku.id, external_sku_code=f"SHIP-{suffix}")
            self.db.session.add(mapping)
            self.db.session.commit()

        with self.client.session_transaction() as session:
            session["_fba_calc_form"] = {
                "multiplier": "1",
                "new_product_multiplier": "1.5",
                "coverage_days": "30",
                "sea_days": "30",
                "marketplace": "US",
                "shipment_store_name": "Amazon-US",
                "shipment_box_count": "5",
                "shipment_remark": "",
            }
            session["_fba_calc_result"] = {
                "rows": [
                    {
                        "external_sku_code": f"SHIP-{suffix}",
                        "warehouse_sku_code": sku_code,
                        "warehouse_qty": 30,
                        "sku_id": sku_id,
                    }
                ],
                "summary": {"sku_count": 1, "suggested_qty": 0, "sellable_qty": 0, "inbound_qty": 0},
            }

        blank_response = self.client.post(
            "/fba-calculator/generate-shipment",
            data={
                "csrf_token": csrf,
                "external_sku_codes": f"SHIP-{suffix}",
                "shipment_store_name": "Amazon-US",
                "shipment_box_count": "5",
                "shipment_remark": "",
                f"shipment_quantity__SHIP-{suffix}": "",
            },
            follow_redirects=False,
        )
        self.assertEqual(blank_response.status_code, 302)
        with self.client.session_transaction() as session:
            flashes = [message for _category, message in session.get("_flashes", [])]
        self.assertIn("请至少填写一个大于 0 的发货数量。填空或填 0 的行会自动跳过。", flashes)
        self.assertFalse(any("发货数量必须大于 0" in message for message in flashes))

        zero_response = self.client.post(
            "/fba-calculator/generate-shipment",
            data={
                "csrf_token": csrf,
                "external_sku_codes": f"SHIP-{suffix}",
                "shipment_store_name": "Amazon-US",
                "shipment_box_count": "5",
                "shipment_remark": "",
                f"shipment_quantity__SHIP-{suffix}": "0",
            },
            follow_redirects=False,
        )
        self.assertEqual(zero_response.status_code, 302)
        with self.client.session_transaction() as session:
            flashes = [message for _category, message in session.get("_flashes", [])]
        self.assertIn("请至少填写一个大于 0 的发货数量。填空或填 0 的行会自动跳过。", flashes)
        self.assertFalse(any("发货数量必须大于 0" in message for message in flashes))

    def test_fba_shipment_record_defaults_collapsed_and_delete_rolls_back_stock(self):
        self.login()
        csrf = self.current_csrf()
        suffix = self.unique_suffix()
        _supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            self.db.session.add(
                self.InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=warehouse.id,
                    quantity=10,
                    transaction_type="initial_inbound",
                    reference_type="test",
                    reference_no=f"SHIP-ROLLBACK-{suffix}",
                    note="shipment rollback test",
                )
            )
            shipment = self.app_module.ShipmentSheet(
                shipment_no=f"SH-ROLL-{suffix}",
                store_name="Amazon-US",
                warehouse_id=warehouse.id,
                operator_user_id=admin.id,
                operator_name=admin.full_name,
                box_count=1,
                status="pending",
            )
            shipment.items.append(
                self.app_module.ShipmentSheetItem(
                    sku_id=sku.id,
                    external_sku_code=f"SHIP-ROLL-{suffix}",
                    quantity=4,
                    per_box_quantity=4,
                )
            )
            self.db.session.add(shipment)
            self.db.session.commit()
            self.app_module.sync_inventory_balances()
            shipment_id = shipment.id

        page = self.client.get("/fba-calculator")
        self.assertEqual(page.status_code, 200)
        body = page.get_data(as_text=True)
        self.assertNotIn('<details class="stack-item order-fold shipment-fold" open>', body)
        self.assertIn("保存修改", body)
        self.assertIn("删除", body)

        confirm_response = self.client.post(
            f"/shipment-sheets/{shipment_id}/confirm",
            data={"csrf_token": csrf},
            follow_redirects=False,
        )
        self.assertEqual(confirm_response.status_code, 302)
        with self.app.app_context():
            balance_after_confirm = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            self.assertEqual(balance_after_confirm.quantity, 6)

        delete_response = self.client.post(
            f"/shipment-sheets/{shipment_id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=False,
        )
        self.assertEqual(delete_response.status_code, 302)
        with self.app.app_context():
            balance_after_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            shipment_after_delete = self.app_module.ShipmentSheet.query.get(shipment_id)
            rollback_txns = self.InventoryTransaction.query.filter_by(reference_type="shipment_sheet", reference_no=f"SH-ROLL-{suffix}").all()
            self.assertEqual(balance_after_delete.quantity, 10)
            self.assertIsNone(shipment_after_delete)
            self.assertEqual(rollback_txns, [])

    def test_inventory_note_legacy_mojibake_is_cleaned_for_display(self):
        self.assertEqual(
            self.app_module.clean_legacy_display_text("杩斾慨鍗曪細DR2026041601"),
            "返修单：DR2026041601",
        )

    def test_shipment_archive_export_hides_without_stock_rollback(self):
        self.login()
        csrf = self.current_csrf()
        suffix = self.unique_suffix()
        _supplier, _customer, warehouse, sku = self.create_master_data(suffix)
        with self.app.app_context():
            admin = self.User.query.filter_by(username="admin").first()
            self.app_module.add_inventory_transaction(
                sku.id,
                warehouse.id,
                10,
                "initial_inbound",
                reference_type="test",
                reference_no=f"SHIP-ARCHIVE-{suffix}",
                note="shipment archive test",
                operator_name=admin.full_name,
            )
            shipment = self.app_module.ShipmentSheet(
                shipment_no=f"SH-ARCH-{suffix}",
                store_name="Amazon-US",
                warehouse_id=warehouse.id,
                operator_user_id=admin.id,
                operator_name=admin.full_name,
                box_count=1,
                status="confirmed",
                warehouse_confirmer_name=admin.full_name,
            )
            shipment.items.append(
                self.app_module.ShipmentSheetItem(
                    sku_id=sku.id,
                    external_sku_code=f"ARCH-{suffix}",
                    quantity=4,
                    per_box_quantity=4,
                )
            )
            self.db.session.add(shipment)
            self.db.session.flush()
            self.app_module.add_inventory_transaction(
                sku.id,
                warehouse.id,
                -4,
                "shipment_outbound",
                reference_type="shipment_sheet",
                reference_no=shipment.shipment_no,
                note="店铺：Amazon-US",
                operator_name=admin.full_name,
            )
            self.db.session.commit()
            shipment_id = shipment.id

        response = self.client.post(
            "/shipment-sheets/archive-export",
            data={"csrf_token": csrf, "shipment_ids": str(shipment_id)},
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        with self.app.app_context():
            archived = self.db.session.get(self.app_module.ShipmentSheet, shipment_id)
            balance = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            outbound_txns = self.InventoryTransaction.query.filter_by(reference_type="shipment_sheet", reference_no=f"SH-ARCH-{suffix}").all()
            self.assertIsNotNone(archived.archived_at)
            self.assertEqual(balance.quantity, 6)
            self.assertEqual(len(outbound_txns), 1)

        active_page = self.client.get("/fba-calculator")
        self.assertNotIn(f"SH-ARCH-{suffix}", active_page.get_data(as_text=True))
        archived_page = self.client.get("/fba-calculator?shipment_archive=archived")
        self.assertIn(f"SH-ARCH-{suffix}", archived_page.get_data(as_text=True))

    def test_completed_defect_repair_cannot_be_deleted(self):
        self.login()
        suffix = self.unique_suffix()
        supplier, customer, warehouse, sku = self.create_master_data(suffix)
        csrf = self.current_csrf()

        inbound_response = self.client.post(
            "/inventory/in",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "supplier_id": str(supplier.id),
                "quantity": "10",
                "reference_type": "defect-delete-lock-test",
                "reference_no": f"DEFECT-LOCK-IN-{suffix}",
                "note": "inventory for protected defect delete",
            },
            follow_redirects=True,
        )
        self.assertEqual(inbound_response.status_code, 200)

        create_response = self.client.post(
            "/defects",
            data={
                "csrf_token": csrf,
                "sku_id": str(sku.id),
                "warehouse_id": str(warehouse.id),
                "quantity": "3",
                "status": "done",
                "defect_reason": "completed defect repair",
                "repair_result": "restored",
            },
            follow_redirects=True,
        )
        self.assertEqual(create_response.status_code, 200)

        with self.app.app_context():
            repair = self.app_module.DefectRepair.query.order_by(self.app_module.DefectRepair.id.desc()).first()
            self.assertIsNotNone(repair)
            repair_id = repair.id
            repair_no = repair.repair_no
            balance_before_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            repair_txns_before = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertEqual(balance_before_delete.quantity, 10)
            self.assertEqual([txn.transaction_type for txn in repair_txns_before], ["defect_hold", "defect_return"])

        delete_response = self.client.post(
            f"/defects/{repair_id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(delete_response.status_code, 200)
        self.assertIn("只有待处理或返修中的返修单才能删除", delete_response.get_data(as_text=True))

        with self.app.app_context():
            still_exists = self.app_module.DefectRepair.query.get(repair_id)
            balance_after_delete = self.InventoryBalance.query.filter_by(sku_id=sku.id, warehouse_id=warehouse.id).first()
            repair_txns_after = self.app_module.InventoryTransaction.query.filter_by(reference_type="defect_repair", reference_no=repair_no).all()
            self.assertIsNotNone(still_exists)
            self.assertEqual(balance_after_delete.quantity, 10)
            self.assertEqual([txn.transaction_type for txn in repair_txns_after], ["defect_hold", "defect_return"])

    def test_freight_compare_imports_rate_book_and_creates_quote(self):
        from openpyxl import Workbook

        self.login()
        csrf = self.current_csrf()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "汇总表"
        sheet.append(
            [
                "产品名称",
                "仓库代码简称",
                "仓库代码",
                "销售产品代码",
                "12KG+",
                "",
                "",
                "",
                "21KG+",
                "",
                "",
                "",
                "51KG+",
                "",
                "",
                "",
                "100KG+",
                "",
                "",
                "",
                "1CBM+不包税（1:363）",
                "",
                "",
                "",
                "参考时效",
            ]
        )
        sheet.append(
            [
                "",
                "",
                "",
                "",
                "华东",
                "华南",
                "福建",
                "青岛",
                "华东",
                "华南",
                "福建",
                "青岛",
                "华东",
                "华南",
                "福建",
                "青岛",
                "华东",
                "华南",
                "福建",
                "青岛",
                "华东",
                "华南",
                "福建",
                "青岛",
                "",
            ]
        )
        sheet.append(
            [
                "正班盈速达卡派",
                "ONT8",
                "ONT8(92551-9534)",
                "包税：CPYSDKP",
                "14",
                "14.5",
                "14.5",
                "14.5",
                "13",
                "13.5",
                "13.5",
                "13.5",
                "12",
                "12.5",
                "12.5",
                "12.5",
                "11",
                "11.5",
                "11.5",
                "11.5",
                "2000",
                "2100",
                "2100",
                "2100",
                "15天",
            ]
        )
        rate_file = BytesIO()
        workbook.save(rate_file)
        rate_file.seek(0)

        upload_response = self.client.post(
            "/freight-compare/rate-books",
            data={
                "csrf_token": csrf,
                "forwarder_name": "盈和",
                "version_label": f"TEST-{self.unique_suffix()}",
                "rate_file": (rate_file, "yinghe-test.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(upload_response.status_code, 200)

        with self.app.app_context():
            rate_count = self.app_module.FreightRate.query.filter_by(forwarder_name="盈和", warehouse_code="ONT8").count()
            self.assertGreater(rate_count, 0)
            postal_mapping = self.app_module.warehouse_postal_mapping()
            self.assertEqual(postal_mapping.get("ONT8"), "92551")
            postal_row = self.WarehousePostalCode.query.filter_by(warehouse_code="ONT8").first()
            self.assertIsNotNone(postal_row)
            self.assertEqual(postal_row.postal_code, "92551")
            rate_book = self.app_module.FreightRateBook.query.filter_by(forwarder_name="盈和").order_by(self.app_module.FreightRateBook.id.desc()).first()
            self.assertIsNotNone(rate_book)
            rate_book_id = rate_book.id
            stored_path = Path(self.app_module.FREIGHT_RATE_BOOK_DIR) / rate_book.stored_filename
            self.assertTrue(stored_path.exists())

        csrf = self.current_csrf()
        quote_response = self.client.post(
            "/freight-compare/quote",
            data={
                "csrf_token": csrf,
                "origin_region": "义乌",
                "item_name_count": "1",
                "quote_warehouse_code": ["ONT8"],
                "quote_postal_code": [""],
                "quote_weight_kg": ["120"],
                "quote_box_count": ["10"],
                "quote_length_cm": ["0"],
                "quote_width_cm": ["0"],
                "quote_height_cm": ["0"],
            },
            follow_redirects=True,
        )
        self.assertEqual(quote_response.status_code, 200)
        quote_html = quote_response.get_data(as_text=True)
        self.assertIn("data-sortable-table", quote_html)
        self.assertIn('<th data-sort-type="number">总费用</th>', quote_html)
        self.assertIn('<th data-sort-type="number">时效</th>', quote_html)
        self.assertIn('data-sort-value="15015"', quote_html)
        self.assertIn('"ONT8"', quote_html)
        self.assertIn('"92551"', quote_html)

        with self.app.app_context():
            batch = self.app_module.FreightQuoteBatch.query.order_by(self.app_module.FreightQuoteBatch.id.desc()).first()
            self.assertIsNotNone(batch)
            self.assertGreater(len(batch.options), 0)
            batch_id = batch.id
            best_option = sorted(batch.options, key=lambda option: option.rank_no)[0]
            self.assertEqual(best_option.forwarder_name, "盈和")
            self.assertEqual(best_option.channel_name, "正班盈速达卡派")
            self.assertEqual(float(best_option.total_cost), 1320.0)
            detail_rows = json.loads(best_option.detail_json)
            self.assertEqual(detail_rows[0]["postal_code"], "92551")

        csrf = self.current_csrf()
        quote_delete_response = self.client.post(
            f"/freight-compare/quotes/{batch_id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(quote_delete_response.status_code, 200)
        self.assertIn("历史比价", quote_delete_response.get_data(as_text=True))

        with self.app.app_context():
            deleted_batch = self.db.session.get(self.app_module.FreightQuoteBatch, batch_id)
            deleted_option_count = self.app_module.FreightQuoteOption.query.filter_by(quote_batch_id=batch_id).count()
            self.assertIsNone(deleted_batch)
            self.assertEqual(deleted_option_count, 0)

        csrf = self.current_csrf()
        delete_response = self.client.post(
            f"/freight-compare/rate-books/{rate_book_id}/delete",
            data={"csrf_token": csrf},
            follow_redirects=True,
        )
        self.assertEqual(delete_response.status_code, 200)

        with self.app.app_context():
            deleted_book = self.db.session.get(self.app_module.FreightRateBook, rate_book_id)
            deleted_rate_count = self.app_module.FreightRate.query.filter_by(rate_book_id=rate_book_id).count()
            self.assertIsNone(deleted_book)
            self.assertEqual(deleted_rate_count, 0)
            self.assertFalse(stored_path.exists())

    def test_freight_transit_parser_ignores_kqd_zip_zone_numbers(self):
        text_value = (
            "美西（8-9）邮编 "
            "预计时效：美西航线12天-13天提取（从开船第二天开始算）。"
        )
        self.assertEqual(self.app_module.parse_transit_days(text_value), (12, 13))
        self.assertEqual(self.app_module.parse_transit_days("预计开船后20-25天送仓 无时效赔付"), (20, 25))
        self.assertEqual(self.app_module.parse_transit_days("预计开船35-40天左右，接受集货的在发"), (35, 40))

    def test_customs_parser_ignores_declaration_mode_codes(self):
        with self.app.app_context():
            rate_book = self.app_module.FreightRateBook(
                forwarder_name="\u76c8\u548c",
                version_label="TEST",
                original_filename="yinghe.xlsx",
                stored_filename="yinghe.xlsx",
            )
            no_fee_rule = self.app_module.extract_customs_rule_from_text(
                rate_book,
                "\u7f8e\u56fd\u6d77\u6d3e",
                "\u5408\u5e76\u62a5\u5173\uff0c\u6ee1\u8db31039\u30019710\u30019810\u7b49\u62a5\u5173\u65b9\u5f0f\u3002",
            )
            yinghe_rule = self.app_module.extract_customs_rule_from_text(
                rate_book,
                "\u5361\u6d3e",
                "\u63a5\u5355\u72ec\u62a5\u5173\u4ef6 150\u5143/\u4e00\u7968\uff0c\u54c1\u540d\u8d85\u8fc75\u4e2a\u52a0\u653635\u5143\u4e00\u4e2a\u3002",
            )
            dingbang_rule = self.app_module.extract_customs_rule_from_text(
                rate_book,
                "\u666e\u8239",
                "\u4e00\u822c\u8d38\u6613\u62a5\u5173200/\u7968\uff0c\u7eed\u9875\u8d3950/\u9875\uff1b\u5355\u79685\u4e2a\u54c1\u540d\u3002",
            )

        self.assertIsNone(no_fee_rule)
        self.assertEqual(yinghe_rule.customs_fee, Decimal("150"))
        self.assertEqual(dingbang_rule.customs_fee, Decimal("200"))

    def test_kqd_card_parser_skips_cbm_price_columns(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "\u7f8e\u68eeCLX\u5361\u6d3e"
        sheet.append(["\u7f8e\u68eeCLX\u5361\u6d3e"])
        sheet.append([""])
        sheet.append(["FBA\u4ed3\u5e93", "\u4e49\u4e4c", "", "\u6cc9\u5dde/\u6df1\u5733", ""])
        sheet.append(
            [
                "",
                "300+",
                "1CBM\u8d77\uff0c1CBM\u6700\u91cd\u4e0d\u8d85363KG",
                "300+",
                "1CBM\u8d77\uff0c1CBM\u6700\u91cd\u4e0d\u8d85363KG",
                "\u65f6\u6548",
            ]
        )
        sheet.append(["ONT8-LGB8", 9.3, 2020, 9.8, 2120, "\u5f00\u8239\u540e14-19\u5929"])

        with self.app.app_context():
            rate_book = self.app_module.FreightRateBook(
                forwarder_name="\u5f00\u6e20\u8fbe",
                version_label="TEST",
                original_filename="kqd.xlsx",
                stored_filename="kqd.xlsx",
            )
            rates = self.app_module.parse_kqd_rates(workbook, rate_book)

        self.assertTrue(rates)
        unit_prices = sorted({Decimal(rate.unit_price) for rate in rates})
        self.assertEqual(unit_prices, [Decimal("9.3"), Decimal("9.8")])
        self.assertNotIn(Decimal("2020"), unit_prices)
        self.assertNotIn(Decimal("2120"), unit_prices)

    def test_hongfan_parser_imports_zip_ranges_and_warehouse_rates(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sea_sheet = workbook.active
        sea_sheet.title = "美国海派"
        sea_sheet.cell(4, 2).value = "下单产品"
        sea_sheet.cell(4, 3).value = "美国分区邮编/重量"
        sea_sheet.cell(4, 4).value = "义乌 (含税)"
        sea_sheet.cell(4, 7).value = "福建 (含税)"
        sea_sheet.cell(4, 10).value = "深圳 (含税)"
        for col_index, weight_label in {
            4: "12KG+",
            5: "51KG+",
            6: "100KG+",
            7: "12KG+",
            8: "51KG+",
            9: "100KG+",
            10: "12KG+",
            11: "51KG+",
            12: "100KG+",
        }.items():
            sea_sheet.cell(5, col_index).value = weight_label
        sea_sheet.cell(6, 2).value = "ZIM 海派"
        sea_sheet.cell(6, 3).value = "美中(USM) (邮编97000-99999 且40000-79999）"
        for col_index, price in {
            4: "12.8",
            5: "11.3",
            6: "10.3",
            7: "13.1",
            8: "11.6",
            9: "10.6",
            10: "13.1",
            11: "11.6",
            12: "10.6",
        }.items():
            sea_sheet.cell(6, col_index).value = price
        sea_sheet.cell(6, 18).value = "16-18天"

        card_sheet = workbook.create_sheet("美国ZIM海卡")
        card_sheet.cell(3, 3).value = "ZIM海卡 合并报关仓点限5个仓，超出分票"
        card_sheet.cell(4, 2).value = "美国分区邮编/重量"
        card_sheet.cell(4, 3).value = "义乌 (按KG含税）"
        card_sheet.cell(4, 4).value = "泉州 (按KG含税）"
        card_sheet.cell(4, 5).value = "深圳 按KG含税）"
        card_sheet.cell(4, 6).value = "参考时效（开船后） 自然日"
        card_sheet.cell(5, 3).value = "12KG+"
        card_sheet.cell(5, 4).value = "12KG+"
        card_sheet.cell(5, 5).value = "12KG+"
        card_sheet.cell(6, 2).value = "ONT8"
        card_sheet.cell(6, 3).value = "4.96"
        card_sheet.cell(6, 4).value = "5.26"
        card_sheet.cell(6, 5).value = "5.26"
        card_sheet.cell(6, 6).value = "17-20天"

        address_sheet = workbook.create_sheet("FBA仓库代码及地址")
        address_sheet.append(["仓库代码", "邮编", "州省", "城市"])
        address_sheet.append(["ONT8", "92551-9534", "CA", "Moreno Valley"])

        with self.app.app_context():
            rate_book = self.app_module.FreightRateBook(
                forwarder_name="鸿帆",
                version_label="TEST",
                original_filename="hongfan.xlsx",
                stored_filename="hongfan.xlsx",
            )
            rates = self.app_module.parse_hongfan_rates(workbook, rate_book)

        warehouse_rates = [rate for rate in rates if rate.destination_type == "warehouse" and rate.warehouse_code == "ONT8"]
        self.assertTrue(warehouse_rates)
        self.assertTrue(any(rate.postal_code == "92551" for rate in warehouse_rates))
        self.assertTrue(any(rate.origin_region == "泉州" and Decimal(rate.unit_price) == Decimal("5.26") for rate in warehouse_rates))
        zip_ranges = {
            (rate.postal_start, rate.postal_end)
            for rate in rates
            if rate.destination_type == "zip_range" and rate.zone_name.startswith("美中")
        }
        self.assertIn((40000, 79999), zip_ranges)
        self.assertIn((97000, 99999), zip_ranges)
        self.assertEqual(self.app_module.normalize_freight_forwarder("", "VIP鸿帆报价.xlsx"), "鸿帆")
        self.assertTrue(self.app_module.freight_origin_matches("福建&深圳", "福建"))

    def test_freight_quote_applies_clothing_surcharge_per_kg(self):
        with self.app.app_context():
            self.app_module.FreightRateBook.query.update({"status": "archived"})
            rate_book = self.app_module.FreightRateBook(
                forwarder_name="\u9f0e\u90a6",
                version_label="TEST",
                original_filename="dingbang.xlsx",
                stored_filename="dingbang.xlsx",
                status="active",
                clothing_surcharge_per_kg=Decimal("1.00"),
            )
            db_rate = self.app_module.FreightRate(
                rate_book=rate_book,
                forwarder_name="\u9f0e\u90a6",
                channel_name="\u7f8e\u68ee\u5361\u6d3e",
                origin_region="\u4e49\u4e4c",
                warehouse_code="ONT8",
                weight_break_kg=Decimal("100"),
                unit_price=Decimal("10.00"),
                min_piece_kg=Decimal("0"),
                divisor=6000,
            )
            self.db.session.add(rate_book)
            self.db.session.add(db_rate)
            self.db.session.commit()

            result = self.app_module.calculate_freight_quote(
                items=[
                    {
                        "warehouse_code": "ONT8",
                        "postal_code": "92551",
                        "actual_weight_kg": Decimal("120"),
                        "box_count": 1,
                    }
                ],
                origin_region="\u4e49\u4e4c",
                clothing_surcharge_enabled=True,
            )
            self.db.session.delete(rate_book)
            self.db.session.commit()

        self.assertEqual(len(result["options"]), 1)
        option = result["options"][0]
        self.assertEqual(option["freight_cost"], Decimal("1200.00"))
        self.assertEqual(option["surcharge_cost"], Decimal("120.00"))
        self.assertEqual(option["total_cost"], Decimal("1320.00"))
        self.assertIn("\u670d\u88c5\u9644\u52a0\u8d39 1.00/kg", option["notes"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
