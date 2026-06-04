import csv
import json
import math
import os
import re
import secrets
import sys
import time
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from functools import wraps
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from flask import Flask, flash, g, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from markupsafe import Markup
from sqlalchemy import UniqueConstraint, case, event, func, or_, text
from sqlalchemy.engine import Engine
from sqlalchemy.exc import IntegrityError
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_IMAGE_RELATIVE_DIR = os.path.join("uploads", "assets")
FBA_CALC_RESULT_DIR = os.path.join(BASE_DIR, "instance", "fba_calc_results")
PURCHASE_CALC_RESULT_DIR = os.path.join(BASE_DIR, "instance", "purchase_calc_results")
FREIGHT_RATE_BOOK_DIR = os.path.join(BASE_DIR, "instance", "freight_rate_books")
ASSET_IMAGE_DIR = os.path.join(BASE_DIR, "static", ASSET_IMAGE_RELATIVE_DIR)
ALLOWED_ASSET_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
ALLOWED_FREIGHT_RATE_EXTENSIONS = {".xlsx", ".xlsm"}
extra_paths = []
wheel_dir = os.path.join(BASE_DIR, ".wheels")
if os.path.isdir(wheel_dir):
    for wheel_name in sorted(os.listdir(wheel_dir)):
        if wheel_name.endswith(".whl"):
            extra_paths.append(os.path.join(wheel_dir, wheel_name))
for vendor_name in (".vendor_pkgs", ".vendor"):
    vendor_dir = os.path.join(BASE_DIR, vendor_name)
    if os.path.isdir(vendor_dir):
        extra_paths.append(vendor_dir)
for extra_path in reversed(extra_paths):
    if extra_path not in sys.path:
        sys.path.insert(0, extra_path)


app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("ERP_SECRET_KEY") or secrets.token_hex(32)
database_uri = os.getenv("ERP_DATABASE_URL", "sqlite:///erp.db")
app.config["SQLALCHEMY_DATABASE_URI"] = database_uri
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
if database_uri.startswith("sqlite"):
    app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
        "connect_args": {"timeout": 30},
        "pool_pre_ping": True,
    }
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("ERP_SESSION_COOKIE_SECURE", "").lower() in {"1", "true", "yes"}
db = SQLAlchemy(app)


@event.listens_for(Engine, "connect")
def configure_sqlite_connection(dbapi_connection, _connection_record):
    if "sqlite" not in dbapi_connection.__class__.__module__:
        return
    cursor = dbapi_connection.cursor()
    try:
        cursor.execute("PRAGMA busy_timeout=30000")
        cursor.execute("PRAGMA journal_mode=WAL")
        cursor.execute("PRAGMA synchronous=NORMAL")
        cursor.execute("PRAGMA foreign_keys=ON")
    finally:
        cursor.close()


try:
    LOCAL_TIMEZONE = ZoneInfo("Asia/Shanghai")
except ZoneInfoNotFoundError:
    LOCAL_TIMEZONE = timezone(timedelta(hours=8))


LOGIN_WINDOW_SECONDS = 300
LOGIN_MAX_ATTEMPTS = 5
LOGIN_ATTEMPTS = defaultdict(list)
SESSION_IDLE_TIMEOUT_SECONDS = 3600
DEFAULT_LOGIN_LOCK_THRESHOLD = 3
DEFAULT_LOGIN_LOCK_DURATION_MINUTES = 7 * 24 * 60
LOGIN_SECURITY_SETTING_KEYS = {
    "login_lock_threshold": str(DEFAULT_LOGIN_LOCK_THRESHOLD),
    "login_lock_duration_minutes": str(DEFAULT_LOGIN_LOCK_DURATION_MINUTES),
    "emergency_admin_ip_whitelist": "127.0.0.1,::1",
}

PURCHASE_STATUSES = {"draft", "submitted", "placed", "partial", "received"}
SALES_STATUSES = {"draft", "confirmed", "shipped"}
DEFECT_STATUSES = {"pending", "repairing", "done", "scrapped"}
PURCHASE_CALC_OPEN_STATUSES = {"submitted", "placed", "partial"}
PURCHASE_CALC_OUTBOUND_TYPES = {"sales_outbound", "manual_outbound", "shipment_outbound"}
EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@")


role_permissions = db.Table(
    "role_permissions",
    db.Column("role_id", db.Integer, db.ForeignKey("role.id"), primary_key=True),
    db.Column("permission_id", db.Integer, db.ForeignKey("permission.id"), primary_key=True),
)

user_roles = db.Table(
    "user_roles",
    db.Column("user_id", db.Integer, db.ForeignKey("user.id"), primary_key=True),
    db.Column("role_id", db.Integer, db.ForeignKey("role.id"), primary_key=True),
)

user_sku_permissions = db.Table(
    "user_sku_permissions",
    db.Column("user_id", db.Integer, db.ForeignKey("user.id"), primary_key=True),
    db.Column("sku_id", db.Integer, db.ForeignKey("sku.id"), primary_key=True),
)


class TimestampMixin:
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)


class Permission(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(64), unique=True, nullable=False)
    name = db.Column(db.String(128), nullable=False)
    description = db.Column(db.String(255), nullable=False, default="")


class Role(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(64), unique=True, nullable=False)
    description = db.Column(db.String(255), nullable=False, default="")
    permissions = db.relationship("Permission", secondary=role_permissions, lazy="joined")


class User(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(64), unique=True, nullable=False)
    full_name = db.Column(db.String(128), nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    is_active = db.Column(db.Boolean, nullable=False, default=True)
    failed_login_attempts = db.Column(db.Integer, nullable=False, default=0)
    locked_until = db.Column(db.DateTime, nullable=True)
    last_failed_login_at = db.Column(db.DateTime, nullable=True)
    is_emergency_account = db.Column(db.Boolean, nullable=False, default=False)
    roles = db.relationship("Role", secondary=user_roles, lazy="joined")
    authorized_skus = db.relationship("SKU", secondary=user_sku_permissions, lazy="select", back_populates="authorized_users")

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

    @property
    def permission_codes(self):
        codes = set()
        for role in self.roles:
            for permission in role.permissions:
                codes.add(permission.code)
        return codes

    def has_permission(self, code):
        return code in self.permission_codes

    def has_role(self, *role_names):
        owned_names = {role.name for role in self.roles}
        return any(name in owned_names for name in role_names)

    @property
    def role_names(self):
        return ", ".join(role.name for role in self.roles) or "-"


class LoginAudit(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    username = db.Column(db.String(64), nullable=False, default="")
    login_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    logout_at = db.Column(db.DateTime, nullable=True)
    login_ip = db.Column(db.String(64), nullable=False, default="")
    logout_reason = db.Column(db.String(64), nullable=False, default="")
    user_agent = db.Column(db.String(255), nullable=False, default="")
    is_success = db.Column(db.Boolean, nullable=False, default=True)
    failure_reason = db.Column(db.String(128), nullable=False, default="")
    locked_until = db.Column(db.DateTime, nullable=True)
    lock_triggered = db.Column(db.Boolean, nullable=False, default=False)
    user = db.relationship("User")


class OperationLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    category = db.Column(db.String(64), nullable=False, default="")
    action = db.Column(db.String(16), nullable=False, default="")
    target_type = db.Column(db.String(64), nullable=False, default="")
    target_label = db.Column(db.String(255), nullable=False, default="")
    detail = db.Column(db.String(255), nullable=False, default="")
    operator_user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    operator_username = db.Column(db.String(64), nullable=False, default="")
    ip_address = db.Column(db.String(64), nullable=False, default="")
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    operator = db.relationship("User")


class SystemSetting(db.Model):
    key = db.Column(db.String(64), primary_key=True)
    value = db.Column(db.String(255), nullable=False, default="")
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)


class Warehouse(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(128), unique=True, nullable=False)
    code = db.Column(db.String(32), unique=True, nullable=False)
    address = db.Column(db.String(255), nullable=False, default="")
    manager = db.Column(db.String(64), nullable=False, default="")
    remark = db.Column(db.String(255), nullable=False, default="")


class Asset(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    asset_code = db.Column(db.String(64), unique=True, nullable=False)
    name = db.Column(db.String(128), nullable=False)
    category = db.Column(db.String(64), nullable=False, default="")
    brand = db.Column(db.String(64), nullable=False, default="")
    model = db.Column(db.String(64), nullable=False, default="")
    serial_no = db.Column(db.String(128), nullable=False, default="")
    status = db.Column(db.String(32), nullable=False, default="in_use")
    keeper = db.Column(db.String(64), nullable=False, default="")
    location = db.Column(db.String(128), nullable=False, default="")
    purchase_date = db.Column(db.Date, nullable=True)
    asset_value = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    image_path = db.Column(db.String(255), nullable=False, default="")
    remark = db.Column(db.String(255), nullable=False, default="")

    @property
    def status_label(self):
        labels = {
            "in_use": "使用中",
            "idle": "闲置",
            "repair": "维修中",
            "scrapped": "已报废",
        }
        return labels.get(self.status, self.status)


class Customer(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(128), unique=True, nullable=False)
    contact_name = db.Column(db.String(64), nullable=False, default="")
    phone = db.Column(db.String(32), nullable=False, default="")
    address = db.Column(db.String(255), nullable=False, default="")
    remark = db.Column(db.String(255), nullable=False, default="")


class Supplier(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(128), unique=True, nullable=False)
    contact_name = db.Column(db.String(64), nullable=False, default="")
    phone = db.Column(db.String(32), nullable=False, default="")
    address = db.Column(db.String(255), nullable=False, default="")
    remark = db.Column(db.String(255), nullable=False, default="")


class SKU(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sku_code = db.Column(db.String(64), unique=True, nullable=False)
    barcode = db.Column(db.String(64), unique=True, nullable=True, default=None)
    name = db.Column(db.String(128), nullable=False)
    category = db.Column(db.String(64), nullable=False, default="服装")
    color = db.Column(db.String(64), nullable=False, default="")
    size = db.Column(db.String(32), nullable=False, default="")
    unit = db.Column(db.String(16), nullable=False, default="件")
    cost_price = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    sale_price = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    safety_stock = db.Column(db.Integer, nullable=False, default=0)
    remark = db.Column(db.String(255), nullable=False, default="")
    authorized_users = db.relationship("User", secondary=user_sku_permissions, lazy="select", back_populates="authorized_skus")


class SKUMapping(TimestampMixin, db.Model):
    __table_args__ = (
        UniqueConstraint("user_id", "external_sku_code", name="uq_sku_mapping_user_external_sku"),
    )

    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    external_sku_code = db.Column(db.String(128), nullable=False)
    remark = db.Column(db.String(255), nullable=False, default="")
    user = db.relationship("User")
    sku = db.relationship("SKU")


class ShipmentSheet(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    shipment_no = db.Column(db.String(64), unique=True, nullable=False)
    store_name = db.Column(db.String(128), nullable=False, default="")
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    operator_user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    box_count = db.Column(db.Integer, nullable=False, default=1)
    status = db.Column(db.String(32), nullable=False, default="pending")
    remark = db.Column(db.String(255), nullable=False, default="")
    archived_at = db.Column(db.DateTime, nullable=True)
    warehouse_confirmed_at = db.Column(db.DateTime, nullable=True)
    warehouse_confirmer_name = db.Column(db.String(128), nullable=False, default="")
    warehouse = db.relationship("Warehouse")
    operator = db.relationship("User")
    items = db.relationship("ShipmentSheetItem", back_populates="shipment_sheet", cascade="all, delete-orphan")

    @property
    def status_label(self):
        labels = {
            "pending": "待备货",
            "confirmed": "已确认",
        }
        return labels.get(self.status, self.status)

    @property
    def total_quantity(self):
        return sum(item.quantity for item in self.items)


class ShipmentSheetItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    shipment_sheet_id = db.Column(db.Integer, db.ForeignKey("shipment_sheet.id"), nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    external_sku_code = db.Column(db.String(128), nullable=False, default="")
    quantity = db.Column(db.Integer, nullable=False, default=0)
    per_box_quantity = db.Column(db.Integer, nullable=False, default=0)
    shipment_sheet = db.relationship("ShipmentSheet", back_populates="items")
    sku = db.relationship("SKU")


class ShipmentDraftLock(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    lock_token = db.Column(db.String(64), nullable=False, index=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    external_sku_code = db.Column(db.String(128), nullable=False, default="")
    quantity = db.Column(db.Integer, nullable=False, default=0)
    expires_at = db.Column(db.DateTime, nullable=False)
    user = db.relationship("User")
    warehouse = db.relationship("Warehouse")
    sku = db.relationship("SKU")


class FBAInboundRecord(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    shipment_ref = db.Column(db.String(128), nullable=False, default="")
    external_sku_code = db.Column(db.String(128), nullable=False, default="")
    quantity = db.Column(db.Integer, nullable=False, default=0)
    status = db.Column(db.String(32), nullable=False, default="active")
    remark = db.Column(db.String(255), nullable=False, default="")
    uploader_user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    uploader = db.relationship("User")

    @property
    def status_label(self):
        return {"active": "在途", "received": "已接收", "deleted": "已删除"}.get(self.status, self.status)


class FBAReplenishmentRule(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    external_sku_code = db.Column(db.String(128), unique=True, nullable=False)
    custom_multiplier = db.Column(db.Numeric(10, 2), nullable=False, default=1)
    is_new_product = db.Column(db.Boolean, nullable=False, default=False)
    remark = db.Column(db.String(255), nullable=False, default="")


class FreightRateBook(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    forwarder_name = db.Column(db.String(128), nullable=False, index=True)
    version_label = db.Column(db.String(64), nullable=False, default="")
    original_filename = db.Column(db.String(255), nullable=False, default="")
    stored_filename = db.Column(db.String(255), nullable=False, default="")
    status = db.Column(db.String(32), nullable=False, default="active")
    parsed_summary = db.Column(db.Text, nullable=False, default="")
    clothing_surcharge_per_kg = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    uploaded_by_user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    uploaded_by = db.relationship("User")
    rates = db.relationship("FreightRate", back_populates="rate_book", cascade="all, delete-orphan")
    customs_rules = db.relationship("FreightCustomsRule", back_populates="rate_book", cascade="all, delete-orphan")
    remote_zips = db.relationship("FreightRemoteZip", back_populates="rate_book", cascade="all, delete-orphan")

    @property
    def status_label(self):
        return {"active": "启用", "archived": "停用", "error": "解析失败"}.get(self.status, self.status)


class FreightRate(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rate_book_id = db.Column(db.Integer, db.ForeignKey("freight_rate_book.id"), nullable=False, index=True)
    forwarder_name = db.Column(db.String(128), nullable=False, default="")
    channel_name = db.Column(db.String(128), nullable=False, default="", index=True)
    origin_region = db.Column(db.String(64), nullable=False, default="")
    destination_type = db.Column(db.String(32), nullable=False, default="warehouse")
    warehouse_code = db.Column(db.String(64), nullable=False, default="", index=True)
    zone_name = db.Column(db.String(64), nullable=False, default="")
    postal_code = db.Column(db.String(16), nullable=False, default="")
    postal_start = db.Column(db.Integer, nullable=True)
    postal_end = db.Column(db.Integer, nullable=True)
    weight_break_kg = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    unit_price = db.Column(db.Numeric(10, 2), nullable=True)
    cbm_price = db.Column(db.Numeric(10, 2), nullable=True)
    min_chargeable_kg = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    min_piece_kg = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    divisor = db.Column(db.Integer, nullable=False, default=6000)
    transit_days_min = db.Column(db.Integer, nullable=True)
    transit_days_max = db.Column(db.Integer, nullable=True)
    tax_mode = db.Column(db.String(64), nullable=False, default="")
    source_sheet = db.Column(db.String(128), nullable=False, default="")
    notes = db.Column(db.Text, nullable=False, default="")
    rate_book = db.relationship("FreightRateBook", back_populates="rates")


class FreightCustomsRule(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rate_book_id = db.Column(db.Integer, db.ForeignKey("freight_rate_book.id"), nullable=False, index=True)
    forwarder_name = db.Column(db.String(128), nullable=False, default="")
    channel_name = db.Column(db.String(128), nullable=False, default="")
    customs_fee = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    supports_merge = db.Column(db.Boolean, nullable=False, default=True)
    merge_scope = db.Column(db.String(32), nullable=False, default="batch")
    item_name_limit = db.Column(db.Integer, nullable=False, default=0)
    item_name_extra_fee = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    inspection_fee = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    formal_inspection_fee = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    notes = db.Column(db.Text, nullable=False, default="")
    rate_book = db.relationship("FreightRateBook", back_populates="customs_rules")


class FreightRemoteZip(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rate_book_id = db.Column(db.Integer, db.ForeignKey("freight_rate_book.id"), nullable=False, index=True)
    forwarder_name = db.Column(db.String(128), nullable=False, default="")
    kind = db.Column(db.String(32), nullable=False, default="remote")
    zip_code = db.Column(db.String(16), nullable=False, index=True)
    rate_book = db.relationship("FreightRateBook", back_populates="remote_zips")


class WarehousePostalCode(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    warehouse_code = db.Column(db.String(32), unique=True, nullable=False, index=True)
    postal_code = db.Column(db.String(16), nullable=False, default="")
    city = db.Column(db.String(128), nullable=False, default="")
    state = db.Column(db.String(64), nullable=False, default="")
    source = db.Column(db.String(128), nullable=False, default="")
    note = db.Column(db.String(255), nullable=False, default="")


class FreightQuoteBatch(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    quote_no = db.Column(db.String(64), unique=True, nullable=False)
    origin_region = db.Column(db.String(64), nullable=False, default="")
    customs_required = db.Column(db.Boolean, nullable=False, default=False)
    item_name_count = db.Column(db.Integer, nullable=False, default=0)
    input_json = db.Column(db.Text, nullable=False, default="")
    summary_json = db.Column(db.Text, nullable=False, default="")
    operator_user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    operator = db.relationship("User")
    options = db.relationship("FreightQuoteOption", back_populates="quote_batch", cascade="all, delete-orphan")


class FreightQuoteOption(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    quote_batch_id = db.Column(db.Integer, db.ForeignKey("freight_quote_batch.id"), nullable=False, index=True)
    forwarder_name = db.Column(db.String(128), nullable=False, default="")
    channel_name = db.Column(db.String(128), nullable=False, default="")
    origin_region = db.Column(db.String(64), nullable=False, default="")
    total_cost = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    freight_cost = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    customs_cost = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    surcharge_cost = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    chargeable_weight = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    transit_days_min = db.Column(db.Integer, nullable=True)
    transit_days_max = db.Column(db.Integer, nullable=True)
    rank_no = db.Column(db.Integer, nullable=False, default=0)
    recommendation = db.Column(db.String(64), nullable=False, default="")
    notes = db.Column(db.Text, nullable=False, default="")
    detail_json = db.Column(db.Text, nullable=False, default="")
    quote_batch = db.relationship("FreightQuoteBatch", back_populates="options")


class PurchaseOrder(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    order_no = db.Column(db.String(64), unique=True, nullable=False)
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"), nullable=False)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    status = db.Column(db.String(32), nullable=False, default="draft")
    expected_date = db.Column(db.Date, nullable=True)
    remark = db.Column(db.String(255), nullable=False, default="")
    supplier = db.relationship("Supplier")
    warehouse = db.relationship("Warehouse")
    items = db.relationship("PurchaseOrderItem", back_populates="purchase_order", cascade="all, delete-orphan")

    @property
    def total_amount(self):
        return sum((item.unit_price or 0) * item.quantity for item in self.items)

    @property
    def status_label(self):
        labels = {
            "draft": "草稿",
            "submitted": "已提交",
            "placed": "已下单",
            "partial": "部分收货",
            "received": "已收货",
        }
        return labels.get(self.status, self.status)


class PurchaseOrderItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    purchase_order_id = db.Column(db.Integer, db.ForeignKey("purchase_order.id"), nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    quantity = db.Column(db.Integer, nullable=False)
    unit_price = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    purchase_order = db.relationship("PurchaseOrder", back_populates="items")
    sku = db.relationship("SKU")


class SalesOrder(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    order_no = db.Column(db.String(64), unique=True, nullable=False)
    customer_id = db.Column(db.Integer, db.ForeignKey("customer.id"), nullable=True)
    customer_name = db.Column(db.String(128), nullable=False)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    phone = db.Column(db.String(32), nullable=False, default="")
    status = db.Column(db.String(32), nullable=False, default="draft")
    shipping_address = db.Column(db.String(255), nullable=False, default="")
    remark = db.Column(db.String(255), nullable=False, default="")
    customer = db.relationship("Customer")
    warehouse = db.relationship("Warehouse")
    items = db.relationship("SalesOrderItem", back_populates="sales_order", cascade="all, delete-orphan")

    @property
    def total_amount(self):
        return sum((item.unit_price or 0) * item.quantity for item in self.items)

    @property
    def status_label(self):
        labels = {"draft": "草稿", "confirmed": "已确认", "shipped": "已发货"}
        return labels.get(self.status, self.status)


class SalesOrderItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sales_order_id = db.Column(db.Integer, db.ForeignKey("sales_order.id"), nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    quantity = db.Column(db.Integer, nullable=False)
    unit_price = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    sales_order = db.relationship("SalesOrder", back_populates="items")
    sku = db.relationship("SKU")


class InventoryTransaction(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"), nullable=True)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    quantity = db.Column(db.Integer, nullable=False)
    transaction_type = db.Column(db.String(32), nullable=False)
    reference_type = db.Column(db.String(32), nullable=False, default="")
    reference_no = db.Column(db.String(64), nullable=False, default="")
    note = db.Column(db.String(255), nullable=False, default="")
    sku = db.relationship("SKU")
    warehouse = db.relationship("Warehouse")
    supplier = db.relationship("Supplier")

    @property
    def transaction_type_label(self):
        labels = {
            "purchase_inbound": "采购入库",
            "manual_inbound": "手工入库",
            "initial_inbound": "期初入库",
            "sales_outbound": "销售出库",
            "manual_outbound": "手工出库",
            "shipment_outbound": "发货表出库",
            "defect_hold": "次品占用",
            "defect_return": "返修回库",
            "defect_scrap": "报废出库",
        }
        return labels.get(self.transaction_type, self.transaction_type)


class InventoryBalance(TimestampMixin, db.Model):
    __table_args__ = (
        UniqueConstraint("sku_id", "warehouse_id", name="uq_inventory_balance_sku_warehouse"),
    )

    id = db.Column(db.Integer, primary_key=True)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=False)
    quantity = db.Column(db.Integer, nullable=False, default=0)
    sku = db.relationship("SKU")
    warehouse = db.relationship("Warehouse")


class DefectRepair(TimestampMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    repair_no = db.Column(db.String(64), unique=True, nullable=False)
    sku_id = db.Column(db.Integer, db.ForeignKey("sku.id"), nullable=False)
    warehouse_id = db.Column(db.Integer, db.ForeignKey("warehouse.id"), nullable=True)
    operator_name = db.Column(db.String(128), nullable=False, default="")
    quantity = db.Column(db.Integer, nullable=False)
    completed_quantity = db.Column(db.Integer, nullable=False, default=0)
    scrapped_quantity = db.Column(db.Integer, nullable=False, default=0)
    defect_reason = db.Column(db.String(255), nullable=False)
    status = db.Column(db.String(32), nullable=False, default="pending")
    repair_result = db.Column(db.String(255), nullable=False, default="")
    sku = db.relationship("SKU")
    warehouse = db.relationship("Warehouse")

    @property
    def status_label(self):
        if self.completed_quantity and self.scrapped_quantity and self.remaining_quantity == 0:
            return "部分完成/部分报废"
        if (self.completed_quantity or self.scrapped_quantity) and self.remaining_quantity > 0:
            return "部分处理"
        labels = {
            "pending": "待处理",
            "repairing": "返修中",
            "done": "已完成",
            "scrapped": "已报废",
        }
        return labels.get(self.status, self.status)

    @property
    def remaining_quantity(self):
        return max((self.quantity or 0) - (self.completed_quantity or 0) - (self.scrapped_quantity or 0), 0)

    @property
    def quantity_summary(self):
        parts = [f"总数 {self.quantity or 0}"]
        if self.completed_quantity:
            parts.append(f"成功 {self.completed_quantity}")
        if self.scrapped_quantity:
            parts.append(f"报废 {self.scrapped_quantity}")
        if self.remaining_quantity:
            parts.append(f"未处理 {self.remaining_quantity}")
        return " / ".join(parts)


ALL_PERMISSIONS = [
    ("dashboard.view", "仪表盘", "查看仪表盘"),
    ("user.manage", "用户管理", "管理用户、角色和权限"),
    ("price.view", "价格查看", "查看单价、售价、成本和库存价值等价格信息"),
    ("sku.manage", "SKU管理", "管理商品主数据"),
    ("sku.mapping.manage", "SKU映射", "维护店铺SKU与仓库SKU的映射关系"),
    ("fba.calc.manage", "FBA发货计算器", "导入亚马逊表格并计算建议发货量"),
    ("fba.inbound.manage", "FBA在途管理", "维护亚马逊在途货件并手工标记已接收"),
    ("freight.compare.manage", "货代比价", "上传货代价格表并进行运费、时效和报关费用比价"),
    ("shipment.manage", "发货表", "创建、编辑、删除和导出发货表"),
    ("shipment.confirm", "发货确认", "查看待备货发货表并由仓管确认"),
    ("supplier.manage", "供应商管理", "管理供应商资料"),
    ("customer.manage", "客户管理", "管理客户资料"),
    ("warehouse.manage", "仓库管理", "管理仓库资料"),
    ("asset.manage", "资产管理", "管理公司固定资产和办公设备"),
    ("purchase.manage", "采购管理", "管理采购订单"),
    ("purchase.calc", "采购计算器", "按销量趋势、库存和采购在途计算采购建议"),
    ("stock.in", "入库管理", "处理入库业务"),
    ("sales.manage", "销售订单", "管理销售订单"),
    ("stock.out", "出库管理", "处理出库业务"),
    ("defect.manage", "次品返修", "管理次品返修记录"),
    ("inventory.view", "库存查询", "查看库存信息"),
    ("inventory.transaction.view", "库存流水", "查看库存流水记录"),
    ("report.view", "报表中心", "查看和导出报表"),
]

NAV_ITEMS = [
    ("dashboard", "仪表盘", "dashboard.view"),
    ("asset_list", "资产管理", "asset.manage"),
    ("user_list", "用户管理", "user.manage"),
    ("sku_list", "SKU管理", "sku.manage"),
    ("sku_mapping_list", "SKU映射", "sku.mapping.manage"),
    ("fba_calculator", "FBA发货计算器", "fba.calc.manage"),
    ("fba_inbound_list", "FBA在途管理", "fba.inbound.manage"),
    ("freight_compare", "货代比价", "freight.compare.manage"),
    ("supplier_list", "供应商管理", "supplier.manage"),
    ("customer_list", "客户管理", "customer.manage"),
    ("warehouse_list", "仓库管理", "warehouse.manage"),
    ("purchase_order_list", "采购订单", "purchase.manage"),
    ("purchase_calculator", "采购计算器", "purchase.calc"),
    ("inventory_in", "入库管理", "stock.in"),
    ("sales_order_list", "销售订单", "sales.manage"),
    ("inventory_out", "出库管理", "stock.out"),
    ("defect_list", "次品返修", "defect.manage"),
    ("inventory", "库存查询", "inventory.view"),
    ("inventory_transactions", "库存流水", "inventory.transaction.view"),
    ("scan_lookup", "扫码查询", "inventory.view"),
    ("reports", "报表中心", "report.view"),
]


class InventoryError(Exception):
    pass


def get_request_ip():
    forwarded = request.headers.get("X-Forwarded-For", "")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.remote_addr or "unknown"


def normalize_ip_value(value):
    return (value or "").split("%")[0].strip().lower()


def get_emergency_admin_ip_whitelist():
    raw_value = get_system_setting("emergency_admin_ip_whitelist", "127.0.0.1,::1")
    values = [normalize_ip_value(item) for item in raw_value.replace("\n", ",").split(",")]
    return [item for item in values if item]


def is_emergency_request_allowed():
    return normalize_ip_value(get_request_ip()) in set(get_emergency_admin_ip_whitelist() + ["localhost"])


def get_system_setting(key, default=""):
    setting = db.session.get(SystemSetting, key)
    return setting.value if setting else default


def set_system_setting(key, value):
    setting = db.session.get(SystemSetting, key)
    if not setting:
        setting = SystemSetting(key=key, value=str(value))
        db.session.add(setting)
    else:
        setting.value = str(value)
    return setting


def get_login_lock_threshold():
    return max(parse_int(get_system_setting("login_lock_threshold", DEFAULT_LOGIN_LOCK_THRESHOLD), DEFAULT_LOGIN_LOCK_THRESHOLD), 1)


def get_login_lock_duration_minutes():
    return max(parse_int(get_system_setting("login_lock_duration_minutes", DEFAULT_LOGIN_LOCK_DURATION_MINUTES), DEFAULT_LOGIN_LOCK_DURATION_MINUTES), 1)


def get_login_lock_duration_label(minutes=None):
    total_minutes = minutes if minutes is not None else get_login_lock_duration_minutes()
    days, remainder = divmod(total_minutes, 24 * 60)
    hours, minutes = divmod(remainder, 60)
    parts = []
    if days:
        parts.append(f"{days}天")
    if hours:
        parts.append(f"{hours}小时")
    if minutes:
        parts.append(f"{minutes}分钟")
    return "".join(parts) or "1分钟"


def reset_user_login_security(user):
    user.failed_login_attempts = 0
    user.locked_until = None
    user.last_failed_login_at = None


def is_user_locked(user, current_time=None):
    if not user or not user.locked_until:
        return False
    return user.locked_until > (current_time or datetime.utcnow())


def record_login_audit_event(username, *, user=None, is_success=False, failure_reason="", locked_until=None, lock_triggered=False):
    db.session.add(
        LoginAudit(
            user_id=user.id if user else None,
            username=username,
            login_ip=get_request_ip(),
            user_agent=(request.headers.get("User-Agent", "") or "")[:255],
            is_success=is_success,
            failure_reason=(failure_reason or "")[:128],
            locked_until=locked_until,
            lock_triggered=lock_triggered,
        )
    )


def get_login_attempt_key(username):
    return f"{get_request_ip()}:{username.lower()}"


def prune_login_attempts(key):
    cutoff = time.time() - LOGIN_WINDOW_SECONDS
    LOGIN_ATTEMPTS[key] = [ts for ts in LOGIN_ATTEMPTS.get(key, []) if ts >= cutoff]
    return LOGIN_ATTEMPTS[key]


def is_login_rate_limited(username):
    attempts = prune_login_attempts(get_login_attempt_key(username))
    return len(attempts) >= LOGIN_MAX_ATTEMPTS


def record_login_failure(username):
    key = get_login_attempt_key(username)
    attempts = prune_login_attempts(key)
    attempts.append(time.time())
    LOGIN_ATTEMPTS[key] = attempts


def clear_login_failures(username):
    LOGIN_ATTEMPTS.pop(get_login_attempt_key(username), None)


def is_session_expired():
    last_seen = session.get("_last_seen_at")
    if last_seen is None:
        return False
    return time.time() - float(last_seen) > SESSION_IDLE_TIMEOUT_SECONDS


def start_login_audit(user):
    audit = LoginAudit(
        user_id=user.id,
        username=user.username,
        login_ip=get_request_ip(),
        user_agent=(request.headers.get("User-Agent", "") or "")[:255],
        is_success=True,
    )
    db.session.add(audit)
    db.session.commit()
    refresh_fba_calc_session_result()
    session["login_audit_id"] = audit.id


def finish_login_audit(reason):
    audit_id = session.get("login_audit_id")
    if not audit_id:
        return
    audit = db.session.get(LoginAudit, audit_id)
    if not audit or audit.logout_at:
        return
    audit.logout_at = datetime.utcnow()
    audit.logout_reason = reason
    db.session.commit()


def is_admin_user(user=None):
    current = user or g.get("user")
    return bool(current and current.has_role("管理员"))


def can_delete_records(user=None):
    current = user or g.get("user")
    return bool(current and (current.has_role("管理员") or current.has_role("仓管员")))


def require_delete_role():
    if can_delete_records():
        return None
    flash("只有管理员和仓管员可以删除记录。", "error")
    return redirect(request.referrer or url_for("dashboard"))


def log_operation(category, action, target_type, target_label, detail=""):
    current = g.get("user")
    db.session.add(
        OperationLog(
            category=category,
            action=action,
            target_type=target_type,
            target_label=(target_label or "")[:255],
            detail=(detail or "")[:255],
            operator_user_id=current.id if current else None,
            operator_username=current.username if current else "",
            ip_address=get_request_ip(),
        )
    )


def get_csrf_token():
    token = session.get("_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["_csrf_token"] = token
    return token


def csrf_field():
    return Markup(f'<input type="hidden" name="csrf_token" value="{get_csrf_token()}">')


def validate_csrf():
    session_token = session.get("_csrf_token")
    request_token = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token", "")
    return bool(session_token and request_token and secrets.compare_digest(session_token, request_token))


@app.before_request
def load_current_user():
    user_id = session.get("user_id")
    if user_id and is_session_expired():
        finish_login_audit("超时退出")
        session.clear()
        g.user = None
        flash("登录已超时，请重新登录。", "error")
        return redirect(url_for("login"))
    g.user = db.session.get(User, user_id) if user_id else None
    if g.user and not g.user.is_active:
        finish_login_audit("璐﹀彿澶辨晥")
        session.clear()
        g.user = None
    elif g.user and is_user_locked(g.user):
        locked_until = g.user.locked_until
        finish_login_audit("账号锁定")
        session.clear()
        g.user = None
        flash(f"账号已锁定，预计 {format_dt(locked_until)} 后可重试。", "error")
        return redirect(url_for("login"))
    elif g.user:
        session["_last_seen_at"] = time.time()
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and not validate_csrf():
        flash("请求已过期或无效，请重试。", "error")
        if request.endpoint == "login":
            return redirect(url_for("login"))
        target = request.referrer or url_for("dashboard" if g.user else "login")
        return redirect(target)


@app.after_request
def add_security_headers(response):
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.context_processor
def inject_helpers():
    return {
        "current_user": g.get("user"),
        "nav_items": NAV_ITEMS,
        "now": to_local_time(datetime.utcnow()),
        "csrf_field": csrf_field,
        "get_csrf_token": get_csrf_token,
        "format_dt": format_dt,
        "can_delete_records": can_delete_records,
        "is_admin_user": is_admin_user,
        "can_view_prices": can_view_prices,
        "can_ship_sales_orders": can_ship_sales_orders,
        "can_access_purchase_calculator": can_access_purchase_calculator,
        "clean_legacy_display_text": clean_legacy_display_text,
        "can_manage_shipment_sheets": can_manage_shipment_sheets,
        "can_confirm_shipment_sheets": can_confirm_shipment_sheets,
        "can_edit_shipment_sheet": can_edit_shipment_sheet,
        "can_delete_shipment_sheet": can_delete_shipment_sheet,
        "can_edit_fba_inbound_record": can_edit_fba_inbound_record,
        "can_edit_fba_inbound_group": can_edit_fba_inbound_group,
        "pending_shipment_count": get_pending_shipment_count() if g.get("user") and can_confirm_shipment_sheets() else 0,
    }


@app.errorhandler(IntegrityError)
def handle_integrity_error(_error):
    db.session.rollback()
    flash("保存失败，数据可能重复或已被其他请求修改。", "error")
    return redirect(request.referrer or url_for("dashboard" if g.get("user") else "login"))


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not g.user:
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def get_default_landing_endpoint(user=None):
    current = user or g.get("user")
    if not current:
        return "login"
    for endpoint, _label, permission in NAV_ITEMS:
        if current.has_permission(permission):
            return endpoint
    if current.has_permission("user.manage"):
        return "role_list"
    return "login"


def can_view_prices(user=None):
    current = user or g.get("user")
    return bool(current and current.has_permission("price.view"))


def can_ship_sales_orders(user=None):
    current = user or g.get("user")
    return bool(current and (current.has_permission("sales.manage") or current.has_permission("stock.out")))


def can_access_purchase_calculator(user=None):
    current = user or g.get("user")
    return bool(current and (current.has_permission("purchase.calc") or current.has_permission("purchase.manage") or current.has_permission("stock.in")))


def can_manage_shipment_sheets(user=None):
    current = user or g.get("user")
    return bool(current and current.has_permission("shipment.manage"))


def can_confirm_shipment_sheets(user=None):
    current = user or g.get("user")
    return bool(current and current.has_permission("shipment.confirm"))


def can_edit_shipment_sheet(shipment, user=None):
    current = user or g.get("user")
    return bool(current and can_manage_shipment_sheets(current) and shipment)


def can_delete_shipment_sheet(shipment=None, user=None):
    current = user or g.get("user")
    return bool(current and can_manage_shipment_sheets(current))


def build_shipment_export_rows(shipments):
    rows = []
    for shipment in shipments:
        for item in shipment.items:
            rows.append(
                [
                    shipment.shipment_no,
                    shipment.store_name,
                    shipment.operator_name,
                    shipment.warehouse.name if shipment.warehouse else "-",
                    shipment.box_count,
                    item.external_sku_code,
                    item.sku.sku_code if item.sku else "-",
                    item.sku.name if item.sku else "-",
                    item.quantity,
                    item.per_box_quantity,
                    shipment.status_label,
                    shipment.warehouse_confirmer_name,
                    format_dt(shipment.created_at),
                    format_dt(shipment.archived_at) if shipment.archived_at else "",
                ]
            )
    return rows


def clean_legacy_display_text(value):
    text = str(value or "")
    replacements = {
        "杩斾慨鍗曪細": "返修单：",
        "渚涘簲鍟嗭細": "供应商：",
        "瀹㈡埛锛": "客户：",
        "搴楅摵锛": "店铺：",
        "浠撶锛": "仓管：",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def can_edit_fba_inbound_record(record, user=None):
    current = user or g.get("user")
    return bool(
        current
        and current.has_permission("fba.inbound.manage")
        and record
        and (record.uploader_user_id is None or record.uploader_user_id == current.id)
    )


def can_edit_fba_inbound_group(records, user=None):
    current = user or g.get("user")
    if not current or not current.has_permission("fba.inbound.manage") or not records:
        return False
    uploader_ids = {record.uploader_user_id for record in records if record.uploader_user_id}
    return not uploader_ids or uploader_ids == {current.id}


def can_access_fba_calculator_page(user=None):
    current = user or g.get("user")
    return bool(
        current
        and (
            current.has_permission("fba.calc.manage")
            or current.has_permission("shipment.manage")
            or current.has_permission("shipment.confirm")
        )
    )


def permission_required(code):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not g.user:
                return redirect(url_for("login"))
            if not g.user.has_permission(code):
                flash("权限不足，无法访问该页面。", "error")
                return redirect(url_for(get_default_landing_endpoint(g.user)))
            return view(*args, **kwargs)

        return wrapped

    return decorator


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not g.user:
            return redirect(url_for("login"))
        if not is_admin_user():
            flash("只有管理员可以访问该页面。", "error")
            return redirect(url_for(get_default_landing_endpoint(g.user)))
        return view(*args, **kwargs)

    return wrapped


def parse_decimal(value, default="0"):
    try:
        return Decimal(value or default)
    except (InvalidOperation, TypeError):
        return Decimal(default)


def parse_int(value, default=0):
    try:
        if value is None:
            return default
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return int(value)
        normalized = str(value).strip().replace(",", "").replace("，", "")
        if not normalized:
            return default
        return int(Decimal(normalized))
    except (InvalidOperation, TypeError, ValueError):
        return default


def parse_float(value, default=0):
    try:
        if value is None:
            return default
        if isinstance(value, bool):
            return float(value)
        if isinstance(value, (int, float)):
            return float(value)
        normalized = str(value).strip().replace(",", "").replace("，", "")
        if not normalized:
            return default
        return float(normalized)
    except (TypeError, ValueError):
        return default


def normalize_external_sku(value):
    text = str(value or "")
    text = text.replace("\ufeff", "").replace("\u200b", "").replace("\xa0", " ")
    return text.strip()


def parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def save_asset_image(file_storage):
    if not file_storage or not (file_storage.filename or "").strip():
        return ""
    extension = os.path.splitext((file_storage.filename or "").lower())[1]
    if extension not in ALLOWED_ASSET_IMAGE_EXTENSIONS:
        raise ValueError("资产图片仅支持 JPG、PNG、WEBP 或 GIF 格式。")
    os.makedirs(ASSET_IMAGE_DIR, exist_ok=True)
    filename = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{secrets.token_hex(4)}{extension}"
    absolute_path = os.path.join(ASSET_IMAGE_DIR, filename)
    file_storage.save(absolute_path)
    return os.path.join(ASSET_IMAGE_RELATIVE_DIR, filename).replace("\\", "/")


def delete_asset_image(image_path):
    if not image_path:
        return
    absolute_path = os.path.normpath(os.path.join(BASE_DIR, "static", image_path))
    static_root = os.path.normpath(os.path.join(BASE_DIR, "static"))
    if not absolute_path.startswith(static_root):
        return
    if os.path.isfile(absolute_path):
        os.remove(absolute_path)


def normalize_status(value, allowed, default):
    return value if value in allowed else default


def get_date_filters():
    start_date = parse_date(request.args.get("start_date", "").strip())
    end_date = parse_date(request.args.get("end_date", "").strip())
    sku_keyword = request.args.get("sku_keyword", "").strip()
    return sku_keyword, start_date, end_date


def apply_datetime_range(query, column, start_date=None, end_date=None):
    if start_date:
        query = query.filter(column >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(column < datetime.combine(end_date, datetime.max.time()))
    return query


def build_sku_keyword_filter(keyword):
    if not keyword:
        return None
    pattern = f"%{keyword}%"
    return or_(SKU.sku_code.ilike(pattern), SKU.name.ilike(pattern), SKU.barcode.ilike(pattern))


def build_unique_code(prefix, model, field_name):
    field = getattr(model, field_name)
    for _ in range(20):
        code = f"{prefix}{datetime.utcnow():%Y%m%d%H%M%S}{secrets.randbelow(10000):04d}"
        if not db.session.query(model.id).filter(field == code).first():
            return code
    raise RuntimeError(f"Unable to generate a unique code for {model.__name__}")


def get_default_warehouse():
    return Warehouse.query.order_by(Warehouse.id.asc()).first()


def get_inventory_map(warehouse_id=None):
    query = db.session.query(InventoryBalance.sku_id, func.coalesce(func.sum(InventoryBalance.quantity), 0))
    if warehouse_id:
        query = query.filter(InventoryBalance.warehouse_id == warehouse_id)
    return {sku_id: qty for sku_id, qty in query.group_by(InventoryBalance.sku_id).all()}


def get_inventory_by_warehouse():
    rows = (
        db.session.query(
            InventoryBalance.warehouse_id,
            InventoryBalance.sku_id,
            func.coalesce(func.sum(InventoryBalance.quantity), 0),
        )
        .group_by(InventoryBalance.warehouse_id, InventoryBalance.sku_id)
        .all()
    )
    return {(warehouse_id, sku_id): qty for warehouse_id, sku_id, qty in rows}


def build_sales_velocity_metrics(*, sales_7_qty, sales_14_qty, sales_30_qty):
    daily_7 = round(sales_7_qty / 7, 2) if sales_7_qty > 0 else 0
    daily_14 = round(sales_14_qty / 14, 2) if sales_14_qty > 0 else 0
    daily_30 = round(sales_30_qty / 30, 2) if sales_30_qty > 0 else 0
    if daily_14 > 0 and daily_30 > 0:
        weighted_daily = round(daily_14 * 0.6 + daily_30 * 0.4, 2)
    else:
        weighted_daily = max(daily_14, daily_30, daily_7)
    base_daily = weighted_daily
    baseline_daily = weighted_daily or max(daily_14, daily_30)
    trend_factor = 1
    if daily_7 > 0 and baseline_daily > 0:
        trend_ratio = daily_7 / baseline_daily
        if trend_ratio > 1.15:
            trend_factor = min(max(trend_ratio, 1.1), 1.3)
        elif trend_ratio < 0.85:
            trend_factor = max(min(trend_ratio, 0.9), 0.7)
    elif daily_7 > 0 and daily_30 > 0:
        trend_factor = min(max(daily_7 / daily_30, 0.7), 1.3)
    return {
        "daily_7": daily_7,
        "daily_14": daily_14,
        "daily_30": daily_30,
        "weighted_daily": weighted_daily,
        "base_daily": base_daily,
        "trend_factor": round(trend_factor, 2),
    }


def get_open_purchase_qty_map(*, warehouse_id=None):
    query = (
        db.session.query(
            PurchaseOrder.warehouse_id,
            PurchaseOrderItem.sku_id,
            func.coalesce(func.sum(PurchaseOrderItem.quantity), 0),
        )
        .join(PurchaseOrderItem, PurchaseOrderItem.purchase_order_id == PurchaseOrder.id)
        .filter(PurchaseOrder.status.in_(PURCHASE_CALC_OPEN_STATUSES))
    )
    if warehouse_id:
        query = query.filter(PurchaseOrder.warehouse_id == warehouse_id)
    rows = query.group_by(PurchaseOrder.warehouse_id, PurchaseOrderItem.sku_id).all()
    return {(warehouse_id, sku_id): parse_int(quantity) for warehouse_id, sku_id, quantity in rows}


def get_outbound_consumption_map(*, warehouse_id=None, days=30):
    since = datetime.utcnow() - timedelta(days=max(parse_int(days, 30), 1))
    quantity = func.coalesce(func.sum(-InventoryTransaction.quantity), 0)
    query = (
        db.session.query(InventoryTransaction.sku_id, quantity)
        .filter(
            InventoryTransaction.transaction_type.in_(PURCHASE_CALC_OUTBOUND_TYPES),
            InventoryTransaction.quantity < 0,
            InventoryTransaction.created_at >= since,
        )
    )
    if warehouse_id:
        query = query.filter(InventoryTransaction.warehouse_id == warehouse_id)
    rows = query.group_by(InventoryTransaction.sku_id).all()
    return {sku_id: parse_int(total) for sku_id, total in rows}


def build_purchase_calculator_result(*, warehouse_id, coverage_days, lead_days, sku_keyword="", suggested_min=0, positive_only=False, user=None, sales_windows=None):
    warehouse = db.session.get(Warehouse, warehouse_id) if warehouse_id else get_default_warehouse()
    warehouse_id = warehouse.id if warehouse else None
    coverage_days = max(parse_int(coverage_days, 30), 1)
    lead_days = max(parse_int(lead_days, 15), 1)
    suggested_min = max(parse_int(suggested_min, 0), 0)
    sales_data_provided = sales_windows is not None
    sales_windows = sales_windows or {}
    sales_7_map = sales_windows.get("sales_7") or {}
    sales_14_map = sales_windows.get("sales_14") or {}
    sales_30_map = sales_windows.get("sales_30") or {}
    sales_sku_ids = {
        sku_id
        for source_map in (sales_7_map, sales_14_map, sales_30_map)
        for sku_id in (parse_int(raw_sku_id) for raw_sku_id in source_map.keys())
        if sku_id
    }
    sku_query = apply_sku_scope(SKU.query, SKU.id, user)
    if sales_data_provided:
        sku_query = sku_query.filter(SKU.id.in_(sales_sku_ids)) if sales_sku_ids else sku_query.filter(False)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        sku_query = sku_query.filter(sku_filter)
    skus = sku_query.order_by(SKU.sku_code.asc()).all()

    inventory_map = get_inventory_by_warehouse()
    open_purchase_map = get_open_purchase_qty_map(warehouse_id=warehouse_id)
    outbound_7_map = get_outbound_consumption_map(warehouse_id=warehouse_id, days=7)
    outbound_14_map = get_outbound_consumption_map(warehouse_id=warehouse_id, days=14)
    outbound_30_map = get_outbound_consumption_map(warehouse_id=warehouse_id, days=30)
    rows = []

    for sku in skus:
        current_stock = parse_int(inventory_map.get((warehouse_id, sku.id))) if warehouse_id else 0
        open_purchase_qty = parse_int(open_purchase_map.get((warehouse_id, sku.id))) if warehouse_id else 0
        outbound_7 = parse_int(outbound_7_map.get(sku.id))
        outbound_14 = parse_int(outbound_14_map.get(sku.id))
        outbound_30 = parse_int(outbound_30_map.get(sku.id))
        outbound_daily = round(max(outbound_7 / 7 if outbound_7 else 0, outbound_14 / 14 if outbound_14 else 0, outbound_30 / 30 if outbound_30 else 0), 2)
        sales_7 = parse_int(sales_7_map.get(sku.id))
        sales_14 = parse_int(sales_14_map.get(sku.id))
        sales_30 = parse_int(sales_30_map.get(sku.id))
        sales_metrics = build_sales_velocity_metrics(
            sales_7_qty=sales_7,
            sales_14_qty=sales_14,
            sales_30_qty=sales_30,
        )
        sales_trend_daily = round(sales_metrics["base_daily"] * sales_metrics["trend_factor"], 2)
        daily_consumption = sales_trend_daily if sales_data_provided else outbound_daily
        if sales_data_provided:
            demand_source = "销量趋势" if sales_7 or sales_14 or sales_30 else "销量表无销量"
        else:
            demand_source = "出库消耗"
        lead_consumption = math.ceil(daily_consumption * lead_days)
        target_stock = math.ceil(parse_int(sku.safety_stock) + daily_consumption * coverage_days)
        projected_available = current_stock + open_purchase_qty - lead_consumption
        suggested_qty = max(target_stock - projected_available, 0)
        if positive_only and suggested_qty <= 0:
            continue
        if suggested_min and suggested_qty < suggested_min:
            continue
        rows.append(
            {
                "sku_id": sku.id,
                "sku_code": sku.sku_code,
                "name": sku.name,
                "color": sku.color,
                "size": sku.size,
                "unit": sku.unit,
                "cost_price": float(sku.cost_price or 0),
                "safety_stock": parse_int(sku.safety_stock),
                "current_stock": current_stock,
                "open_purchase_qty": open_purchase_qty,
                "outbound_7": outbound_7,
                "outbound_14": outbound_14,
                "outbound_30": outbound_30,
                "outbound_daily": outbound_daily,
                "sales_7": sales_7,
                "sales_14": sales_14,
                "sales_30": sales_30,
                "sales_daily_7": sales_metrics["daily_7"],
                "sales_daily_14": sales_metrics["daily_14"],
                "sales_daily_30": sales_metrics["daily_30"],
                "sales_weighted_daily": sales_metrics["weighted_daily"],
                "sales_trend_factor": sales_metrics["trend_factor"],
                "sales_trend_daily": sales_trend_daily,
                "demand_source": demand_source,
                "daily_consumption": daily_consumption,
                "lead_consumption": lead_consumption,
                "target_stock": target_stock,
                "projected_available": projected_available,
                "suggested_qty": suggested_qty,
            }
        )

    rows.sort(key=lambda row: (-row["suggested_qty"], row["sku_code"].casefold()))
    return {
        "rows": rows,
        "summary": {
            "sku_count": len(rows),
            "suggested_sku_count": sum(1 for row in rows if row["suggested_qty"] > 0),
            "suggested_qty": sum(row["suggested_qty"] for row in rows),
            "current_stock": sum(row["current_stock"] for row in rows),
            "open_purchase_qty": sum(row["open_purchase_qty"] for row in rows),
            "sales_30": sum(row["sales_30"] for row in rows),
        },
        "params": {
            "warehouse_id": warehouse_id,
            "coverage_days": coverage_days,
            "lead_days": lead_days,
            "sku_keyword": sku_keyword,
            "suggested_min": suggested_min,
            "positive_only": bool(positive_only),
            "sales_data_provided": bool(sales_data_provided),
        },
    }


def get_dashboard_sku_rankings(*, transaction_types, quantity_expression, sku_keyword="", limit=10, user=None):
    total_quantity = func.coalesce(func.sum(quantity_expression), 0).label("total_quantity")
    query = (
        db.session.query(
            SKU.id,
            SKU.sku_code,
            SKU.name,
            SKU.color,
            SKU.size,
            total_quantity,
        )
        .join(InventoryTransaction, InventoryTransaction.sku_id == SKU.id)
        .filter(InventoryTransaction.transaction_type.in_(transaction_types))
    )
    query = apply_sku_scope(query, InventoryTransaction.sku_id, user)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        query = query.filter(sku_filter)
    rows = (
        query.group_by(SKU.id, SKU.sku_code, SKU.name, SKU.color, SKU.size)
        .order_by(total_quantity.desc(), SKU.sku_code.asc())
        .limit(limit)
        .all()
    )
    return [
        {
            "sku_id": sku_id,
            "sku_code": sku_code,
            "name": name,
            "color": color,
            "size": size,
            "quantity": int(quantity or 0),
        }
        for sku_id, sku_code, name, color, size, quantity in rows
    ]


def get_defect_status_map(*, warehouse_id=None, sku_filter=None, start_date=None, end_date=None):
    query = DefectRepair.query
    if warehouse_id:
        query = query.filter(DefectRepair.warehouse_id == warehouse_id)
    if sku_filter is not None:
        query = query.join(DefectRepair.sku).filter(sku_filter)
    query = apply_datetime_range(query, DefectRepair.created_at, start_date, end_date)
    defect_status_map = {}
    for repair in query.all():
        pending_or_repairing_quantity = repair.remaining_quantity
        if pending_or_repairing_quantity:
            status = "repairing" if repair.status == "repairing" else "pending"
            defect_status_map[(repair.sku_id, status)] = (
                defect_status_map.get((repair.sku_id, status), 0) + pending_or_repairing_quantity
            )
        if repair.completed_quantity:
            defect_status_map[(repair.sku_id, "done")] = (
                defect_status_map.get((repair.sku_id, "done"), 0) + repair.completed_quantity
            )
        if repair.scrapped_quantity:
            defect_status_map[(repair.sku_id, "scrapped")] = (
                defect_status_map.get((repair.sku_id, "scrapped"), 0) + repair.scrapped_quantity
            )
    return defect_status_map


def sync_inventory_balances(commit=True):
    summary = (
        db.session.query(
            InventoryTransaction.sku_id,
            InventoryTransaction.warehouse_id,
            func.coalesce(func.sum(InventoryTransaction.quantity), 0),
        )
        .group_by(InventoryTransaction.sku_id, InventoryTransaction.warehouse_id)
        .all()
    )
    InventoryBalance.query.delete()
    for sku_id, warehouse_id, quantity in summary:
        db.session.add(InventoryBalance(sku_id=sku_id, warehouse_id=warehouse_id, quantity=quantity))
    if commit:
        db.session.commit()


def ensure_barcode_optional_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    columns = db.session.execute(text("PRAGMA table_info(sku)")).mappings().all()
    if not columns:
        return

    barcode_column = next((column for column in columns if column["name"] == "barcode"), None)
    if not barcode_column or barcode_column["notnull"] == 0:
        return

    db.session.execute(text("PRAGMA foreign_keys=OFF"))
    db.session.execute(
        text(
            """
            CREATE TABLE sku_new (
                id INTEGER NOT NULL PRIMARY KEY,
                created_at DATETIME NOT NULL,
                updated_at DATETIME NOT NULL,
                sku_code VARCHAR(64) NOT NULL UNIQUE,
                barcode VARCHAR(64) UNIQUE,
                name VARCHAR(128) NOT NULL,
                category VARCHAR(64) NOT NULL DEFAULT '',
                color VARCHAR(64) NOT NULL DEFAULT '',
                size VARCHAR(32) NOT NULL DEFAULT '',
                unit VARCHAR(16) NOT NULL DEFAULT '',
                cost_price NUMERIC(10, 2) NOT NULL DEFAULT 0,
                sale_price NUMERIC(10, 2) NOT NULL DEFAULT 0,
                safety_stock INTEGER NOT NULL DEFAULT 0,
                remark VARCHAR(255) NOT NULL DEFAULT ''
            )
            """
        )
    )
    db.session.execute(
        text(
            """
            INSERT INTO sku_new (
                id, created_at, updated_at, sku_code, barcode, name, category, color, size, unit,
                cost_price, sale_price, safety_stock, remark
            )
            SELECT
                id, created_at, updated_at, sku_code, NULLIF(barcode, ''), name, category, color, size, unit,
                cost_price, sale_price, safety_stock, remark
            FROM sku
            """
        )
    )
    db.session.execute(text("DROP TABLE sku"))
    db.session.execute(text("ALTER TABLE sku_new RENAME TO sku"))
    db.session.commit()
    db.session.execute(text("PRAGMA foreign_keys=ON"))


def ensure_defect_repair_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    columns = db.session.execute(text("PRAGMA table_info(defect_repair)")).mappings().all()
    if not columns:
        return

    column_names = {column["name"] for column in columns}

    if "warehouse_id" not in column_names:
        default_warehouse = get_default_warehouse()
        default_warehouse_id = default_warehouse.id if default_warehouse else None
        db.session.execute(text("ALTER TABLE defect_repair ADD COLUMN warehouse_id INTEGER"))
        if default_warehouse_id:
            db.session.execute(
                text("UPDATE defect_repair SET warehouse_id = :warehouse_id WHERE warehouse_id IS NULL"),
                {"warehouse_id": default_warehouse_id},
            )
    if "completed_quantity" not in column_names:
        db.session.execute(text("ALTER TABLE defect_repair ADD COLUMN completed_quantity INTEGER NOT NULL DEFAULT 0"))
        db.session.execute(
            text("UPDATE defect_repair SET completed_quantity = quantity WHERE status = 'done'")
        )
    if "scrapped_quantity" not in column_names:
        db.session.execute(text("ALTER TABLE defect_repair ADD COLUMN scrapped_quantity INTEGER NOT NULL DEFAULT 0"))
        db.session.execute(
            text("UPDATE defect_repair SET scrapped_quantity = quantity WHERE status = 'scrapped'")
        )
    db.session.commit()


def ensure_operator_name_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    table_names = ("purchase_order", "sales_order", "inventory_transaction", "defect_repair")
    for table_name in table_names:
        columns = db.session.execute(text(f"PRAGMA table_info({table_name})")).mappings().all()
        if not columns or any(column["name"] == "operator_name" for column in columns):
            continue
        db.session.execute(text(f"ALTER TABLE {table_name} ADD COLUMN operator_name VARCHAR(128) NOT NULL DEFAULT ''"))
    db.session.commit()


def ensure_inventory_supplier_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    columns = db.session.execute(text("PRAGMA table_info(inventory_transaction)")).mappings().all()
    if not columns:
        return
    if not any(column["name"] == "supplier_id" for column in columns):
        db.session.execute(text("ALTER TABLE inventory_transaction ADD COLUMN supplier_id INTEGER"))
        db.session.commit()


def ensure_login_security_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    user_columns = db.session.execute(text("PRAGMA table_info(user)")).mappings().all()
    if user_columns:
        if not any(column["name"] == "failed_login_attempts" for column in user_columns):
            db.session.execute(text("ALTER TABLE user ADD COLUMN failed_login_attempts INTEGER NOT NULL DEFAULT 0"))
        if not any(column["name"] == "locked_until" for column in user_columns):
            db.session.execute(text("ALTER TABLE user ADD COLUMN locked_until DATETIME"))
        if not any(column["name"] == "last_failed_login_at" for column in user_columns):
            db.session.execute(text("ALTER TABLE user ADD COLUMN last_failed_login_at DATETIME"))
        if not any(column["name"] == "is_emergency_account" for column in user_columns):
            db.session.execute(text("ALTER TABLE user ADD COLUMN is_emergency_account BOOLEAN NOT NULL DEFAULT 0"))

    login_audit_columns = db.session.execute(text("PRAGMA table_info(login_audit)")).mappings().all()
    if login_audit_columns:
        if not any(column["name"] == "is_success" for column in login_audit_columns):
            db.session.execute(text("ALTER TABLE login_audit ADD COLUMN is_success BOOLEAN NOT NULL DEFAULT 1"))
        if not any(column["name"] == "failure_reason" for column in login_audit_columns):
            db.session.execute(text("ALTER TABLE login_audit ADD COLUMN failure_reason VARCHAR(128) NOT NULL DEFAULT ''"))
        if not any(column["name"] == "locked_until" for column in login_audit_columns):
            db.session.execute(text("ALTER TABLE login_audit ADD COLUMN locked_until DATETIME"))
        if not any(column["name"] == "lock_triggered" for column in login_audit_columns):
            db.session.execute(text("ALTER TABLE login_audit ADD COLUMN lock_triggered BOOLEAN NOT NULL DEFAULT 0"))

    db.session.commit()


def ensure_asset_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    asset_columns = db.session.execute(text("PRAGMA table_info(asset)")).mappings().all()
    if not asset_columns:
        return
    if not any(column["name"] == "asset_value" for column in asset_columns):
        db.session.execute(text("ALTER TABLE asset ADD COLUMN asset_value NUMERIC(12, 2) NOT NULL DEFAULT 0"))
    if not any(column["name"] == "image_path" for column in asset_columns):
        db.session.execute(text("ALTER TABLE asset ADD COLUMN image_path VARCHAR(255) NOT NULL DEFAULT ''"))
    db.session.commit()


def ensure_fba_inbound_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    inbound_columns = db.session.execute(text("PRAGMA table_info(fba_inbound_record)")).mappings().all()
    if not inbound_columns:
        return
    if not any(column["name"] == "uploader_user_id" for column in inbound_columns):
        db.session.execute(text("ALTER TABLE fba_inbound_record ADD COLUMN uploader_user_id INTEGER"))
    db.session.commit()


def ensure_shipment_archive_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    shipment_columns = db.session.execute(text("PRAGMA table_info(shipment_sheet)")).mappings().all()
    if not shipment_columns:
        return
    if not any(column["name"] == "archived_at" for column in shipment_columns):
        db.session.execute(text("ALTER TABLE shipment_sheet ADD COLUMN archived_at DATETIME"))
    db.session.commit()


def ensure_freight_compare_schema():
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        return

    rate_book_columns = db.session.execute(text("PRAGMA table_info(freight_rate_book)")).mappings().all()
    if rate_book_columns and not any(column["name"] == "clothing_surcharge_per_kg" for column in rate_book_columns):
        db.session.execute(text("ALTER TABLE freight_rate_book ADD COLUMN clothing_surcharge_per_kg NUMERIC(10, 2) NOT NULL DEFAULT 0"))
        db.session.execute(
            text("UPDATE freight_rate_book SET clothing_surcharge_per_kg = 1 WHERE forwarder_name = :forwarder_name"),
            {"forwarder_name": "鼎邦"},
        )
        db.session.commit()


def validate_entity_ids(model, ids):
    clean_ids = sorted({value for value in ids if value})
    if not clean_ids:
        return {}
    entities = model.query.filter(model.id.in_(clean_ids)).all()
    return {entity.id: entity for entity in entities}


def parse_order_items(prefix, *, allow_price=True):
    sku_ids = request.form.getlist(f"{prefix}_sku_id")
    quantities = request.form.getlist(f"{prefix}_quantity")
    prices = request.form.getlist(f"{prefix}_price") if allow_price else []
    if len(prices) < len(sku_ids):
        prices.extend(["0"] * (len(sku_ids) - len(prices)))
    items = []
    for sku_id, qty, price in zip(sku_ids, quantities, prices):
        parsed_sku_id = parse_int(sku_id)
        parsed_qty = parse_int(qty)
        parsed_price = parse_decimal(price)
        if not parsed_sku_id or parsed_qty <= 0 or parsed_price < 0:
            continue
        items.append({"sku_id": parsed_sku_id, "quantity": parsed_qty, "unit_price": parsed_price})
    return items


def parse_shipment_sheet_items(*, box_count, warehouse_id, user=None, exclude_shipment_id=None):
    external_codes = request.form.getlist("shipment_external_sku_code")
    quantities = request.form.getlist("shipment_quantity")
    items = []
    errors = []
    current = user or g.get("user")
    mapping_query = SKUMapping.query
    if current and not can_manage_all_sku_mappings(current):
        mapping_query = mapping_query.filter(SKUMapping.user_id == current.id)
    mapping_rows = mapping_query.all()
    mapping_by_external_code = {row.external_sku_code.strip().lower(): row for row in mapping_rows if row.external_sku_code.strip()}
    available_map = get_shipment_available_inventory_map(warehouse_id=warehouse_id, exclude_shipment_id=exclude_shipment_id)

    for index, (external_code_raw, quantity_raw) in enumerate(zip(external_codes, quantities), start=1):
        external_code = (external_code_raw or "").strip()
        quantity = parse_int(quantity_raw)
        if not external_code and quantity <= 0:
            continue
        if not external_code:
            errors.append(f"第 {index} 行：请选择店铺SKU。")
            continue
        if quantity <= 0:
            errors.append(f"第 {index} 行：数量必须大于 0。")
            continue
        if box_count <= 0:
            errors.append("箱数必须大于 0。")
            break
        if quantity % box_count != 0:
            errors.append(f"第 {index} 行：数量 {quantity} 不能被箱数 {box_count} 整除。")
            continue
        mapping = mapping_by_external_code.get(external_code.lower())
        if not mapping:
            errors.append(f"第 {index} 行：店铺SKU {external_code} 未找到映射。")
            continue
        available_quantity = int(available_map.get((warehouse_id, mapping.sku_id), 0) or 0)
        if quantity > available_quantity:
            errors.append(f"第 {index} 行：店铺SKU {external_code} 对应仓库SKU {mapping.sku.sku_code} 库存不足，实时可用库存仅 {available_quantity}。")
            continue
        items.append(
            {
                "sku_id": mapping.sku_id,
                "external_sku_code": external_code,
                "quantity": quantity,
                "per_box_quantity": quantity // box_count,
                "sku_code": mapping.sku.sku_code,
                "sku_name": mapping.sku.name,
                "available_quantity": available_quantity,
            }
        )

    return items, errors


def parse_uploaded_table(file_storage):
    filename = (file_storage.filename or "").strip()
    extension = os.path.splitext(filename.lower())[1]
    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(file_storage, data_only=True, read_only=True)
        worksheet = workbook.active
        return [[str(cell).strip() if cell is not None else "" for cell in row] for row in worksheet.iter_rows(values_only=True)]

    raw = file_storage.read()
    text_content = None

    def score_decoded_text(decoded, *, csv_like=False):
        replacement_count = decoded.count("\ufffd")
        null_count = decoded.count("\x00")
        control_count = sum(
            1
            for char in decoded
            if (ord(char) < 32 or 0x7F <= ord(char) <= 0x9F) and char not in "\r\n\t"
        )
        line_count = len(decoded.splitlines())
        sample_lines = decoded.splitlines()[:5]
        delimiter_hits = 0
        if csv_like:
            delimiter_hits = max(sum(line.count(candidate) for line in sample_lines) for candidate in ("\t", ",", "|", ";"))
        printable_count = sum(1 for char in decoded if char.isprintable() or char in "\r\n\t")
        suspicious_ratio = 0 if not decoded else 1 - (printable_count / len(decoded))
        return (
            replacement_count,
            null_count,
            control_count,
            suspicious_ratio,
            -delimiter_hits,
            -line_count,
            len(decoded),
        )

    csv_like_extension = extension in {".txt", ".csv", ".tsv"}
    encodings_to_try = [
        "utf-8-sig",
        "utf-8",
        "utf-16",
        "utf-16le",
        "utf-16be",
        "utf-32",
        "utf-32le",
        "utf-32be",
        "gb18030",
        "gbk",
        "gb2312",
        "cp936",
        "mbcs",
        "big5",
        "cp1252",
        "latin1",
    ]
    decode_candidates = []
    for encoding in encodings_to_try:
        try:
            decoded = raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        decode_candidates.append((score_decoded_text(decoded, csv_like=csv_like_extension), encoding, decoded))

    if decode_candidates:
        decode_candidates.sort(key=lambda item: item[0])
        best_score, _best_encoding, best_text = decode_candidates[0]
        if best_score[0] <= 2 and best_score[1] == 0:
            text_content = best_text

    if text_content is None and csv_like_extension:
        fallback_candidates = []
        for encoding in ("utf-8", "utf-8-sig", "utf-16le", "utf-16be", "gb18030", "gbk", "cp936", "big5", "cp1252", "latin1"):
            try:
                decoded = raw.decode(encoding, errors="replace")
            except LookupError:
                continue
            fallback_candidates.append((score_decoded_text(decoded, csv_like=True), encoding, decoded))
        if fallback_candidates:
            fallback_candidates.sort(key=lambda item: item[0])
            best_score, _best_encoding, best_text = fallback_candidates[0]
            if best_score[0] <= 2 and best_score[1] == 0:
                text_content = best_text

    if text_content is None:
        raise ValueError("文件编码无法识别。7天销量和30天销量的 .txt 表建议直接使用亚马逊原始导出文件，或另存为 UTF-8 / UTF-16 / GBK / GB18030。")

    sample = "\n".join(text_content.splitlines()[:5])
    delimiter = "\t" if extension == ".tsv" else ","
    if extension in {".txt", ".csv", ".tsv"}:
        delimiter_candidates = ["\t", ",", "|", ";"]
        delimiter = max(delimiter_candidates, key=lambda item: sample.count(item))
    reader = csv.reader(StringIO(text_content), delimiter=delimiter)
    return [[str(cell).strip() for cell in row] for row in reader]


def clean_freight_cell(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:g}"
    return str(value).replace("\ufeff", "").replace("\u200b", "").replace("\u200c", "").replace("\xa0", " ").replace("\n", " ").replace("\r", " ").strip()


def freight_number(value, default=None):
    text_value = clean_freight_cell(value).replace(",", "").replace("，", "")
    if not text_value or text_value in {"-", "——", "单询", "详询"}:
        return default
    match = re.search(r"-?\d+(?:\.\d+)?", text_value)
    if not match:
        return default
    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return default


def normalize_freight_forwarder(value="", filename=""):
    text_value = f"{value or ''} {filename or ''}"
    if "盈和" in text_value:
        return "盈和"
    if "开渠达" in text_value:
        return "开渠达"
    if "鼎邦" in text_value:
        return "鼎邦"
    if "鸿帆" in text_value or "hongfan" in text_value.lower() or "hongvan" in text_value.lower():
        return "鸿帆"
    return (value or "").strip() or "未命名货代"


def normalize_zip_code(value):
    match = re.search(r"\d{4,5}", clean_freight_cell(value))
    if not match:
        return ""
    return match.group(0).zfill(5)


def extract_zip_code(value):
    return normalize_zip_code(value)


def extract_warehouse_codes(value):
    text_value = clean_freight_cell(value).upper()
    ignored = {
        "FBA",
        "UPS",
        "USPS",
        "FEDEX",
        "DHL",
        "LTL",
        "RMB",
        "KG",
        "KGS",
        "CBM",
        "MAX",
        "CLX",
        "YHE",
        "PO",
        "POD",
        "OA",
        "SZ",
        "USA",
        "DDP",
        "WALMART",
        "TIKTOK",
        "WAYFAIR",
        "ORANGE",
        "AMAZON",
        "SERVICES",
        "INC",
    }
    codes = []
    for match in re.finditer(r"(?<![A-Z0-9])[A-Z]{2,}[A-Z0-9]{1,5}(?![A-Z0-9])", text_value):
        code = match.group(0).strip()
        if code in ignored or len(code) < 3 or len(code) > 7:
            continue
        if code not in codes:
            codes.append(code)
    return codes


def parse_transit_days(value):
    text_value = clean_freight_cell(value)
    if "天" not in text_value:
        return None, None
    normalized = re.sub(r"[\u200b-\u200f\ufeff]", "", text_value)

    def clean_transit_pair(first, second=None):
        first_value = parse_int(first)
        second_value = parse_int(second) if second else first_value
        if first_value <= 0 or second_value <= 0:
            return None, None
        if first_value > 90 or second_value > 90:
            return None, None
        return min(first_value, second_value), max(first_value, second_value)

    semantic_patterns = (
        r"(?:预计|标准)?时效[^0-9]{0,40}(\d{1,2})\s*天\s*[-~—至到]?\s*(\d{1,2})?\s*天?",
        r"(?:预计|标准)?时效[^0-9]{0,40}(\d{1,2})\s*[-~—至到]\s*(\d{1,2})\s*天",
        r"开船后?\s*(\d{1,2})\s*[-~—至到]\s*(\d{1,2})\s*天",
        r"开船后?\s*(\d{1,2})\s*天\s*[-~—至到]?\s*(\d{1,2})?\s*天?",
        r"预计开船\s*(\d{1,2})\s*[-~—至到]\s*(\d{1,2})\s*天",
        r"预计开船\s*(\d{1,2})\s*天\s*[-~—至到]?\s*(\d{1,2})?\s*天?",
        r"(\d{1,2})\s*[-~—至到]\s*(\d{1,2})\s*天(?:左右)?(?:送仓|提取|到仓|入仓|派送|签收)",
        r"(\d{1,2})\s*天\s*[-~—至到]?\s*(\d{1,2})?\s*天?(?:左右)?(?:送仓|提取|到仓|入仓|派送|签收)",
        r"(\d{1,2})\s*天(?:极速达|限时达|快船|美森)",
    )
    for pattern in semantic_patterns:
        match = re.search(pattern, normalized)
        if match:
            return clean_transit_pair(match.group(1), match.group(2) if len(match.groups()) > 1 else None)

    if "邮编" in normalized or "赔" in normalized or "延迟" in normalized or "超过" in normalized or "申请调查" in normalized:
        return None, None
    numbers = [parse_int(item) for item in re.findall(r"\d{1,2}", normalized)]
    numbers = [item for item in numbers if 0 < item <= 90]
    if not numbers:
        return None, None
    if len(numbers) == 1:
        return numbers[0], numbers[0]
    return min(numbers[:2]), max(numbers[:2])


def zip_range_from_text(value):
    text_value = clean_freight_cell(value)
    range_match = re.search(r"(\d{4,5})\s*[-~—至]\s*(\d{4,5})", text_value)
    if range_match:
        start = parse_int(range_match.group(1).zfill(5))
        end = parse_int(range_match.group(2).zfill(5))
        return min(start, end), max(start, end)
    if "美西" in text_value or "8-9" in text_value or "8-9）" in text_value:
        return 80000, 99999
    if "美中" in text_value or "4-7" in text_value:
        return 40000, 79999
    if "美东" in text_value or "0-3" in text_value:
        return 0, 39999
    return None, None


def zip_ranges_from_text(value):
    text_value = clean_freight_cell(value)
    ranges = []
    for match in re.finditer(r"(\d{4,5})\s*[-~—至]\s*(\d{4,5})", text_value):
        start = parse_int(match.group(1).zfill(5))
        end = parse_int(match.group(2).zfill(5))
        ranges.append((min(start, end), max(start, end)))
    has_explicit_range = bool(ranges)
    if not has_explicit_range and ("美西" in text_value or "8-9" in text_value or "8-9）" in text_value):
        ranges.append((80000, 99999))
    if "美中" in text_value or "4-7" in text_value:
        ranges.append((40000, 79999))
        if not has_explicit_range:
            ranges.append((97000, 99999))
    if not has_explicit_range and ("美东" in text_value or "0-3" in text_value):
        ranges.append((0, 39999))
    if not ranges:
        zip_code = normalize_zip_code(text_value)
        if zip_code:
            zip_number = parse_int(zip_code)
            ranges.append((zip_number, zip_number))
    deduped = []
    seen = set()
    for start, end in ranges:
        if start is None or end is None:
            continue
        key = (start, end)
        if key not in seen:
            deduped.append(key)
            seen.add(key)
    return deduped


def normalize_origin_region(value):
    text_value = clean_freight_cell(value)
    if not text_value:
        return ""
    if ("深圳" in text_value or "华南" in text_value) and ("泉州" in text_value or "福建" in text_value):
        return "泉州/深圳"
    if text_value.upper().startswith("SZ") or "深圳" in text_value or "华南" in text_value:
        return "深圳" if "深圳" in text_value or text_value.upper().startswith("SZ") else "华南"
    if "义乌" in text_value or "华东" in text_value or "宁波" in text_value or "上海" in text_value:
        return "义乌" if "义乌" in text_value else "华东"
    if "泉州" in text_value or "福建" in text_value:
        return "泉州"
    if "青岛" in text_value:
        return "青岛"
    return text_value


def freight_origin_matches(rate_origin, selected_origin):
    selected = normalize_origin_region(selected_origin)
    origin = normalize_origin_region(rate_origin)
    if not selected or selected in {"全部", "不限"} or not origin:
        return True
    aliases = {
        "义乌": {"义乌", "华东", "宁波", "上海"},
        "华东": {"义乌", "华东", "宁波", "上海"},
        "深圳": {"深圳", "华南", "泉州/深圳"},
        "华南": {"深圳", "华南", "泉州/深圳"},
        "泉州": {"泉州", "福建", "泉州/深圳"},
        "福建": {"泉州", "福建", "泉州/深圳"},
        "泉州/深圳": {"深圳", "华南", "泉州", "福建", "泉州/深圳"},
        "青岛": {"青岛"},
    }
    return origin in aliases.get(selected, {selected}) or selected in aliases.get(origin, {origin})


def freight_sheet_by_keywords(workbook, *keywords):
    for worksheet in workbook.worksheets:
        if all(keyword in worksheet.title for keyword in keywords):
            return worksheet
    return None


def iter_freight_row_values(worksheet, row_index, max_col=None):
    column_count = min(worksheet.max_column, max_col or worksheet.max_column)
    return [clean_freight_cell(worksheet.cell(row_index, column_index).value) for column_index in range(1, column_count + 1)]


def build_freight_rate(
    rate_book,
    *,
    channel_name,
    origin_region="",
    destination_type="warehouse",
    warehouse_code="",
    zone_name="",
    postal_code="",
    postal_start=None,
    postal_end=None,
    weight_break_kg=0,
    unit_price=None,
    cbm_price=None,
    min_chargeable_kg=0,
    min_piece_kg=0,
    divisor=6000,
    transit_days_min=None,
    transit_days_max=None,
    tax_mode="",
    source_sheet="",
    notes="",
):
    if unit_price is None and cbm_price is None:
        return None
    return FreightRate(
        rate_book=rate_book,
        forwarder_name=rate_book.forwarder_name,
        channel_name=(channel_name or "").strip(),
        origin_region=normalize_origin_region(origin_region),
        destination_type=destination_type,
        warehouse_code=(warehouse_code or "").upper().strip(),
        zone_name=(zone_name or "").strip(),
        postal_code=postal_code or "",
        postal_start=postal_start,
        postal_end=postal_end,
        weight_break_kg=weight_break_kg or 0,
        unit_price=unit_price,
        cbm_price=cbm_price,
        min_chargeable_kg=min_chargeable_kg or weight_break_kg or 0,
        min_piece_kg=min_piece_kg or 0,
        divisor=divisor or 6000,
        transit_days_min=transit_days_min,
        transit_days_max=transit_days_max,
        tax_mode=tax_mode or "",
        source_sheet=source_sheet or "",
        notes=(notes or "")[:2000],
    )


def extract_customs_rule_from_text(rate_book, channel_name, text_value):
    text_value = clean_freight_cell(text_value)
    if "报关" not in text_value and "品名" not in text_value and "商检" not in text_value:
        return None
    fee_matches = []
    for pattern in (
        r"(?:一般(?:贸易)?报关|单独报关|报关费|接单独报关件)[：:，,、\s]*(\d{2,4})\s*(?:元|RMB|CNY)?\s*/?\s*(?:一?票|票)",
        r"(?:一般(?:贸易)?报关|单独报关|报关费|接单独报关件)[^\d]{0,12}(\d{2,4})\s*(?:元|RMB|CNY)?\s*/?\s*(?:一?票|票)",
    ):
        fee_matches.extend(re.findall(pattern, text_value, flags=re.IGNORECASE))
    customs_fee = Decimal(fee_matches[0]) if fee_matches else Decimal("0")
    item_limit = 0
    for pattern in (r"限\s*(\d+)\s*个品名", r"单票\s*(\d+)\s*个品名", r"超过\s*(\d+)\s*个品名", r"超出\s*(\d+)\s*个"):
        match = re.search(pattern, text_value)
        if match:
            item_limit = parse_int(match.group(1))
            break
    extra_fee_match = re.search(r"加收\s*(?:RMB)?\s*(\d{1,4})\s*元?\s*/?\s*个", text_value)
    item_extra_fee = Decimal(extra_fee_match.group(1)) if extra_fee_match else Decimal("0")
    inspection_matches = re.findall(r"普通商检[^\d]{0,10}(\d{2,4})", text_value)
    formal_matches = re.findall(r"正式报关商检[^\d]{0,12}(\d{2,4})", text_value)
    supports_merge = "不支持合并报关" not in text_value and "不支持跨区合并报关" not in text_value
    merge_scope = "batch"
    if not supports_merge:
        merge_scope = "zone" if "分区" in text_value or "跨区" in text_value else "item"
    if not any([customs_fee, item_limit, item_extra_fee, inspection_matches, formal_matches, "不支持合并报关" in text_value]):
        return None
    return FreightCustomsRule(
        rate_book=rate_book,
        forwarder_name=rate_book.forwarder_name,
        channel_name=channel_name or "",
        customs_fee=customs_fee,
        supports_merge=supports_merge,
        merge_scope=merge_scope,
        item_name_limit=item_limit,
        item_name_extra_fee=item_extra_fee,
        inspection_fee=Decimal(inspection_matches[0]) if inspection_matches else Decimal("0"),
        formal_inspection_fee=Decimal(formal_matches[0]) if formal_matches else Decimal("0"),
        notes=text_value[:2000],
    )


def collect_customs_rules(workbook, rate_book):
    rules = []
    seen = set()
    for worksheet in workbook.worksheets:
        snippets = []
        for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 400), max_col=min(worksheet.max_column, 80), values_only=True):
            row_text = " ".join(clean_freight_cell(value) for value in row if clean_freight_cell(value))
            if "报关" in row_text or "品名" in row_text or "商检" in row_text:
                snippets.append(row_text)
        if not snippets:
            continue
        combined = " ".join(snippets)
        rule = extract_customs_rule_from_text(rate_book, worksheet.title, combined)
        if not rule:
            continue
        key = (
            rule.channel_name,
            str(rule.customs_fee),
            rule.supports_merge,
            rule.item_name_limit,
            str(rule.item_name_extra_fee),
        )
        if key not in seen:
            rules.append(rule)
            seen.add(key)
    global_text = " ".join(rule.notes for rule in rules)
    global_rule = extract_customs_rule_from_text(rate_book, "", global_text)
    if global_rule:
        rules.append(global_rule)
    return rules


def collect_remote_zips(workbook, rate_book):
    rows = []
    seen = set()
    for worksheet in workbook.worksheets:
        title = worksheet.title
        if not any(token in title for token in ("偏远", "Remote", "DAS")):
            continue
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, max_col=min(worksheet.max_column, 20), values_only=True):
            for index, value in enumerate(row):
                zip_code = normalize_zip_code(value)
                if not zip_code:
                    continue
                kind = "super_remote" if ("超" in title or "Ext" in title or index == 1 and "偏远" in title) else "remote"
                key = (kind, zip_code)
                if key in seen:
                    continue
                rows.append(FreightRemoteZip(rate_book=rate_book, forwarder_name=rate_book.forwarder_name, kind=kind, zip_code=zip_code))
                seen.add(key)
                if len(rows) >= 30000:
                    return rows
    return rows


def parse_yinghe_rates(workbook, rate_book):
    rates = []
    summary_sheet = freight_sheet_by_keywords(workbook, "汇总")
    if summary_sheet:
        weight_blocks = [
            (5, 8, Decimal("12"), False),
            (9, 12, Decimal("21"), False),
            (13, 16, Decimal("51"), False),
            (17, 20, Decimal("100"), False),
            (21, 24, Decimal("300"), True),
        ]
        for row_index in range(3, summary_sheet.max_row + 1):
            channel_name = clean_freight_cell(summary_sheet.cell(row_index, 1).value)
            warehouse_code = (clean_freight_cell(summary_sheet.cell(row_index, 2).value) or "").upper()
            warehouse_text = clean_freight_cell(summary_sheet.cell(row_index, 3).value)
            if not channel_name or not warehouse_code:
                continue
            postal_code = extract_zip_code(warehouse_text)
            transit_min, transit_max = parse_transit_days(summary_sheet.cell(row_index, 25).value)
            tax_mode = clean_freight_cell(summary_sheet.cell(row_index, 4).value)
            for start_col, end_col, weight_break, is_cbm in weight_blocks:
                for col_index in range(start_col, end_col + 1):
                    origin_region = clean_freight_cell(summary_sheet.cell(2, col_index).value)
                    price = freight_number(summary_sheet.cell(row_index, col_index).value)
                    rate = build_freight_rate(
                        rate_book,
                        channel_name=channel_name,
                        origin_region=origin_region,
                        warehouse_code=warehouse_code,
                        postal_code=postal_code,
                        weight_break_kg=weight_break,
                        unit_price=None if is_cbm else price,
                        cbm_price=price if is_cbm else None,
                        min_piece_kg=Decimal("12"),
                        divisor=6000,
                        transit_days_min=transit_min,
                        transit_days_max=transit_max,
                        tax_mode=tax_mode,
                        source_sheet=summary_sheet.title,
                    )
                    if rate:
                        rates.append(rate)

    sea_sheet = freight_sheet_by_keywords(workbook, "美国海派")
    if sea_sheet:
        channel_name = ""
        origin_row = None
        for row_index in range(1, sea_sheet.max_row + 1):
            row_values = iter_freight_row_values(sea_sheet, row_index, 20)
            if row_values[1] and not any(freight_number(value) for value in row_values[3:16]):
                channel_name = row_values[1]
            if row_values[4:8] and any(value in {"华东", "华南", "福建", "青岛"} for value in row_values[4:8]):
                origin_row = row_index
                continue
            postal_start, postal_end = zip_range_from_text(row_values[2] if len(row_values) > 2 else "")
            if not channel_name or postal_start is None:
                continue
            transit_min, transit_max = parse_transit_days(row_values[16] if len(row_values) > 16 else "")
            for start_col, end_col, weight_break in [(5, 8, Decimal("12")), (9, 12, Decimal("51")), (13, 16, Decimal("100"))]:
                for col_index in range(start_col, end_col + 1):
                    origin_region = clean_freight_cell(sea_sheet.cell(origin_row or row_index - 1, col_index).value)
                    price = freight_number(sea_sheet.cell(row_index, col_index).value)
                    rate = build_freight_rate(
                        rate_book,
                        channel_name=channel_name,
                        origin_region=origin_region,
                        destination_type="zip_range",
                        zone_name=row_values[1],
                        postal_start=postal_start,
                        postal_end=postal_end,
                        weight_break_kg=weight_break,
                        unit_price=price,
                        min_piece_kg=Decimal("12"),
                        divisor=6000,
                        transit_days_min=transit_min,
                        transit_days_max=transit_max,
                        tax_mode="包税",
                        source_sheet=sea_sheet.title,
                    )
                    if rate:
                        rates.append(rate)
    return rates


def parse_kqd_rates(workbook, rate_book):
    rates = []
    skip_tokens = ("目录", "分区", "偏远", "下单", "地址", "注意", "模板")
    for worksheet in workbook.worksheets:
        if any(token in worksheet.title for token in skip_tokens):
            continue
        if worksheet.max_row < 5:
            continue
        origin_by_col = {}
        current_origin = ""
        for col_index in range(2, min(worksheet.max_column, 12) + 1):
            origin_value = clean_freight_cell(worksheet.cell(3, col_index).value)
            if origin_value:
                current_origin = origin_value
            origin_by_col[col_index] = current_origin
        for row_index in range(5, worksheet.max_row + 1):
            destination_text = clean_freight_cell(worksheet.cell(row_index, 1).value)
            warehouse_codes = extract_warehouse_codes(destination_text)
            if not warehouse_codes:
                continue
            row_values = iter_freight_row_values(worksheet, row_index, min(worksheet.max_column, 14))
            transit_text = " ".join(value for value in row_values if "天" in value)
            transit_min, transit_max = parse_transit_days(transit_text)
            for col_index in range(2, min(worksheet.max_column, 12) + 1):
                weight_label = clean_freight_cell(worksheet.cell(4, col_index).value)
                if "CBM" in weight_label.upper() or "方" in weight_label:
                    continue
                weight_break = freight_number(weight_label)
                price = freight_number(worksheet.cell(row_index, col_index).value)
                if not weight_break or price is None:
                    continue
                for warehouse_code in warehouse_codes:
                    rate = build_freight_rate(
                        rate_book,
                        channel_name=worksheet.title,
                        origin_region=origin_by_col.get(col_index, ""),
                        warehouse_code=warehouse_code,
                        weight_break_kg=weight_break,
                        unit_price=price,
                        min_piece_kg=Decimal("8") if "卡派" in worksheet.title else Decimal("12"),
                        divisor=6000,
                        transit_days_min=transit_min,
                        transit_days_max=transit_max,
                        tax_mode="包税",
                        source_sheet=worksheet.title,
                    )
                    if rate:
                        rates.append(rate)

    sea_sheet = freight_sheet_by_keywords(workbook, "美国海派")
    if sea_sheet:
        channel_name = ""
        for row_index in range(1, sea_sheet.max_row + 1):
            first_cell = clean_freight_cell(sea_sheet.cell(row_index, 1).value)
            if first_cell and "邮编" not in first_cell and not freight_number(sea_sheet.cell(row_index, 2).value):
                channel_name = first_cell
            postal_start, postal_end = zip_range_from_text(first_cell)
            if not channel_name or postal_start is None:
                continue
            transit_min, transit_max = parse_transit_days(" ".join(iter_freight_row_values(sea_sheet, row_index, 8)))
            for col_index, weight_break in [(2, Decimal("12")), (3, Decimal("51")), (4, Decimal("101"))]:
                price = freight_number(sea_sheet.cell(row_index, col_index).value)
                rate = build_freight_rate(
                    rate_book,
                    channel_name=channel_name,
                    origin_region="义乌",
                    destination_type="zip_range",
                    zone_name=first_cell,
                    postal_start=postal_start,
                    postal_end=postal_end,
                    weight_break_kg=weight_break,
                    unit_price=price,
                    min_piece_kg=Decimal("12"),
                    divisor=6000,
                    transit_days_min=transit_min,
                    transit_days_max=transit_max,
                    tax_mode="包税",
                    source_sheet=sea_sheet.title,
                )
                if rate:
                    rates.append(rate)
    return rates


def parse_dingbang_pair_table(workbook, rate_book, worksheet):
    rates = []
    current_channel_yw = ""
    current_channel_sz = ""
    current_transit = ""
    for row_index in range(1, worksheet.max_row + 1):
        row_values = iter_freight_row_values(worksheet, row_index, min(worksheet.max_column, 12))
        row_text = " ".join(row_values)
        if "渠道代码" in row_text:
            current_channel_yw = clean_freight_cell(worksheet.cell(row_index, 4).value) or clean_freight_cell(worksheet.cell(row_index, 3).value)
            current_channel_sz = clean_freight_cell(worksheet.cell(row_index, 6).value) or clean_freight_cell(worksheet.cell(row_index, 5).value)
            current_transit = clean_freight_cell(worksheet.cell(row_index, 8).value) or clean_freight_cell(worksheet.cell(row_index, 10).value)
            continue
        warehouse_text = clean_freight_cell(worksheet.cell(row_index, 3).value) or clean_freight_cell(worksheet.cell(row_index, 2).value)
        warehouse_codes = extract_warehouse_codes(warehouse_text)
        if not warehouse_codes:
            continue
        transit_min, transit_max = parse_transit_days(clean_freight_cell(worksheet.cell(row_index, 8).value) or current_transit)
        for col_index, origin_region, channel_name, weight_break in [
            (4, "义乌", current_channel_yw, Decimal("12")),
            (5, "义乌", current_channel_yw, Decimal("100")),
            (6, "深圳", current_channel_sz, Decimal("12")),
            (7, "深圳", current_channel_sz, Decimal("100")),
        ]:
            price = freight_number(worksheet.cell(row_index, col_index).value)
            if not channel_name or price is None:
                continue
            for warehouse_code in warehouse_codes:
                rate = build_freight_rate(
                    rate_book,
                    channel_name=channel_name,
                    origin_region=origin_region,
                    warehouse_code=warehouse_code,
                    zone_name=clean_freight_cell(worksheet.cell(row_index, 2).value),
                    weight_break_kg=weight_break,
                    unit_price=price,
                    min_piece_kg=Decimal("12"),
                    divisor=6000,
                    transit_days_min=transit_min,
                    transit_days_max=transit_max,
                    tax_mode="包税",
                    source_sheet=worksheet.title,
                )
                if rate:
                    rates.append(rate)
    return rates


def parse_dingbang_rates(workbook, rate_book):
    rates = []
    sea_sheet = workbook.worksheets[2] if len(workbook.worksheets) > 2 else None
    if sea_sheet:
        current_channel_yw = ""
        current_channel_sz = ""
        current_transit = ""
        for row_index in range(1, sea_sheet.max_row + 1):
            row_values = iter_freight_row_values(sea_sheet, row_index, 12)
            row_text = " ".join(row_values)
            if "渠道代码" in row_text:
                current_channel_yw = row_values[0].split("：")[-1].strip() if row_values else ""
                current_channel_sz = row_values[5].split("：")[-1].strip() if len(row_values) > 5 and row_values[5] else ""
                current_transit = row_values[9] if len(row_values) > 9 else ""
                continue
            postal_start, postal_end = zip_range_from_text(row_values[0] if row_values else "")
            if postal_start is None:
                continue
            transit_min, transit_max = parse_transit_days(current_transit)
            for col_index, origin_region, channel_name, weight_break in [
                (2, "义乌", current_channel_yw, Decimal("12")),
                (3, "义乌", current_channel_yw, Decimal("51")),
                (4, "义乌", current_channel_yw, Decimal("100")),
                (6, "深圳", current_channel_sz, Decimal("12")),
                (7, "深圳", current_channel_sz, Decimal("51")),
                (8, "深圳", current_channel_sz, Decimal("100")),
            ]:
                price = freight_number(sea_sheet.cell(row_index, col_index).value)
                if not channel_name or price is None:
                    continue
                rate = build_freight_rate(
                    rate_book,
                    channel_name=channel_name,
                    origin_region=origin_region,
                    destination_type="zip_range",
                    zone_name=row_values[0],
                    postal_start=postal_start,
                    postal_end=postal_end,
                    weight_break_kg=weight_break,
                    unit_price=price,
                    min_piece_kg=Decimal("12"),
                    divisor=6000,
                    transit_days_min=transit_min,
                    transit_days_max=transit_max,
                    tax_mode="包税",
                    source_sheet=sea_sheet.title,
                )
                if rate:
                    rates.append(rate)

    if len(workbook.worksheets) > 3:
        card_sheet = workbook.worksheets[3]
        channel_by_col = {}
        for col_index in range(3, min(card_sheet.max_column, 32) + 1):
            channel = clean_freight_cell(card_sheet.cell(5, col_index).value)
            if channel:
                channel_by_col[col_index] = channel
            else:
                channel_by_col[col_index] = channel_by_col.get(col_index - 1, "")
        for row_index in range(8, min(card_sheet.max_row, 24) + 1):
            warehouse_codes = extract_warehouse_codes(card_sheet.cell(row_index, 2).value)
            if not warehouse_codes:
                continue
            for col_index in range(3, min(card_sheet.max_column, 32) + 1):
                channel_name = channel_by_col.get(col_index, "")
                price = freight_number(card_sheet.cell(row_index, col_index).value)
                weight_break = freight_number(card_sheet.cell(7, col_index).value)
                if not channel_name or price is None or not weight_break:
                    continue
                origin_region = "深圳" if channel_name.upper().startswith("SZ") else "义乌"
                for warehouse_code in warehouse_codes:
                    rate = build_freight_rate(
                        rate_book,
                        channel_name=channel_name,
                        origin_region=origin_region,
                        warehouse_code=warehouse_code,
                        zone_name=clean_freight_cell(card_sheet.cell(row_index, 1).value),
                        weight_break_kg=weight_break,
                        unit_price=price,
                        min_piece_kg=Decimal("12"),
                        divisor=6000,
                        transit_days_min=14,
                        transit_days_max=25,
                        tax_mode="包税",
                        source_sheet=card_sheet.title,
                    )
                    if rate:
                        rates.append(rate)

    for index in (4, 5, 6):
        if len(workbook.worksheets) > index:
            rates.extend(parse_dingbang_pair_table(workbook, rate_book, workbook.worksheets[index]))
    return rates


def hongfan_merged_cell_value(worksheet, row_index, col_index):
    value = clean_freight_cell(worksheet.cell(row_index, col_index).value)
    if value:
        return value
    for merged_range in worksheet.merged_cells.ranges:
        if (
            merged_range.min_row <= row_index <= merged_range.max_row
            and merged_range.min_col <= col_index <= merged_range.max_col
        ):
            return clean_freight_cell(worksheet.cell(merged_range.min_row, merged_range.min_col).value)
    return ""


def hongfan_effective_max_col(worksheet):
    return min(max(worksheet.max_column, 20), 80)


def hongfan_channel_name(worksheet, header_row, warehouse_col, price_col):
    candidates = [
        hongfan_merged_cell_value(worksheet, header_row - 1, price_col) if header_row > 1 else "",
        hongfan_merged_cell_value(worksheet, header_row - 1, warehouse_col) if header_row > 1 else "",
        worksheet.title,
    ]
    for candidate in candidates:
        text_value = clean_freight_cell(candidate)
        if not text_value or text_value == "#REF!":
            continue
        for marker in (" 合并报关", "合并报关仓点", "合并报关"):
            marker_index = text_value.find(marker)
            if marker_index > 0:
                text_value = text_value[:marker_index].strip()
                break
        return text_value[:128]
    return worksheet.title[:128]


def hongfan_sheet_has_small_weight_surcharge(worksheet):
    max_col = min(hongfan_effective_max_col(worksheet), 20)
    snippets = []
    for row_index in range(1, min(worksheet.max_row, 6) + 1):
        snippets.extend(hongfan_merged_cell_value(worksheet, row_index, col_index) for col_index in range(1, max_col + 1))
    text_value = " ".join(snippets).upper().replace(" ", "")
    return bool(re.search(r"12\s*[-~—至]\s*50\s*KG\s*\+?1", text_value) or ("12-50KG" in text_value and "+1" in text_value))


def hongfan_extract_warehouse_codes(value):
    skipped = {"WALMART", "TIKTOK", "WAYFAIR", "ORANGE", "AMAZON", "SERVICES", "INC"}
    codes = []
    for code in extract_warehouse_codes(value):
        normalized = clean_freight_cell(code).upper()
        if normalized in skipped:
            continue
        if normalized not in codes:
            codes.append(normalized)
    return codes


def collect_hongfan_warehouse_postal_codes(workbook):
    mapping = {}
    for worksheet in workbook.worksheets:
        max_col = hongfan_effective_max_col(worksheet)
        header_row = None
        warehouse_col = None
        postal_col = None
        city_col = None
        state_col = None
        for row_index in range(1, min(worksheet.max_row, 12) + 1):
            row_values = [hongfan_merged_cell_value(worksheet, row_index, col_index) for col_index in range(1, max_col + 1)]
            row_text = " ".join(row_values).lower()
            if "仓库" not in row_text or ("邮编" not in row_text and "zip" not in row_text and "postal" not in row_text):
                continue
            header_row = row_index
            for col_index, header in enumerate(row_values, start=1):
                normalized = header.lower()
                if warehouse_col is None and ("仓库" in normalized or "warehouse" in normalized or "code" in normalized):
                    warehouse_col = col_index
                if postal_col is None and ("邮编" in normalized or "zip" in normalized or "postal" in normalized):
                    postal_col = col_index
                if city_col is None and ("城市" in normalized or "city" in normalized):
                    city_col = col_index
                if state_col is None and ("州" in normalized or "省" in normalized or "state" in normalized):
                    state_col = col_index
            break
        if header_row is None or warehouse_col is None or postal_col is None:
            continue
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            warehouse_text = hongfan_merged_cell_value(worksheet, row_index, warehouse_col)
            postal_code = normalize_zip_code(hongfan_merged_cell_value(worksheet, row_index, postal_col))
            if not warehouse_text or not postal_code:
                continue
            for warehouse_code in hongfan_extract_warehouse_codes(warehouse_text):
                mapping[warehouse_code] = {
                    "postal_code": postal_code,
                    "city": hongfan_merged_cell_value(worksheet, row_index, city_col) if city_col else "",
                    "state": hongfan_merged_cell_value(worksheet, row_index, state_col) if state_col else "",
                }
    return mapping


def hongfan_destinations_from_text(value, postal_by_code):
    text_value = clean_freight_cell(value)
    warehouse_codes = hongfan_extract_warehouse_codes(text_value)
    postal_code = normalize_zip_code(text_value)
    destinations = []
    for warehouse_code in warehouse_codes:
        mapped_postal = postal_by_code.get(warehouse_code, {}).get("postal_code", "")
        destinations.append(
            {
                "destination_type": "warehouse",
                "warehouse_code": warehouse_code,
                "postal_code": mapped_postal or (postal_code if len(warehouse_codes) == 1 else ""),
                "postal_start": None,
                "postal_end": None,
            }
        )
    if not destinations and postal_code:
        zip_number = parse_int(postal_code)
        destinations.append(
            {
                "destination_type": "zip_range",
                "warehouse_code": "",
                "postal_code": postal_code,
                "postal_start": zip_number,
                "postal_end": zip_number,
            }
        )
    return destinations


def parse_hongfan_zone_rates(workbook, rate_book):
    rates = []
    worksheet = freight_sheet_by_keywords(workbook, "美国海派")
    if not worksheet:
        return rates
    price_columns = [
        (4, "义乌", Decimal("12")),
        (5, "义乌", Decimal("51")),
        (6, "义乌", Decimal("100")),
        (7, "泉州", Decimal("12")),
        (8, "泉州", Decimal("51")),
        (9, "泉州", Decimal("100")),
        (10, "深圳", Decimal("12")),
        (11, "深圳", Decimal("51")),
        (12, "深圳", Decimal("100")),
    ]
    current_channel = ""
    for row_index in range(6, worksheet.max_row + 1):
        channel_cell = clean_freight_cell(worksheet.cell(row_index, 2).value)
        if channel_cell:
            current_channel = channel_cell
        zone_text = clean_freight_cell(worksheet.cell(row_index, 3).value)
        zip_ranges = zip_ranges_from_text(zone_text)
        if not current_channel or not zone_text or not zip_ranges:
            continue
        transit_min, transit_max = parse_transit_days(worksheet.cell(row_index, 18).value)
        notes = " ".join(
            clean_freight_cell(worksheet.cell(row_index, col_index).value)
            for col_index in (13, 17)
            if clean_freight_cell(worksheet.cell(row_index, col_index).value)
        )
        for col_index, origin_region, weight_break in price_columns:
            price = freight_number(worksheet.cell(row_index, col_index).value)
            if price is None:
                continue
            for postal_start, postal_end in zip_ranges:
                rate = build_freight_rate(
                    rate_book,
                    channel_name=current_channel,
                    origin_region=origin_region,
                    destination_type="zip_range",
                    zone_name=zone_text,
                    postal_start=postal_start,
                    postal_end=postal_end,
                    weight_break_kg=weight_break,
                    unit_price=price,
                    min_piece_kg=Decimal("12"),
                    divisor=6000,
                    transit_days_min=transit_min,
                    transit_days_max=transit_max,
                    tax_mode="包税",
                    source_sheet=worksheet.title,
                    notes=notes,
                )
                if rate:
                    rates.append(rate)
    return rates


def parse_hongfan_warehouse_rates(workbook, rate_book, postal_by_code):
    rates = []
    for worksheet in workbook.worksheets:
        if "仓库代码及地址" in worksheet.title:
            continue
        max_col = hongfan_effective_max_col(worksheet)
        parsed_sheet = False
        for header_row in range(1, min(worksheet.max_row, 10) + 1):
            warehouse_cols = []
            for col_index in range(1, max_col + 1):
                header = hongfan_merged_cell_value(worksheet, header_row, col_index).replace(" ", "")
                if any(keyword in header for keyword in ("仓库代码", "仓点", "FBA仓点")) or (
                    "美国分区邮编" in header and "海派" not in worksheet.title
                ):
                    warehouse_cols.append(col_index)
            if not warehouse_cols:
                continue
            next_warehouse_cols = warehouse_cols[1:] + [max_col + 1]
            for warehouse_col, next_warehouse_col in zip(warehouse_cols, next_warehouse_cols):
                segment_end = next_warehouse_col - 1
                transit_cols = []
                specs = []
                for col_index in range(warehouse_col + 1, segment_end + 1):
                    header_text = hongfan_merged_cell_value(worksheet, header_row, col_index)
                    weight_text = hongfan_merged_cell_value(worksheet, header_row + 1, col_index)
                    combined_header = f"{header_text} {weight_text}"
                    if "时效" in combined_header or "签收" in combined_header:
                        transit_cols.append(col_index)
                        continue
                    weight_break = freight_number(weight_text)
                    if not weight_break or "CBM" in weight_text.upper() or "方" in weight_text:
                        continue
                    if not any(token in header_text for token in ("义乌", "深圳", "福建", "泉州", "华东", "华南")):
                        continue
                    specs.append(
                        {
                            "col": col_index,
                            "origin_region": header_text,
                            "weight_break": weight_break,
                            "channel_name": hongfan_channel_name(worksheet, header_row, warehouse_col, col_index),
                            "tax_mode": "包税" if "含税" in header_text else "",
                        }
                    )
                if not specs:
                    continue
                parsed_sheet = True
                add_small_weight_surcharge = (
                    hongfan_sheet_has_small_weight_surcharge(worksheet)
                    and not any(spec["weight_break"] == Decimal("12") for spec in specs)
                )
                for spec in specs:
                    spec["transit_col"] = next((col_index for col_index in transit_cols if col_index > spec["col"]), None)
                for row_index in range(header_row + 2, worksheet.max_row + 1):
                    warehouse_text = hongfan_merged_cell_value(worksheet, row_index, warehouse_col)
                    if not warehouse_text:
                        continue
                    destinations = hongfan_destinations_from_text(warehouse_text, postal_by_code)
                    if not destinations:
                        continue
                    row_values = iter_freight_row_values(worksheet, row_index, min(segment_end, max_col))
                    row_transit_text = " ".join(value for value in row_values if "天" in value)
                    for spec in specs:
                        price = freight_number(worksheet.cell(row_index, spec["col"]).value)
                        if price is None:
                            continue
                        price_breaks = [(spec["weight_break"], price)]
                        if add_small_weight_surcharge and spec["weight_break"] == Decimal("51"):
                            price_breaks.append((Decimal("12"), price + Decimal("1")))
                        transit_text = (
                            hongfan_merged_cell_value(worksheet, row_index, spec["transit_col"])
                            if spec.get("transit_col")
                            else row_transit_text
                        )
                        transit_min, transit_max = parse_transit_days(transit_text)
                        for weight_break, unit_price in price_breaks:
                            for destination in destinations:
                                rate = build_freight_rate(
                                    rate_book,
                                    channel_name=spec["channel_name"],
                                    origin_region=spec["origin_region"],
                                    destination_type=destination["destination_type"],
                                    warehouse_code=destination["warehouse_code"],
                                    postal_code=destination["postal_code"],
                                    postal_start=destination["postal_start"],
                                    postal_end=destination["postal_end"],
                                    weight_break_kg=weight_break,
                                    unit_price=unit_price,
                                    min_piece_kg=Decimal("12"),
                                    divisor=6000,
                                    transit_days_min=transit_min,
                                    transit_days_max=transit_max,
                                    tax_mode=spec["tax_mode"],
                                    source_sheet=worksheet.title,
                                )
                                if rate:
                                    rates.append(rate)
            if parsed_sheet:
                break
    return rates


def dedupe_freight_rates(rates):
    deduped = []
    seen = set()
    for rate in rates:
        key = (
            rate.forwarder_name,
            rate.channel_name,
            rate.origin_region,
            rate.destination_type,
            rate.warehouse_code,
            rate.postal_code,
            rate.postal_start,
            rate.postal_end,
            str(rate.weight_break_kg),
            str(rate.unit_price),
            str(rate.cbm_price),
            rate.source_sheet,
        )
        if key in seen:
            continue
        deduped.append(rate)
        seen.add(key)
    return deduped


def parse_hongfan_rates(workbook, rate_book):
    postal_by_code = collect_hongfan_warehouse_postal_codes(workbook)
    rates = []
    rates.extend(parse_hongfan_zone_rates(workbook, rate_book))
    rates.extend(parse_hongfan_warehouse_rates(workbook, rate_book, postal_by_code))
    return dedupe_freight_rates(rates)


def parse_freight_rate_workbook(file_path, rate_book):
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, data_only=True, read_only=False)
    if rate_book.forwarder_name == "盈和":
        rates = parse_yinghe_rates(workbook, rate_book)
    elif rate_book.forwarder_name == "开渠达":
        rates = parse_kqd_rates(workbook, rate_book)
    elif rate_book.forwarder_name == "鼎邦":
        rates = parse_dingbang_rates(workbook, rate_book)
    elif rate_book.forwarder_name == "鸿帆":
        rates = parse_hongfan_rates(workbook, rate_book)
    else:
        rates = []
    customs_rules = collect_customs_rules(workbook, rate_book)
    remote_zips = collect_remote_zips(workbook, rate_book)
    return {
        "rates": rates,
        "customs_rules": customs_rules,
        "remote_zips": remote_zips,
        "summary": {
            "sheet_count": len(workbook.worksheets),
            "rate_count": len(rates),
            "customs_rule_count": len(customs_rules),
            "remote_zip_count": len(remote_zips),
        },
    }


def decimal_to_float(value):
    if value is None:
        return 0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0


def freight_json_dumps(payload):
    return json.dumps(payload, ensure_ascii=False, default=decimal_to_float)


def load_freight_rate_book_summary(rate_book):
    try:
        return json.loads(rate_book.parsed_summary or "{}")
    except json.JSONDecodeError:
        return {}


def warehouse_postal_mapping():
    mapping = {}
    rate_rows = (
        db.session.query(FreightRate.warehouse_code, FreightRate.postal_code)
        .filter(FreightRate.warehouse_code != "", FreightRate.postal_code != "")
        .distinct()
        .all()
    )
    for warehouse_code, postal_code in rate_rows:
        code = clean_freight_cell(warehouse_code).upper()
        zip_code = normalize_zip_code(postal_code)
        if code and zip_code and code not in mapping:
            mapping[code] = zip_code
    rows = WarehousePostalCode.query.order_by(WarehousePostalCode.warehouse_code.asc()).all()
    for row in rows:
        code = clean_freight_cell(row.warehouse_code).upper()
        zip_code = normalize_zip_code(row.postal_code)
        if code and zip_code:
            mapping[code] = zip_code
    return mapping


def resolve_warehouse_postal_code(warehouse_code, mapping=None):
    code = clean_freight_cell(warehouse_code).upper()
    if not code:
        return ""
    mapping = mapping if mapping is not None else warehouse_postal_mapping()
    return normalize_zip_code(mapping.get(code, ""))


def upsert_warehouse_postal_code(warehouse_code, postal_code, *, city="", state="", source="", note=""):
    code = clean_freight_cell(warehouse_code).upper()
    zip_code = normalize_zip_code(postal_code)
    if not code or not zip_code:
        return False
    row = WarehousePostalCode.query.filter_by(warehouse_code=code).first()
    if not row:
        row = WarehousePostalCode(warehouse_code=code)
        db.session.add(row)
    row.postal_code = zip_code
    if city:
        row.city = city[:128]
    if state:
        row.state = state[:64]
    if source:
        row.source = source[:128]
    if note:
        row.note = note[:255]
    return True


def learn_warehouse_postal_codes_from_rates(rates, source):
    learned_count = 0
    seen = set()
    for rate in rates:
        warehouse_code = clean_freight_cell(getattr(rate, "warehouse_code", "")).upper()
        if not warehouse_code or warehouse_code in seen:
            continue
        postal_code = normalize_zip_code(getattr(rate, "postal_code", ""))
        if not postal_code and getattr(rate, "postal_start", None) is not None:
            start = parse_int(rate.postal_start)
            end = parse_int(rate.postal_end)
            if 0 < start <= 99999 and (not end or start == end or str(start).zfill(5)[:3] == str(end).zfill(5)[:3]):
                postal_code = str(start).zfill(5)
        if upsert_warehouse_postal_code(warehouse_code, postal_code, source=source, note="从货代价格表自动学习"):
            learned_count += 1
            seen.add(warehouse_code)
    return learned_count


def parse_warehouse_postal_upload(file_storage):
    rows = parse_uploaded_table(file_storage)
    header_index = None
    header_cells = []
    for index, row in enumerate(rows[:20]):
        cells = [clean_freight_cell(value) for value in row]
        row_text = " ".join(cells).lower()
        if ("仓" in row_text or "warehouse" in row_text or "fba" in row_text) and ("邮编" in row_text or "zip" in row_text or "postal" in row_text):
            header_index = index
            header_cells = cells
            break
    if header_index is None:
        raise ValueError("仓库邮编表未识别到表头，请至少包含仓库代码和邮编列。")

    def find_col(*keywords):
        for keyword in keywords:
            for col_index, header in enumerate(header_cells):
                if keyword.lower() in header.lower():
                    return col_index
        return None

    warehouse_col = find_col("仓库", "fba", "warehouse", "code")
    postal_col = find_col("邮编", "zip", "postal")
    city_col = find_col("城市", "city")
    state_col = find_col("州", "state")
    source_col = find_col("来源", "source")
    note_col = find_col("备注", "note")
    if warehouse_col is None or postal_col is None:
        raise ValueError("仓库邮编表缺少仓库代码或邮编列。")

    imported = []
    for row in rows[header_index + 1 :]:
        cells = [clean_freight_cell(value) for value in row]

        def value_at(col_index):
            if col_index is None or col_index >= len(cells):
                return ""
            return cells[col_index]

        warehouse_codes = extract_warehouse_codes(value_at(warehouse_col))
        warehouse_code = warehouse_codes[0] if warehouse_codes else clean_freight_cell(value_at(warehouse_col)).upper()
        postal_code = normalize_zip_code(value_at(postal_col))
        if not warehouse_code and not postal_code:
            continue
        if not warehouse_code or not postal_code:
            continue
        imported.append(
            {
                "warehouse_code": warehouse_code,
                "postal_code": postal_code,
                "city": value_at(city_col),
                "state": value_at(state_col),
                "source": value_at(source_col),
                "note": value_at(note_col),
            }
        )
    if not imported:
        raise ValueError("仓库邮编表没有可导入的数据。")
    return imported


def parse_freight_quote_upload(file_storage):
    if not file_storage or not (file_storage.filename or "").strip():
        return []
    rows = parse_uploaded_table(file_storage)
    postal_mapping = warehouse_postal_mapping()
    header_index = None
    header_cells = []
    for index, row in enumerate(rows[:20]):
        cells = [clean_freight_cell(value) for value in row]
        row_text = " ".join(cells)
        if ("仓" in row_text or "FBA" in row_text.upper()) and ("重量" in row_text or "KG" in row_text.upper()):
            header_index = index
            header_cells = cells
            break
    if header_index is None:
        raise ValueError("导入货件表未识别到表头，请至少包含仓库代码和重量列。")

    def find_col(*keywords):
        for keyword in keywords:
            for col_index, header in enumerate(header_cells):
                if keyword.lower() in header.lower():
                    return col_index
        return None

    warehouse_col = find_col("仓库", "FBA", "warehouse")
    zip_col = find_col("邮编", "zip")
    weight_col = find_col("重量", "实重", "kg")
    box_col = find_col("箱数", "件数", "boxes")
    length_col = find_col("长", "length")
    width_col = find_col("宽", "width")
    height_col = find_col("高", "height")
    if warehouse_col is None and zip_col is None:
        raise ValueError("导入货件表缺少仓库代码或邮编列。")
    if weight_col is None:
        raise ValueError("导入货件表缺少重量列。")

    items = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        cells = [clean_freight_cell(value) for value in row]

        def value_at(col_index):
            if col_index is None or col_index >= len(cells):
                return ""
            return cells[col_index]

        warehouse_codes = extract_warehouse_codes(value_at(warehouse_col))
        warehouse_code = warehouse_codes[0] if warehouse_codes else clean_freight_cell(value_at(warehouse_col)).upper()
        postal_code = normalize_zip_code(value_at(zip_col))
        if not postal_code:
            postal_code = resolve_warehouse_postal_code(warehouse_code, postal_mapping)
        weight = freight_number(value_at(weight_col), Decimal("0"))
        if not warehouse_code and not postal_code:
            continue
        if weight <= 0:
            raise ValueError(f"导入货件表第 {row_number} 行重量必须大于 0。")
        items.append(
            {
                "warehouse_code": warehouse_code,
                "postal_code": postal_code,
                "actual_weight_kg": weight,
                "box_count": max(parse_int(value_at(box_col), 1), 1),
                "length_cm": freight_number(value_at(length_col), Decimal("0")) or Decimal("0"),
                "width_cm": freight_number(value_at(width_col), Decimal("0")) or Decimal("0"),
                "height_cm": freight_number(value_at(height_col), Decimal("0")) or Decimal("0"),
            }
        )
    return items


def parse_freight_quote_form_items():
    warehouse_values = request.form.getlist("quote_warehouse_code")
    postal_values = request.form.getlist("quote_postal_code")
    weight_values = request.form.getlist("quote_weight_kg")
    box_values = request.form.getlist("quote_box_count")
    length_values = request.form.getlist("quote_length_cm")
    width_values = request.form.getlist("quote_width_cm")
    height_values = request.form.getlist("quote_height_cm")
    row_count = max(len(warehouse_values), len(postal_values), len(weight_values), 6)
    postal_mapping = warehouse_postal_mapping()
    items = []
    errors = []
    for index in range(row_count):
        warehouse_raw = warehouse_values[index] if index < len(warehouse_values) else ""
        postal_raw = postal_values[index] if index < len(postal_values) else ""
        weight_raw = weight_values[index] if index < len(weight_values) else ""
        warehouse_codes = extract_warehouse_codes(warehouse_raw)
        warehouse_code = warehouse_codes[0] if warehouse_codes else clean_freight_cell(warehouse_raw).upper()
        postal_code = normalize_zip_code(postal_raw)
        if not postal_code:
            postal_code = resolve_warehouse_postal_code(warehouse_code, postal_mapping)
        weight = freight_number(weight_raw, Decimal("0")) or Decimal("0")
        if not warehouse_code and not postal_code and weight <= 0:
            continue
        if not warehouse_code and not postal_code:
            errors.append(f"第 {index + 1} 行请填写 FBA 仓库代码或邮编。")
            continue
        if weight <= 0:
            errors.append(f"第 {index + 1} 行重量必须大于 0。")
            continue
        items.append(
            {
                "warehouse_code": warehouse_code,
                "postal_code": postal_code,
                "actual_weight_kg": weight,
                "box_count": max(parse_int(box_values[index] if index < len(box_values) else "1", 1), 1),
                "length_cm": freight_number(length_values[index] if index < len(length_values) else "", Decimal("0")) or Decimal("0"),
                "width_cm": freight_number(width_values[index] if index < len(width_values) else "", Decimal("0")) or Decimal("0"),
                "height_cm": freight_number(height_values[index] if index < len(height_values) else "", Decimal("0")) or Decimal("0"),
            }
        )
    if errors:
        raise ValueError("\n".join(errors[:8]))
    return items


def freight_item_volume_weight(item, divisor=6000):
    length = Decimal(item.get("length_cm") or 0)
    width = Decimal(item.get("width_cm") or 0)
    height = Decimal(item.get("height_cm") or 0)
    boxes = Decimal(item.get("box_count") or 1)
    if length <= 0 or width <= 0 or height <= 0:
        return Decimal("0")
    return (length * width * height * boxes / Decimal(divisor or 6000)).quantize(Decimal("0.01"))


def freight_base_weight_for_item(item, rate):
    actual_weight = Decimal(item.get("actual_weight_kg") or 0)
    volume_weight = freight_item_volume_weight(item, rate.divisor or 6000)
    min_piece_total = Decimal(rate.min_piece_kg or 0) * Decimal(item.get("box_count") or 1)
    return max(actual_weight, volume_weight, min_piece_total)


def freight_rate_destination_matches(rate, item):
    warehouse_code = (item.get("warehouse_code") or "").upper().strip()
    postal_code = normalize_zip_code(item.get("postal_code") or "")
    if rate.destination_type == "warehouse":
        return bool(warehouse_code and rate.warehouse_code and rate.warehouse_code.upper() == warehouse_code)
    if rate.destination_type == "zip_range":
        if not postal_code or rate.postal_start is None or rate.postal_end is None:
            return False
        zip_number = parse_int(postal_code)
        return rate.postal_start <= zip_number <= rate.postal_end
    return False


def freight_quote_key(rate):
    return (rate.rate_book_id, rate.forwarder_name, rate.channel_name, normalize_origin_region(rate.origin_region))


def build_freight_candidates_for_item(item, rates, selected_origin):
    grouped = defaultdict(list)
    for rate in rates:
        if rate.unit_price is None:
            continue
        if not freight_origin_matches(rate.origin_region, selected_origin):
            continue
        if not freight_rate_destination_matches(rate, item):
            continue
        grouped[freight_quote_key(rate)].append(rate)

    selected = {}
    for key, option_rates in grouped.items():
        option_rates.sort(key=lambda row: Decimal(row.weight_break_kg or 0))
        base_weight = max(freight_base_weight_for_item(item, row) for row in option_rates)
        applicable = [row for row in option_rates if base_weight >= Decimal(row.weight_break_kg or 0)]
        chosen = applicable[-1] if applicable else option_rates[0]
        chargeable_weight = max(base_weight, Decimal(chosen.min_chargeable_kg or 0), Decimal(chosen.weight_break_kg or 0))
        freight_cost = (chargeable_weight * Decimal(chosen.unit_price or 0)).quantize(Decimal("0.01"))
        selected[key] = {
            "rate": chosen,
            "chargeable_weight": chargeable_weight,
            "freight_cost": freight_cost,
            "base_weight": base_weight,
            "volume_weight": freight_item_volume_weight(item, chosen.divisor or 6000),
        }
    return selected


def select_freight_customs_rule(rate_book_id, channel_name, rules_by_book, source_sheets=None):
    source_sheets = set(source_sheets or [])
    rules = rules_by_book.get(rate_book_id, [])
    exact = [
        rule
        for rule in rules
        if rule.channel_name
        and (
            rule.channel_name == channel_name
            or rule.channel_name in channel_name
            or channel_name in rule.channel_name
        )
    ]
    if exact:
        return max(exact, key=lambda rule: Decimal(rule.customs_fee or 0))
    sheet_rules = [rule for rule in rules if rule.channel_name and rule.channel_name in source_sheets]
    if sheet_rules:
        return max(sheet_rules, key=lambda rule: Decimal(rule.customs_fee or 0))
    global_rules = [rule for rule in rules if not rule.channel_name]
    if global_rules:
        return max(global_rules, key=lambda rule: Decimal(rule.customs_fee or 0))
    return max(rules, key=lambda rule: Decimal(rule.customs_fee or 0), default=None)


def freight_remote_kind(rate_book_id, postal_code, remote_zips_by_book):
    postal_code = normalize_zip_code(postal_code)
    if not postal_code:
        return ""
    rows = remote_zips_by_book.get(rate_book_id, {})
    return rows.get(postal_code, "")


def calculate_freight_quote(*, items, origin_region, customs_required=False, inspection_required=False, item_name_count=0, clothing_surcharge_enabled=True):
    active_books = FreightRateBook.query.filter_by(status="active").order_by(FreightRateBook.created_at.desc()).all()
    active_books_by_id = {book.id: book for book in active_books}
    active_book_ids = [book.id for book in active_books]
    if not active_book_ids:
        raise ValueError("还没有启用的货代价格表，请先上传价格表。")
    rates = FreightRate.query.filter(FreightRate.rate_book_id.in_(active_book_ids)).all()
    if not rates:
        raise ValueError("启用的价格表里没有可用价格行，请重新上传或检查解析结果。")
    rules_by_book = defaultdict(list)
    for rule in FreightCustomsRule.query.filter(FreightCustomsRule.rate_book_id.in_(active_book_ids)).all():
        rules_by_book[rule.rate_book_id].append(rule)
    remote_zips_by_book = defaultdict(dict)
    for remote in FreightRemoteZip.query.filter(FreightRemoteZip.rate_book_id.in_(active_book_ids)).all():
        remote_zips_by_book[remote.rate_book_id][remote.zip_code] = remote.kind

    option_accumulators = {}
    missing_items = []
    for item_index, item in enumerate(items, start=1):
        candidates = build_freight_candidates_for_item(item, rates, origin_region)
        if not candidates:
            missing_items.append(item_index)
            continue
        for key, payload in candidates.items():
            if key not in option_accumulators:
                option_accumulators[key] = {
                    "rate_book_id": key[0],
                    "forwarder_name": key[1],
                    "channel_name": key[2],
                    "origin_region": key[3],
                    "items": [],
                    "freight_cost": Decimal("0"),
                    "surcharge_cost": Decimal("0"),
                    "chargeable_weight": Decimal("0"),
                    "transit_min": None,
                    "transit_max": None,
                    "notes": [],
                    "zones": set(),
                    "source_sheets": set(),
                }
            option = option_accumulators[key]
            rate = payload["rate"]
            postal_code = item.get("postal_code") or rate.postal_code
            remote_kind = freight_remote_kind(rate.rate_book_id, postal_code, remote_zips_by_book)
            surcharge_cost = Decimal("0")
            if remote_kind:
                option["notes"].append(f"{item.get('warehouse_code') or postal_code} 命中{'超偏远' if remote_kind == 'super_remote' else '偏远'}邮编，附加费需复核")
                if rate.forwarder_name == "开渠达":
                    surcharge_cost = (payload["chargeable_weight"] * Decimal("2")).quantize(Decimal("0.01"))
            rate_book = active_books_by_id.get(rate.rate_book_id)
            clothing_surcharge_per_kg = Decimal(rate_book.clothing_surcharge_per_kg or 0) if rate_book else Decimal("0")
            if clothing_surcharge_enabled and clothing_surcharge_per_kg > 0:
                clothing_surcharge_cost = (payload["chargeable_weight"] * clothing_surcharge_per_kg).quantize(Decimal("0.01"))
                surcharge_cost += clothing_surcharge_cost
                option["notes"].append(f"服装附加费 {clothing_surcharge_per_kg}/kg")
            option["freight_cost"] += payload["freight_cost"]
            option["surcharge_cost"] += surcharge_cost
            option["chargeable_weight"] += payload["chargeable_weight"]
            option["zones"].add(rate.zone_name or rate.warehouse_code or postal_code or f"item-{item_index}")
            if rate.transit_days_min:
                option["transit_min"] = rate.transit_days_min if option["transit_min"] is None else min(option["transit_min"], rate.transit_days_min)
            if rate.transit_days_max:
                option["transit_max"] = rate.transit_days_max if option["transit_max"] is None else max(option["transit_max"], rate.transit_days_max)
            if rate.source_sheet:
                option["source_sheets"].add(rate.source_sheet)
            option["items"].append(
                {
                    "index": item_index,
                    "warehouse_code": item.get("warehouse_code", ""),
                    "postal_code": postal_code,
                    "unit_price": decimal_to_float(rate.unit_price),
                    "weight_break_kg": decimal_to_float(rate.weight_break_kg),
                    "actual_weight_kg": decimal_to_float(item.get("actual_weight_kg")),
                    "volume_weight_kg": decimal_to_float(payload["volume_weight"]),
                    "chargeable_weight_kg": decimal_to_float(payload["chargeable_weight"]),
                    "freight_cost": decimal_to_float(payload["freight_cost"]),
                    "surcharge_cost": decimal_to_float(surcharge_cost),
                    "source_sheet": rate.source_sheet,
                }
            )

    complete_options = []
    for option in option_accumulators.values():
        if len(option["items"]) != len(items):
            continue
        customs_cost = Decimal("0")
        rule = select_freight_customs_rule(option["rate_book_id"], option["channel_name"], rules_by_book, option.get("source_sheets"))
        if customs_required:
            if not rule:
                option["notes"].append("未识别到报关规则，报关费未计入")
            else:
                declaration_count = 1
                if not rule.supports_merge:
                    declaration_count = max(len(option["zones"]) if rule.merge_scope == "zone" else len(items), 1)
                customs_cost += Decimal(rule.customs_fee or 0) * declaration_count
                if item_name_count and rule.item_name_limit and item_name_count > rule.item_name_limit:
                    customs_cost += Decimal(item_name_count - rule.item_name_limit) * Decimal(rule.item_name_extra_fee or 0) * declaration_count
                if inspection_required:
                    customs_cost += Decimal(rule.inspection_fee or rule.formal_inspection_fee or 0) * declaration_count
                if not rule.supports_merge:
                    option["notes"].append("该渠道不支持完全合并报关，已按分票规则计费")
        elif rule and not rule.supports_merge:
            option["notes"].append("如需报关，该渠道存在分票/分区限制")
        total_cost = (option["freight_cost"] + option["surcharge_cost"] + customs_cost).quantize(Decimal("0.01"))
        complete_options.append(
            {
                "rate_book_id": option["rate_book_id"],
                "forwarder_name": option["forwarder_name"],
                "channel_name": option["channel_name"],
                "origin_region": option["origin_region"],
                "freight_cost": option["freight_cost"].quantize(Decimal("0.01")),
                "customs_cost": customs_cost.quantize(Decimal("0.01")),
                "surcharge_cost": option["surcharge_cost"].quantize(Decimal("0.01")),
                "total_cost": total_cost,
                "chargeable_weight": option["chargeable_weight"].quantize(Decimal("0.01")),
                "unit_cost": (total_cost / option["chargeable_weight"]).quantize(Decimal("0.01")) if option["chargeable_weight"] else Decimal("0"),
                "transit_days_min": option["transit_min"],
                "transit_days_max": option["transit_max"],
                "notes": "；".join(dict.fromkeys(option["notes"]))[:2000],
                "items": option["items"],
            }
        )

    complete_options.sort(key=lambda row: (row["total_cost"], row["transit_days_max"] or 999))
    fastest_days = min((row["transit_days_max"] or 999 for row in complete_options), default=None)
    for index, option in enumerate(complete_options, start=1):
        labels = []
        if index == 1:
            labels.append("最低价")
        if fastest_days is not None and (option["transit_days_max"] or 999) == fastest_days:
            labels.append("最快")
        if not labels and index <= 3:
            labels.append("备选")
        option["rank_no"] = index
        option["recommendation"] = "/".join(labels)

    return {
        "items": items,
        "clothing_surcharge_enabled": clothing_surcharge_enabled,
        "options": complete_options,
        "missing_items": missing_items,
        "summary": {
            "item_count": len(items),
            "option_count": len(complete_options),
            "best_total_cost": decimal_to_float(complete_options[0]["total_cost"]) if complete_options else 0,
            "missing_item_count": len(missing_items),
        },
    }


def save_freight_quote_batch(result, *, origin_region, customs_required, inspection_required, item_name_count):
    batch = FreightQuoteBatch(
        quote_no=build_unique_code("FQ", FreightQuoteBatch, "quote_no"),
        origin_region=origin_region,
        customs_required=customs_required,
        item_name_count=item_name_count,
        input_json=freight_json_dumps(
            {
                "items": result.get("items", []),
                "inspection_required": inspection_required,
                "clothing_surcharge_enabled": result.get("clothing_surcharge_enabled", True),
            }
        ),
        summary_json=freight_json_dumps(result.get("summary", {})),
        operator_user_id=g.user.id if g.get("user") else None,
        operator_name=get_current_operator_name(),
    )
    db.session.add(batch)
    db.session.flush()
    for option in result.get("options", []):
        batch.options.append(
            FreightQuoteOption(
                forwarder_name=option["forwarder_name"],
                channel_name=option["channel_name"],
                origin_region=option["origin_region"],
                total_cost=option["total_cost"],
                freight_cost=option["freight_cost"],
                customs_cost=option["customs_cost"],
                surcharge_cost=option.get("surcharge_cost", 0),
                chargeable_weight=option["chargeable_weight"],
                transit_days_min=option.get("transit_days_min"),
                transit_days_max=option.get("transit_days_max"),
                rank_no=option.get("rank_no", 0),
                recommendation=option.get("recommendation", ""),
                notes=option.get("notes", ""),
                detail_json=freight_json_dumps(option.get("items", [])),
            )
        )
    log_operation("货代比价", "新增", "freight_quote", batch.quote_no, f"报价选项：{len(result.get('options', []))}")
    db.session.commit()
    return batch


def get_sheet_b1_value(rows):
    if not rows:
        return ""
    first_row = rows[0] if rows else []
    if len(first_row) < 2:
        return ""
    return (first_row[1] or "").strip()


def is_fba_shipment_table(rows):
    b1_value = get_sheet_b1_value(rows)
    normalized = re.sub(r"\s+", " ", b1_value or "").strip().upper()
    return bool(re.search(r"(?<![A-Z0-9])FBA[A-Z0-9_-]*", normalized))


def resolve_import_sku_codes(raw_sku_codes, *, user=None):
    normalized_codes = [code.strip() for code in raw_sku_codes if (code or "").strip()]
    if not normalized_codes:
        return {}, []

    exact_skus = SKU.query.filter(SKU.sku_code.in_(normalized_codes)).all()
    resolved_map = {sku.sku_code: sku for sku in exact_skus}
    missing_codes = [code for code in normalized_codes if code not in resolved_map]

    current = user or g.get("user")
    if current and missing_codes:
        mapping_rows = SKUMapping.query.join(SKUMapping.sku).filter(
            SKUMapping.user_id == current.id,
            SKUMapping.external_sku_code.in_(missing_codes),
        ).all()
        for mapping in mapping_rows:
            resolved_map[mapping.external_sku_code] = mapping.sku
        missing_codes = [code for code in normalized_codes if code not in resolved_map]

    return resolved_map, missing_codes


def build_missing_mapping_error(missing_codes):
    preview = ", ".join(missing_codes[:10])
    suffix = " 等" if len(missing_codes) > 10 else ""
    return f"以下卖家SKU未匹配到仓库SKU或SKU映射：{preview}{suffix}。请先到“SKU映射”或“SKU管理”中补充后重试。"


def parse_shipment_import(file_storage, *, include_customer, user=None):
    filename = (file_storage.filename or "").strip()
    rows = parse_uploaded_table(file_storage)
    metadata = {}
    header_row = None
    seller_sku_col = None
    shipped_qty_col = None

    for index, row in enumerate(rows):
        cells = [cell.strip() for cell in row]
        non_empty = [cell for cell in cells if cell]
        if len(non_empty) >= 2 and "卖家 SKU" in cells and "已发货" in cells:
            header_row = index
            seller_sku_col = cells.index("卖家 SKU")
            shipped_qty_col = cells.index("已发货")
            break
        if len(non_empty) >= 2:
            metadata[cells[0]] = cells[1]

    if header_row is None or seller_sku_col is None or shipped_qty_col is None:
        raise ValueError("未识别到导入模板中的“卖家 SKU”或“已发货”列。")

    order_no = (metadata.get("货件编号") or "").strip()
    delivery_value = (metadata.get("配送地址") or "").strip()
    customer_name = "亚马逊" if include_customer and is_fba_shipment_table(rows) else ""
    grouped_quantities = defaultdict(int)

    for row in rows[header_row + 1 :]:
        cells = [cell.strip() for cell in row]
        if not any(cells):
            continue
        sku_code = cells[seller_sku_col] if seller_sku_col < len(cells) else ""
        qty_raw = cells[shipped_qty_col] if shipped_qty_col < len(cells) else ""
        quantity = parse_int(qty_raw)
        if not sku_code or quantity <= 0:
            continue
        grouped_quantities[sku_code] += quantity

    if not grouped_quantities:
        raise ValueError("导入文件里没有可用的 SKU 和已发货数量。")

    sku_codes = sorted(grouped_quantities.keys())
    sku_map, missing_codes = resolve_import_sku_codes(sku_codes, user=user)
    if missing_codes:
        raise ValueError(build_missing_mapping_error(missing_codes))

    merged_items = defaultdict(lambda: {"quantity": 0, "sku": None})
    for raw_sku_code in sku_codes:
        mapped_sku = sku_map[raw_sku_code]
        merged_items[mapped_sku.sku_code]["quantity"] += grouped_quantities[raw_sku_code]
        merged_items[mapped_sku.sku_code]["sku"] = mapped_sku

    items = [
        {
            "sku_id": payload["sku"].id,
            "sku_code": warehouse_sku_code,
            "quantity": payload["quantity"],
            "unit_price": "0",
        }
        for warehouse_sku_code, payload in sorted(merged_items.items())
    ]

    customer = Customer.query.filter_by(name=customer_name).first() if include_customer and customer_name else None
    return {
        "order_no": order_no,
        "customer_id": customer.id if customer else "",
        "customer_name": customer.name if customer else customer_name,
        "shipping_address": delivery_value if include_customer else "",
        "reference_no": order_no,
        "reference_type": "FBA货件",
        "items": items,
        "source_name": filename,
    }


def parse_sku_mapping_import(file_storage, *, allowed_user_ids=None, current_user=None, overwrite_existing=False):
    rows = parse_uploaded_table(file_storage)
    if not rows:
        raise ValueError("导入文件为空。")

    header_index = None
    headers = []
    for index, row in enumerate(rows):
        normalized_row = [(cell or "").strip() for cell in row]
        if "店铺SKU" in normalized_row and "仓库SKU" in normalized_row:
            header_index = index
            headers = normalized_row
            break
    if header_index is None:
        raise ValueError("未识别到模板表头，请先下载模板后填写。")

    def get_value(cells, column_name):
        if column_name not in headers:
            return ""
        column_index = headers.index(column_name)
        return cells[column_index].strip() if column_index < len(cells) else ""

    user_lookup = {}
    for user in User.query.filter(User.is_active.is_(True)).all():
        user_lookup[user.username.strip().lower()] = user
        if user.full_name:
            user_lookup[user.full_name.strip().lower()] = user

    allowed_user_ids = set(allowed_user_ids or [])
    seen_pairs = set()
    create_rows = []
    update_rows = []
    errors = []

    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        cells = [(cell or "").strip() for cell in row]
        if not any(cells):
            continue

        user_key = (get_value(cells, "运营账号") or get_value(cells, "运营姓名")).strip()
        external_sku_code = get_value(cells, "店铺SKU")
        warehouse_sku_code = get_value(cells, "仓库SKU")
        remark = get_value(cells, "备注")

        user = user_lookup.get(user_key.lower()) if user_key else current_user
        sku = SKU.query.filter_by(sku_code=warehouse_sku_code).first() if warehouse_sku_code else None

        if not user:
            errors.append(f"第 {row_number} 行：运营不存在或未填写。")
            continue
        if allowed_user_ids and user.id not in allowed_user_ids:
            errors.append(f"第 {row_number} 行：无权为运营 {user.full_name} 导入映射。")
            continue
        if not external_sku_code:
            errors.append(f"第 {row_number} 行：店铺SKU不能为空。")
            continue
        if not sku:
            errors.append(f"第 {row_number} 行：仓库SKU {warehouse_sku_code or '-'} 不存在。")
            continue
        if not validate_sku_access([sku.id], user=current_user):
            errors.append(f"第 {row_number} 行：无权映射仓库SKU {warehouse_sku_code}。")
            continue
        if (user.id, external_sku_code) in seen_pairs:
            errors.append(f"第 {row_number} 行：同一运营下的店铺SKU {external_sku_code} 在文件中重复。")
            continue
        existing_mapping = SKUMapping.query.filter_by(user_id=user.id, external_sku_code=external_sku_code).first()
        if existing_mapping and not overwrite_existing:
            errors.append(f"第 {row_number} 行：运营 {user.full_name} 的店铺SKU {external_sku_code} 已存在映射。")
            continue

        seen_pairs.add((user.id, external_sku_code))
        row_payload = {
            "user_id": user.id,
            "sku_id": sku.id,
            "external_sku_code": external_sku_code,
            "remark": remark,
        }
        if existing_mapping:
            update_rows.append((existing_mapping, row_payload))
        else:
            create_rows.append(row_payload)

    if errors:
        raise ValueError("\n".join(errors[:20]))
    if not create_rows and not update_rows:
        raise ValueError("导入文件中没有可导入的SKU映射数据。")
    return {"create_rows": create_rows, "update_rows": update_rows}


def sanitize_excel_cell(value):
    if isinstance(value, str) and value.startswith(EXCEL_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def export_workbook(filename, sheet_name, headers, rows):
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError:
        flash("当前 Python 环境缺少 openpyxl，暂时无法导出报表。", "error")
        return redirect(url_for("reports"))

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append([sanitize_excel_cell(value) for value in headers])
    for row in rows:
        ws.append([sanitize_excel_cell(value) for value in row])
    for column in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 28)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def export_purchase_order_workbook(order):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError:
        flash("当前 Python 环境缺少 openpyxl，暂时无法导出采购订单。", "error")
        return redirect(url_for("purchase_order_list"))

    show_prices = can_view_prices()
    headers = ["SKU", "商品", "规格", "数量", "单位"]
    if show_prices:
        headers.extend(["单价", "金额"])

    wb = Workbook()
    ws = wb.active
    ws.title = "采购订单"
    last_column = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = ws.cell(row=1, column=1, value=f"采购订单 {order.order_no}")
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center")

    info_rows = [
        ("订单号", order.order_no),
        ("供应商", order.supplier.name if order.supplier else ""),
        ("目标仓库", order.warehouse.name if order.warehouse else ""),
        ("订单状态", order.status_label),
        ("预计到货", order.expected_date.strftime("%Y-%m-%d") if order.expected_date else ""),
        ("操作人", order.operator_name or ""),
        ("备注", order.remark or ""),
    ]
    if show_prices:
        info_rows.append(("订单金额", float(order.total_amount or 0)))

    for row_index, (label, value) in enumerate(info_rows, start=3):
        ws.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_index, column=2, value=sanitize_excel_cell(value))

    table_start = len(info_rows) + 5
    header_fill = PatternFill("solid", fgColor="EAF0F8")
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=table_start, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_index, item in enumerate(order.items, start=table_start + 1):
        row_values = [
            item.sku.sku_code,
            item.sku.name,
            f"{item.sku.color or '-'} / {item.sku.size or '-'}",
            item.quantity,
            item.sku.unit,
        ]
        if show_prices:
            unit_price = float(item.unit_price or 0)
            row_values.extend([unit_price, unit_price * item.quantity])
        for column_index, value in enumerate(row_values, start=1):
            ws.cell(row=row_index, column=column_index, value=sanitize_excel_cell(value))

    for column_index in range(1, ws.max_column + 1):
        max_length = max(len(str(ws.cell(row=row_index, column=column_index).value or "")) for row_index in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 32)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=f"{order.order_no}-采购订单.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def count_active_users_with_permission(code, exclude_user_id=None):
    users = User.query.filter_by(is_active=True).all()
    count = 0
    for user in users:
        if exclude_user_id and user.id == exclude_user_id:
            continue
        if user.has_permission(code):
            count += 1
    return count


def count_active_users_with_role(role_id, exclude_role_id=None):
    users = User.query.filter_by(is_active=True).all()
    count = 0
    for user in users:
        role_ids = {role.id for role in user.roles}
        if exclude_role_id and exclude_role_id in role_ids:
            role_ids.discard(exclude_role_id)
        if role_id in role_ids:
            count += 1
    return count


def would_remove_last_user_manager(user, roles, is_active):
    if not user.is_active or not user.has_permission("user.manage"):
        return False
    resulting_permissions = {permission.code for role in roles for permission in role.permissions}
    if is_active and "user.manage" in resulting_permissions:
        return False
    return count_active_users_with_permission("user.manage", exclude_user_id=user.id) == 0


def would_role_update_remove_last_user_manager(role, permissions):
    resulting_permission_codes = {permission.code for permission in permissions}
    if "user.manage" in resulting_permission_codes:
        return False
    affected_users = [user for user in User.query.filter_by(is_active=True).all() if any(item.id == role.id for item in user.roles)]
    for user in affected_users:
        remaining_permission_codes = {
            permission.code
            for current_role in user.roles
            if current_role.id != role.id
            for permission in current_role.permissions
        }
        if "user.manage" not in remaining_permission_codes and count_active_users_with_permission("user.manage", exclude_user_id=user.id) == 0:
            return True
    return False


def would_role_delete_remove_last_user_manager(role):
    if "user.manage" not in {permission.code for permission in role.permissions}:
        return False
    affected_users = [user for user in User.query.filter_by(is_active=True).all() if any(item.id == role.id for item in user.roles)]
    for user in affected_users:
        remaining_permission_codes = {
            permission.code
            for current_role in user.roles
            if current_role.id != role.id
            for permission in current_role.permissions
        }
        if "user.manage" not in remaining_permission_codes and count_active_users_with_permission("user.manage", exclude_user_id=user.id) == 0:
            return True
    return False


def get_current_operator_name():
    user = g.get("user")
    if not user:
        return ""
    return (user.full_name or user.username or "").strip()


def user_requires_sku_scope(user=None):
    current = user or g.get("user")
    if not current or is_admin_user(current):
        return False
    if not current.has_permission("sales.manage"):
        return False
    # SKU 授权范围只限制销售侧账号，仓库/采购/管理类账号不应被误过滤为空。
    if any(
        current.has_permission(code)
        for code in ("stock.in", "purchase.manage", "warehouse.manage", "inventory.transaction.view", "shipment.confirm")
    ):
        return False
    return True


def get_sku_assignable_users(include_user_ids=None):
    query = User.query.filter(
        User.is_active.is_(True),
        or_(
            User.roles.any(Role.permissions.any(Permission.code == "sales.manage")),
            User.id.in_(include_user_ids or []),
        ),
    )
    return query.order_by(User.full_name.asc(), User.username.asc()).all()


def can_manage_all_sku_mappings(user=None):
    current = user or g.get("user")
    return bool(current and (is_admin_user(current) or current.has_permission("sku.manage") or current.has_permission("user.manage")))


def get_mapping_manageable_users(include_user_ids=None, user=None):
    current = user or g.get("user")
    if not current:
        return []
    included_ids = {value for value in (include_user_ids or []) if value}
    if can_manage_all_sku_mappings(current):
        return get_sku_assignable_users(included_ids)
    return (
        User.query.filter(User.is_active.is_(True), User.id.in_(included_ids | {current.id}))
        .order_by(User.full_name.asc(), User.username.asc())
        .all()
    )


def get_shipment_mapping_options(user=None):
    current = user or g.get("user")
    if not current:
        return []
    query = SKUMapping.query.join(SKUMapping.sku)
    if not can_manage_all_sku_mappings(current):
        query = query.filter(SKUMapping.user_id == current.id)
    return query.order_by(SKUMapping.external_sku_code.asc(), SKU.sku_code.asc()).all()


def get_shipment_lock_token():
    token = session.get("_shipment_lock_token")
    if not token:
        token = secrets.token_urlsafe(24)
        session["_shipment_lock_token"] = token
    return token


def cleanup_shipment_draft_locks():
    ShipmentDraftLock.query.filter(ShipmentDraftLock.expires_at <= datetime.utcnow()).delete()
    db.session.commit()


def get_active_shipment_draft_lock_map(*, warehouse_id=None, exclude_lock_token=None):
    query = db.session.query(
        ShipmentDraftLock.warehouse_id,
        ShipmentDraftLock.sku_id,
        func.coalesce(func.sum(ShipmentDraftLock.quantity), 0),
    ).filter(ShipmentDraftLock.expires_at > datetime.utcnow())
    if warehouse_id:
        query = query.filter(ShipmentDraftLock.warehouse_id == warehouse_id)
    if exclude_lock_token:
        query = query.filter(ShipmentDraftLock.lock_token != exclude_lock_token)
    rows = query.group_by(ShipmentDraftLock.warehouse_id, ShipmentDraftLock.sku_id).all()
    return {(warehouse_id, sku_id): quantity for warehouse_id, sku_id, quantity in rows}


def get_current_shipment_draft_locks():
    return ShipmentDraftLock.query.filter(
        ShipmentDraftLock.lock_token == get_shipment_lock_token(),
        ShipmentDraftLock.expires_at > datetime.utcnow(),
    ).all()


def get_pending_shipment_reserved_map(*, warehouse_id=None, exclude_shipment_id=None):
    query = (
        db.session.query(
            ShipmentSheet.warehouse_id,
            ShipmentSheetItem.sku_id,
            func.coalesce(func.sum(ShipmentSheetItem.quantity), 0),
        )
        .join(ShipmentSheetItem, ShipmentSheetItem.shipment_sheet_id == ShipmentSheet.id)
        .filter(ShipmentSheet.status == "pending")
    )
    if warehouse_id:
        query = query.filter(ShipmentSheet.warehouse_id == warehouse_id)
    if exclude_shipment_id:
        query = query.filter(ShipmentSheet.id != exclude_shipment_id)
    rows = query.group_by(ShipmentSheet.warehouse_id, ShipmentSheetItem.sku_id).all()
    return {(warehouse_id, sku_id): quantity for warehouse_id, sku_id, quantity in rows}


def get_shipment_available_inventory_map(*, warehouse_id=None, exclude_shipment_id=None):
    inventory_map = get_inventory_by_warehouse()
    reserved_map = get_pending_shipment_reserved_map(warehouse_id=warehouse_id, exclude_shipment_id=exclude_shipment_id)
    draft_lock_map = get_active_shipment_draft_lock_map(warehouse_id=warehouse_id, exclude_lock_token=get_shipment_lock_token())
    keys = set(inventory_map) | set(reserved_map) | set(draft_lock_map)
    available = {}
    for key in keys:
        available[key] = max(
            int(inventory_map.get(key, 0) or 0)
            - int(reserved_map.get(key, 0) or 0)
            - int(draft_lock_map.get(key, 0) or 0),
            0,
        )
    return available


def get_pending_shipment_count():
    return ShipmentSheet.query.filter_by(status="pending").count()


def parse_amazon_inventory_report(file_storage):
    rows = parse_uploaded_table(file_storage)
    result = {}
    for row in rows:
        if len(row) <= 10:
            continue
        external_sku = normalize_external_sku(row[0])
        if not external_sku or external_sku.lower() in {"店铺sku", "sku", "merchant sku"}:
            continue
        result[external_sku] = {
            "received_qty": parse_int(row[9]),
            "sellable_qty": parse_int(row[10]),
        }
    return result


def normalize_report_header(value):
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", (value or "").strip().lower())


def find_report_header_index(normalized_row, priorities):
    for candidate_group in priorities:
        for candidate in candidate_group:
            for index, cell in enumerate(normalized_row):
                if cell and candidate in cell:
                    return index
    return None


def find_report_header_row(rows, required_groups):
    for row_index, row in enumerate(rows[:10]):
        normalized_row = [normalize_report_header(cell) for cell in row]
        matched_indexes = []
        for priorities in required_groups:
            match_index = find_report_header_index(normalized_row, priorities)
            if match_index is None:
                matched_indexes = []
                break
            matched_indexes.append(match_index)
        if matched_indexes:
            return row_index, matched_indexes
    return None, []


def parse_reserved_inventory_report(file_storage):
    rows = parse_uploaded_table(file_storage)
    result = defaultdict(int)
    header_row_index, matched_indexes = find_report_header_row(
        rows,
        [
            [
                ("店铺sku", "merchantsku", "sellersku", "msku"),
                ("sku",),
            ],
            [
                ("reservedfctransfers",),
                ("预留库存", "预留数量", "afnreservedquantity"),
                ("reserved",),
            ],
        ],
    )
    if matched_indexes:
        sku_index, reserved_index = matched_indexes
        data_rows = rows[header_row_index + 1 :]
    else:
        sku_index, reserved_index = 0, 6
        data_rows = rows

    for row in data_rows:
        if len(row) <= max(sku_index, reserved_index):
            continue
        external_sku = normalize_external_sku(row[sku_index])
        normalized_sku = normalize_report_header(external_sku)
        if not external_sku or normalized_sku in {"店铺sku", "sku", "merchantsku", "sellersku", "msku"}:
            continue
        result[external_sku] += parse_int(row[reserved_index])
    return dict(result)


def parse_sales_velocity_report(file_storage, *, marketplace="US"):
    return parse_sales_velocity_windows(file_storage, marketplace=marketplace, require_dates=False)["sales_30"]


def parse_sales_report_datetime(value):
    raw_value = str(value or "").strip()
    if not raw_value:
        return None
    if re.fullmatch(r"\d+(\.\d+)?", raw_value):
        try:
            serial_value = float(raw_value)
        except ValueError:
            serial_value = 0
        if serial_value > 20000:
            return datetime(1899, 12, 30) + timedelta(days=serial_value)

    normalized = raw_value.replace("T", " ").replace("Z", "")
    normalized = re.sub(r"\s+[+-]\d{2}:?\d{2}$", "", normalized)
    normalized = re.sub(r"\.\d+", "", normalized)
    try:
        return datetime.fromisoformat(raw_value.replace("Z", "+00:00"))
    except ValueError:
        pass
    for date_format in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%y %H:%M:%S",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m/%d/%y",
    ):
        try:
            return datetime.strptime(normalized[: len(datetime.now().strftime(date_format))], date_format)
        except ValueError:
            continue
    match = re.search(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", raw_value)
    if match:
        return parse_sales_report_datetime(match.group(0))
    return None


def parse_sales_report_date(value, *, target_timezone=None):
    parsed_datetime = parse_sales_report_datetime(value)
    if not parsed_datetime:
        return None
    if target_timezone and parsed_datetime.tzinfo:
        return parsed_datetime.astimezone(target_timezone).date()
    return parsed_datetime.date()


def find_sales_velocity_columns(rows):
    header_row_index, matched_indexes = find_report_header_row(
        rows,
        [
            [
                ("店铺sku", "卖家sku", "merchantsku", "sellersku", "msku"),
                ("sku",),
            ],
            [
                ("quantitypurchased", "itemquantity", "orderedquantity", "unitsordered", "商品数量", "购买数量", "已订购商品数量"),
                ("销量", "数量", "quantity"),
            ],
        ],
    )
    if not matched_indexes:
        return {
            "header_row_index": None,
            "sku_index": 11,
            "quantity_index": 14,
            "marketplace_index": 27,
            "date_index": None,
        }

    normalized_header = [normalize_report_header(cell) for cell in rows[header_row_index]]
    marketplace_index = find_report_header_index(
        normalized_header,
        [
            ("marketplacename", "marketplace", "saleschannel", "国家店铺", "站点", "商城"),
        ],
    )
    country_index = find_report_header_index(
        normalized_header,
        [
            ("shipcountry", "shiptocountry", "countrycode", "country", "收货国家", "配送国家"),
        ],
    )
    date_index = find_report_header_index(
        normalized_header,
        [
            ("purchasedate", "orderdate", "paymentsdate", "date", "购买日期", "订单日期", "付款时间", "日期"),
        ],
    )
    return {
        "header_row_index": header_row_index,
        "sku_index": matched_indexes[0],
        "quantity_index": matched_indexes[1],
        "marketplace_index": marketplace_index,
        "country_index": country_index,
        "date_index": date_index,
    }


def normalize_marketplace_code(value):
    text = str(value or "").strip().upper()
    if not text or text in {"ALL", "全部"}:
        return ""
    marketplace_aliases = {
        "AMAZON.COM": "US",
        "AMAZON.CA": "CA",
        "AMAZON.COM.MX": "MX",
        "AMAZON.CO.UK": "UK",
        "AMAZON.DE": "DE",
        "AMAZON.FR": "FR",
        "AMAZON.IT": "IT",
        "AMAZON.ES": "ES",
        "AMAZON.JP": "JP",
        "AMAZON.COM.AU": "AU",
    }
    if text in marketplace_aliases:
        return marketplace_aliases[text]
    if text.startswith("AMAZON."):
        suffix = text.rsplit(".", 1)[-1]
        return "US" if suffix == "COM" else suffix
    return text


MARKETPLACE_TIMEZONE_NAMES = {
    "US": "America/Los_Angeles",
    "CA": "America/Los_Angeles",
    "MX": "America/Mexico_City",
    "UK": "Europe/London",
    "DE": "Europe/Berlin",
    "FR": "Europe/Paris",
    "IT": "Europe/Rome",
    "ES": "Europe/Madrid",
    "JP": "Asia/Tokyo",
    "AU": "Australia/Sydney",
}


def get_marketplace_timezone(marketplace):
    code = normalize_marketplace_code(marketplace)
    timezone_name = MARKETPLACE_TIMEZONE_NAMES.get(code) or MARKETPLACE_TIMEZONE_NAMES["US"]
    try:
        return ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        return timezone.utc


def get_sales_complete_window_end_date(marketplace, *, as_of_datetime=None):
    current_time = as_of_datetime or datetime.now(timezone.utc)
    if current_time.tzinfo is None:
        current_time = current_time.replace(tzinfo=timezone.utc)
    marketplace_today = current_time.astimezone(get_marketplace_timezone(marketplace)).date()
    return marketplace_today - timedelta(days=1)


def marketplace_matches_filter(report_value, selected_marketplace):
    selected = normalize_marketplace_code(selected_marketplace)
    if not selected:
        return True
    report_marketplace = normalize_marketplace_code(report_value)
    return not report_marketplace or report_marketplace == selected


def parse_sales_velocity_windows(file_storage, *, marketplace="US", require_dates=True, as_of_datetime=None):
    rows = parse_uploaded_table(file_storage)
    columns = find_sales_velocity_columns(rows)
    result = defaultdict(int)
    daily_result = defaultdict(lambda: defaultdict(int))
    max_sale_date = None
    data_rows = rows[(columns["header_row_index"] + 1) :] if columns["header_row_index"] is not None else rows

    for row in data_rows:
        required_index = max(
            index
            for index in (
                columns["sku_index"],
                columns["quantity_index"],
                columns["marketplace_index"] if columns["marketplace_index"] is not None else -1,
                columns["country_index"] if columns.get("country_index") is not None else -1,
                columns["date_index"] if columns["date_index"] is not None else -1,
            )
            if index is not None
        )
        if len(row) <= required_index:
            continue
        external_sku = normalize_external_sku(row[columns["sku_index"]])
        marketplace_code = ""
        if columns["marketplace_index"] is not None and columns["marketplace_index"] < len(row):
            marketplace_code = (row[columns["marketplace_index"]] or "").strip().upper()
        elif columns.get("country_index") is not None and columns["country_index"] < len(row):
            marketplace_code = (row[columns["country_index"]] or "").strip().upper()
        elif len(row) > 27:
            marketplace_code = (row[27] or "").strip().upper()
        quantity = parse_int(row[columns["quantity_index"]])
        if not external_sku or external_sku.lower() in {"店铺sku", "sku", "merchant sku"}:
            continue
        if not marketplace_matches_filter(marketplace_code, marketplace):
            continue
        if quantity <= 0:
            continue
        sale_date = None
        if columns["date_index"] is not None and columns["date_index"] < len(row):
            sale_date = parse_sales_report_date(
                row[columns["date_index"]],
                target_timezone=get_marketplace_timezone(marketplace_code or marketplace),
            )
        if sale_date:
            max_sale_date = max(max_sale_date, sale_date) if max_sale_date else sale_date
            daily_result[external_sku][sale_date.isoformat()] += quantity
        result[external_sku] += quantity

    if require_dates and result and not max_sale_date:
        raise ValueError("30天销量表需要包含订单日期/购买日期/付款时间等日期列，系统才能自动拆分最近7天、14天和30天销量。")

    if not max_sale_date:
        return {
            "sales_7": dict(result),
            "sales_14": dict(result),
            "sales_30": dict(result),
            "daily_sales": {},
            "window_end_date": "",
            "latest_sale_date": "",
            "excluded_latest_date": "",
            "date_column_found": False,
        }

    window_end_date = get_sales_complete_window_end_date(marketplace, as_of_datetime=as_of_datetime)

    def build_window(days):
        start_date = window_end_date - timedelta(days=days - 1)
        window_result = defaultdict(int)
        for external_sku, date_map in daily_result.items():
            for date_text, quantity in date_map.items():
                sale_date = parse_sales_report_date(date_text)
                if sale_date and start_date <= sale_date <= window_end_date:
                    window_result[external_sku] += quantity
        return dict(window_result)

    return {
        "sales_7": build_window(7),
        "sales_14": build_window(14),
        "sales_30": build_window(30),
        "daily_sales": {sku: dict(date_map) for sku, date_map in daily_result.items()},
        "window_end_date": window_end_date.isoformat(),
        "latest_sale_date": max_sale_date.isoformat(),
        "excluded_latest_date": (window_end_date + timedelta(days=1)).isoformat(),
        "date_column_found": True,
    }


def build_purchase_sales_windows(file_storage, *, user=None, marketplace="ALL", as_of_datetime=None):
    raw_windows = parse_sales_velocity_windows(file_storage, marketplace=marketplace, as_of_datetime=as_of_datetime)
    sku_lookup = {}
    sku_query = apply_sku_scope(SKU.query, SKU.id, user)
    for sku in sku_query.all():
        for value in (sku.sku_code, sku.barcode):
            key = normalize_external_sku(value).casefold()
            if key:
                sku_lookup.setdefault(key, sku.id)

    mapping_lookup = {}
    mapping_query = apply_sku_scope(SKUMapping.query, SKUMapping.sku_id, user)
    for mapping in mapping_query.all():
        key = normalize_external_sku(mapping.external_sku_code).casefold()
        if key:
            mapping_lookup.setdefault(key, mapping.sku_id)

    normalized_windows = {"sales_7": defaultdict(int), "sales_14": defaultdict(int), "sales_30": defaultdict(int)}
    unmatched_codes = set()
    all_codes = set()
    for window_key in normalized_windows:
        all_codes.update((raw_windows.get(window_key) or {}).keys())

    for code in sorted(all_codes):
        match_key = normalize_external_sku(code).casefold()
        sku_id = sku_lookup.get(match_key) or mapping_lookup.get(match_key)
        if not sku_id:
            unmatched_codes.add(code)
            continue
        for window_key in normalized_windows:
            normalized_windows[window_key][sku_id] += parse_int((raw_windows.get(window_key) or {}).get(code))

    return {
        "sales_7": dict(normalized_windows["sales_7"]),
        "sales_14": dict(normalized_windows["sales_14"]),
        "sales_30": dict(normalized_windows["sales_30"]),
        "unmatched_codes": sorted(unmatched_codes),
        "window_end_date": raw_windows.get("window_end_date", ""),
        "latest_sale_date": raw_windows.get("latest_sale_date", ""),
        "excluded_latest_date": raw_windows.get("excluded_latest_date", ""),
        "date_column_found": raw_windows.get("date_column_found", False),
    }


def parse_amazon_inbound_report(file_storage):
    rows = parse_uploaded_table(file_storage)
    result = defaultdict(int)
    for row in rows:
        if len(row) <= 9:
            continue
        external_sku = normalize_external_sku(row[0])
        quantity = parse_int(row[9])
        if not external_sku or external_sku.lower() in {"店铺sku", "sku", "merchant sku"}:
            continue
        if quantity <= 0:
            continue
        result[external_sku] += quantity
    return dict(result)


def parse_amazon_inbound_records(file_storage):
    rows = parse_uploaded_table(file_storage)
    shipment_ref = os.path.splitext((file_storage.filename or "").strip())[0] or build_unique_code("FBA", FBAInboundRecord, "shipment_ref")
    result = defaultdict(int)
    for row in rows:
        if len(row) <= 9:
            continue
        external_sku = normalize_external_sku(row[0])
        quantity = parse_int(row[9])
        if not external_sku or external_sku.lower() in {"店铺sku", "sku", "merchant sku"}:
            continue
        if quantity <= 0:
            continue
        result[external_sku] += quantity
    return shipment_ref, dict(result)


def get_active_fba_inbound_totals():
    rows = (
        db.session.query(FBAInboundRecord.external_sku_code, func.coalesce(func.sum(FBAInboundRecord.quantity), 0))
        .filter(FBAInboundRecord.status == "active")
        .group_by(FBAInboundRecord.external_sku_code)
        .all()
    )
    result = defaultdict(int)
    for external_sku_code, quantity in rows:
        normalized_code = normalize_external_sku(external_sku_code)
        if not normalized_code:
            continue
        result[normalized_code] += parse_int(quantity)
    return dict(result)


FBA_SIZE_ORDER = {
    "XXXS": 0,
    "XXS": 1,
    "XS": 2,
    "S": 3,
    "M": 4,
    "L": 5,
    "XL": 6,
    "XXL": 7,
    "XXXL": 8,
    "4XL": 9,
    "5XL": 10,
}


def parse_fba_sort_parts(code):
    normalized = (code or "").strip()
    if not normalized or normalized == "-":
        return ("~", "~", 999, normalized)
    parts = [part.strip() for part in normalized.split("-") if part.strip()]
    if not parts:
        return (normalized.casefold(), "", 999, "")
    style = parts[0].casefold()
    size_token = parts[-1].upper()
    if size_token in FBA_SIZE_ORDER:
        color = "-".join(parts[1:-1]).casefold()
        return (style, color, FBA_SIZE_ORDER[size_token], size_token)
    color = "-".join(parts[1:]).casefold()
    return (style, color, 999, size_token.casefold())


def sort_fba_result_rows(rows):
    rows.sort(
        key=lambda row: (
            *parse_fba_sort_parts(row.get("warehouse_sku_code") or row.get("external_sku_code") or ""),
            (row.get("external_sku_code") or "").casefold(),
        )
    )
    return rows


def get_fba_calc_result_path(result_key):
    safe_key = re.sub(r"[^A-Za-z0-9_-]", "", str(result_key or ""))
    if not safe_key:
        return ""
    return os.path.join(FBA_CALC_RESULT_DIR, f"{safe_key}.json")


def set_fba_calc_result(calc_result):
    os.makedirs(FBA_CALC_RESULT_DIR, exist_ok=True)
    result_key = session.get("_fba_calc_result_key") or secrets.token_urlsafe(24)
    result_path = get_fba_calc_result_path(result_key)
    with open(result_path, "w", encoding="utf-8") as handle:
        json.dump(calc_result or {}, handle, ensure_ascii=False, separators=(",", ":"))
    session["_fba_calc_result_key"] = result_key
    # Older versions stored the full result in the signed cookie; remove it so large
    # calculations do not overflow the browser cookie on the next filter/update request.
    session.pop("_fba_calc_result", None)
    session.modified = True
    return calc_result or {}


def get_fba_calc_result():
    legacy_result = session.get("_fba_calc_result")
    if legacy_result:
        return set_fba_calc_result(legacy_result)
    result_path = get_fba_calc_result_path(session.get("_fba_calc_result_key"))
    if not result_path or not os.path.exists(result_path):
        return {}
    try:
        with open(result_path, "r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return {}


def get_purchase_calc_result_path(result_key):
    safe_key = re.sub(r"[^A-Za-z0-9_-]", "", str(result_key or ""))
    if not safe_key:
        return ""
    return os.path.join(PURCHASE_CALC_RESULT_DIR, f"{safe_key}.json")


def set_purchase_calc_result(calc_result):
    os.makedirs(PURCHASE_CALC_RESULT_DIR, exist_ok=True)
    result_key = session.get("_purchase_calc_result_key") or secrets.token_urlsafe(24)
    result_path = get_purchase_calc_result_path(result_key)
    with open(result_path, "w", encoding="utf-8") as handle:
        json.dump(calc_result or {}, handle, ensure_ascii=False, separators=(",", ":"))
    session["_purchase_calc_result_key"] = result_key
    session.pop("_purchase_calc_result", None)
    session.modified = True
    return calc_result or {}


def get_purchase_calc_result():
    legacy_result = session.get("_purchase_calc_result")
    if legacy_result:
        return set_purchase_calc_result(legacy_result)
    result_path = get_purchase_calc_result_path(session.get("_purchase_calc_result_key"))
    if not result_path or not os.path.exists(result_path):
        return {}
    try:
        with open(result_path, "r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return {}


def build_filtered_fba_calc_result(calc_result, *, sku_keyword="", suggested_state="", suggested_min=""):
    source_rows = calc_result.get("rows") or []
    sku_keyword = (sku_keyword or "").strip().casefold()
    suggested_state = (suggested_state or "").strip()
    suggested_min_value = parse_int(suggested_min, None) if str(suggested_min).strip() else None
    filtered_rows = []

    for raw_row in source_rows:
        row = dict(raw_row)
        if sku_keyword:
            haystacks = [
                row.get("external_sku_code") or "",
                row.get("warehouse_sku_code") or "",
                row.get("warehouse_sku_name") or "",
            ]
            if not any(sku_keyword in value.casefold() for value in haystacks):
                continue
        suggested_qty = parse_int(row.get("suggested_qty"))
        if suggested_state == "positive" and suggested_qty <= 0:
            continue
        if suggested_state == "zero" and suggested_qty != 0:
            continue
        if suggested_min_value is not None and suggested_qty < suggested_min_value:
            continue
        filtered_rows.append(row)

    return {
        "rows": filtered_rows,
        "unmatched_codes": calc_result.get("unmatched_codes") or [],
        "summary": {
            "sku_count": len(filtered_rows),
            "suggested_qty": sum(parse_int(row.get("suggested_qty")) for row in filtered_rows),
            "sellable_qty": sum(parse_int(row.get("sellable_qty")) for row in filtered_rows),
            "inbound_qty": sum(parse_int(row.get("inbound_qty")) for row in filtered_rows),
        },
    }


def build_shipment_page_context(*, shipment_form=None):
    cleanup_shipment_draft_locks()
    selected_status = request.args.get("shipment_status", "").strip()
    keyword = request.args.get("shipment_q", "").strip()
    selected_store = request.args.get("shipment_store_name", "").strip()
    selected_archive = request.args.get("shipment_archive", "active").strip() or "active"
    shipment_query = ShipmentSheet.query.join(ShipmentSheet.warehouse)
    if selected_archive == "archived":
        shipment_query = shipment_query.filter(ShipmentSheet.archived_at.isnot(None))
    elif selected_archive != "all":
        selected_archive = "active"
        shipment_query = shipment_query.filter(ShipmentSheet.archived_at.is_(None))
    if selected_status:
        shipment_query = shipment_query.filter(ShipmentSheet.status == selected_status)
    if selected_store:
        shipment_query = shipment_query.filter(ShipmentSheet.store_name == selected_store)
    if keyword:
        keyword_like = f"%{keyword}%"
        shipment_query = shipment_query.filter(
            or_(
                ShipmentSheet.shipment_no.ilike(keyword_like),
                ShipmentSheet.store_name.ilike(keyword_like),
                ShipmentSheet.operator_name.ilike(keyword_like),
                ShipmentSheetItem.external_sku_code.ilike(keyword_like),
                SKU.sku_code.ilike(keyword_like),
                SKU.name.ilike(keyword_like),
            )
        ).outerjoin(ShipmentSheet.items).outerjoin(ShipmentSheetItem.sku)

    shipments = shipment_query.order_by(
        case((ShipmentSheet.status == "pending", 0), else_=1),
        ShipmentSheet.created_at.desc(),
    ).distinct().all()
    shipment_summary = {
        "count": len(shipments),
        "pending_count": sum(1 for row in shipments if row.status == "pending"),
        "confirmed_count": sum(1 for row in shipments if row.status == "confirmed"),
        "quantity": sum(row.total_quantity for row in shipments),
    }
    return {
        "shipments": shipments,
        "shipment_summary": shipment_summary,
        "shipment_stores": sorted({row[0] for row in ShipmentSheet.query.with_entities(ShipmentSheet.store_name).all() if row[0]}),
        "shipment_mapping_options": get_shipment_mapping_options(user=g.user),
        "selected_shipment_status": selected_status,
        "selected_shipment_store": selected_store,
        "selected_shipment_archive": selected_archive,
        "shipment_q": keyword,
        "shipment_form": shipment_form or {},
    }


def build_fba_sales_metrics(*, sales_7_qty, sales_14_qty, sales_30_qty):
    return build_sales_velocity_metrics(
        sales_7_qty=sales_7_qty,
        sales_14_qty=sales_14_qty,
        sales_30_qty=sales_30_qty,
    )


def build_fba_target_days(*, coverage_days=None, production_days=None, sea_days=None, inbound_processing_days=None, safety_stock_days=None):
    cycle_values = [production_days, sea_days, inbound_processing_days, safety_stock_days]
    if any(value is not None for value in cycle_values):
        return max(
            parse_int(production_days)
            + parse_int(sea_days)
            + parse_int(inbound_processing_days)
            + parse_int(safety_stock_days),
            1,
        )
    return max(parse_int(coverage_days, 30), 1)


def classify_fba_sales_health(*, sellable_qty, reserved_qty, sales_7_qty, sales_14_qty, sales_30_qty):
    daily_7 = sales_7_qty / 7 if sales_7_qty > 0 else 0
    daily_14 = sales_14_qty / 14 if sales_14_qty > 0 else 0
    daily_30 = sales_30_qty / 30 if sales_30_qty > 0 else 0
    recent_baseline = max(daily_14, daily_30)
    recent_drop = recent_baseline > 0 and daily_7 < recent_baseline * 0.5
    stock_pressure_qty = sellable_qty + reserved_qty
    stock_pressure_limit = max(1, math.ceil(max(daily_30, daily_14) * 3))

    if sales_30_qty > 0 and sellable_qty <= 0 and (recent_drop or sales_7_qty <= 0):
        return {
            "status": "疑似断货",
            "note": "当前可售库存为0，近30天仍有销量；近7天低销量优先按断货处理。",
            "stockout_suspected": True,
        }
    if sales_30_qty > 0 and recent_drop and stock_pressure_qty <= stock_pressure_limit:
        return {
            "status": "疑似缺货影响",
            "note": f"近7天日均低于历史基准，且可售/预留合计不高于{stock_pressure_limit}件。",
            "stockout_suspected": True,
        }
    if sales_30_qty > 0 and recent_drop:
        return {
            "status": "销量下滑",
            "note": "近7天日均明显低于14/30天基准，当前仍有可售库存。",
            "stockout_suspected": False,
        }
    if recent_baseline > 0 and daily_7 > recent_baseline * 1.3:
        return {
            "status": "销量上升",
            "note": "近7天日均高于14/30天基准。",
            "stockout_suspected": False,
        }
    return {
        "status": "正常",
        "note": "近7/14/30天销量节奏未出现明显异常。",
        "stockout_suspected": False,
    }


def calculate_fba_replenishment(
    *,
    inventory_data,
    reserved_data,
    sales_7_data,
    sales_14_data,
    sales_30_data,
    inbound_data,
    multiplier,
    new_product_multiplier,
    coverage_days,
    production_days=None,
    sea_days=None,
    inbound_processing_days=None,
    safety_stock_days=None,
    user,
):
    mapping_rows = get_shipment_mapping_options(user=user)
    mapping_by_external_sku = {}
    for row in mapping_rows:
        mapping_by_external_sku[normalize_external_sku(row.external_sku_code)] = row
    rules_by_external_sku = {
        normalize_external_sku(row.external_sku_code): row
        for row in FBAReplenishmentRule.query.filter(
            FBAReplenishmentRule.external_sku_code.in_(
                list(set(inventory_data) | set(reserved_data) | set(sales_7_data) | set(sales_14_data) | set(sales_30_data) | set(inbound_data))
            )
        ).all()
    }

    all_external_codes = sorted(
        set(inventory_data)
        | set(reserved_data)
        | set(sales_7_data)
        | set(sales_14_data)
        | set(sales_30_data)
        | set(inbound_data)
    )
    warehouse = get_default_warehouse()
    warehouse_available = get_inventory_by_warehouse()
    result_rows = []
    unmatched_codes = []
    target_days = build_fba_target_days(
        coverage_days=coverage_days,
        production_days=production_days,
        sea_days=sea_days,
        inbound_processing_days=inbound_processing_days,
        safety_stock_days=safety_stock_days,
    )

    for external_code in all_external_codes:
        mapping = mapping_by_external_sku.get(external_code)
        inventory_row = inventory_data.get(external_code, {})
        sellable_qty = parse_int(inventory_row.get("sellable_qty"))
        received_qty = parse_int(inventory_row.get("received_qty"))
        reserved_qty = parse_int(reserved_data.get(external_code))
        sales_7_qty = parse_int(sales_7_data.get(external_code))
        sales_14_qty = parse_int(sales_14_data.get(external_code))
        sales_30_qty = parse_int(sales_30_data.get(external_code))
        inbound_qty = parse_int(inbound_data.get(external_code))
        rule = rules_by_external_sku.get(external_code)
        sales_metrics = build_fba_sales_metrics(
            sales_7_qty=sales_7_qty,
            sales_14_qty=sales_14_qty,
            sales_30_qty=sales_30_qty,
        )
        daily_7 = sales_metrics["daily_7"]
        daily_14 = sales_metrics["daily_14"]
        daily_30 = sales_metrics["daily_30"]
        weighted_daily = sales_metrics["weighted_daily"]
        base_daily = sales_metrics["base_daily"]
        trend_factor = sales_metrics["trend_factor"]
        sales_health = classify_fba_sales_health(
            sellable_qty=sellable_qty,
            reserved_qty=reserved_qty,
            sales_7_qty=sales_7_qty,
            sales_14_qty=sales_14_qty,
            sales_30_qty=sales_30_qty,
        )
        if sales_health["stockout_suspected"]:
            trend_factor = max(trend_factor, 1)
        applied_multiplier = float(rule.custom_multiplier) if rule else max(multiplier, 0)
        if rule and rule.is_new_product:
            applied_multiplier *= max(new_product_multiplier, 0)
        forecast_qty = base_daily * target_days * applied_multiplier
        suggested_qty = max(math.ceil(forecast_qty * trend_factor - sellable_qty - inbound_qty), 0)
        if not mapping:
            unmatched_codes.append(external_code)
            warehouse_sku_code = "-"
            warehouse_sku_name = "未映射"
            warehouse_qty = 0
            sku_id = None
        else:
            sku_id = mapping.sku_id
            warehouse_sku_code = mapping.sku.sku_code
            warehouse_sku_name = mapping.sku.name
            warehouse_qty = parse_int(warehouse_available.get((warehouse.id, mapping.sku_id), 0) if warehouse else 0)

        result_rows.append(
            {
                "external_sku_code": external_code,
                "sku_id": sku_id,
                "warehouse_sku_code": warehouse_sku_code,
                "warehouse_sku_name": warehouse_sku_name,
                "sellable_qty": sellable_qty,
                "received_qty": received_qty,
                "reserved_qty": reserved_qty,
                "sales_7_qty": sales_7_qty,
                "sales_14_qty": sales_14_qty,
                "sales_30_qty": sales_30_qty,
                "daily_7": daily_7,
                "daily_14": daily_14,
                "daily_30": daily_30,
                "weighted_daily": weighted_daily,
                "target_days": target_days,
                "inbound_qty": inbound_qty,
                "warehouse_qty": warehouse_qty,
                "trend_factor": round(trend_factor, 2),
                "applied_multiplier": round(applied_multiplier, 2),
                "is_new_product": bool(rule.is_new_product) if rule else False,
                "custom_multiplier": float(rule.custom_multiplier) if rule else round(max(multiplier, 0), 2),
                "rule_remark": rule.remark if rule else "",
                "sales_health_status": sales_health["status"],
                "sales_health_note": sales_health["note"],
                "stockout_suspected": sales_health["stockout_suspected"],
                "suggested_qty": suggested_qty,
            }
        )

    sort_fba_result_rows(result_rows)

    return {
        "rows": result_rows,
        "unmatched_codes": unmatched_codes,
        "summary": {
            "sku_count": len(result_rows),
            "suggested_qty": sum(row["suggested_qty"] for row in result_rows),
            "sellable_qty": sum(row["sellable_qty"] for row in result_rows),
            "inbound_qty": sum(row["inbound_qty"] for row in result_rows),
        },
    }


def refresh_fba_calc_session_result():
    calc_result = get_fba_calc_result()
    rows = calc_result.get("rows") or []
    if not rows:
        return

    calc_form = session.get("_fba_calc_form") or {}
    global_multiplier = max(parse_float(calc_form.get("multiplier"), 1), 0)
    new_product_multiplier = max(parse_float(calc_form.get("new_product_multiplier"), 1.5), 0)
    target_days = build_fba_target_days(
        coverage_days=calc_form.get("coverage_days"),
        production_days=calc_form.get("production_days") if any(key in calc_form for key in ("production_days", "inbound_processing_days", "safety_stock_days")) else None,
        sea_days=calc_form.get("sea_days") if any(key in calc_form for key in ("production_days", "inbound_processing_days", "safety_stock_days")) else None,
        inbound_processing_days=calc_form.get("inbound_processing_days") if any(key in calc_form for key in ("production_days", "inbound_processing_days", "safety_stock_days")) else None,
        safety_stock_days=calc_form.get("safety_stock_days") if any(key in calc_form for key in ("production_days", "inbound_processing_days", "safety_stock_days")) else None,
    )
    inbound_totals = get_active_fba_inbound_totals()
    external_codes = [row.get("external_sku_code", "").strip() for row in rows if row.get("external_sku_code")]
    rules_by_external_sku = {
        row.external_sku_code: row
        for row in FBAReplenishmentRule.query.filter(FBAReplenishmentRule.external_sku_code.in_(external_codes)).all()
    }

    for row in rows:
        external_code = normalize_external_sku(row.get("external_sku_code"))
        rule = rules_by_external_sku.get(external_code)
        custom_multiplier = float(rule.custom_multiplier) if rule else round(global_multiplier, 2)
        is_new_product = bool(rule.is_new_product) if rule else False
        applied_multiplier = custom_multiplier * (new_product_multiplier if is_new_product else 1)
        row["inbound_qty"] = parse_int(inbound_totals.get(external_code))
        sales_metrics = build_fba_sales_metrics(
            sales_7_qty=parse_int(row.get("sales_7_qty")),
            sales_14_qty=parse_int(row.get("sales_14_qty", row.get("sales_8_14_qty"))),
            sales_30_qty=parse_int(row.get("sales_30_qty")),
        )
        row["daily_7"] = sales_metrics["daily_7"]
        row["sales_14_qty"] = parse_int(row.get("sales_14_qty", row.get("sales_8_14_qty")))
        row["daily_14"] = sales_metrics["daily_14"]
        row["daily_30"] = sales_metrics["daily_30"]
        row["weighted_daily"] = sales_metrics["weighted_daily"]
        row["target_days"] = target_days
        row["trend_factor"] = sales_metrics["trend_factor"]
        base_daily = sales_metrics["base_daily"]
        trend_factor = sales_metrics["trend_factor"]
        sales_health = classify_fba_sales_health(
            sellable_qty=parse_int(row.get("sellable_qty")),
            reserved_qty=parse_int(row.get("reserved_qty")),
            sales_7_qty=parse_int(row.get("sales_7_qty")),
            sales_14_qty=parse_int(row.get("sales_14_qty", row.get("sales_8_14_qty"))),
            sales_30_qty=parse_int(row.get("sales_30_qty")),
        )
        if sales_health["stockout_suspected"]:
            trend_factor = max(trend_factor, 1)
            row["trend_factor"] = round(trend_factor, 2)
        forecast_qty = base_daily * target_days * applied_multiplier
        row["custom_multiplier"] = round(custom_multiplier, 2)
        row["is_new_product"] = is_new_product
        row["rule_remark"] = rule.remark if rule else ""
        row["sales_health_status"] = sales_health["status"]
        row["sales_health_note"] = sales_health["note"]
        row["stockout_suspected"] = sales_health["stockout_suspected"]
        row["applied_multiplier"] = round(applied_multiplier, 2)
        row["suggested_qty"] = max(
            math.ceil(forecast_qty * trend_factor - parse_int(row.get("sellable_qty")) - parse_int(row.get("inbound_qty"))),
            0,
        )

    sort_fba_result_rows(rows)
    calc_result["summary"] = {
        "sku_count": len(rows),
        "suggested_qty": sum(parse_int(row.get("suggested_qty")) for row in rows),
        "sellable_qty": sum(parse_int(row.get("sellable_qty")) for row in rows),
        "inbound_qty": sum(parse_int(row.get("inbound_qty")) for row in rows),
    }
    set_fba_calc_result(calc_result)


def get_authorized_sku_ids(user=None):
    current = user or g.get("user")
    if not user_requires_sku_scope(current):
        return None
    return {sku.id for sku in current.authorized_skus}


def apply_sku_scope(query, sku_column, user=None):
    authorized_sku_ids = get_authorized_sku_ids(user)
    if authorized_sku_ids is None:
        return query
    if not authorized_sku_ids:
        return query.filter(False)
    return query.filter(sku_column.in_(authorized_sku_ids))


def validate_sku_access(sku_ids, user=None):
    authorized_sku_ids = get_authorized_sku_ids(user)
    if authorized_sku_ids is None:
        return True
    requested_ids = {parse_int(value) for value in sku_ids if parse_int(value)}
    return requested_ids.issubset(authorized_sku_ids)


def apply_order_query_sku_scope(query, item_model, relation_attr, user=None):
    authorized_sku_ids = get_authorized_sku_ids(user)
    if authorized_sku_ids is None:
        return query
    if not authorized_sku_ids:
        return query.filter(False)
    return query.filter(relation_attr.any(item_model.sku_id.in_(authorized_sku_ids)))


def add_inventory_transaction(sku_id, warehouse_id, quantity, txn_type, reference_type="", reference_no="", note="", operator_name=None, supplier_id=None):
    if quantity == 0:
        raise InventoryError("数量不能为 0。")

    updated_at = datetime.utcnow()
    if quantity > 0:
        updated = (
            InventoryBalance.query.filter_by(sku_id=sku_id, warehouse_id=warehouse_id)
            .update(
                {
                    InventoryBalance.quantity: InventoryBalance.quantity + quantity,
                    InventoryBalance.updated_at: updated_at,
                },
                synchronize_session=False,
            )
        )
        if not updated:
            db.session.add(InventoryBalance(sku_id=sku_id, warehouse_id=warehouse_id, quantity=quantity))
            db.session.flush()
    else:
        updated = (
            InventoryBalance.query.filter_by(sku_id=sku_id, warehouse_id=warehouse_id)
            .filter(InventoryBalance.quantity >= -quantity)
            .update(
                {
                    InventoryBalance.quantity: InventoryBalance.quantity + quantity,
                    InventoryBalance.updated_at: updated_at,
                },
                synchronize_session=False,
            )
        )
        if not updated:
            raise InventoryError("库存不足。")

    db.session.add(
        InventoryTransaction(
            sku_id=sku_id,
            warehouse_id=warehouse_id,
            operator_name=(operator_name if operator_name is not None else get_current_operator_name()),
            quantity=quantity,
            transaction_type=txn_type,
            reference_type=reference_type,
            reference_no=reference_no,
            note=note,
            supplier_id=supplier_id,
        )
    )


def sync_defect_repair_inventory(repair):
    delete_inventory_transactions(reference_type="defect_repair", reference_no=repair.repair_no)

    if not repair.warehouse_id:
        raise InventoryError("返修单缺少仓库，无法同步库存。")

    note = f"返修单：{repair.repair_no}"
    operator_name = get_current_operator_name() or repair.operator_name
    if repair.status in {"pending", "repairing", "done"} or (repair.completed_quantity and repair.scrapped_quantity):
        add_inventory_transaction(
            repair.sku_id,
            repair.warehouse_id,
            -repair.quantity,
            "defect_hold",
            reference_type="defect_repair",
            reference_no=repair.repair_no,
            note=note,
            operator_name=operator_name,
        )
    if repair.completed_quantity:
        add_inventory_transaction(
            repair.sku_id,
            repair.warehouse_id,
            repair.completed_quantity,
            "defect_return",
            reference_type="defect_repair",
            reference_no=repair.repair_no,
            note=note,
            operator_name=operator_name,
        )
    elif repair.status == "scrapped":
        add_inventory_transaction(
            repair.sku_id,
            repair.warehouse_id,
            -repair.quantity,
            "defect_scrap",
            reference_type="defect_repair",
            reference_no=repair.repair_no,
            note=note,
            operator_name=operator_name,
        )


def apply_defect_repair_progress(repair, requested_status, *, completed_quantity=None, scrapped_quantity=None):
    quantity = repair.quantity or 0
    has_explicit_progress = completed_quantity is not None or scrapped_quantity is not None

    if has_explicit_progress:
        completed = parse_int(completed_quantity, repair.completed_quantity or 0)
        scrapped = parse_int(scrapped_quantity, repair.scrapped_quantity or 0)
        if completed == 0 and scrapped == 0 and requested_status == "done":
            completed = quantity
            has_explicit_progress = False
        elif completed == 0 and scrapped == 0 and requested_status == "scrapped":
            scrapped = quantity
            has_explicit_progress = False
    else:
        completed = repair.completed_quantity or 0
        scrapped = repair.scrapped_quantity or 0
        if requested_status == "done":
            completed = quantity
            scrapped = 0
        elif requested_status == "scrapped":
            completed = 0
            scrapped = quantity
        elif requested_status in {"pending", "repairing"}:
            completed = 0
            scrapped = 0

    if completed < 0 or scrapped < 0:
        raise ValueError("成功数量和报废数量不能小于 0。")
    if completed + scrapped > quantity:
        raise ValueError("成功数量和报废数量之和不能超过返修单总数量。")

    repair.completed_quantity = completed
    repair.scrapped_quantity = scrapped

    if has_explicit_progress:
        processed_quantity = completed + scrapped
        if processed_quantity >= quantity:
            if scrapped and not completed:
                repair.status = "scrapped"
            else:
                repair.status = "done"
        elif processed_quantity:
            repair.status = "repairing" if requested_status not in {"pending", "repairing"} else requested_status
        else:
            repair.status = requested_status
    else:
        repair.status = requested_status


def delete_inventory_transactions(*, transaction_ids=None, reference_type=None, reference_no=None):
    query = InventoryTransaction.query
    if transaction_ids is not None:
        clean_ids = sorted({parse_int(value) for value in transaction_ids if parse_int(value)})
        if not clean_ids:
            return 0
        query = query.filter(InventoryTransaction.id.in_(clean_ids))
    else:
        query = query.filter_by(reference_type=reference_type, reference_no=reference_no)

    transactions = query.all()
    if not transactions:
        return 0

    for transaction in transactions:
        db.session.delete(transaction)
    db.session.flush()
    sync_inventory_balances(commit=False)
    return len(transactions)


def entity_has_rows(model, **filters):
    return db.session.query(model.id).filter_by(**filters).first() is not None


def matches_sku_keyword_for_item(sku, keyword):
    if not keyword:
        return True
    keyword = keyword.lower()
    return any(
        keyword in (value or "").lower()
        for value in (sku.sku_code, sku.name, sku.barcode)
    )


def group_manual_outbound_transactions(transactions):
    groups = {}
    ordered_groups = []
    for txn in transactions:
        group_key = txn.reference_no or f"txn-{txn.id}"
        group = groups.get(group_key)
        if not group:
            group = {
                "group_key": group_key,
                "reference_type": txn.reference_type or "manual_outbound",
                "reference_no": txn.reference_no or f"TXN-{txn.id}",
                "warehouse": txn.warehouse,
                "created_at": txn.created_at,
                "note": txn.note or "",
                "lines": [],
                "total_quantity": 0,
            }
            groups[group_key] = group
            ordered_groups.append(group)
        group["lines"].append(txn)
        group["total_quantity"] += abs(txn.quantity)
        if txn.created_at > group["created_at"]:
            group["created_at"] = txn.created_at
        if not group["note"] and txn.note:
            group["note"] = txn.note
    return ordered_groups


def get_package_version(package_name, fallback="-"):
    try:
        return version(package_name)
    except PackageNotFoundError:
        return fallback


def get_deployment_mode():
    if os.path.exists("/.dockerenv"):
        return "Docker 容器"
    if app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite:////app/"):
        return "Docker 挂载模式"
    return "本地或其他环境"


def get_database_location():
    database_path = db.engine.url.database or ""
    if not database_path:
        return app.config["SQLALCHEMY_DATABASE_URI"]
    return os.path.abspath(database_path)


def to_local_time(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(LOCAL_TIMEZONE)


def format_dt(value, fmt="%Y-%m-%d %H:%M"):
    local_value = to_local_time(value)
    if not local_value:
        return "-"
    return local_value.strftime(fmt)


def seed_data():
    for key, value in LOGIN_SECURITY_SETTING_KEYS.items():
        if not db.session.get(SystemSetting, key):
            db.session.add(SystemSetting(key=key, value=value))

    for code, name, description in ALL_PERMISSIONS:
        if not Permission.query.filter_by(code=code).first():
            db.session.add(Permission(code=code, name=name, description=description))
    db.session.commit()

    def assign_permissions(role_name, description, codes):
        role = Role.query.filter_by(name=role_name).first()
        is_new_role = role is None
        if not role:
            role = Role(name=role_name, description=description)
            db.session.add(role)
        else:
            role.description = description
        if is_new_role or role_name == "管理员":
            role.permissions = Permission.query.filter(Permission.code.in_(codes)).all()
        db.session.commit()
        return role

    admin_role = assign_permissions("管理员", "全部权限", [item[0] for item in ALL_PERMISSIONS])
    assign_permissions("仓管员", "仓库作业", ["dashboard.view", "stock.in", "stock.out", "purchase.calc", "inventory.view", "inventory.transaction.view", "shipment.confirm"])
    assign_permissions("采购员", "采购作业", ["dashboard.view", "supplier.manage", "warehouse.manage", "asset.manage", "purchase.manage", "purchase.calc", "stock.in", "report.view"])
    assign_permissions("销售员", "销售作业", ["dashboard.view", "customer.manage", "sales.manage", "stock.out", "inventory.view", "report.view", "sku.mapping.manage", "shipment.manage", "fba.calc.manage", "fba.inbound.manage", "freight.compare.manage"])

    if not User.query.filter_by(username="admin").first():
        admin_password = os.getenv("ERP_INIT_ADMIN_PASSWORD", "").strip()
        if admin_password and len(admin_password) >= 8:
            admin = User(username="admin", full_name="系统管理员", is_active=True)
            admin.set_password(admin_password)
            admin.roles = [admin_role]
            db.session.add(admin)
            db.session.commit()
        elif not User.query.first():
            print("未创建管理员账号。请设置 ERP_INIT_ADMIN_PASSWORD（至少 8 位）后重新初始化。")

    emergency_user = User.query.filter_by(is_emergency_account=True).first()
    if not emergency_user:
        emergency_username = os.getenv("ERP_EMERGENCY_ADMIN_USERNAME", "local_admin").strip() or "local_admin"
        emergency_password = os.getenv("ERP_EMERGENCY_ADMIN_PASSWORD", "").strip()
        generated_password = ""
        if len(emergency_password) < 8:
            generated_password = secrets.token_urlsafe(12)
            emergency_password = generated_password
        if User.query.filter_by(username=emergency_username).first():
            emergency_username = f"local_admin_{secrets.token_hex(2)}"
        emergency_user = User(
            username=emergency_username,
            full_name="本地应急管理员",
            is_active=True,
            is_emergency_account=True,
        )
        emergency_user.set_password(emergency_password)
        emergency_user.roles = [admin_role]
        db.session.add(emergency_user)
        db.session.commit()
        if generated_password:
            print(
                f"本地应急管理员已创建：用户名 {emergency_username}，初始密码 {generated_password}。"
                "该账号仅允许从应急白名单 IP 登录，请尽快修改密码。"
            )

    if not Warehouse.query.first():
        db.session.add(Warehouse(name="主仓库", code="MAIN", address="默认地址", manager="仓库主管", remark="默认仓库"))
        db.session.commit()

    if not Supplier.query.first():
        db.session.add(Supplier(name="示例供应商", contact_name="供应商负责人", phone="13800000000", address="广州", remark="初始化数据"))
    if not Customer.query.first():
        db.session.add(Customer(name="示例客户", contact_name="采购员", phone="13900000000", address="广州", remark="初始化数据"))
    if not SKU.query.first():
        db.session.add_all(
            [
                SKU(sku_code="TSHIRT-BLACK-M", barcode="690000000001", name="基础T恤", category="上衣", color="黑色", size="M", unit="件", cost_price=Decimal("35.00"), sale_price=Decimal("89.00"), safety_stock=20),
                SKU(sku_code="HOODIE-GRAY-L", barcode="690000000002", name="连帽卫衣", category="外套", color="灰色", size="L", unit="件", cost_price=Decimal("78.00"), sale_price=Decimal("169.00"), safety_stock=10),
            ]
        )
    db.session.commit()

    default_warehouse = get_default_warehouse()
    if default_warehouse and not InventoryTransaction.query.first():
        for sku in SKU.query.all():
            db.session.add(
                InventoryTransaction(
                    sku_id=sku.id,
                    warehouse_id=default_warehouse.id,
                    quantity=30,
                    transaction_type="initial_inbound",
                    reference_type="system",
                    reference_no="INIT",
                    note="期初库存",
                )
            )
        db.session.commit()

    sync_inventory_balances()


@app.route("/")
def index():
    if g.user:
        return redirect(url_for(get_default_landing_endpoint(g.user)))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        now = datetime.utcnow()

        if not username or not password:
            flash("用户名和密码不能为空。", "error")
            return render_template("login.html")

        user = User.query.filter_by(username=username).first()

        if user and user.locked_until and user.locked_until <= now:
            reset_user_login_security(user)
            db.session.commit()

        if user and is_user_locked(user, now):
            record_login_audit_event(
                username,
                user=user,
                is_success=False,
                failure_reason="账号已锁定",
                locked_until=user.locked_until,
            )
            db.session.commit()
            flash(f"账号已锁定，预计 {format_dt(user.locked_until)} 后可重试。", "error")
            return render_template("login.html")

        if not user:
            record_login_audit_event(username, is_success=False, failure_reason="用户不存在")
            db.session.commit()
            flash("用户名或密码错误。", "error")
            return render_template("login.html")

        if not user.is_active:
            record_login_audit_event(username, user=user, is_success=False, failure_reason="账号已停用")
            db.session.commit()
            flash("账号已停用。", "error")
            return render_template("login.html")

        if user.is_emergency_account and not is_emergency_request_allowed():
            record_login_audit_event(username, user=user, is_success=False, failure_reason="IP 不在应急白名单")
            db.session.commit()
            flash("应急管理员仅允许从白名单 IP 登录。", "error")
            return render_template("login.html")

        if user.check_password(password):
            clear_login_failures(username)
            reset_user_login_security(user)
            session.clear()
            session["user_id"] = user.id
            session["_csrf_token"] = secrets.token_urlsafe(32)
            session["_last_seen_at"] = time.time()
            start_login_audit(user)
            flash("登录成功。", "success")
            return redirect(url_for(get_default_landing_endpoint(user)))

        user.failed_login_attempts = (user.failed_login_attempts or 0) + 1
        user.last_failed_login_at = now
        lock_triggered = False
        if user.failed_login_attempts >= get_login_lock_threshold():
            user.locked_until = now + timedelta(minutes=get_login_lock_duration_minutes())
            lock_triggered = True
        record_login_failure(username)
        record_login_audit_event(
            username,
            user=user,
            is_success=False,
            failure_reason="密码错误",
            locked_until=user.locked_until,
            lock_triggered=lock_triggered,
        )
        db.session.commit()
        if lock_triggered:
            flash(f"密码错误次数已达上限，账号已锁定至 {format_dt(user.locked_until)}。", "error")
        else:
            remaining_attempts = max(get_login_lock_threshold() - user.failed_login_attempts, 0)
            flash(f"用户名或密码错误。再错 {remaining_attempts} 次将锁定账号。", "error")

    return render_template("login.html")


@app.post("/logout")
def logout():
    finish_login_audit("主动退出")
    session.clear()
    flash("已安全退出登录。", "success")
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
@permission_required("dashboard.view")
def dashboard():
    sku_keyword = request.args.get("sku_keyword", "").strip()
    inventory_map = get_inventory_map()
    inventory_value = Decimal("0")
    low_stock = []
    for sku in SKU.query.order_by(SKU.name.asc()).all():
        current_qty = inventory_map.get(sku.id, 0)
        inventory_value += (sku.cost_price or 0) * current_qty
        if current_qty <= sku.safety_stock:
            low_stock.append({"sku": sku, "quantity": current_qty})
    fixed_asset_value = db.session.query(func.coalesce(func.sum(Asset.asset_value), 0)).scalar() or Decimal("0")
    outbound_rankings = get_dashboard_sku_rankings(
        transaction_types=("sales_outbound", "manual_outbound"),
        quantity_expression=-InventoryTransaction.quantity,
        sku_keyword=sku_keyword,
        user=g.user,
    )
    inbound_rankings = get_dashboard_sku_rankings(
        transaction_types=("purchase_inbound", "manual_inbound", "initial_inbound"),
        quantity_expression=InventoryTransaction.quantity,
        sku_keyword=sku_keyword,
        user=g.user,
    )
    return render_template(
        "dashboard.html",
        stats={
            "sku_count": SKU.query.count(),
            "purchase_count": PurchaseOrder.query.count(),
            "sales_count": SalesOrder.query.count(),
            "stock_qty": sum(inventory_map.values()),
            "inventory_value": inventory_value,
            "fixed_asset_value": fixed_asset_value,
        },
        recent_purchases=PurchaseOrder.query.order_by(PurchaseOrder.created_at.desc()).limit(5).all(),
        recent_sales=SalesOrder.query.order_by(SalesOrder.created_at.desc()).limit(5).all(),
        recent_repairs=DefectRepair.query.order_by(DefectRepair.created_at.desc()).limit(5).all(),
        low_stock=low_stock[:8],
        sku_keyword=sku_keyword,
        outbound_rankings=outbound_rankings,
        inbound_rankings=inbound_rankings,
    )


@app.route("/system-info")
@login_required
@permission_required("user.manage")
def system_info():
    return render_template(
        "system_info.html",
        login_security_config={
            "login_lock_threshold": get_login_lock_threshold(),
            "login_lock_duration_minutes": get_login_lock_duration_minutes(),
            "login_lock_duration_label": get_login_lock_duration_label(),
            "emergency_admin_ip_whitelist": ", ".join(get_emergency_admin_ip_whitelist()),
        },
        system_info_items=[
            ("系统名称", "苏客 ERP"),
            ("当前版本", os.getenv("ERP_APP_VERSION", "2026.04.03")),
            ("部署方式", get_deployment_mode()),
            ("数据库位置", get_database_location()),
            ("数据库连接", app.config["SQLALCHEMY_DATABASE_URI"]),
            ("Python 版本", sys.version.split()[0]),
            ("Flask 版本", get_package_version("Flask")),
            ("SQLAlchemy 版本", get_package_version("SQLAlchemy")),
            ("渲染时间", format_dt(datetime.utcnow(), "%Y-%m-%d %H:%M:%S")),
            ("根目录", BASE_DIR),
        ],
    )


@app.post("/system-info/login-security")
@login_required
@permission_required("user.manage")
def system_info_update_login_security():
    login_lock_threshold = max(parse_int(request.form.get("login_lock_threshold"), DEFAULT_LOGIN_LOCK_THRESHOLD), 1)
    login_lock_duration_minutes = max(parse_int(request.form.get("login_lock_duration_minutes"), DEFAULT_LOGIN_LOCK_DURATION_MINUTES), 1)
    emergency_admin_ip_whitelist = ",".join(
        item for item in [normalize_ip_value(value) for value in request.form.get("emergency_admin_ip_whitelist", "").replace("\n", ",").split(",")] if item
    ) or "127.0.0.1,::1"
    set_system_setting("login_lock_threshold", login_lock_threshold)
    set_system_setting("login_lock_duration_minutes", login_lock_duration_minutes)
    set_system_setting("emergency_admin_ip_whitelist", emergency_admin_ip_whitelist)
    db.session.commit()
    log_operation(
        "绯荤粺閰嶇疆",
        "鏇存柊",
        "system_setting",
        "login_security",
        f"失败阈值：{login_lock_threshold}；锁定时长：{login_lock_duration_minutes} 分钟；应急白名单：{emergency_admin_ip_whitelist}",
    )
    flash("登录安全配置已更新。", "success")
    return redirect(url_for("system_info"))


def render_sales_order_page(*, sales_form=None):
    sku_keyword, start_date, end_date = get_date_filters()
    orders_query = SalesOrder.query
    orders_query = apply_order_query_sku_scope(orders_query, SalesOrderItem, SalesOrder.items)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        orders_query = (
            orders_query.join(SalesOrder.items)
            .join(SalesOrderItem.sku)
            .filter(or_(sku_filter, SalesOrder.order_no.ilike(f"%{sku_keyword}%")))
            .distinct()
        )
    orders_query = apply_datetime_range(orders_query, SalesOrder.created_at, start_date, end_date)
    visible_skus = apply_sku_scope(SKU.query, SKU.id).order_by(SKU.name.asc()).all()
    orders = orders_query.order_by(SalesOrder.created_at.desc()).all()
    sales_summary = {
        "order_count": len(orders),
        "sku_count": sum(len(order.items) for order in orders),
        "quantity": sum(item.quantity for order in orders for item in order.items),
        "amount": float(sum(item.quantity * (item.unit_price or 0) for order in orders for item in order.items)),
        "shipped_count": sum(1 for order in orders if order.status == "shipped"),
    }
    return render_template(
        "sales_orders.html",
        orders=orders,
        customers=Customer.query.order_by(Customer.name.asc()).all(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        skus=visible_skus,
        inventory_by_warehouse=get_inventory_by_warehouse(),
        sales_summary=sales_summary,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
        sales_form=sales_form or {},
    )


def render_inventory_out_page(*, outbound_form=None):
    sku_keyword, start_date, end_date = get_date_filters()
    txns_query = (
        InventoryTransaction.query.join(InventoryTransaction.sku)
        .filter(InventoryTransaction.quantity < 0, InventoryTransaction.transaction_type == "manual_outbound")
    )
    txns_query = apply_sku_scope(txns_query, InventoryTransaction.sku_id)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        txns_query = txns_query.filter(or_(sku_filter, InventoryTransaction.reference_no.ilike(f"%{sku_keyword}%")))
    txns_query = apply_datetime_range(txns_query, InventoryTransaction.created_at, start_date, end_date)
    recent_txns = txns_query.order_by(InventoryTransaction.created_at.desc()).limit(200).all()
    visible_skus = apply_sku_scope(SKU.query, SKU.id).order_by(SKU.name.asc()).all()
    txn_groups = group_manual_outbound_transactions(recent_txns)
    outbound_summary = {
        "order_count": len(txn_groups),
        "line_count": sum(len(group["lines"]) for group in txn_groups),
        "quantity": abs(sum(txn.quantity for txn in recent_txns)),
        "warehouse_count": len({txn.warehouse_id for txn in recent_txns}),
    }
    return render_template(
        "inventory_out.html",
        skus=visible_skus,
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        txns=recent_txns,
        txn_groups=txn_groups,
        outbound_summary=outbound_summary,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
        outbound_form=outbound_form or {},
    )


@app.route("/login-audits")
@login_required
@permission_required("user.manage")
def login_audit_list():
    login_audits = LoginAudit.query.order_by(LoginAudit.login_at.desc()).limit(500).all()
    return render_template("login_audits.html", login_audits=login_audits)


@app.route("/operation-logs")
@login_required
@admin_required
def operation_log_list():
    category = request.args.get("category", "").strip()
    logs_query = OperationLog.query
    if category:
        logs_query = logs_query.filter(OperationLog.category == category)
    logs = logs_query.order_by(OperationLog.created_at.desc()).limit(500).all()
    categories = [row[0] for row in db.session.query(OperationLog.category).distinct().order_by(OperationLog.category.asc()).all() if row[0]]
    return render_template("operation_logs.html", logs=logs, categories=categories, selected_category=category)


@app.route("/permissions")
@login_required
@permission_required("user.manage")
def permission_list():
    return render_template("permissions.html", permissions=Permission.query.order_by(Permission.code.asc()).all())


@app.route("/roles", methods=["GET", "POST"])
@login_required
@permission_required("user.manage")
def role_list():
    if request.method == "POST":
        role_name = request.form.get("name", "").strip()
        description = request.form.get("description", "").strip()
        permission_ids = [parse_int(value) for value in request.form.getlist("permission_ids") if value]
        permissions = Permission.query.filter(Permission.id.in_(permission_ids)).all() if permission_ids else []

        if not role_name:
            flash("角色名称不能为空。", "error")
        elif Role.query.filter_by(name=role_name).first():
            flash("角色名称已存在。", "error")
        else:
            role = Role(name=role_name, description=description)
            role.permissions = permissions
            db.session.add(role)
            log_operation("角色", "新增", "role", role.name, f"权限数：{len(permissions)}")
            db.session.commit()
            flash("角色创建成功。", "success")
        return redirect(url_for("role_list"))

    return render_template(
        "roles.html",
        roles=Role.query.order_by(Role.created_at.desc()).all(),
        permissions=Permission.query.order_by(Permission.code.asc()).all(),
    )


@app.post("/roles/<int:role_id>/update")
@login_required
@permission_required("user.manage")
def role_update(role_id):
    role = db.session.get(Role, role_id)
    if not role:
        flash("角色不存在。", "error")
        return redirect(url_for("role_list"))

    role_name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()
    permission_ids = [parse_int(value) for value in request.form.getlist("permission_ids") if value]
    permissions = Permission.query.filter(Permission.id.in_(permission_ids)).all() if permission_ids else []
    existing = Role.query.filter(Role.name == role_name, Role.id != role_id).first()

    if not role_name:
        flash("角色名称不能为空。", "error")
    elif existing:
        flash("角色名称已存在。", "error")
    elif would_role_update_remove_last_user_manager(role, permissions):
        flash("必须至少保留一个具备用户管理权限的角色，当前修改会移除最后一个用户管理权限。", "error")
    else:
        role.name = role_name
        role.description = description
        role.permissions = permissions
        db.session.commit()
        log_operation("角色", "更新", "role", role.name, f"权限数：{len(permissions)}")
        flash("角色已更新。", "success")
    return redirect(url_for("role_list"))


@app.post("/roles/<int:role_id>/delete")
@login_required
@permission_required("user.manage")
def role_delete(role_id):
    denied = require_delete_role()
    if denied:
        return denied
    role = db.session.get(Role, role_id)
    if not role:
        flash("角色不存在。", "error")
    elif User.query.filter(User.roles.any(Role.id == role.id)).first():
        flash("该角色仍被用户使用，请先从相关用户身上移除。", "error")
    elif would_role_delete_remove_last_user_manager(role):
        flash("必须至少保留一个具备用户管理权限的角色，当前删除会移除最后一个用户管理权限。", "error")
    else:
        log_operation("角色", "删除", "role", role.name, f"权限数：{len(role.permissions)}")
        db.session.delete(role)
        db.session.commit()
        flash("角色已删除。", "success")
    return redirect(url_for("role_list"))


@app.route("/users", methods=["GET", "POST"])
@login_required
@permission_required("user.manage")
def user_list():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        full_name = request.form.get("full_name", "").strip()
        is_active = request.form.get("is_active") == "on"
        role_ids = [parse_int(value) for value in request.form.getlist("role_ids") if value]
        roles = Role.query.filter(Role.id.in_(role_ids)).all() if role_ids else []
        authorized_skus = None
        if "authorized_sku_ids" in request.form:
            authorized_sku_ids = [parse_int(value) for value in request.form.getlist("authorized_sku_ids") if parse_int(value)]
            authorized_skus = SKU.query.filter(SKU.id.in_(authorized_sku_ids)).all() if authorized_sku_ids else []

        if not username or not full_name or not password:
            flash("用户名、姓名和密码不能为空。", "error")
        elif len(password) < 8:
            flash("密码长度不能少于 8 位。", "error")
        elif not roles:
            flash("请至少分配一个角色。", "error")
        elif User.query.filter_by(username=username).first():
            flash("用户名已存在。", "error")
        else:
            user = User(username=username, full_name=full_name, is_active=is_active)
            user.set_password(password)
            user.roles = roles
            if authorized_skus is not None:
                user.authorized_skus = authorized_skus
            db.session.add(user)
            log_operation("用户", "新增", "user", username, f"姓名：{full_name}；角色：{user.role_names}")
            db.session.commit()
            flash("用户创建成功。", "success")
        return redirect(url_for("user_list"))

    return render_template(
        "users.html",
        users=User.query.order_by(User.created_at.desc()).all(),
        roles=Role.query.order_by(Role.name.asc()).all(),
    )


@app.post("/users/<int:user_id>/update")
@login_required
@permission_required("user.manage")
def user_update(user_id):
    user = db.session.get(User, user_id)
    if not user:
        flash("用户不存在。", "error")
        return redirect(url_for("user_list"))

    username = request.form.get("username", "").strip()
    full_name = request.form.get("full_name", "").strip()
    is_active = request.form.get("is_active") == "on"
    role_ids = [parse_int(value) for value in request.form.getlist("role_ids") if value]
    roles = Role.query.filter(Role.id.in_(role_ids)).all() if role_ids else []
    authorized_skus = None
    if "authorized_sku_ids" in request.form:
        authorized_sku_ids = [parse_int(value) for value in request.form.getlist("authorized_sku_ids") if parse_int(value)]
        authorized_skus = SKU.query.filter(SKU.id.in_(authorized_sku_ids)).all() if authorized_sku_ids else []
    existing = User.query.filter(User.username == username, User.id != user_id).first()

    if not username or not full_name:
        flash("用户名和姓名不能为空。", "error")
    elif not roles:
        flash("请至少分配一个角色。", "error")
    elif existing:
        flash("用户名已存在。", "error")
    elif user.is_emergency_account and not is_active:
        flash("本地应急管理员不能被停用。", "error")
    elif would_remove_last_user_manager(user, roles, is_active):
        flash("必须至少保留一个启用状态且拥有用户管理权限的账号。", "error")
    else:
        user.username = username
        user.full_name = full_name
        user.is_active = is_active
        user.roles = roles
        if authorized_skus is not None:
            user.authorized_skus = authorized_skus
        db.session.commit()
        flash("用户信息已更新。", "success")
    return redirect(url_for("user_list"))


@app.post("/users/<int:user_id>/password")
@login_required
@permission_required("user.manage")
def user_change_password(user_id):
    user = db.session.get(User, user_id)
    new_password = request.form.get("new_password", "").strip()
    if not user:
        flash("用户不存在。", "error")
    elif len(new_password) < 8:
        flash("密码长度不能少于 8 位。", "error")
    else:
        user.set_password(new_password)
        reset_user_login_security(user)
        clear_login_failures(user.username)
        db.session.commit()
        flash("密码重置成功。", "success")
    return redirect(url_for("user_list"))


@app.post("/users/<int:user_id>/unlock")
@login_required
@permission_required("user.manage")
def user_unlock(user_id):
    user = db.session.get(User, user_id)
    if not user:
        flash("用户不存在。", "error")
    else:
        reset_user_login_security(user)
        clear_login_failures(user.username)
        db.session.commit()
        log_operation("用户", "解锁", "user", user.username, f"姓名：{user.full_name}")
        flash("账号已解锁。", "success")
    return redirect(request.referrer or url_for("user_list"))


@app.post("/users/<int:user_id>/delete")
@login_required
@permission_required("user.manage")
def user_delete(user_id):
    denied = require_delete_role()
    if denied:
        return denied
    user = db.session.get(User, user_id)
    if not user:
        flash("用户不存在。", "error")
    elif g.user and g.user.id == user.id:
        flash("不能删除当前登录用户。", "error")
    elif user.is_emergency_account:
        flash("本地应急管理员不能删除。", "error")
    elif user.is_active and user.has_permission("user.manage") and count_active_users_with_permission("user.manage", exclude_user_id=user.id) == 0:
        flash("必须至少保留一个启用状态且拥有用户管理权限的账号。", "error")
    else:
        log_operation("用户", "删除", "user", user.username, f"姓名：{user.full_name}；角色：{user.role_names}")
        db.session.delete(user)
        db.session.commit()
        flash("用户已删除。", "success")
    return redirect(url_for("user_list"))


def get_freight_quote_batch_payload(batch):
    if not batch:
        return None
    try:
        input_payload = json.loads(batch.input_json or "{}")
    except json.JSONDecodeError:
        input_payload = {}
    try:
        summary = json.loads(batch.summary_json or "{}")
    except json.JSONDecodeError:
        summary = {}
    return {
        "batch": batch,
        "items": input_payload.get("items", []),
        "summary": summary,
        "options": batch.options,
    }


def render_freight_compare_page(*, quote_batch=None, quote_form=None):
    if quote_batch is None:
        batch_id = parse_int(session.get("_freight_quote_batch_id"))
        quote_batch = db.session.get(FreightQuoteBatch, batch_id) if batch_id else None
    rate_books = FreightRateBook.query.order_by(FreightRateBook.created_at.desc()).limit(20).all()
    active_rate_books = [book for book in rate_books if book.status == "active"]
    recent_quotes = FreightQuoteBatch.query.order_by(FreightQuoteBatch.created_at.desc()).limit(10).all()
    rate_book_summaries = {book.id: load_freight_rate_book_summary(book) for book in rate_books}
    warehouse_postal_map = warehouse_postal_mapping()
    warehouse_postal_rows = {
        code: {"warehouse_code": code, "postal_code": postal_code, "source": "已导入价格表", "city": "", "state": "", "note": ""}
        for code, postal_code in warehouse_postal_map.items()
    }
    for row in WarehousePostalCode.query.order_by(WarehousePostalCode.warehouse_code.asc()).all():
        warehouse_postal_rows[row.warehouse_code.upper()] = row
    warehouse_postal_codes = list(warehouse_postal_rows.values())[:300]
    return render_template(
        "freight_compare.html",
        rate_books=rate_books,
        active_rate_books=active_rate_books,
        rate_book_summaries=rate_book_summaries,
        warehouse_postal_codes=warehouse_postal_codes,
        warehouse_postal_map=warehouse_postal_map,
        quote_payload=get_freight_quote_batch_payload(quote_batch),
        recent_quotes=recent_quotes,
        quote_form=quote_form or {},
    )


@app.get("/freight-compare")
@login_required
@permission_required("freight.compare.manage")
def freight_compare():
    return render_freight_compare_page()


@app.post("/freight-compare/rate-books")
@login_required
@permission_required("freight.compare.manage")
def freight_rate_book_upload():
    upload = request.files.get("rate_file")
    if not upload or not (upload.filename or "").strip():
        flash("请选择要上传的货代价格表。", "error")
        return redirect(url_for("freight_compare"))
    original_filename = upload.filename.strip()
    extension = os.path.splitext(original_filename.lower())[1]
    if extension not in ALLOWED_FREIGHT_RATE_EXTENSIONS:
        flash("价格表仅支持 .xlsx 或 .xlsm 文件。", "error")
        return redirect(url_for("freight_compare"))

    forwarder_name = normalize_freight_forwarder(request.form.get("forwarder_name", ""), original_filename)
    version_label = request.form.get("version_label", "").strip() or datetime.utcnow().strftime("%Y-%m-%d")
    os.makedirs(FREIGHT_RATE_BOOK_DIR, exist_ok=True)
    safe_name = secure_filename(original_filename) or f"rate-book{extension}"
    stored_filename = f"{datetime.utcnow():%Y%m%d%H%M%S}-{secrets.token_hex(4)}-{safe_name}"
    stored_path = os.path.join(FREIGHT_RATE_BOOK_DIR, stored_filename)
    upload.save(stored_path)

    rate_book = FreightRateBook(
        forwarder_name=forwarder_name,
        version_label=version_label,
        original_filename=original_filename,
        stored_filename=stored_filename,
        clothing_surcharge_per_kg=Decimal("1") if forwarder_name == "鼎邦" else Decimal("0"),
        uploaded_by_user_id=g.user.id if g.user else None,
    )
    try:
        parsed = parse_freight_rate_workbook(stored_path, rate_book)
        if not parsed["rates"]:
            raise ValueError("未解析到可用价格行，请检查是否为盈和、开渠达、鼎邦或鸿帆的美线价格表。")
        FreightRateBook.query.filter(
            FreightRateBook.forwarder_name == forwarder_name,
            FreightRateBook.status == "active",
        ).update({"status": "archived"})
        learned_postal_count = learn_warehouse_postal_codes_from_rates(parsed["rates"], f"{forwarder_name} {version_label}")
        db.session.add(rate_book)
        db.session.add_all(parsed["rates"] + parsed["customs_rules"] + parsed["remote_zips"])
        rate_book.parsed_summary = freight_json_dumps(parsed["summary"])
        rate_book.status = "active"
        log_operation("货代比价", "导入", "freight_rate_book", forwarder_name, f"版本：{version_label}；价格行：{parsed['summary']['rate_count']}；仓库邮编：{learned_postal_count}")
        db.session.commit()
        flash(
            f"{forwarder_name} 价格表已导入：{parsed['summary']['rate_count']} 条价格、{parsed['summary']['customs_rule_count']} 条报关规则，自动学习 {learned_postal_count} 个仓库邮编。",
            "success",
        )
    except Exception as error:
        db.session.rollback()
        try:
            os.remove(stored_path)
        except OSError:
            pass
        flash(f"价格表解析失败：{error}", "error")
    return redirect(url_for("freight_compare"))


@app.post("/freight-compare/rate-books/<int:rate_book_id>/archive")
@login_required
@permission_required("freight.compare.manage")
def freight_rate_book_archive(rate_book_id):
    rate_book = db.session.get(FreightRateBook, rate_book_id)
    if not rate_book:
        flash("价格表版本不存在。", "error")
    else:
        rate_book.status = "archived" if rate_book.status == "active" else "active"
        log_operation("货代比价", "更新", "freight_rate_book", rate_book.forwarder_name, f"版本：{rate_book.version_label}；状态：{rate_book.status}")
        db.session.commit()
        flash(f"{rate_book.forwarder_name} {rate_book.version_label} 已{rate_book.status_label}。", "success")
    return redirect(url_for("freight_compare"))


@app.post("/freight-compare/rate-books/<int:rate_book_id>/surcharge")
@login_required
@permission_required("freight.compare.manage")
def freight_rate_book_surcharge(rate_book_id):
    rate_book = db.session.get(FreightRateBook, rate_book_id)
    if not rate_book:
        flash("价格表版本不存在。", "error")
        return redirect(url_for("freight_compare"))

    surcharge = parse_decimal(request.form.get("clothing_surcharge_per_kg"), "0").quantize(Decimal("0.01"))
    if surcharge < 0:
        surcharge = Decimal("0")
    rate_book.clothing_surcharge_per_kg = surcharge
    log_operation("货代比价", "更新", "freight_rate_book", rate_book.forwarder_name, f"版本：{rate_book.version_label}；服装附加费：{surcharge}/kg")
    db.session.commit()
    flash(f"{rate_book.forwarder_name} {rate_book.version_label or '-'} 服装附加费已更新为 {surcharge}/kg。", "success")
    return redirect(url_for("freight_compare"))


@app.post("/freight-compare/rate-books/<int:rate_book_id>/delete")
@login_required
@permission_required("freight.compare.manage")
def freight_rate_book_delete(rate_book_id):
    denied = require_delete_role()
    if denied:
        return denied
    rate_book = db.session.get(FreightRateBook, rate_book_id)
    if not rate_book:
        flash("价格表版本不存在。", "error")
        return redirect(url_for("freight_compare"))

    forwarder_name = rate_book.forwarder_name
    version_label = rate_book.version_label or "-"
    stored_path = os.path.join(FREIGHT_RATE_BOOK_DIR, rate_book.stored_filename) if rate_book.stored_filename else ""
    rate_count = FreightRate.query.filter_by(rate_book_id=rate_book.id).count()
    customs_rule_count = FreightCustomsRule.query.filter_by(rate_book_id=rate_book.id).count()
    remote_zip_count = FreightRemoteZip.query.filter_by(rate_book_id=rate_book.id).count()
    try:
        FreightRate.query.filter_by(rate_book_id=rate_book.id).delete(synchronize_session=False)
        FreightCustomsRule.query.filter_by(rate_book_id=rate_book.id).delete(synchronize_session=False)
        FreightRemoteZip.query.filter_by(rate_book_id=rate_book.id).delete(synchronize_session=False)
        log_operation(
            "货代比价",
            "删除",
            "freight_rate_book",
            forwarder_name,
            f"版本：{version_label}；价格行：{rate_count}；报关规则：{customs_rule_count}；偏远邮编：{remote_zip_count}",
        )
        db.session.delete(rate_book)
        db.session.commit()
        if stored_path and os.path.isfile(stored_path):
            try:
                os.remove(stored_path)
            except OSError:
                pass
        flash(f"{forwarder_name} {version_label} 价格表已删除。", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"价格表删除失败：{error}", "error")
    return redirect(url_for("freight_compare"))


@app.get("/freight-compare/warehouse-postal-template")
@login_required
@permission_required("freight.compare.manage")
def freight_warehouse_postal_template():
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(["仓库代码", "邮编", "城市", "州", "来源", "备注"])
    rows = WarehousePostalCode.query.order_by(WarehousePostalCode.warehouse_code.asc()).limit(1000).all()
    if rows:
        for row in rows:
            writer.writerow([row.warehouse_code, row.postal_code, row.city, row.state, row.source, row.note])
    else:
        writer.writerow(["ONT8", "92551", "Moreno Valley", "CA", "示例", "可删除示例行"])
    payload = BytesIO(stream.getvalue().encode("utf-8-sig"))
    return send_file(
        payload,
        as_attachment=True,
        download_name="fba_warehouse_postal_template.csv",
        mimetype="text/csv; charset=utf-8",
    )


@app.post("/freight-compare/warehouse-postals")
@login_required
@permission_required("freight.compare.manage")
def freight_warehouse_postal_upload():
    upload = request.files.get("warehouse_postal_file")
    if not upload or not (upload.filename or "").strip():
        flash("请选择要上传的仓库邮编表。", "error")
        return redirect(url_for("freight_compare"))
    try:
        rows = parse_warehouse_postal_upload(upload)
        imported_count = 0
        for row in rows:
            if upsert_warehouse_postal_code(
                row["warehouse_code"],
                row["postal_code"],
                city=row.get("city", ""),
                state=row.get("state", ""),
                source=row.get("source", "") or "手动上传",
                note=row.get("note", ""),
            ):
                imported_count += 1
        log_operation("货代比价", "导入", "warehouse_postal_code", "仓库邮编库", f"导入/更新：{imported_count}")
        db.session.commit()
        flash(f"仓库邮编库已导入/更新 {imported_count} 条。", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"仓库邮编表导入失败：{error}", "error")
    return redirect(url_for("freight_compare"))


@app.post("/freight-compare/quote")
@login_required
@permission_required("freight.compare.manage")
def freight_compare_quote():
    quote_form = {
        "origin_region": request.form.get("origin_region", "义乌").strip() or "义乌",
        "customs_required": request.form.get("customs_required") == "on",
        "inspection_required": request.form.get("inspection_required") == "on",
        "clothing_surcharge_enabled": request.form.get("clothing_surcharge_enabled") == "on",
        "item_name_count": request.form.get("item_name_count", "1").strip() or "1",
    }
    try:
        items = parse_freight_quote_form_items()
        items.extend(parse_freight_quote_upload(request.files.get("quote_file")))
        if not items:
            raise ValueError("请至少填写一行货件，或上传一份包含仓库代码和重量的货件表。")
        item_name_count = max(parse_int(quote_form["item_name_count"], 1), 1)
        result = calculate_freight_quote(
            items=items,
            origin_region=quote_form["origin_region"],
            customs_required=quote_form["customs_required"],
            inspection_required=quote_form["inspection_required"],
            item_name_count=item_name_count,
            clothing_surcharge_enabled=quote_form["clothing_surcharge_enabled"],
        )
        if not result["options"]:
            raise ValueError("没有找到能覆盖全部目的仓的渠道。可以补充邮编、确认起运仓，或上传更多货代价格表。")
        batch = save_freight_quote_batch(
            result,
            origin_region=quote_form["origin_region"],
            customs_required=quote_form["customs_required"],
            inspection_required=quote_form["inspection_required"],
            item_name_count=item_name_count,
        )
        session["_freight_quote_batch_id"] = batch.id
        if result["missing_items"]:
            flash(f"有 {len(result['missing_items'])} 行货件没有匹配到任何价格，已从本次比价中排除。", "error")
        flash(f"比价完成，生成 {len(result['options'])} 个报价选项。", "success")
        return render_freight_compare_page(quote_batch=batch, quote_form=quote_form)
    except ValueError as error:
        flash(str(error), "error")
        return render_freight_compare_page(quote_form=quote_form)


@app.get("/freight-compare/export")
@login_required
@permission_required("freight.compare.manage")
def freight_compare_export():
    batch_id = parse_int(request.args.get("batch_id")) or parse_int(session.get("_freight_quote_batch_id"))
    batch = db.session.get(FreightQuoteBatch, batch_id) if batch_id else None
    if not batch or not batch.options:
        flash("当前没有可导出的比价结果。", "error")
        return redirect(url_for("freight_compare"))
    rows = []
    for option in sorted(batch.options, key=lambda row: row.rank_no):
        rows.append(
            [
                batch.quote_no,
                option.rank_no,
                option.recommendation,
                option.forwarder_name,
                option.channel_name,
                option.origin_region,
                float(option.total_cost or 0),
                float(option.freight_cost or 0),
                float(option.surcharge_cost or 0),
                float(option.customs_cost or 0),
                float(option.chargeable_weight or 0),
                float((option.total_cost or 0) / option.chargeable_weight) if option.chargeable_weight else 0,
                f"{option.transit_days_min or '-'}-{option.transit_days_max or '-'}天",
                option.notes,
            ]
        )
    return export_workbook(
        f"货代比价结果-{batch.quote_no}.xlsx",
        "货代比价",
        ["比价单号", "排名", "推荐", "货代", "渠道", "起运仓", "总费用", "基础运费", "附加费", "报关费用", "计费重KG", "单公斤费用", "参考时效", "备注"],
        rows,
    )


@app.post("/freight-compare/quotes/<int:batch_id>/delete")
@login_required
@permission_required("freight.compare.manage")
def freight_quote_batch_delete(batch_id):
    denied = require_delete_role()
    if denied:
        return denied
    batch = db.session.get(FreightQuoteBatch, batch_id)
    if not batch:
        flash("历史比价记录不存在。", "error")
        return redirect(url_for("freight_compare"))
    quote_no = batch.quote_no
    option_count = len(batch.options)
    if parse_int(session.get("_freight_quote_batch_id")) == batch.id:
        session.pop("_freight_quote_batch_id", None)
    db.session.delete(batch)
    log_operation("货代比价", "删除", "freight_quote_batch", quote_no, f"报价选项：{option_count}")
    db.session.commit()
    flash(f"历史比价 {quote_no} 已删除。", "success")
    return redirect(url_for("freight_compare"))


@app.route("/skus", methods=["GET", "POST"])
@login_required
@permission_required("sku.manage")
def sku_list():
    if request.method == "POST":
        barcode = request.form.get("barcode", "").strip() or None
        allow_price = can_view_prices()
        sku = SKU(
            sku_code=request.form.get("sku_code", "").strip(),
            barcode=barcode,
            name=request.form.get("name", "").strip(),
            category=request.form.get("category", "").strip() or "Apparel",
            color=request.form.get("color", "").strip(),
            size=request.form.get("size", "").strip(),
            unit=request.form.get("unit", "").strip() or "pcs",
            cost_price=parse_decimal(request.form.get("cost_price")) if allow_price else Decimal("0"),
            sale_price=parse_decimal(request.form.get("sale_price")) if allow_price else Decimal("0"),
            safety_stock=max(parse_int(request.form.get("safety_stock"), 0), 0),
            remark=request.form.get("remark", "").strip(),
        )
        if not sku.sku_code or not sku.name:
            flash("SKU 编码和商品名称不能为空。", "error")
        elif sku.cost_price < 0 or sku.sale_price < 0:
            flash("价格不能为负数。", "error")
        elif SKU.query.filter_by(sku_code=sku.sku_code).first():
            flash("SKU 编码已存在。", "error")
        elif sku.barcode and SKU.query.filter_by(barcode=sku.barcode).first():
            flash("条形码已存在。", "error")
        else:
            db.session.add(sku)
            log_operation("SKU", "新增", "sku", sku.sku_code, f"{sku.name} / {sku.color or '-'} / {sku.size or '-'}")
            db.session.commit()
            flash("SKU 创建成功。", "success")
        return redirect(url_for("sku_list"))

    sku_keyword = request.args.get("sku_keyword", "").strip()
    selected_sales_user_id = parse_int(request.args.get("sales_user_id"))
    authorization_state = request.args.get("authorization_state", "").strip()
    sku_query = SKU.query
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        sku_query = sku_query.filter(sku_filter)
    if authorization_state == "unassigned":
        sku_query = sku_query.filter(~SKU.authorized_users.any())
    if selected_sales_user_id:
        sku_query = sku_query.filter(SKU.authorized_users.any(User.id == selected_sales_user_id))
    skus = sku_query.order_by(SKU.created_at.desc()).all()
    assignable_user_ids = {
        user.id
        for sku in skus
        for user in sku.authorized_users
    }
    return render_template(
        "skus.html",
        skus=skus,
        inventory_map=get_inventory_map(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        inventory_by_warehouse=get_inventory_by_warehouse(),
        sku_keyword=sku_keyword,
        selected_sales_user_id=selected_sales_user_id,
        authorization_state=authorization_state,
        sales_users=get_sku_assignable_users(assignable_user_ids),
    )


@app.post("/skus/<int:sku_id>/delete")
@login_required
@permission_required("sku.manage")
def sku_delete(sku_id):
    denied = require_delete_role()
    if denied:
        return denied
    sku = db.session.get(SKU, sku_id)
    if not sku:
        flash("SKU 不存在。", "error")
    elif entity_has_rows(PurchaseOrderItem, sku_id=sku.id) or entity_has_rows(SalesOrderItem, sku_id=sku.id):
        flash("该 SKU 已被采购单或销售单引用，请先删除相关单据。", "error")
    elif entity_has_rows(DefectRepair, sku_id=sku.id):
        flash("该 SKU 已被返修记录引用，请先删除相关返修记录。", "error")
    elif entity_has_rows(InventoryTransaction, sku_id=sku.id):
        flash("该 SKU 仍存在库存流水，请先删除相关入库或出库记录。", "error")
    else:
        InventoryBalance.query.filter_by(sku_id=sku.id).delete()
        log_operation("SKU", "删除", "sku", sku.sku_code, f"{sku.name} / {sku.color or '-'} / {sku.size or '-'}")
        db.session.delete(sku)
        db.session.commit()
        flash("SKU 已删除。", "success")
    return redirect(url_for("sku_list"))


@app.post("/skus/<int:sku_id>/update")
@login_required
@permission_required("sku.manage")
def sku_update(sku_id):
    sku = db.session.get(SKU, sku_id)
    if not sku:
        flash("SKU 不存在。", "error")
        return redirect(
            url_for(
                "sku_list",
                sku_keyword=request.args.get("sku_keyword", "").strip(),
                sales_user_id=parse_int(request.args.get("sales_user_id")),
                authorization_state=request.args.get("authorization_state", "").strip(),
            )
        )

    barcode = request.form.get("barcode", "").strip() or None
    sku_code = request.form.get("sku_code", "").strip()
    name = request.form.get("name", "").strip()
    allow_price = can_view_prices()
    cost_price = parse_decimal(request.form.get("cost_price")) if allow_price else sku.cost_price
    sale_price = parse_decimal(request.form.get("sale_price")) if allow_price else sku.sale_price
    existing_sku_code = SKU.query.filter(SKU.sku_code == sku_code, SKU.id != sku.id).first()
    existing_barcode = SKU.query.filter(SKU.barcode == barcode, SKU.id != sku.id).first() if barcode else None

    if not sku_code or not name:
        flash("SKU 编码和商品名称不能为空。", "error")
    elif cost_price < 0 or sale_price < 0:
        flash("价格不能为负数。", "error")
    elif existing_sku_code:
        flash("SKU 编码已存在。", "error")
    elif existing_barcode:
        flash("条形码已存在。", "error")
    else:
        sku.sku_code = sku_code
        sku.barcode = barcode
        sku.name = name
        sku.category = request.form.get("category", "").strip() or "Apparel"
        sku.color = request.form.get("color", "").strip()
        sku.size = request.form.get("size", "").strip()
        sku.unit = request.form.get("unit", "").strip() or "pcs"
        if allow_price:
            sku.cost_price = cost_price
            sku.sale_price = sale_price
        sku.safety_stock = max(parse_int(request.form.get("safety_stock"), 0), 0)
        sku.remark = request.form.get("remark", "").strip()
        db.session.commit()
        flash("SKU 信息已更新。", "success")
    return redirect(
        url_for(
            "sku_list",
            sku_keyword=request.args.get("sku_keyword", "").strip(),
            sales_user_id=parse_int(request.args.get("sales_user_id")),
            authorization_state=request.args.get("authorization_state", "").strip(),
        )
    )


@app.post("/skus/<int:sku_id>/permissions")
@login_required
@permission_required("sku.manage")
def sku_update_permissions(sku_id):
    sku = db.session.get(SKU, sku_id)
    if not sku:
        flash("SKU 不存在。", "error")
        return redirect(url_for("sku_list"))

    user_ids = [parse_int(value) for value in request.form.getlist("authorized_user_ids") if parse_int(value)]
    sales_users = get_sku_assignable_users(user_ids) if user_ids else []
    if user_ids:
        allowed_user_ids = {user.id for user in sales_users}
        sales_users = [user for user in sales_users if user.id in allowed_user_ids and user.id in set(user_ids)]
    sku.authorized_users = sales_users
    db.session.commit()
    flash("SKU 销售授权已更新。", "success")
    return redirect(
        url_for(
            "sku_list",
            sku_keyword=request.args.get("sku_keyword", "").strip(),
            sales_user_id=parse_int(request.args.get("sales_user_id")),
            authorization_state=request.args.get("authorization_state", "").strip(),
        )
    )


@app.post("/skus/permissions/batch")
@login_required
@permission_required("sku.manage")
def sku_batch_update_permissions():
    sku_ids = [parse_int(value) for value in request.form.getlist("selected_sku_ids") if parse_int(value)]
    user_ids = [parse_int(value) for value in request.form.getlist("authorized_user_ids") if parse_int(value)]
    if not sku_ids:
        flash("请先选择需要批量授权的 SKU。", "error")
    elif not user_ids:
        flash("请先选择至少一位销售员。", "error")
    else:
        skus = SKU.query.filter(SKU.id.in_(sku_ids)).all()
        sales_users = [user for user in get_sku_assignable_users(user_ids) if user.id in set(user_ids)]
        sales_user_ids = {user.id for user in sales_users}
        updated_count = 0
        for sku in skus:
            merged_users = {user.id: user for user in sku.authorized_users}
            for user in sales_users:
                merged_users[user.id] = user
            if sales_user_ids - {user.id for user in sku.authorized_users}:
                updated_count += 1
            sku.authorized_users = list(merged_users.values())
        db.session.commit()
        flash(f"已为 {updated_count} 个 SKU 追加销售授权。", "success")
    return redirect(
        url_for(
            "sku_list",
            sku_keyword=request.args.get("sku_keyword", "").strip(),
            sales_user_id=parse_int(request.args.get("sales_user_id")),
            authorization_state=request.args.get("authorization_state", "").strip(),
        )
    )


@app.route("/sku-mappings", methods=["GET", "POST"])
@login_required
@permission_required("sku.mapping.manage")
def sku_mapping_list():
    mapping_keyword = request.args.get("mapping_keyword", "").strip()
    selected_user_id = parse_int(request.args.get("mapping_user_id"))
    allowed_users = get_mapping_manageable_users([selected_user_id] if selected_user_id else None)
    allowed_user_ids = {user.id for user in allowed_users}
    if allowed_user_ids and selected_user_id and selected_user_id not in allowed_user_ids:
        selected_user_id = 0

    if request.method == "POST":
        user_id = parse_int(request.form.get("user_id"))
        sku_id = parse_int(request.form.get("sku_id"))
        external_sku_code = request.form.get("external_sku_code", "").strip()
        remark = request.form.get("remark", "").strip()

        if user_id not in allowed_user_ids:
            flash("请选择有权限的运营账号。", "error")
        elif not external_sku_code:
            flash("店铺 SKU 不能为空。", "error")
        elif not sku_id:
            flash("请选择仓库 SKU。", "error")
        elif not db.session.get(SKU, sku_id):
            flash("仓库 SKU 不存在。", "error")
        elif not validate_sku_access([sku_id]):
            flash("无权映射该仓库 SKU。", "error")
        elif SKUMapping.query.filter_by(user_id=user_id, external_sku_code=external_sku_code).first():
            flash("该运营下的店铺 SKU 已存在映射。", "error")
        else:
            mapping = SKUMapping(user_id=user_id, sku_id=sku_id, external_sku_code=external_sku_code, remark=remark)
            db.session.add(mapping)
            log_operation("SKU映射", "新增", "sku_mapping", external_sku_code, f"运营ID：{user_id}；仓库SKU：{sku_id}")
            db.session.commit()
            flash("SKU 映射已创建。", "success")
        return redirect(
            url_for(
                "sku_mapping_list",
                mapping_keyword=mapping_keyword,
                mapping_user_id=selected_user_id or "",
            )
        )

    mappings_query = SKUMapping.query.join(SKUMapping.user).join(SKUMapping.sku)
    if allowed_user_ids:
        mappings_query = mappings_query.filter(SKUMapping.user_id.in_(allowed_user_ids))
    if selected_user_id:
        mappings_query = mappings_query.filter(SKUMapping.user_id == selected_user_id)
    if mapping_keyword:
        sku_filter = build_sku_keyword_filter(mapping_keyword)
        keyword_like = f"%{mapping_keyword}%"
        mapping_filters = [
            SKUMapping.external_sku_code.ilike(keyword_like),
            User.full_name.ilike(keyword_like),
            User.username.ilike(keyword_like),
        ]
        if sku_filter is not None:
            mapping_filters.append(sku_filter)
        mappings_query = mappings_query.filter(or_(*mapping_filters))
    mappings_query = apply_sku_scope(mappings_query, SKUMapping.sku_id)
    available_skus = apply_sku_scope(SKU.query, SKU.id).order_by(SKU.name.asc(), SKU.sku_code.asc()).all()
    return render_template(
        "sku_mappings.html",
        mappings=mappings_query.order_by(SKUMapping.created_at.desc()).all(),
        mapping_users=allowed_users,
        skus=available_skus,
        mapping_keyword=mapping_keyword,
        selected_mapping_user_id=selected_user_id,
        default_mapping_user_id=(g.user.id if not can_manage_all_sku_mappings() else (selected_user_id or "")),
    )


@app.get("/sku-mappings/template")
@login_required
@permission_required("sku.mapping.manage")
def sku_mapping_template():
    rows = [
        ["运营账号", "店铺SKU", "仓库SKU", "备注"],
        [g.user.username if g.user else "", "E1006T-Green-S", "1006-Green-S", "店铺A示例"],
        [g.user.username if g.user else "", "STY1006-Green-S", "1006-Green-S", "店铺B示例"],
    ]
    return export_workbook("SKU映射导入模板.xlsx", "SKU映射模板", rows[0], rows[1:])


@app.post("/sku-mappings/import")
@login_required
@permission_required("sku.mapping.manage")
def sku_mapping_import():
    upload = request.files.get("import_file")
    overwrite_existing = request.form.get("overwrite_existing") == "on"
    mapping_keyword = request.args.get("mapping_keyword", "").strip()
    selected_user_id = parse_int(request.args.get("mapping_user_id"))
    allowed_users = get_mapping_manageable_users([selected_user_id] if selected_user_id else None)
    allowed_user_ids = {user.id for user in allowed_users}

    if not upload or not (upload.filename or "").strip():
        flash("请选择要导入的 CSV、TSV 或 Excel 文件。", "error")
    else:
        try:
            import_result = parse_sku_mapping_import(
                upload,
                allowed_user_ids=allowed_user_ids,
                current_user=g.user,
                overwrite_existing=overwrite_existing,
            )
            create_rows = import_result["create_rows"]
            update_rows = import_result["update_rows"]
            for row in create_rows:
                db.session.add(SKUMapping(**row))
            for mapping, row in update_rows:
                mapping.sku_id = row["sku_id"]
                mapping.remark = row["remark"]
            total_count = len(create_rows) + len(update_rows)
            log_operation("SKU映射", "导入", "sku_mapping", upload.filename, f"新增：{len(create_rows)}；覆盖：{len(update_rows)}；合计：{total_count}")
            db.session.commit()
            flash(f"SKU 映射导入成功，新增 {len(create_rows)} 条，覆盖 {len(update_rows)} 条。", "success")
        except ValueError as error:
            db.session.rollback()
            flash(str(error), "error")

    return redirect(
        url_for(
            "sku_mapping_list",
            mapping_keyword=mapping_keyword,
            mapping_user_id=selected_user_id or "",
        )
    )


def render_shipment_sheet_page(*, shipment_form=None):
    if not (can_manage_shipment_sheets() or can_confirm_shipment_sheets()):
        flash("你没有访问发货单记录的权限。", "error")
        return redirect(url_for("dashboard"))

    cleanup_shipment_draft_locks()
    selected_status = request.args.get("status", "").strip()
    keyword = request.args.get("q", "").strip()
    selected_store = request.args.get("store_name", "").strip()
    shipment_query = ShipmentSheet.query.join(ShipmentSheet.warehouse)
    if selected_status:
        shipment_query = shipment_query.filter(ShipmentSheet.status == selected_status)
    if selected_store:
        shipment_query = shipment_query.filter(ShipmentSheet.store_name == selected_store)
    if keyword:
        keyword_like = f"%{keyword}%"
        shipment_query = shipment_query.filter(
            or_(
                ShipmentSheet.shipment_no.ilike(keyword_like),
                ShipmentSheet.store_name.ilike(keyword_like),
                ShipmentSheet.operator_name.ilike(keyword_like),
                ShipmentSheetItem.external_sku_code.ilike(keyword_like),
                SKU.sku_code.ilike(keyword_like),
                SKU.name.ilike(keyword_like),
            )
        ).outerjoin(ShipmentSheet.items).outerjoin(ShipmentSheetItem.sku)

    shipments = shipment_query.order_by(
        case((ShipmentSheet.status == "pending", 0), else_=1),
        ShipmentSheet.created_at.desc(),
    ).distinct().all()
    warehouse = get_default_warehouse()
    mapping_options = get_shipment_mapping_options()
    available_map = get_shipment_available_inventory_map(warehouse_id=warehouse.id if warehouse else None)
    current_lock_map = {row.external_sku_code: row.quantity for row in get_current_shipment_draft_locks()}
    available_by_external_code = {}
    for mapping in mapping_options:
        available_by_external_code[mapping.external_sku_code] = int(available_map.get((warehouse.id, mapping.sku_id), 0) if warehouse else 0)
    shipment_summary = {
        "count": len(shipments),
        "pending_count": sum(1 for row in shipments if row.status == "pending"),
        "confirmed_count": sum(1 for row in shipments if row.status == "confirmed"),
        "quantity": sum(row.total_quantity for row in shipments),
    }
    shipment_form = shipment_form or {
        "store_name": "",
        "box_count": 1,
        "warehouse_id": warehouse.id if warehouse else "",
        "warehouse_name": warehouse.name if warehouse else "-",
        "operator_name": get_current_operator_name(),
        "items": [{}],
        "remark": "",
    }
    return render_template(
        "shipment_sheets.html",
        shipments=shipments,
        shipment_form=shipment_form,
        shipment_summary=shipment_summary,
        stores=sorted({row[0] for row in ShipmentSheet.query.with_entities(ShipmentSheet.store_name).all() if row[0]}),
        selected_status=selected_status,
        selected_store=selected_store,
        q=keyword,
        warehouse=warehouse,
        mapping_options=mapping_options,
        available_by_external_code=available_by_external_code,
        current_shipment_lock_map=current_lock_map,
    )


@app.route("/shipment-sheets", methods=["GET", "POST"])
@login_required
def shipment_sheet_list():
    if request.method == "GET":
        return redirect(
            url_for(
                "fba_calculator",
                shipment_q=request.args.get("q", "").strip(),
                shipment_status=request.args.get("status", "").strip(),
                shipment_store_name=request.args.get("store_name", "").strip(),
            )
        )
    if not (can_manage_shipment_sheets() or can_confirm_shipment_sheets()):
        flash("你没有访问发货单记录的权限。", "error")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        if not can_manage_shipment_sheets():
            flash("你没有创建发货单的权限。", "error")
            return redirect(url_for("shipment_sheet_list"))
        warehouse = get_default_warehouse()
        store_name = request.form.get("store_name", "").strip()
        box_count = parse_int(request.form.get("box_count"), 1)
        remark = request.form.get("remark", "").strip()
        items, errors = parse_shipment_sheet_items(
            box_count=box_count,
            warehouse_id=warehouse.id if warehouse else 0,
            user=g.user,
        )
        shipment_form = {
            "store_name": store_name,
            "box_count": box_count,
            "warehouse_id": warehouse.id if warehouse else "",
            "warehouse_name": warehouse.name if warehouse else "-",
            "operator_name": get_current_operator_name(),
            "items": [{"external_sku_code": row.get("external_sku_code", ""), "quantity": row.get("quantity", "")} for row in items]
            or [{"external_sku_code": code, "quantity": qty} for code, qty in zip(request.form.getlist("shipment_external_sku_code"), request.form.getlist("shipment_quantity"))],
            "remark": remark,
        }
        if not warehouse:
            flash("系统还没有可用仓库，暂时无法创建发货单。", "error")
            return render_shipment_sheet_page(shipment_form=shipment_form)
        if not store_name:
            errors.append("店铺不能为空。")
        if box_count <= 0:
            errors.append("箱数必须大于 0。")
        if not items:
            errors.append("请至少添加一条有效的发货SKU。")
        if errors:
            for error in errors:
                flash(error, "error")
            return render_shipment_sheet_page(shipment_form=shipment_form)

        shipment = ShipmentSheet(
            shipment_no=build_unique_code("SH", ShipmentSheet, "shipment_no"),
            store_name=store_name,
            warehouse_id=warehouse.id,
            operator_user_id=g.user.id if g.user else None,
            operator_name=get_current_operator_name(),
            box_count=box_count,
            remark=remark,
        )
        for item in items:
            shipment.items.append(
                ShipmentSheetItem(
                    sku_id=item["sku_id"],
                    external_sku_code=item["external_sku_code"],
                    quantity=item["quantity"],
                    per_box_quantity=item["per_box_quantity"],
                )
            )
        db.session.add(shipment)
        ShipmentDraftLock.query.filter_by(lock_token=get_shipment_lock_token()).delete()
        log_operation("发货表", "新增", "shipment_sheet", shipment.shipment_no, f"店铺：{shipment.store_name}；箱数：{shipment.box_count}；明细：{len(items)}")
        db.session.commit()
        flash("发货单已确认并预锁库存，仓管员可以进入发货单记录进行确认。", "success")
        return redirect(url_for("shipment_sheet_list"))

    return render_fba_calculator_page()


@app.post("/shipment-sheets/<int:shipment_id>/update")
@login_required
def shipment_sheet_update(shipment_id):
    if not can_manage_shipment_sheets():
        flash("你没有编辑发货单的权限。", "error")
        return redirect(url_for("fba_calculator"))
    shipment = db.session.get(ShipmentSheet, shipment_id)
    if not shipment:
        flash("发货单不存在。", "error")
        return redirect(url_for("fba_calculator"))
    if shipment.status != "pending":
        flash("只有待备货的发货单可以编辑。", "error")
        return redirect(url_for("fba_calculator"))
    if not can_edit_shipment_sheet(shipment):
        flash("只有创建人可以编辑这张发货单。", "error")
        return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))

    store_name = request.form.get("store_name", "").strip()
    box_count = parse_int(request.form.get("box_count"), 1)
    remark = request.form.get("remark", "").strip()
    items, errors = parse_shipment_sheet_items(
        box_count=box_count,
        warehouse_id=shipment.warehouse_id,
        user=g.user,
        exclude_shipment_id=shipment.id,
    )
    if not store_name:
        errors.append("店铺不能为空。")
    if box_count <= 0:
        errors.append("箱数必须大于 0。")
    if not items:
        errors.append("请至少保留一条有效的发货SKU。")
    if errors:
        for error in errors:
            flash(error, "error")
        return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))

    shipment.store_name = store_name
    shipment.box_count = box_count
    shipment.remark = remark
    shipment.items.clear()
    for item in items:
        shipment.items.append(
            ShipmentSheetItem(
                sku_id=item["sku_id"],
                external_sku_code=item["external_sku_code"],
                quantity=item["quantity"],
                per_box_quantity=item["per_box_quantity"],
            )
        )
    ShipmentDraftLock.query.filter_by(lock_token=get_shipment_lock_token()).delete()
    log_operation("发货表", "更新", "shipment_sheet", shipment.shipment_no, f"店铺：{shipment.store_name}；箱数：{shipment.box_count}；明细：{len(items)}")
    db.session.commit()
    flash("发货单已更新。", "success")
    return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))


@app.post("/shipment-sheets/<int:shipment_id>/confirm")
@login_required
def shipment_sheet_confirm(shipment_id):
    if not can_confirm_shipment_sheets():
        flash("你没有仓管确认发货单的权限。", "error")
        return redirect(url_for("fba_calculator"))
    shipment = db.session.get(ShipmentSheet, shipment_id)
    if not shipment:
        flash("发货单不存在。", "error")
        return redirect(url_for("fba_calculator"))
    if shipment.status == "confirmed":
        flash("该发货单已经确认过了。", "error")
        return redirect(url_for("fba_calculator"))
    try:
        for item in shipment.items:
            add_inventory_transaction(
                item.sku_id,
                shipment.warehouse_id,
                -item.quantity,
                "shipment_outbound",
                reference_type="shipment_sheet",
                reference_no=shipment.shipment_no,
                note=f"店铺：{shipment.store_name}",
                operator_name=get_current_operator_name(),
            )
        shipment.status = "confirmed"
        shipment.warehouse_confirmed_at = datetime.utcnow()
        shipment.warehouse_confirmer_name = get_current_operator_name()
        log_operation("发货表", "确认", "shipment_sheet", shipment.shipment_no, f"仓管确认：{shipment.warehouse_confirmer_name}")
        db.session.commit()
        flash("仓管确认成功，库存已正式扣减。", "success")
    except InventoryError as error:
        db.session.rollback()
        flash(f"确认失败：{error}", "error")
    return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))


@app.post("/shipment-sheets/<int:shipment_id>/delete")
@login_required
def shipment_sheet_delete(shipment_id):
    if not can_manage_shipment_sheets():
        flash("你没有删除发货单的权限。", "error")
        return redirect(url_for("fba_calculator"))
    shipment = db.session.get(ShipmentSheet, shipment_id)
    if not shipment:
        flash("发货单不存在。", "error")
        return redirect(url_for("fba_calculator"))
    if not can_delete_shipment_sheet(shipment):
        flash("你没有删除这张发货单的权限。", "error")
        return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))
    if shipment.status == "confirmed":
        deleted_txn_count = delete_inventory_transactions(reference_type="shipment_sheet", reference_no=shipment.shipment_no)
        log_operation("发货表", "删除", "shipment_sheet", shipment.shipment_no, f"已回滚流水：{deleted_txn_count}")
    else:
        log_operation("发货表", "删除", "shipment_sheet", shipment.shipment_no, f"店铺：{shipment.store_name}")
    db.session.delete(shipment)
    db.session.commit()
    if shipment.status == "confirmed":
        flash("发货单已删除，已同步回滚扣减库存。", "success")
    else:
        flash("发货单已删除。", "success")
    return redirect(url_for("fba_calculator", shipment_q=request.args.get("shipment_q", "").strip(), shipment_status=request.args.get("shipment_status", "").strip(), shipment_store_name=request.args.get("shipment_store_name", "").strip()))


@app.post("/shipment-sheets/prelock")
@login_required
def shipment_sheet_prelock():
    if not can_manage_shipment_sheets():
        return jsonify({"ok": False, "message": "你没有预锁发货库存的权限。"}), 403
    cleanup_shipment_draft_locks()
    warehouse = get_default_warehouse()
    external_sku_code = request.form.get("external_sku_code", "").strip()
    quantity = parse_int(request.form.get("quantity"))
    if not warehouse:
        return jsonify({"ok": False, "message": "系统还没有可用仓库。"}), 400
    if not external_sku_code:
        return jsonify({"ok": False, "message": "请先选择店铺SKU。"}), 400
    if quantity <= 0:
        return jsonify({"ok": False, "message": "请先填写有效数量。"}), 400
    mapping_query = SKUMapping.query.filter(func.lower(SKUMapping.external_sku_code) == external_sku_code.lower())
    if not can_manage_all_sku_mappings():
        mapping_query = mapping_query.filter(SKUMapping.user_id == g.user.id)
    mapping = mapping_query.first()
    if not mapping:
        return jsonify({"ok": False, "message": f"店铺SKU {external_sku_code} 未找到映射。"}), 400
    available_map = get_shipment_available_inventory_map(warehouse_id=warehouse.id)
    available_quantity = int(available_map.get((warehouse.id, mapping.sku_id), 0) or 0)
    if quantity > available_quantity:
        return jsonify({"ok": False, "message": f"库存不足：{external_sku_code} 当前实时可用库存只有 {available_quantity}", "available": available_quantity}), 400

    ShipmentDraftLock.query.filter_by(lock_token=get_shipment_lock_token(), external_sku_code=external_sku_code).delete()
    db.session.add(
        ShipmentDraftLock(
            lock_token=get_shipment_lock_token(),
            user_id=g.user.id if g.user else None,
            warehouse_id=warehouse.id,
            sku_id=mapping.sku_id,
            external_sku_code=external_sku_code,
            quantity=quantity,
            expires_at=datetime.utcnow() + timedelta(hours=2),
        )
    )
    db.session.commit()
    return jsonify({"ok": True, "message": f"已预锁 {external_sku_code} 数量 {quantity}", "locked_quantity": quantity})


@app.post("/shipment-sheets/prelock/release")
@login_required
def shipment_sheet_prelock_release():
    if not can_manage_shipment_sheets():
        return jsonify({"ok": False, "message": "你没有操作预锁的权限。"}), 403
    external_sku_code = request.form.get("external_sku_code", "").strip()
    if not external_sku_code:
        return jsonify({"ok": False, "message": "缺少店铺SKU。"}), 400
    ShipmentDraftLock.query.filter_by(lock_token=get_shipment_lock_token(), external_sku_code=external_sku_code).delete()
    db.session.commit()
    return jsonify({"ok": True, "message": f"已取消预锁 {external_sku_code}"})


@app.get("/shipment-sheets/export")
@login_required
def shipment_sheet_export():
    if not (can_manage_shipment_sheets() or can_confirm_shipment_sheets()):
        flash("你没有导出发货单的权限。", "error")
        return redirect(url_for("dashboard"))
    selected_status = request.args.get("status", "").strip()
    keyword = request.args.get("q", "").strip()
    selected_store = request.args.get("store_name", "").strip()
    selected_archive = request.args.get("archive", "active").strip() or "active"
    query = ShipmentSheet.query
    if selected_archive == "archived":
        query = query.filter(ShipmentSheet.archived_at.isnot(None))
    elif selected_archive != "all":
        query = query.filter(ShipmentSheet.archived_at.is_(None))
    if selected_status:
        query = query.filter(ShipmentSheet.status == selected_status)
    if selected_store:
        query = query.filter(ShipmentSheet.store_name == selected_store)
    if keyword:
        keyword_like = f"%{keyword}%"
        query = query.filter(
            or_(
                ShipmentSheet.shipment_no.ilike(keyword_like),
                ShipmentSheet.store_name.ilike(keyword_like),
                ShipmentSheet.operator_name.ilike(keyword_like),
            )
        )
    rows = build_shipment_export_rows(query.order_by(ShipmentSheet.created_at.desc()).all())
    return export_workbook(
        "发货表.xlsx",
        "发货表",
        ["发货单号", "店铺", "运营", "仓库", "箱数", "店铺SKU", "仓库SKU", "商品名称", "总数量", "每箱数量", "状态", "仓管确认人", "创建时间", "归档时间"],
        rows,
    )


@app.post("/shipment-sheets/archive-export")
@login_required
def shipment_sheet_archive_export():
    if not can_manage_shipment_sheets():
        flash("你没有归档发货单的权限。", "error")
        return redirect(url_for("fba_calculator"))
    shipment_ids = sorted({parse_int(value) for value in request.form.getlist("shipment_ids") if parse_int(value)})
    if not shipment_ids:
        flash("请先勾选要归档并导出的发货单。", "error")
        return redirect(url_for("fba_calculator"))
    shipments = ShipmentSheet.query.filter(ShipmentSheet.id.in_(shipment_ids)).order_by(ShipmentSheet.created_at.desc()).all()
    if not shipments:
        flash("未找到可归档的发货单。", "error")
        return redirect(url_for("fba_calculator"))
    archived_at = datetime.utcnow()
    archived_count = 0
    for shipment in shipments:
        if shipment.archived_at is None:
            shipment.archived_at = archived_at
            archived_count += 1
    log_operation("发货表", "归档并导出", "shipment_sheet", f"{len(shipments)} 张", f"新增归档：{archived_count} 张")
    db.session.commit()
    return export_workbook(
        "发货表归档.xlsx",
        "归档发货表",
        ["发货单号", "店铺", "运营", "仓库", "箱数", "店铺SKU", "仓库SKU", "商品名称", "总数量", "每箱数量", "状态", "仓管确认人", "创建时间", "归档时间"],
        build_shipment_export_rows(shipments),
    )


def render_fba_calculator_page(*, calc_form=None, calc_result=None, shipment_form=None):
    calc_form = calc_form or session.get("_fba_calc_form") or {
        "multiplier": "1",
        "new_product_multiplier": "1.5",
        "coverage_days": "37",
        "production_days": "0",
        "sea_days": "30",
        "inbound_processing_days": "0",
        "safety_stock_days": "7",
        "marketplace": "ALL",
        "shipment_store_name": "",
        "shipment_box_count": "1",
        "shipment_remark": "",
    }
    if calc_result is None and (session.get("_fba_calc_result") or session.get("_fba_calc_result_key")):
        refresh_fba_calc_session_result()
    calc_result = calc_result or get_fba_calc_result()
    result_filters = {
        "sku_keyword": request.args.get("sku_keyword", "").strip(),
        "suggested_state": request.args.get("suggested_state", "").strip(),
        "suggested_min": request.args.get("suggested_min", "").strip(),
    }
    filtered_calc_result = build_filtered_fba_calc_result(
        calc_result,
        sku_keyword=result_filters["sku_keyword"],
        suggested_state=result_filters["suggested_state"],
        suggested_min=result_filters["suggested_min"],
    )
    shipment_context = build_shipment_page_context(shipment_form=shipment_form)
    return render_template(
        "fba_calculator.html",
        calc_form=calc_form,
        calc_result=filtered_calc_result,
        full_calc_result=calc_result,
        result_filters=result_filters,
        active_inbound_count=FBAInboundRecord.query.filter_by(status="active").count(),
        shipments=shipment_context["shipments"],
        shipment_summary=shipment_context["shipment_summary"],
        shipment_stores=shipment_context["shipment_stores"],
        shipment_mapping_options=shipment_context["shipment_mapping_options"],
        selected_shipment_status=shipment_context["selected_shipment_status"],
        selected_shipment_store=shipment_context["selected_shipment_store"],
        selected_shipment_archive=shipment_context["selected_shipment_archive"],
        shipment_q=shipment_context["shipment_q"],
        shipment_form=shipment_context["shipment_form"],
    )


@app.route("/fba-calculator", methods=["GET", "POST"])
@login_required
def fba_calculator():
    if not can_access_fba_calculator_page():
        flash("你没有访问 FBA发货计算器 的权限。", "error")
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        if not g.user.has_permission("fba.calc.manage"):
            flash("你没有执行 FBA 计算的权限。", "error")
            return render_fba_calculator_page()
        calc_form = {
            "multiplier": request.form.get("multiplier", "1").strip() or "1",
            "new_product_multiplier": request.form.get("new_product_multiplier", "1.5").strip() or "1.5",
            "coverage_days": request.form.get("coverage_days", "30").strip() or "30",
            "production_days": request.form.get("production_days", "0").strip() or "0",
            "sea_days": request.form.get("sea_days", "30").strip() or "30",
            "inbound_processing_days": request.form.get("inbound_processing_days", "0").strip() or "0",
            "safety_stock_days": request.form.get("safety_stock_days", "7").strip() or "7",
            "marketplace": request.form.get("marketplace", "ALL").strip().upper() or "ALL",
        }
        inventory_file = request.files.get("inventory_file")
        reserved_file = request.files.get("reserved_file")
        sales_30_file = request.files.get("sales_30_file")
        missing_files = [
            name
            for name, upload in [
                ("亚马逊物流库存表", inventory_file),
                ("预留库存表", reserved_file),
                ("最近30天销量明细", sales_30_file),
            ]
            if not upload or not (upload.filename or "").strip()
        ]
        if missing_files:
            flash(f"请补齐这些文件后再计算：{'、'.join(missing_files)}。", "error")
            return render_fba_calculator_page(calc_form=calc_form)

        try:
            multiplier = float(calc_form["multiplier"] or "1")
            new_product_multiplier = float(calc_form["new_product_multiplier"] or "1.5")
            coverage_days = parse_int(calc_form["coverage_days"], 30)
            inventory_data = parse_amazon_inventory_report(inventory_file)
            reserved_data = parse_reserved_inventory_report(reserved_file)
            sales_windows = parse_sales_velocity_windows(sales_30_file, marketplace=calc_form["marketplace"])
            sales_7_data = sales_windows["sales_7"]
            sales_14_data = sales_windows["sales_14"]
            sales_30_data = sales_windows["sales_30"]
            inbound_data = defaultdict(int, get_active_fba_inbound_totals())
            calc_result = calculate_fba_replenishment(
                inventory_data=inventory_data,
                reserved_data=reserved_data,
                sales_7_data=sales_7_data,
                sales_14_data=sales_14_data,
                sales_30_data=sales_30_data,
                inbound_data=dict(inbound_data),
                multiplier=multiplier,
                new_product_multiplier=new_product_multiplier,
                coverage_days=coverage_days,
                production_days=calc_form["production_days"],
                sea_days=calc_form["sea_days"],
                inbound_processing_days=calc_form["inbound_processing_days"],
                safety_stock_days=calc_form["safety_stock_days"],
                user=g.user,
            )
            calc_result["sales_window"] = sales_windows
            set_fba_calc_result(calc_result)
            session["_fba_calc_form"] = calc_form
            if calc_result["unmatched_codes"]:
                flash(f"以下店铺SKU还没有映射，已按未映射显示：{'、'.join(calc_result['unmatched_codes'][:8])}", "error")
            if sales_windows.get("window_end_date"):
                flash(
                    f"已排除最新未完整日 {sales_windows.get('excluded_latest_date') or '-'}，按统计截止日 {sales_windows['window_end_date']} 自动汇总最近7天、14天和30天销量。",
                    "success",
                )
            flash(f"计算完成，共生成 {calc_result['summary']['sku_count']} 条SKU建议。", "success")
            return render_fba_calculator_page(calc_form=calc_form, calc_result=calc_result)
        except ValueError as error:
            flash(str(error), "error")
            return render_fba_calculator_page(calc_form=calc_form)

    return render_fba_calculator_page()


@app.post("/fba-calculator/rules")
@login_required
@permission_required("fba.calc.manage")
def fba_calculator_rule_update():
    external_sku_code = request.form.get("external_sku_code", "").strip()
    custom_multiplier = parse_decimal(request.form.get("custom_multiplier"), "1")
    is_new_product = request.form.get("is_new_product") == "on"
    remark = request.form.get("remark", "").strip()
    filter_args = {
        "sku_keyword": request.form.get("sku_keyword", "").strip(),
        "suggested_state": request.form.get("suggested_state", "").strip(),
        "suggested_min": request.form.get("suggested_min", "").strip(),
    }
    if not external_sku_code:
        flash("缺少店铺SKU，无法保存补货规则。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    rule = FBAReplenishmentRule.query.filter_by(external_sku_code=external_sku_code).first()
    if not rule:
        rule = FBAReplenishmentRule(external_sku_code=external_sku_code)
        db.session.add(rule)
    rule.custom_multiplier = custom_multiplier
    rule.is_new_product = is_new_product
    rule.remark = remark
    db.session.commit()
    refresh_fba_calc_session_result()
    flash(f"{external_sku_code} 的补货规则已保存。", "success")
    return redirect(url_for("fba_calculator", **filter_args))


@app.post("/fba-calculator/rules/batch")
@login_required
@permission_required("fba.calc.manage")
def fba_calculator_rule_batch_update():
    external_codes = sorted({(value or "").strip() for value in request.form.getlist("external_sku_codes") if (value or "").strip()})
    action = request.form.get("batch_action", "").strip()
    calc_form = session.get("_fba_calc_form") or {}
    calc_form["multiplier"] = request.form.get("multiplier", str(calc_form.get("multiplier", "1"))).strip() or "1"
    calc_form["new_product_multiplier"] = request.form.get(
        "new_product_multiplier",
        str(calc_form.get("new_product_multiplier", "1.5")),
    ).strip() or "1.5"
    calc_form["coverage_days"] = request.form.get("coverage_days", str(calc_form.get("coverage_days", "30"))).strip() or "30"
    calc_form["production_days"] = request.form.get("production_days", str(calc_form.get("production_days", "0"))).strip() or "0"
    calc_form["sea_days"] = request.form.get("sea_days", str(calc_form.get("sea_days", "30"))).strip() or "30"
    calc_form["inbound_processing_days"] = request.form.get(
        "inbound_processing_days",
        str(calc_form.get("inbound_processing_days", "0")),
    ).strip() or "0"
    calc_form["safety_stock_days"] = request.form.get(
        "safety_stock_days",
        str(calc_form.get("safety_stock_days", "7")),
    ).strip() or "7"
    calc_form["shipment_store_name"] = request.form.get("shipment_store_name", "").strip()
    calc_form["shipment_box_count"] = request.form.get("shipment_box_count", "1").strip() or "1"
    calc_form["shipment_remark"] = request.form.get("shipment_remark", "").strip()
    session["_fba_calc_form"] = calc_form
    default_multiplier = max(parse_float(calc_form.get("multiplier"), 1), 0)
    filter_args = {
        "sku_keyword": request.form.get("sku_keyword", "").strip(),
        "suggested_state": request.form.get("suggested_state", "").strip(),
        "suggested_min": request.form.get("suggested_min", "").strip(),
    }

    if action == "refresh_calc":
        refresh_fba_calc_session_result()
        flash("计算参数已更新。", "success")
        return redirect(url_for("fba_calculator", **filter_args))

    if not external_codes:
        flash("请先勾选要批量处理的 SKU。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if action not in {"mark_new", "clear_new"}:
        flash("批量操作无效。", "error")
        return redirect(url_for("fba_calculator", **filter_args))

    rules_by_external_sku = {
        row.external_sku_code: row
        for row in FBAReplenishmentRule.query.filter(FBAReplenishmentRule.external_sku_code.in_(external_codes)).all()
    }
    updated_count = 0
    for external_code in external_codes:
        rule = rules_by_external_sku.get(external_code)
        if action == "mark_new":
            if not rule:
                rule = FBAReplenishmentRule(external_sku_code=external_code, custom_multiplier=default_multiplier)
                db.session.add(rule)
            rule.is_new_product = True
            updated_count += 1
        elif rule:
            rule.is_new_product = False
            updated_count += 1

    db.session.commit()
    refresh_fba_calc_session_result()
    flash(
        f"已批量为 {updated_count} 个SKU{'标记为新品' if action == 'mark_new' else '解除新品状态'}。",
        "success",
    )
    return redirect(url_for("fba_calculator", **filter_args))


@app.post("/fba-calculator/generate-shipment")
@login_required
@permission_required("fba.calc.manage")
def fba_calculator_generate_shipment():
    filter_args = {
        "sku_keyword": request.form.get("sku_keyword", "").strip(),
        "suggested_state": request.form.get("suggested_state", "").strip(),
        "suggested_min": request.form.get("suggested_min", "").strip(),
    }
    if not can_manage_shipment_sheets():
        flash("你没有把 FBA 结果生成发货表的权限。", "error")
        return redirect(url_for("fba_calculator", **filter_args))

    calc_result = get_fba_calc_result()
    rows = calc_result.get("rows") or []
    row_by_external_sku = {row.get("external_sku_code", "").strip(): row for row in rows if row.get("external_sku_code")}
    external_codes = [(value or "").strip() for value in request.form.getlist("external_sku_codes") if (value or "").strip()]
    store_name = request.form.get("shipment_store_name", "").strip()
    remark = request.form.get("shipment_remark", "").strip()
    box_count = parse_int(request.form.get("shipment_box_count"), 1)
    calc_form = session.get("_fba_calc_form") or {}
    calc_form["shipment_store_name"] = store_name
    calc_form["shipment_box_count"] = request.form.get("shipment_box_count", "1").strip() or "1"
    calc_form["shipment_remark"] = remark
    session["_fba_calc_form"] = calc_form
    warehouse = get_default_warehouse()

    if not rows:
        flash("当前还没有可用的 FBA 计算结果，请先完成一次计算。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if not external_codes:
        flash("请先勾选你要生成发货表的 SKU。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if not store_name:
        flash("请先填写发货表的店铺名称。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if box_count <= 0:
        flash("箱数必须大于 0。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if not warehouse:
        flash("系统还没有可用仓库，暂时无法生成发货表草稿。", "error")
        return redirect(url_for("fba_calculator", **filter_args))

    items = []
    errors = []
    requested_by_sku = defaultdict(int)
    available_by_sku = {}
    for external_code in external_codes:
        row = row_by_external_sku.get(external_code)
        if not row:
            errors.append(f"{external_code} 不在当前计算结果里。")
            continue
        raw_quantity = request.form.get(f"shipment_quantity__{external_code}", "").strip()
        shipment_quantity = parse_int(raw_quantity, 0)
        if shipment_quantity <= 0:
            continue
        if row.get("warehouse_sku_code") == "-":
            errors.append(f"{external_code} 还没有对应的仓库SKU映射。")
            continue
        if shipment_quantity % box_count != 0:
            errors.append(f"{external_code} 的发货数量 {shipment_quantity} 不能被箱数 {box_count} 整除。")
            continue
        available_quantity = parse_int(row.get("warehouse_qty"))
        sku_id = parse_int(row.get("sku_id"))
        requested_by_sku[sku_id] += shipment_quantity
        available_by_sku[sku_id] = available_quantity
        items.append(
            {
                "external_sku_code": external_code,
                "quantity": shipment_quantity,
                "per_box_quantity": shipment_quantity // box_count,
                "sku_id": sku_id,
                "warehouse_sku_code": row.get("warehouse_sku_code"),
            }
        )

    for sku_id, total_quantity in requested_by_sku.items():
        if total_quantity > parse_int(available_by_sku.get(sku_id)):
            sample_item = next((item for item in items if item["sku_id"] == sku_id), None)
            if sample_item:
                errors.append(f"仓库SKU {sample_item['warehouse_sku_code']} 合计发货数量 {total_quantity} 超过当前可用库存 {available_by_sku.get(sku_id, 0)}。")

    if errors:
        for error in errors[:8]:
            flash(error, "error")
        if len(errors) > 8:
            flash(f"还有 {len(errors) - 8} 条错误未展示，请调整后再重试。", "error")
        return redirect(url_for("fba_calculator", **filter_args))
    if not items:
        flash("请至少填写一个大于 0 的发货数量。填空或填 0 的行会自动跳过。", "error")
        return redirect(url_for("fba_calculator", **filter_args))

    shipment = ShipmentSheet(
        shipment_no=build_unique_code("SH", ShipmentSheet, "shipment_no"),
        store_name=store_name,
        warehouse_id=warehouse.id,
        operator_user_id=g.user.id if g.user else None,
        operator_name=get_current_operator_name(),
        box_count=box_count,
        remark=remark or "FBA发货计算器生成",
    )
    for item in items:
        shipment.items.append(
            ShipmentSheetItem(
                sku_id=item["sku_id"],
                external_sku_code=item["external_sku_code"],
                quantity=item["quantity"],
                per_box_quantity=item["per_box_quantity"],
            )
        )
    db.session.add(shipment)
    log_operation("发货表", "新增", "shipment_sheet", shipment.shipment_no, f"来源：FBA发货计算器；店铺：{shipment.store_name}；箱数：{shipment.box_count}；明细：{len(items)}")
    db.session.commit()
    flash(f"已创建发货单 {shipment.shipment_no}，共 {len(items)} 个SKU，等待仓管确认。", "success")
    return redirect(url_for("fba_calculator", shipment_status="pending", **filter_args))


@app.get("/fba-calculator/export")
@login_required
@permission_required("fba.calc.manage")
def fba_calculator_export():
    calc_result = get_fba_calc_result()
    filtered_calc_result = build_filtered_fba_calc_result(
        calc_result,
        sku_keyword=request.args.get("sku_keyword", "").strip(),
        suggested_state=request.args.get("suggested_state", "").strip(),
        suggested_min=request.args.get("suggested_min", "").strip(),
    )
    rows = filtered_calc_result.get("rows") or []
    if not rows:
        flash("当前没有可导出的计算结果，请先完成计算或调整筛选条件。", "error")
        return redirect(url_for("fba_calculator"))
    export_rows = []
    for row in rows:
        export_rows.append(
            [
                row.get("external_sku_code", ""),
                row.get("warehouse_sku_code", ""),
                row.get("warehouse_sku_name", ""),
                row.get("sellable_qty", 0),
                row.get("received_qty", 0),
                row.get("reserved_qty", 0),
                row.get("sales_7_qty", 0),
                row.get("sales_14_qty", row.get("sales_8_14_qty", 0)),
                row.get("sales_30_qty", 0),
                row.get("daily_7", 0),
                row.get("daily_14", row.get("daily_8_14", 0)),
                row.get("daily_30", 0),
                row.get("weighted_daily", row.get("daily_30", 0)),
                row.get("target_days", 0),
                row.get("trend_factor", 0),
                row.get("applied_multiplier", 0),
                row.get("inbound_qty", 0),
                row.get("warehouse_qty", 0),
                row.get("sales_health_status", ""),
                row.get("sales_health_note", ""),
                row.get("suggested_qty", 0),
            ]
        )
    return export_workbook(
        "FBA发货计算结果.xlsx",
        "FBA计算结果",
        ["店铺SKU", "仓库SKU", "商品", "可售库存", "已接收库存", "预留库存", "7天销量", "14天销量", "30天销量", "7天日均", "14天日均", "30天日均", "加权日均", "目标天数", "趋势系数", "应用倍率", "在途数量", "仓库库存", "销量判断", "判断说明", "建议发货量"],
        export_rows,
    )


@app.route("/fba-inbounds", methods=["GET", "POST"])
@login_required
@permission_required("fba.inbound.manage")
def fba_inbound_list():
    if request.method == "POST":
        upload = request.files.get("inbound_file")
        if not upload or not (upload.filename or "").strip():
            flash("请选择要导入的亚马逊发货单。", "error")
            return redirect(url_for("fba_inbound_list"))
        try:
            shipment_ref, rows = parse_amazon_inbound_records(upload)
            if not rows:
                flash("这份发货单里没有识别到有效的在途数据。", "error")
                return redirect(url_for("fba_inbound_list"))
            existing_records = FBAInboundRecord.query.filter_by(shipment_ref=shipment_ref).all()
            conflicting_record = next(
                (
                    row
                    for row in existing_records
                    if row.uploader_user_id and g.user and row.uploader_user_id != g.user.id
                ),
                None,
            )
            if conflicting_record:
                owner = db.session.get(User, conflicting_record.uploader_user_id)
                owner_name = owner.full_name if owner else "其他用户"
                flash(f"在途货件 {shipment_ref} 由 {owner_name} 上传，只有上传人可以覆盖或更新。", "error")
                return redirect(url_for("fba_inbound_list"))
            for external_code, quantity in rows.items():
                record = FBAInboundRecord.query.filter_by(shipment_ref=shipment_ref, external_sku_code=external_code, status="active").first()
                if record:
                    record.quantity = quantity
                    if not record.uploader_user_id and g.user:
                        record.uploader_user_id = g.user.id
                else:
                    db.session.add(
                        FBAInboundRecord(
                            shipment_ref=shipment_ref,
                            external_sku_code=external_code,
                            quantity=quantity,
                            status="active",
                            uploader_user_id=g.user.id if g.user else None,
                        )
                    )
            log_operation("FBA在途", "导入", "fba_inbound", shipment_ref, f"SKU数：{len(rows)}")
            db.session.commit()
            flash(f"已导入在途货件 {shipment_ref}，共 {len(rows)} 条。", "success")
        except ValueError as error:
            flash(str(error), "error")
        return redirect(url_for("fba_inbound_list"))

    keyword = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()
    query = FBAInboundRecord.query
    if keyword:
        keyword_like = f"%{keyword}%"
        query = query.filter(or_(FBAInboundRecord.shipment_ref.ilike(keyword_like), FBAInboundRecord.external_sku_code.ilike(keyword_like)))
    if status:
        query = query.filter(FBAInboundRecord.status == status)
    records = query.order_by(
        case((FBAInboundRecord.status == "active", 0), else_=1),
        FBAInboundRecord.shipment_ref.desc(),
        FBAInboundRecord.created_at.desc(),
    ).all()
    record_groups = []
    current_group = None
    for record in records:
        if not current_group or current_group["shipment_ref"] != record.shipment_ref:
            current_group = {
                "shipment_ref": record.shipment_ref,
                "records": [],
                "active_qty": 0,
                "active_count": 0,
                "can_edit": False,
            }
            record_groups.append(current_group)
        current_group["records"].append(record)
        if record.status == "active":
            current_group["active_qty"] += record.quantity
            current_group["active_count"] += 1
        current_group["can_edit"] = can_edit_fba_inbound_group(current_group["records"])
    summary = {
        "count": len(records),
        "active_count": sum(1 for row in records if row.status == "active"),
        "active_qty": sum(row.quantity for row in records if row.status == "active"),
    }
    return render_template(
        "fba_inbounds.html",
        record_groups=record_groups,
        summary=summary,
        q=keyword,
        selected_status=status,
    )


@app.post("/fba-inbounds/<int:record_id>/receive")
@login_required
@permission_required("fba.inbound.manage")
def fba_inbound_receive(record_id):
    record = db.session.get(FBAInboundRecord, record_id)
    if not record:
        flash("在途记录不存在。", "error")
        return redirect(url_for("fba_inbound_list"))
    if not can_edit_fba_inbound_record(record):
        flash("只有上传人可以标记这条在途记录。", "error")
        return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))
    record.status = "received"
    log_operation("FBA在途", "标记已接收", "fba_inbound", record.shipment_ref, f"{record.external_sku_code} / {record.quantity}")
    db.session.commit()
    flash("该在途记录已标记为已接收。", "success")
    return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))


@app.post("/fba-inbounds/<int:record_id>/delete")
@login_required
@permission_required("fba.inbound.manage")
def fba_inbound_delete(record_id):
    record = db.session.get(FBAInboundRecord, record_id)
    if not record:
        flash("在途记录不存在。", "error")
        return redirect(url_for("fba_inbound_list"))
    if not can_edit_fba_inbound_record(record):
        flash("只有上传人可以删除这条在途记录。", "error")
        return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))
    shipment_ref = record.shipment_ref
    log_operation("FBA在途", "删除", "fba_inbound", record.shipment_ref, f"{record.external_sku_code} / {record.quantity}")
    db.session.delete(record)
    db.session.commit()
    flash("该在途记录已彻底删除。", "success")
    return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))


@app.post("/fba-inbounds/delete-group")
@login_required
@permission_required("fba.inbound.manage")
def fba_inbound_group_delete():
    shipment_ref = request.form.get("shipment_ref", "").strip()
    if not shipment_ref:
        flash("缺少货件号。", "error")
        return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))
    records = FBAInboundRecord.query.filter_by(shipment_ref=shipment_ref).all()
    if not records:
        flash("在途货件不存在。", "error")
        return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))
    if not can_edit_fba_inbound_group(records):
        flash("只有上传人可以整单删除这票在途货件。", "error")
        return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))
    deleted_count = len(records)
    for record in records:
        db.session.delete(record)
    log_operation("FBA在途", "整单删除", "fba_inbound", shipment_ref, f"删除记录：{deleted_count} 条")
    db.session.commit()
    flash(f"在途货件 {shipment_ref} 已彻底删除，共删除 {deleted_count} 条记录。", "success")
    return redirect(url_for("fba_inbound_list", q=request.args.get("q", "").strip(), status=request.args.get("status", "").strip()))


@app.post("/sku-mappings/<int:mapping_id>/update")
@login_required
@permission_required("sku.mapping.manage")
def sku_mapping_update(mapping_id):
    mapping = db.session.get(SKUMapping, mapping_id)
    if not mapping:
        flash("SKU 映射不存在。", "error")
        return redirect(url_for("sku_mapping_list"))

    allowed_users = get_mapping_manageable_users([mapping.user_id])
    allowed_user_ids = {user.id for user in allowed_users}
    if mapping.user_id not in allowed_user_ids:
        flash("无权编辑该 SKU 映射。", "error")
        return redirect(url_for("sku_mapping_list"))

    user_id = parse_int(request.form.get("user_id"))
    sku_id = parse_int(request.form.get("sku_id"))
    external_sku_code = request.form.get("external_sku_code", "").strip()
    remark = request.form.get("remark", "").strip()
    duplicate = (
        SKUMapping.query.filter(
            SKUMapping.user_id == user_id,
            SKUMapping.external_sku_code == external_sku_code,
            SKUMapping.id != mapping.id,
        ).first()
    )

    if user_id not in allowed_user_ids:
        flash("请选择有权限的运营账号。", "error")
    elif not external_sku_code:
        flash("店铺 SKU 不能为空。", "error")
    elif not sku_id or not db.session.get(SKU, sku_id):
        flash("请选择有效的仓库 SKU。", "error")
    elif not validate_sku_access([sku_id]):
        flash("无权映射该仓库 SKU。", "error")
    elif duplicate:
        flash("该运营下的店铺 SKU 已存在映射。", "error")
    else:
        previous_external_sku_code = mapping.external_sku_code
        mapping.user_id = user_id
        mapping.sku_id = sku_id
        mapping.external_sku_code = external_sku_code
        mapping.remark = remark
        log_operation("SKU映射", "更新", "sku_mapping", external_sku_code, f"原店铺SKU：{previous_external_sku_code}；运营ID：{user_id}；仓库SKU：{sku_id}")
        db.session.commit()
        flash("SKU 映射已更新。", "success")
    return redirect(
        url_for(
            "sku_mapping_list",
            mapping_keyword=request.args.get("mapping_keyword", "").strip(),
            mapping_user_id=parse_int(request.args.get("mapping_user_id")) or "",
        )
    )


@app.post("/sku-mappings/<int:mapping_id>/delete")
@login_required
@permission_required("sku.mapping.manage")
def sku_mapping_delete(mapping_id):
    mapping = db.session.get(SKUMapping, mapping_id)
    if not mapping:
        flash("SKU 映射不存在。", "error")
        return redirect(url_for("sku_mapping_list"))

    allowed_users = get_mapping_manageable_users([mapping.user_id])
    allowed_user_ids = {user.id for user in allowed_users}
    if mapping.user_id not in allowed_user_ids:
        flash("无权删除该 SKU 映射。", "error")
    else:
        log_operation("SKU映射", "删除", "sku_mapping", mapping.external_sku_code, f"运营：{mapping.user.full_name}；仓库SKU：{mapping.sku.sku_code}")
        db.session.delete(mapping)
        db.session.commit()
        flash("SKU 映射已删除。", "success")
    return redirect(
        url_for(
            "sku_mapping_list",
            mapping_keyword=request.args.get("mapping_keyword", "").strip(),
            mapping_user_id=parse_int(request.args.get("mapping_user_id")) or "",
        )
    )


@app.route("/suppliers", methods=["GET", "POST"])
@login_required
@permission_required("supplier.manage")
def supplier_list():
    if request.method == "POST":
        supplier = Supplier(
            name=request.form.get("name", "").strip(),
            contact_name=request.form.get("contact_name", "").strip(),
            phone=request.form.get("phone", "").strip(),
            address=request.form.get("address", "").strip(),
            remark=request.form.get("remark", "").strip(),
        )
        if not supplier.name:
            flash("供应商名称不能为空。", "error")
        elif Supplier.query.filter_by(name=supplier.name).first():
            flash("供应商名称已存在。", "error")
        else:
            db.session.add(supplier)
            log_operation("供应商", "新增", "supplier", supplier.name, f"联系人：{supplier.contact_name or '-'}")
            db.session.commit()
            flash("供应商创建成功。", "success")
        return redirect(url_for("supplier_list"))
    return render_template("suppliers.html", suppliers=Supplier.query.order_by(Supplier.created_at.desc()).all())


@app.post("/suppliers/<int:supplier_id>/delete")
@login_required
@permission_required("supplier.manage")
def supplier_delete(supplier_id):
    denied = require_delete_role()
    if denied:
        return denied
    supplier = db.session.get(Supplier, supplier_id)
    if not supplier:
        flash("供应商不存在。", "error")
    elif entity_has_rows(PurchaseOrder, supplier_id=supplier.id):
        flash("该供应商已被采购订单引用，请先删除相关采购单。", "error")
    elif entity_has_rows(InventoryTransaction, supplier_id=supplier.id):
        flash("该供应商已被入库记录引用，请先删除相关入库记录。", "error")
    else:
        log_operation("供应商", "删除", "supplier", supplier.name, f"联系人：{supplier.contact_name or '-'}")
        db.session.delete(supplier)
        db.session.commit()
        flash("供应商已删除。", "success")
    return redirect(url_for("supplier_list"))


@app.route("/customers", methods=["GET", "POST"])
@login_required
@permission_required("customer.manage")
def customer_list():
    if request.method == "POST":
        customer = Customer(
            name=request.form.get("name", "").strip(),
            contact_name=request.form.get("contact_name", "").strip(),
            phone=request.form.get("phone", "").strip(),
            address=request.form.get("address", "").strip(),
            remark=request.form.get("remark", "").strip(),
        )
        if not customer.name:
            flash("客户名称不能为空。", "error")
        elif Customer.query.filter_by(name=customer.name).first():
            flash("客户名称已存在。", "error")
        else:
            db.session.add(customer)
            log_operation("客户", "新增", "customer", customer.name, f"联系人：{customer.contact_name or '-'}")
            db.session.commit()
            flash("客户创建成功。", "success")
        return redirect(url_for("customer_list"))
    return render_template("customers.html", customers=Customer.query.order_by(Customer.created_at.desc()).all())


@app.post("/customers/<int:customer_id>/delete")
@login_required
@permission_required("customer.manage")
def customer_delete(customer_id):
    denied = require_delete_role()
    if denied:
        return denied
    customer = db.session.get(Customer, customer_id)
    if not customer:
        flash("客户不存在。", "error")
    elif entity_has_rows(SalesOrder, customer_id=customer.id):
        flash("该客户已被销售单引用，请先删除相关销售单。", "error")
    else:
        log_operation("客户", "删除", "customer", customer.name, f"联系人：{customer.contact_name or '-'}")
        db.session.delete(customer)
        db.session.commit()
        flash("客户已删除。", "success")
    return redirect(url_for("customer_list"))


@app.route("/warehouses", methods=["GET", "POST"])
@login_required
@permission_required("warehouse.manage")
def warehouse_list():
    if request.method == "POST":
        warehouse = Warehouse(
            name=request.form.get("name", "").strip(),
            code=request.form.get("code", "").strip().upper(),
            address=request.form.get("address", "").strip(),
            manager=request.form.get("manager", "").strip(),
            remark=request.form.get("remark", "").strip(),
        )
        if not warehouse.name or not warehouse.code:
            flash("仓库名称和编码不能为空。", "error")
        elif Warehouse.query.filter(or_(Warehouse.name == warehouse.name, Warehouse.code == warehouse.code)).first():
            flash("仓库名称或编码已存在。", "error")
        else:
            db.session.add(warehouse)
            log_operation("仓库", "新增", "warehouse", warehouse.code, warehouse.name)
            db.session.commit()
            flash("仓库创建成功。", "success")
        return redirect(url_for("warehouse_list"))

    warehouse_stock = defaultdict(int)
    for (warehouse_id, _sku_id), qty in get_inventory_by_warehouse().items():
        warehouse_stock[warehouse_id] += qty
    return render_template(
        "warehouses.html",
        warehouses=Warehouse.query.order_by(Warehouse.created_at.desc()).all(),
        warehouse_stock=warehouse_stock,
    )


@app.route("/assets", methods=["GET", "POST"])
@login_required
@permission_required("asset.manage")
def asset_list():
    if request.method == "POST":
        asset_code = request.form.get("asset_code", "").strip()
        name = request.form.get("name", "").strip()
        category = request.form.get("category", "").strip()
        status = request.form.get("status", "in_use").strip() or "in_use"
        purchase_date_raw = request.form.get("purchase_date", "").strip()
        purchase_date = parse_date(purchase_date_raw)
        asset_value = parse_decimal(request.form.get("asset_value"))
        asset_image = request.files.get("image_file")

        if not asset_code or not name:
            flash("资产编号和资产名称不能为空。", "error")
        elif Asset.query.filter_by(asset_code=asset_code).first():
            flash("资产编号已存在。", "error")
        elif purchase_date_raw and not purchase_date:
            flash("购入日期格式不正确。", "error")
        else:
            try:
                image_path = save_asset_image(asset_image)
                asset = Asset(
                    asset_code=asset_code,
                    name=name,
                    category=category,
                    brand=request.form.get("brand", "").strip(),
                    model=request.form.get("model", "").strip(),
                    serial_no=request.form.get("serial_no", "").strip(),
                    status=status,
                    keeper=request.form.get("keeper", "").strip(),
                    location=request.form.get("location", "").strip(),
                    purchase_date=purchase_date,
                    asset_value=asset_value,
                    image_path=image_path,
                    remark=request.form.get("remark", "").strip(),
                )
                db.session.add(asset)
                log_operation("资产", "新增", "asset", asset.asset_code, f"{asset.name} / {asset.category or '-'} / {asset.status_label}")
                db.session.commit()
                flash("资产已创建。", "success")
            except ValueError as error:
                flash(str(error), "error")
        return redirect(url_for("asset_list"))

    asset_keyword = request.args.get("asset_keyword", "").strip()
    status = request.args.get("status", "").strip()
    category = request.args.get("category", "").strip()
    assets_query = Asset.query
    if asset_keyword:
        keyword_like = f"%{asset_keyword}%"
        assets_query = assets_query.filter(
            or_(
                Asset.asset_code.ilike(keyword_like),
                Asset.name.ilike(keyword_like),
                Asset.brand.ilike(keyword_like),
                Asset.model.ilike(keyword_like),
                Asset.serial_no.ilike(keyword_like),
                Asset.keeper.ilike(keyword_like),
                Asset.location.ilike(keyword_like),
            )
        )
    if status:
        assets_query = assets_query.filter(Asset.status == status)
    if category:
        assets_query = assets_query.filter(Asset.category == category)

    categories = [
        row[0]
        for row in db.session.query(Asset.category).filter(Asset.category != "").distinct().order_by(Asset.category.asc()).all()
        if row[0]
    ]
    return render_template(
        "assets.html",
        assets=assets_query.order_by(Asset.created_at.desc()).all(),
        asset_keyword=asset_keyword,
        selected_status=status,
        selected_category=category,
        categories=categories,
    )


@app.post("/assets/<int:asset_id>/update")
@login_required
@permission_required("asset.manage")
def asset_update(asset_id):
    asset = db.session.get(Asset, asset_id)
    if not asset:
        flash("资产不存在。", "error")
        return redirect(url_for("asset_list"))

    asset_code = request.form.get("asset_code", "").strip()
    name = request.form.get("name", "").strip()
    purchase_date_raw = request.form.get("purchase_date", "").strip()
    purchase_date = parse_date(purchase_date_raw)
    asset_value = parse_decimal(request.form.get("asset_value"))
    asset_image = request.files.get("image_file")
    duplicate = Asset.query.filter(Asset.asset_code == asset_code, Asset.id != asset.id).first()

    if not asset_code or not name:
        flash("资产编号和资产名称不能为空。", "error")
    elif duplicate:
        flash("资产编号已存在。", "error")
    elif purchase_date_raw and not purchase_date:
        flash("购入日期格式不正确。", "error")
    else:
        try:
            new_image_path = save_asset_image(asset_image)
            previous_image_path = asset.image_path
            asset.asset_code = asset_code
            asset.name = name
            asset.category = request.form.get("category", "").strip()
            asset.brand = request.form.get("brand", "").strip()
            asset.model = request.form.get("model", "").strip()
            asset.serial_no = request.form.get("serial_no", "").strip()
            asset.status = request.form.get("status", "in_use").strip() or "in_use"
            asset.keeper = request.form.get("keeper", "").strip()
            asset.location = request.form.get("location", "").strip()
            asset.purchase_date = purchase_date
            asset.asset_value = asset_value
            if new_image_path:
                asset.image_path = new_image_path
            asset.remark = request.form.get("remark", "").strip()
            db.session.commit()
            if new_image_path and previous_image_path and previous_image_path != new_image_path:
                delete_asset_image(previous_image_path)
            flash("资产信息已更新。", "success")
        except ValueError as error:
            flash(str(error), "error")
    return redirect(
        url_for(
            "asset_list",
            asset_keyword=request.args.get("asset_keyword", "").strip(),
            status=request.args.get("status", "").strip(),
            category=request.args.get("category", "").strip(),
        )
    )


@app.post("/assets/<int:asset_id>/delete")
@login_required
@permission_required("asset.manage")
def asset_delete(asset_id):
    denied = require_delete_role()
    if denied:
        return denied
    asset = db.session.get(Asset, asset_id)
    if not asset:
        flash("资产不存在。", "error")
    else:
        image_path = asset.image_path
        log_operation("资产", "删除", "asset", asset.asset_code, f"{asset.name} / {asset.status_label}")
        db.session.delete(asset)
        db.session.commit()
        delete_asset_image(image_path)
        flash("资产已删除。", "success")
    return redirect(
        url_for(
            "asset_list",
            asset_keyword=request.args.get("asset_keyword", "").strip(),
            status=request.args.get("status", "").strip(),
            category=request.args.get("category", "").strip(),
        )
    )


@app.post("/warehouses/<int:warehouse_id>/delete")
@login_required
@permission_required("warehouse.manage")
def warehouse_delete(warehouse_id):
    denied = require_delete_role()
    if denied:
        return denied
    warehouse = db.session.get(Warehouse, warehouse_id)
    if not warehouse:
        flash("仓库不存在。", "error")
    elif entity_has_rows(PurchaseOrder, warehouse_id=warehouse.id) or entity_has_rows(SalesOrder, warehouse_id=warehouse.id):
        flash("该仓库已被采购单或销售单引用，请先删除相关单据。", "error")
    elif entity_has_rows(InventoryTransaction, warehouse_id=warehouse.id):
        flash("该仓库仍存在库存流水，请先删除相关入库或出库记录。", "error")
    else:
        InventoryBalance.query.filter_by(warehouse_id=warehouse.id).delete()
        log_operation("仓库", "删除", "warehouse", warehouse.code, warehouse.name)
        db.session.delete(warehouse)
        db.session.commit()
        flash("仓库已删除。", "success")
    return redirect(url_for("warehouse_list"))


@app.route("/purchase-calculator", methods=["GET", "POST"])
@login_required
def purchase_calculator():
    if not can_access_purchase_calculator():
        flash("你没有访问采购计算器的权限。", "error")
        return redirect(url_for("dashboard"))
    default_warehouse = get_default_warehouse()
    form = {
        "warehouse_id": request.values.get("warehouse_id", str(default_warehouse.id if default_warehouse else "")).strip(),
        "coverage_days": request.values.get("coverage_days", "30").strip() or "30",
        "lead_days": request.values.get("lead_days", "15").strip() or "15",
        "sku_keyword": request.values.get("sku_keyword", "").strip(),
        "suggested_min": request.values.get("suggested_min", "").strip(),
        "positive_only": request.values.get("positive_only", "1").strip(),
    }
    calc_result = get_purchase_calc_result()
    if request.method == "POST":
        warehouse_id = parse_int(form["warehouse_id"])
        warehouse = db.session.get(Warehouse, warehouse_id)
        sales_file = request.files.get("sales_30_file")
        if not warehouse:
            flash("请选择有效仓库后再计算。", "error")
        elif not sales_file or not (sales_file.filename or "").strip():
            flash("请上传最近30天销量表后再计算。", "error")
        else:
            try:
                sales_windows = build_purchase_sales_windows(sales_file, user=g.user)
                calc_result = build_purchase_calculator_result(
                    warehouse_id=warehouse.id,
                    coverage_days=form["coverage_days"],
                    lead_days=form["lead_days"],
                    sku_keyword=form["sku_keyword"],
                    suggested_min=form["suggested_min"],
                    positive_only=form["positive_only"] in {"1", "true", "on"},
                    user=g.user,
                    sales_windows=sales_windows,
                )
                calc_result["sales_window"] = {
                    "window_end_date": sales_windows.get("window_end_date", ""),
                    "latest_sale_date": sales_windows.get("latest_sale_date", ""),
                    "excluded_latest_date": sales_windows.get("excluded_latest_date", ""),
                    "unmatched_codes": sales_windows.get("unmatched_codes", []),
                }
                set_purchase_calc_result(calc_result)
                session["_purchase_calc_form"] = form
                if sales_windows.get("window_end_date"):
                    flash(
                        f"销量表已解析，按统计截止日 {sales_windows['window_end_date']} 汇总最近7天、14天和30天销量。",
                        "success",
                    )
                if sales_windows.get("unmatched_codes"):
                    flash(f"这些销量表 SKU 未匹配到仓库 SKU：{'、'.join(sales_windows['unmatched_codes'][:8])}", "error")
                flash(f"采购建议已生成，共 {calc_result['summary']['suggested_sku_count']} 个 SKU 建议采购。", "success")
            except ValueError as error:
                flash(str(error), "error")
    else:
        form = session.get("_purchase_calc_form") or form

    return render_template(
        "purchase_calculator.html",
        form=form,
        result=calc_result,
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        suppliers=Supplier.query.order_by(Supplier.name.asc()).all(),
    )


@app.get("/purchase-calculator/export")
@login_required
def purchase_calculator_export():
    if not can_access_purchase_calculator():
        flash("你没有导出采购计算结果的权限。", "error")
        return redirect(url_for("dashboard"))
    calc_result = get_purchase_calc_result()
    rows = calc_result.get("rows") or []
    if not rows:
        flash("当前没有可导出的采购计算结果，请先点击开始计算。", "error")
        return redirect(url_for("purchase_calculator"))
    export_rows = [
        [
            row.get("sku_code", ""),
            row.get("name", ""),
            f"{row.get('color', '')} / {row.get('size', '')}",
            row.get("current_stock", 0),
            row.get("safety_stock", 0),
            row.get("open_purchase_qty", 0),
            row.get("outbound_7", 0),
            row.get("outbound_14", 0),
            row.get("outbound_30", 0),
            row.get("outbound_daily", 0),
            row.get("sales_7", 0),
            row.get("sales_14", 0),
            row.get("sales_30", 0),
            row.get("sales_weighted_daily", 0),
            row.get("sales_trend_factor", 1),
            row.get("sales_trend_daily", 0),
            row.get("demand_source", ""),
            row.get("daily_consumption", 0),
            row.get("lead_consumption", 0),
            row.get("target_stock", 0),
            row.get("projected_available", 0),
            row.get("suggested_qty", 0),
        ]
        for row in rows
    ]
    return export_workbook(
        "采购计算结果.xlsx",
        "采购建议",
        ["SKU编码", "商品", "规格", "当前库存", "安全库存", "采购在途", "7天出库", "14天出库", "30天出库", "出库日均", "7天销量", "14天销量", "30天销量", "销量加权日均", "趋势系数", "销量趋势日均", "需求来源", "需求日均", "到货前消耗", "目标库存", "预计可用", "建议采购量"],
        export_rows,
    )


@app.post("/purchase-calculator/generate-order")
@login_required
@permission_required("purchase.manage")
def purchase_calculator_generate_order():
    calc_result = get_purchase_calc_result()
    row_by_sku_id = {str(row.get("sku_id")): row for row in calc_result.get("rows") or []}
    supplier_id = parse_int(request.form.get("supplier_id"))
    warehouse_id = parse_int(request.form.get("warehouse_id") or (calc_result.get("params") or {}).get("warehouse_id"))
    expected_date_raw = request.form.get("expected_date", "").strip()
    expected_date = parse_date(expected_date_raw)
    supplier = db.session.get(Supplier, supplier_id)
    warehouse = db.session.get(Warehouse, warehouse_id)
    if not supplier or not warehouse:
        flash("请选择供应商和仓库后再生成采购订单草稿。", "error")
        return redirect(url_for("purchase_calculator"))
    if expected_date_raw and not expected_date:
        flash("预计到货日期格式不正确。", "error")
        return redirect(url_for("purchase_calculator"))

    items = []
    for sku_id in request.form.getlist("purchase_calc_sku_id"):
        sku_id = str(parse_int(sku_id) or "")
        row = row_by_sku_id.get(sku_id)
        if not row:
            continue
        quantity = parse_int(request.form.get(f"purchase_calc_quantity__{sku_id}"))
        if quantity <= 0:
            continue
        sku = db.session.get(SKU, parse_int(sku_id))
        if not sku:
            continue
        items.append({"sku_id": sku.id, "quantity": quantity, "unit_price": sku.cost_price or Decimal("0")})
    if not items:
        flash("请至少勾选一条建议采购量大于 0 的 SKU。", "error")
        return redirect(url_for("purchase_calculator"))
    if not validate_sku_access([item["sku_id"] for item in items]):
        flash("存在未授权的 SKU，无法生成采购订单。", "error")
        return redirect(url_for("purchase_calculator"))

    order = PurchaseOrder(
        order_no=build_unique_code("PO", PurchaseOrder, "order_no"),
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        operator_name=get_current_operator_name(),
        status="draft",
        expected_date=expected_date,
        remark=request.form.get("remark", "").strip() or "由采购计算器生成的草稿",
    )
    for item in items:
        order.items.append(PurchaseOrderItem(**item))
    db.session.add(order)
    log_operation("采购计算器", "生成草稿", "purchase_order", order.order_no, f"供应商：{supplier.name}；仓库：{warehouse.name}；明细：{len(items)}")
    db.session.commit()
    flash(f"已生成采购订单草稿 {order.order_no}，请在采购订单页复核后提交。", "success")
    return redirect(url_for("purchase_order_list"))


@app.route("/purchase-orders", methods=["GET", "POST"])
@login_required
@permission_required("purchase.manage")
def purchase_order_list():
    if request.method == "POST":
        items = parse_order_items("purchase", allow_price=can_view_prices())
        supplier_id = parse_int(request.form.get("supplier_id"))
        warehouse_id = parse_int(request.form.get("warehouse_id"))
        supplier = db.session.get(Supplier, supplier_id)
        warehouse = db.session.get(Warehouse, warehouse_id)
        item_skus = validate_entity_ids(SKU, [item["sku_id"] for item in items])
        expected_date_raw = request.form.get("expected_date", "").strip()
        expected_date = parse_date(expected_date_raw)

        if not items:
            flash("请至少添加一条采购明细。", "error")
            return redirect(url_for("purchase_order_list"))
        if not supplier or not warehouse:
            flash("供应商和仓库不能为空。", "error")
            return redirect(url_for("purchase_order_list"))
        if len(item_skus) != len({item["sku_id"] for item in items}):
            flash("所选 SKU 中存在无效项。", "error")
            return redirect(url_for("purchase_order_list"))
        if expected_date_raw and not expected_date:
            flash("预计到货日期格式不正确。", "error")
            return redirect(url_for("purchase_order_list"))

        order = PurchaseOrder(
            order_no=request.form.get("order_no", "").strip() or build_unique_code("PO", PurchaseOrder, "order_no"),
            supplier_id=supplier_id,
            warehouse_id=warehouse_id,
            operator_name=get_current_operator_name(),
            status=normalize_status(request.form.get("status", "submitted"), PURCHASE_STATUSES - {"received"}, "submitted"),
            remark=request.form.get("remark", "").strip(),
            expected_date=expected_date,
        )
        for item in items:
            order.items.append(PurchaseOrderItem(**item))
        db.session.add(order)
        log_operation("采购订单", "新增", "purchase_order", order.order_no, f"供应商：{supplier.name}；仓库：{warehouse.name}；明细：{len(items)}")
        db.session.commit()
        flash("采购订单创建成功。", "success")
        return redirect(url_for("purchase_order_list"))

    sku_keyword, start_date, end_date = get_date_filters()
    orders_query = PurchaseOrder.query
    orders_query = apply_order_query_sku_scope(orders_query, PurchaseOrderItem, PurchaseOrder.items)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        orders_query = orders_query.join(PurchaseOrder.items).join(PurchaseOrderItem.sku).filter(sku_filter).distinct()
    orders_query = apply_datetime_range(orders_query, PurchaseOrder.created_at, start_date, end_date)

    return render_template(
        "purchase_orders.html",
        orders=orders_query.order_by(PurchaseOrder.created_at.desc()).all(),
        suppliers=Supplier.query.order_by(Supplier.name.asc()).all(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        skus=SKU.query.order_by(SKU.name.asc()).all(),
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.get("/purchase-orders/<int:order_id>/export")
@login_required
@permission_required("purchase.manage")
def purchase_order_export(order_id):
    order = db.session.get(PurchaseOrder, order_id)
    if not order:
        flash("采购订单不存在。", "error")
        return redirect(url_for("purchase_order_list"))
    if not validate_sku_access([item.sku_id for item in order.items]):
        flash("无权导出该采购订单。", "error")
        return redirect(url_for("purchase_order_list"))
    return export_purchase_order_workbook(order)


@app.post("/purchase-orders/<int:order_id>/delete")
@login_required
@permission_required("purchase.manage")
def purchase_order_delete(order_id):
    denied = require_delete_role()
    if denied:
        return denied
    order = db.session.get(PurchaseOrder, order_id)
    if not order:
        flash("采购订单不存在。", "error")
        return redirect(url_for("purchase_order_list"))

    deleted_txn_count = delete_inventory_transactions(reference_type="purchase_order", reference_no=order.order_no)
    log_operation("采购订单", "删除", "purchase_order", order.order_no, f"供应商：{order.supplier.name}；回滚流水：{deleted_txn_count}")
    db.session.delete(order)
    db.session.commit()
    if deleted_txn_count:
        flash("采购订单已删除，关联入库记录和库存余额已回滚。", "success")
    else:
        flash("采购订单已删除。", "success")
    return redirect(url_for("purchase_order_list"))


@app.post("/purchase-orders/<int:order_id>/receive")
@login_required
@permission_required("stock.in")
def purchase_order_receive(order_id):
    order = db.session.get(PurchaseOrder, order_id)
    if not order:
        flash("采购订单不存在。", "error")
        return redirect(url_for("purchase_order_list"))
    if order.status == "received":
        flash("该采购单已经收货，不能重复操作。", "error")
        return redirect(url_for("purchase_order_list"))

    try:
        operator_name = get_current_operator_name()
        for item in order.items:
            add_inventory_transaction(
                item.sku_id,
                order.warehouse_id,
                item.quantity,
                "purchase_inbound",
                reference_type="purchase_order",
                reference_no=order.order_no,
                note=f"供应商：{order.supplier.name}",
                operator_name=operator_name,
                supplier_id=order.supplier_id,
            )
        order.status = "received"
        db.session.commit()
        flash("采购订单已收货，库存已更新。", "success")
    except InventoryError as error:
        db.session.rollback()
        flash(str(error), "error")
    return redirect(url_for("purchase_order_list"))


@app.route("/inventory/in", methods=["GET", "POST"])
@login_required
@permission_required("stock.in")
def inventory_in():
    if request.method == "POST":
        items = parse_order_items("manual_inbound", allow_price=can_view_prices())
        warehouse_id = parse_int(request.form.get("warehouse_id"))
        supplier_id = parse_int(request.form.get("supplier_id"))
        if not items:
            legacy_sku_id = parse_int(request.form.get("sku_id"))
            legacy_quantity = parse_int(request.form.get("quantity"))
            if legacy_sku_id and legacy_quantity > 0:
                items = [{"sku_id": legacy_sku_id, "quantity": legacy_quantity, "unit_price": Decimal("0")}]
        warehouse = db.session.get(Warehouse, warehouse_id)
        supplier = db.session.get(Supplier, supplier_id)

        if not items:
            flash("请至少添加一条入库明细。", "error")
            return redirect(url_for("inventory_in"))
        if not warehouse:
            flash("请选择有效的 SKU、仓库、供应商，并填写正确数量。", "error")
            return redirect(url_for("inventory_in"))
        if not supplier:
            flash("请选择供应商。", "error")
            return redirect(url_for("inventory_in"))

        reference_type = request.form.get("reference_type", "").strip() or "manual_inbound_order"
        reference_no = request.form.get("reference_no", "").strip() or build_unique_code("MI", InventoryTransaction, "reference_no")
        note = request.form.get("note", "").strip()
        operator_name = get_current_operator_name()
        for item in items:
            if not db.session.get(SKU, item["sku_id"]):
                flash("存在无效的 SKU，无法完成入库。", "error")
                return redirect(url_for("inventory_in"))
            add_inventory_transaction(
                item["sku_id"],
                warehouse_id,
                item["quantity"],
                "manual_inbound",
                reference_type=reference_type,
                reference_no=reference_no,
                note=note,
                operator_name=operator_name,
                supplier_id=supplier.id,
            )
        sku = db.session.get(SKU, items[0]["sku_id"])
        quantity = sum(item["quantity"] for item in items)
        log_operation("手工入库", "新增", "inventory_transaction", sku.sku_code, f"供应商：{supplier.name}；仓库：{warehouse.name}；数量：+{quantity}")
        db.session.commit()
        flash("入库操作成功。", "success")
        return redirect(url_for("inventory_in"))

    sku_keyword, start_date, end_date = get_date_filters()
    txns_query = InventoryTransaction.query.join(InventoryTransaction.sku).filter(InventoryTransaction.quantity > 0)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        txns_query = txns_query.filter(sku_filter)
    txns_query = apply_datetime_range(txns_query, InventoryTransaction.created_at, start_date, end_date)
    txns = txns_query.order_by(InventoryTransaction.created_at.desc()).limit(100).all()
    inbound_summary = {
        "record_count": len(txns),
        "quantity": sum(txn.quantity for txn in txns),
        "sku_count": len({txn.sku_id for txn in txns}),
        "warehouse_count": len({txn.warehouse_id for txn in txns}),
    }
    return render_template(
        "inventory_in.html",
        skus=SKU.query.order_by(SKU.name.asc()).all(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        suppliers=Supplier.query.order_by(Supplier.name.asc()).all(),
        txns=txns,
        inbound_summary=inbound_summary,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.post("/inventory/transactions/<int:txn_id>/delete")
@login_required
@permission_required("stock.in")
def inventory_transaction_delete(txn_id):
    transaction = db.session.get(InventoryTransaction, txn_id)
    if not transaction:
        flash("库存流水不存在。", "error")
    elif not validate_sku_access([transaction.sku_id]):
        flash("无权操作该 SKU 的库存流水。", "error")
    elif transaction.reference_type in {"purchase_order", "sales_order"}:
        flash("这条库存流水来自业务单据，请到对应采购单或销售单中处理。", "error")
    else:
        target_label = transaction.reference_no or f"TXN-{transaction.id}"
        warehouse_name = transaction.warehouse.name if transaction.warehouse else "-"
        group_delete_reference_types = {"shipment_sheet", "defect_repair"}
        if transaction.reference_no and transaction.reference_type in group_delete_reference_types:
            deleted_count = delete_inventory_transactions(
                reference_type=transaction.reference_type,
                reference_no=transaction.reference_no,
            )
        else:
            deleted_count = delete_inventory_transactions(transaction_ids=[txn_id])
        log_operation("库存流水", "删除", "inventory_transaction", target_label, f"仓库：{warehouse_name}；数量：{deleted_count} 条")
        db.session.commit()
        sync_inventory_balances()
        db.session.expire_all()
        if deleted_count > 1:
            flash("库存单据已删除，库存余额已重新计算。", "success")
        else:
            flash("库存流水已删除，库存余额已重新计算。", "success")
    return redirect(request.referrer or url_for("inventory_transactions"))


@app.route("/sales-orders", methods=["GET", "POST"])
@login_required
@permission_required("sales.manage")
def sales_order_list():
    if request.method == "POST":
        if request.form.get("form_action") == "import":
            upload = request.files.get("import_file")
            if not upload or not (upload.filename or "").strip():
                flash("请选择要导入的 CSV、TSV 或 Excel 文件。", "error")
                return render_sales_order_page()
            try:
                sales_form = parse_shipment_import(upload, include_customer=True, user=g.user)
                if not validate_sku_access([item["sku_id"] for item in sales_form.get("items", [])]):
                    flash("导入数据中存在未授权的 SKU。", "error")
                    return render_sales_order_page()
                flash(f"已导入文件：{sales_form['source_name']}，请确认仓库和明细后再保存。", "success")
                return render_sales_order_page(sales_form=sales_form)
            except ValueError as error:
                flash(str(error), "error")
                return render_sales_order_page()

        items = parse_order_items("sales", allow_price=can_view_prices())
        if not validate_sku_access([item["sku_id"] for item in items]):
            flash("存在未授权的 SKU，无法创建销售单。", "error")
            return redirect(url_for("sales_order_list"))
        customer_id = parse_int(request.form.get("customer_id"))
        warehouse_id = parse_int(request.form.get("warehouse_id"))
        customer = db.session.get(Customer, customer_id) if customer_id else None
        warehouse = db.session.get(Warehouse, warehouse_id)
        item_skus = validate_entity_ids(SKU, [item["sku_id"] for item in items])

        if not items:
            flash("请至少添加一条销售明细。", "error")
            return redirect(url_for("sales_order_list"))
        if not warehouse:
            flash("请选择有效仓库。", "error")
            return redirect(url_for("sales_order_list"))
        if len(item_skus) != len({item["sku_id"] for item in items}):
            flash("所选 SKU 中存在无效项。", "error")
            return redirect(url_for("sales_order_list"))

        customer_name = customer.name if customer else request.form.get("customer_name", "").strip()
        phone = customer.phone if customer else request.form.get("phone", "").strip()
        shipping_address = customer.address if customer else request.form.get("shipping_address", "").strip()
        if not customer_name:
            flash("客户名称不能为空。", "error")
            return redirect(url_for("sales_order_list"))

        if not customer:
            customer = Customer.query.filter_by(name=customer_name).first()
            if not customer:
                customer = Customer(
                    name=customer_name,
                    contact_name="",
                    phone=phone,
                    address=shipping_address,
                    remark="Auto-created from sales import",
                )
                db.session.add(customer)
                db.session.flush()
                log_operation("customer", "create", "customer", customer.name, "Auto-created from sales import")
            customer_id = customer.id

        order = SalesOrder(
            order_no=request.form.get("order_no", "").strip() or build_unique_code("SO", SalesOrder, "order_no"),
            customer_id=customer_id if customer_id else None,
            customer_name=customer_name,
            phone=phone,
            warehouse_id=warehouse_id,
            operator_name=get_current_operator_name(),
            status=normalize_status(request.form.get("status", "confirmed"), SALES_STATUSES - {"shipped"}, "confirmed"),
            shipping_address=shipping_address,
            remark=request.form.get("remark", "").strip(),
        )
        for item in items:
            order.items.append(SalesOrderItem(**item))
        db.session.add(order)
        log_operation("销售订单", "新增", "sales_order", order.order_no, f"客户：{customer_name}；仓库：{warehouse.name}；明细：{len(items)}")
        db.session.commit()
        flash("销售单创建成功。", "success")
        return redirect(url_for("sales_order_list"))

    return render_sales_order_page()


@app.post("/sales-orders/<int:order_id>/delete")
@login_required
@permission_required("sales.manage")
def sales_order_delete(order_id):
    denied = require_delete_role()
    if denied:
        return denied
    order = db.session.get(SalesOrder, order_id)
    if not order:
        flash("销售单不存在。", "error")
        return redirect(url_for("sales_order_list"))
    if not validate_sku_access([item.sku_id for item in order.items]):
        flash("无权操作该销售单。", "error")
        return redirect(url_for("sales_order_list"))

    deleted_txn_count = delete_inventory_transactions(reference_type="sales_order", reference_no=order.order_no)
    log_operation("销售订单", "删除", "sales_order", order.order_no, f"客户：{order.customer_name}；回滚流水：{deleted_txn_count}")
    db.session.delete(order)
    db.session.commit()
    if deleted_txn_count:
        flash("销售单已删除，关联出库记录和库存余额已回滚。", "success")
    else:
        flash("销售单已删除。", "success")
    return redirect(url_for("sales_order_list"))


@app.post("/sales-orders/<int:order_id>/ship")
@login_required
def sales_order_ship(order_id):
    if not can_ship_sales_orders():
        flash("您没有确认销售订单发货的权限。", "error")
        return redirect(url_for("sales_order_list"))
    order = db.session.get(SalesOrder, order_id)
    if not order:
        flash("销售单不存在。", "error")
        return redirect(url_for("sales_order_list"))
    if not validate_sku_access([item.sku_id for item in order.items]):
        flash("无权操作该销售单。", "error")
        return redirect(url_for("sales_order_list"))
    if order.status == "shipped":
        flash("该销售单已经发货，不能重复操作。", "error")
        return redirect(url_for("sales_order_list"))

    try:
        operator_name = get_current_operator_name()
        for item in order.items:
            add_inventory_transaction(
                item.sku_id,
                order.warehouse_id,
                -item.quantity,
                "sales_outbound",
                reference_type="sales_order",
                reference_no=order.order_no,
                note=f"客户：{order.customer_name}",
                operator_name=operator_name,
            )
        order.status = "shipped"
        db.session.commit()
        flash("销售单已发货，库存已更新。", "success")
    except InventoryError:
        db.session.rollback()
        flash("该仓库中一个或多个商品库存不足，无法发货。", "error")
    return redirect(url_for("sales_order_list"))


@app.route("/inventory/out", methods=["GET", "POST"])
@login_required
@permission_required("stock.out")
def inventory_out():
    if request.method == "POST":
        if request.form.get("form_action") == "import":
            upload = request.files.get("import_file")
            if not upload or not (upload.filename or "").strip():
                flash("请选择要导入的 CSV、TSV 或 Excel 文件。", "error")
                return render_inventory_out_page()
            try:
                outbound_form = parse_shipment_import(upload, include_customer=False, user=g.user)
                if not validate_sku_access([item["sku_id"] for item in outbound_form.get("items", [])]):
                    flash("导入数据中存在未授权的 SKU。", "error")
                    return render_inventory_out_page()
                flash(f"已导入文件：{outbound_form['source_name']}，请选择仓库后再保存。", "success")
                return render_inventory_out_page(outbound_form=outbound_form)
            except ValueError as error:
                flash(str(error), "error")
                return render_inventory_out_page()

        warehouse_id = parse_int(request.form.get("warehouse_id"))
        items = parse_order_items("manual_outbound", allow_price=can_view_prices())
        if not items:
            legacy_sku_id = parse_int(request.form.get("sku_id"))
            legacy_quantity = parse_int(request.form.get("quantity"))
            if legacy_sku_id and legacy_quantity > 0:
                items = [{"sku_id": legacy_sku_id, "quantity": legacy_quantity, "unit_price": Decimal("0")}]
        warehouse = db.session.get(Warehouse, warehouse_id)
        if not validate_sku_access([item["sku_id"] for item in items]):
            flash("存在未授权的 SKU，无法创建出库单。", "error")
            return redirect(url_for("inventory_out"))

        if not items:
            flash("请至少添加一条出库明细。", "error")
            return redirect(url_for("inventory_out"))
        if not warehouse:
            flash("请选择有效的仓库。", "error")
            return redirect(url_for("inventory_out"))

        try:
            reference_type = request.form.get("reference_type", "").strip() or "manual_outbound_order"
            reference_no = request.form.get("reference_no", "").strip() or build_unique_code("MO", InventoryTransaction, "reference_no")
            note = request.form.get("note", "").strip()
            operator_name = get_current_operator_name()
            for item in items:
                if not db.session.get(SKU, item["sku_id"]):
                    raise InventoryError("存在无效的 SKU，无法完成出库。")
                add_inventory_transaction(
                    item["sku_id"],
                    warehouse_id,
                    -item["quantity"],
                    "manual_outbound",
                    reference_type=reference_type,
                    reference_no=reference_no,
                    note=note,
                    operator_name=operator_name,
                )
            log_operation("手工出库", "新增", "inventory_transaction", reference_no, f"仓库：{warehouse.name}；明细：{len(items)}")
            db.session.commit()
            flash("手工出库单创建成功。", "success")
        except InventoryError as error:
            db.session.rollback()
            flash(str(error) if str(error) else "所选仓库库存不足。", "error")
        return redirect(url_for("inventory_out"))

    return render_inventory_out_page()


@app.route("/defects", methods=["GET", "POST"])
@login_required
@permission_required("defect.manage")
def defect_list():
    if request.method == "POST":
        sku_id = parse_int(request.form.get("sku_id"))
        if not validate_sku_access([sku_id]):
            flash("该 SKU 未授权给当前销售员。", "error")
            return redirect(url_for("defect_list"))
        warehouse_id = parse_int(request.form.get("warehouse_id"))
        warehouse = db.session.get(Warehouse, warehouse_id) if warehouse_id else None
        repair = DefectRepair(
            repair_no=request.form.get("repair_no", "").strip() or build_unique_code("DR", DefectRepair, "repair_no"),
            sku_id=sku_id,
            warehouse_id=warehouse_id if warehouse else None,
            operator_name=get_current_operator_name(),
            quantity=parse_int(request.form.get("quantity")),
            defect_reason=request.form.get("defect_reason", "").strip(),
            repair_result=request.form.get("repair_result", "").strip(),
        )
        new_status = normalize_status(request.form.get("status", "pending"), DEFECT_STATUSES, "pending")
        if not db.session.get(SKU, sku_id):
            flash("请选择有效的 SKU。", "error")
        elif not warehouse:
            flash("请选择有效的仓库。", "error")
        elif repair.quantity <= 0 or not repair.defect_reason:
            flash("请完整填写返修单信息。", "error")
        else:
            try:
                apply_defect_repair_progress(
                    repair,
                    new_status,
                    completed_quantity=request.form.get("completed_quantity") if "completed_quantity" in request.form else None,
                    scrapped_quantity=request.form.get("scrapped_quantity") if "scrapped_quantity" in request.form else None,
                )
                db.session.add(repair)
                db.session.flush()
                sync_defect_repair_inventory(repair)
                log_operation("返修单", "新增", "defect_repair", repair.repair_no, f"SKU ID：{sku_id}；仓库：{warehouse.name}；数量：{repair.quantity}；状态：{repair.status_label}")
                db.session.commit()
                flash("返修单创建成功。", "success")
            except InventoryError as error:
                db.session.rollback()
                flash(str(error) if str(error) else "库存不足，无法保存返修单。", "error")
            except ValueError as error:
                db.session.rollback()
                flash(str(error), "error")
        return redirect(url_for("defect_list"))

    sku_keyword, start_date, end_date = get_date_filters()
    repairs_query = DefectRepair.query.join(DefectRepair.sku)
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if sku_filter is not None:
        repairs_query = repairs_query.filter(sku_filter)
    repairs_query = apply_datetime_range(repairs_query, DefectRepair.created_at, start_date, end_date)
    page = max(parse_int(request.args.get("page"), 1), 1)
    repairs_pagination = repairs_query.order_by(DefectRepair.created_at.desc()).paginate(
        page=page,
        per_page=25,
        error_out=False,
    )
    return render_template(
        "defects.html",
        repairs=repairs_pagination.items,
        repairs_pagination=repairs_pagination,
        skus=SKU.query.order_by(SKU.name.asc()).all(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.post("/defects/<int:repair_id>/update")
@login_required
@permission_required("defect.manage")
def defect_update(repair_id):
    repair = db.session.get(DefectRepair, repair_id)
    if not repair:
        flash("返修单不存在。", "error")
        return redirect(url_for("defect_list"))
    if not validate_sku_access([repair.sku_id]):
        flash("无权操作该返修单。", "error")
        return redirect(url_for("defect_list"))

    new_status = normalize_status(request.form.get("status", repair.status), DEFECT_STATUSES, repair.status)
    repair_result = request.form.get("repair_result", "").strip()
    previous_status = repair.status
    previous_completed_quantity = repair.completed_quantity
    previous_scrapped_quantity = repair.scrapped_quantity

    try:
        apply_defect_repair_progress(
            repair,
            new_status,
            completed_quantity=request.form.get("completed_quantity") if "completed_quantity" in request.form else None,
            scrapped_quantity=request.form.get("scrapped_quantity") if "scrapped_quantity" in request.form else None,
        )
        repair.repair_result = repair_result
        sync_defect_repair_inventory(repair)
        log_operation("返修单", "更新", "defect_repair", repair.repair_no, f"状态：{previous_status} -> {repair.status}；成功：{previous_completed_quantity} -> {repair.completed_quantity}；报废：{previous_scrapped_quantity} -> {repair.scrapped_quantity}；仓库：{repair.warehouse.name if repair.warehouse else '-'}")
        db.session.commit()
        flash("返修单状态已更新。", "success")
    except InventoryError as error:
        db.session.rollback()
        flash(str(error) if str(error) else "库存不足，无法更新返修单状态。", "error")
    except ValueError as error:
        db.session.rollback()
        flash(str(error), "error")
    return redirect(url_for("defect_list"))


@app.post("/defects/<int:repair_id>/delete")
@login_required
@permission_required("defect.manage")
def defect_delete(repair_id):
    denied = require_delete_role()
    if denied:
        return denied
    repair = db.session.get(DefectRepair, repair_id)
    if not repair:
        flash("返修单不存在。", "error")
        return redirect(url_for("defect_list"))
    if not validate_sku_access([repair.sku_id]):
        flash("无权操作该返修单。", "error")
        return redirect(url_for("defect_list"))
    if repair.status not in {"pending", "repairing"}:
        flash("只有待处理或返修中的返修单才能删除；已完成或已报废单据保留记录。", "error")
        return redirect(url_for("defect_list"))

    deleted_txn_count = delete_inventory_transactions(reference_type="defect_repair", reference_no=repair.repair_no)
    warehouse_name = repair.warehouse.name if repair.warehouse else "-"
    log_operation("返修单", "删除", "defect_repair", repair.repair_no, f"仓库：{warehouse_name}；回滚流水：{deleted_txn_count}")
    db.session.delete(repair)
    db.session.commit()
    if deleted_txn_count:
        flash("返修单已删除，关联库存流水和库存余额已回滚。", "success")
    else:
        flash("返修单已删除。", "success")
    return redirect(url_for("defect_list"))


@app.route("/inventory")
@login_required
@permission_required("inventory.view")
def inventory():
    selected_warehouse_id = parse_int(request.args.get("warehouse_id"))
    sku_keyword, start_date, end_date = get_date_filters()
    sku_filter = build_sku_keyword_filter(sku_keyword)
    stock_rows = []
    sku_query = apply_sku_scope(SKU.query, SKU.id)
    if sku_filter is not None:
        sku_query = sku_query.filter(sku_filter)

    if start_date or end_date:
        summary_query = db.session.query(InventoryTransaction.sku_id, func.coalesce(func.sum(InventoryTransaction.quantity), 0))
        if selected_warehouse_id:
            summary_query = summary_query.filter(InventoryTransaction.warehouse_id == selected_warehouse_id)
        if sku_filter is not None:
            summary_query = summary_query.join(InventoryTransaction.sku).filter(sku_filter)
        summary_query = apply_datetime_range(summary_query, InventoryTransaction.created_at, start_date, end_date)
        inventory_map = {sku_id: qty for sku_id, qty in summary_query.group_by(InventoryTransaction.sku_id).all()}
    else:
        inventory_map = get_inventory_map(selected_warehouse_id if selected_warehouse_id else None)

    defect_status_map = get_defect_status_map(
        warehouse_id=selected_warehouse_id if selected_warehouse_id else None,
        sku_filter=sku_filter,
        start_date=start_date,
        end_date=end_date,
    )

    for sku in sku_query.order_by(SKU.category.asc(), SKU.name.asc()).all():
        current_qty = inventory_map.get(sku.id, 0)
        stock_rows.append(
            {
                "sku": sku,
                "quantity": current_qty,
                "pending_quantity": defect_status_map.get((sku.id, "pending"), 0),
                "repairing_quantity": defect_status_map.get((sku.id, "repairing"), 0),
                "done_quantity": defect_status_map.get((sku.id, "done"), 0),
                "scrapped_quantity": defect_status_map.get((sku.id, "scrapped"), 0),
                "stock_value": current_qty * float(sku.cost_price or 0),
                "low_stock": current_qty <= sku.safety_stock,
            }
        )
    inventory_summary = {
        "sku_count": len(stock_rows),
        "quantity": sum(row["quantity"] for row in stock_rows),
        "pending_quantity": sum(row["pending_quantity"] for row in stock_rows),
        "repairing_quantity": sum(row["repairing_quantity"] for row in stock_rows),
        "done_quantity": sum(row["done_quantity"] for row in stock_rows),
        "scrapped_quantity": sum(row["scrapped_quantity"] for row in stock_rows),
        "safety_stock": sum(row["sku"].safety_stock for row in stock_rows),
        "stock_value": sum(row["stock_value"] for row in stock_rows),
        "low_stock_count": sum(1 for row in stock_rows if row["low_stock"]),
    }
    return render_template(
        "inventory.html",
        stock_rows=stock_rows,
        inventory_summary=inventory_summary,
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        selected_warehouse_id=selected_warehouse_id,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.route("/inventory/transactions")
@login_required
@permission_required("inventory.transaction.view")
def inventory_transactions():
    selected_warehouse_id = parse_int(request.args.get("warehouse_id"))
    sku_keyword, start_date, end_date = get_date_filters()
    sku_filter = build_sku_keyword_filter(sku_keyword)

    txns_query = apply_sku_scope(InventoryTransaction.query.join(InventoryTransaction.sku), InventoryTransaction.sku_id)
    if selected_warehouse_id:
        txns_query = txns_query.filter(InventoryTransaction.warehouse_id == selected_warehouse_id)
    if sku_filter is not None:
        txns_query = txns_query.filter(sku_filter)
    txns_query = apply_datetime_range(txns_query, InventoryTransaction.created_at, start_date, end_date)

    return render_template(
        "inventory_transactions.html",
        recent_txns=txns_query.order_by(InventoryTransaction.created_at.desc()).limit(200).all(),
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        selected_warehouse_id=selected_warehouse_id,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.route("/scan")
@login_required
@permission_required("inventory.view")
def scan_lookup():
    q = request.args.get("q", "").strip()
    sku = None
    warehouse_rows = []
    if q:
        sku_query = apply_sku_scope(SKU.query, SKU.id)
        sku = sku_query.filter(or_(SKU.barcode == q, SKU.sku_code == q, SKU.name.ilike(f"%{q}%"))).first()
        if sku:
            stock_map = get_inventory_by_warehouse()
            for warehouse in Warehouse.query.order_by(Warehouse.name.asc()).all():
                warehouse_rows.append({"warehouse": warehouse, "quantity": stock_map.get((warehouse.id, sku.id), 0)})
    return render_template("scan.html", q=q, sku=sku, warehouse_rows=warehouse_rows)


@app.route("/reports")
@login_required
@permission_required("report.view")
def reports():
    show_prices = can_view_prices()
    selected_warehouse_id = parse_int(request.args.get("warehouse_id"))
    sku_keyword, start_date, end_date = get_date_filters()
    sku_filter = build_sku_keyword_filter(sku_keyword)
    inventory_map = get_inventory_by_warehouse()
    purchase_rows = (
        db.session.query(
            PurchaseOrder.order_no,
            Supplier.name,
            Warehouse.name,
            PurchaseOrder.status,
            func.sum(PurchaseOrderItem.quantity * PurchaseOrderItem.unit_price),
        )
        .join(Supplier, PurchaseOrder.supplier_id == Supplier.id)
        .join(Warehouse, PurchaseOrder.warehouse_id == Warehouse.id)
        .join(PurchaseOrderItem, PurchaseOrderItem.purchase_order_id == PurchaseOrder.id)
    )
    if selected_warehouse_id:
        purchase_rows = purchase_rows.filter(PurchaseOrder.warehouse_id == selected_warehouse_id)
    if sku_filter is not None:
        purchase_rows = purchase_rows.join(PurchaseOrderItem.sku).filter(sku_filter)
    purchase_rows = (
        apply_datetime_range(purchase_rows, PurchaseOrder.created_at, start_date, end_date)
        .group_by(PurchaseOrder.id, Supplier.name, Warehouse.name)
        .order_by(PurchaseOrder.created_at.desc())
        .limit(10)
        .all()
    )
    sales_rows = (
        db.session.query(
            SalesOrder.order_no,
            SalesOrder.customer_name,
            Warehouse.name,
            SalesOrder.status,
            func.sum(SalesOrderItem.quantity * SalesOrderItem.unit_price),
        )
        .join(Warehouse, SalesOrder.warehouse_id == Warehouse.id)
        .join(SalesOrderItem, SalesOrderItem.sales_order_id == SalesOrder.id)
    )
    if selected_warehouse_id:
        sales_rows = sales_rows.filter(SalesOrder.warehouse_id == selected_warehouse_id)
    if sku_filter is not None:
        sales_rows = sales_rows.join(SalesOrderItem.sku).filter(sku_filter)
    sales_rows = (
        apply_datetime_range(sales_rows, SalesOrder.created_at, start_date, end_date)
        .group_by(SalesOrder.id, Warehouse.name)
        .order_by(SalesOrder.created_at.desc())
        .limit(10)
        .all()
    )
    report_summary = {
        "purchase_count": len(purchase_rows),
        "purchase_amount": float(sum(row[4] or 0 for row in purchase_rows)),
        "sales_count": len(sales_rows),
        "sales_amount": float(sum(row[4] or 0 for row in sales_rows)),
    }
    inventory_value = 0
    warehouse_stats = defaultdict(lambda: {"qty": 0, "value": 0})
    if start_date or end_date or sku_filter is not None:
        inventory_summary_query = db.session.query(
            InventoryTransaction.warehouse_id,
            InventoryTransaction.sku_id,
            func.coalesce(func.sum(InventoryTransaction.quantity), 0),
        ).join(InventoryTransaction.sku)
        if selected_warehouse_id:
            inventory_summary_query = inventory_summary_query.filter(InventoryTransaction.warehouse_id == selected_warehouse_id)
        if sku_filter is not None:
            inventory_summary_query = inventory_summary_query.filter(sku_filter)
        inventory_summary_query = apply_datetime_range(inventory_summary_query, InventoryTransaction.created_at, start_date, end_date)
        inventory_map = {
            (warehouse_id, sku_id): qty
            for warehouse_id, sku_id, qty in inventory_summary_query.group_by(InventoryTransaction.warehouse_id, InventoryTransaction.sku_id).all()
        }
    skus = SKU.query.filter(sku_filter).all() if sku_filter is not None else SKU.query.all()
    warehouse_query = Warehouse.query
    if selected_warehouse_id:
        warehouse_query = warehouse_query.filter(Warehouse.id == selected_warehouse_id)
    for warehouse in warehouse_query.all():
        for sku in skus:
            qty = inventory_map.get((warehouse.id, sku.id), 0)
            value = qty * float(sku.cost_price or 0)
            if show_prices:
                inventory_value += value
            warehouse_stats[warehouse.name]["qty"] += qty
            if show_prices:
                warehouse_stats[warehouse.name]["value"] += value
    return render_template(
        "reports.html",
        purchase_rows=purchase_rows,
        sales_rows=sales_rows,
        report_summary=report_summary,
        inventory_value=inventory_value,
        warehouse_stats=warehouse_stats,
        warehouses=Warehouse.query.order_by(Warehouse.name.asc()).all(),
        selected_warehouse_id=selected_warehouse_id,
        sku_keyword=sku_keyword,
        start_date=start_date.isoformat() if start_date else "",
        end_date=end_date.isoformat() if end_date else "",
    )


@app.route("/reports/export/<string:report_type>")
@login_required
@permission_required("report.view")
def report_export(report_type):
    show_prices = can_view_prices()
    sku_keyword, start_date, end_date = get_date_filters()
    sku_filter = build_sku_keyword_filter(sku_keyword)
    if report_type == "inventory":
        rows = []
        if start_date or end_date or sku_filter is not None:
            inventory_summary_query = db.session.query(
                InventoryTransaction.warehouse_id,
                InventoryTransaction.sku_id,
                func.coalesce(func.sum(InventoryTransaction.quantity), 0),
            ).join(InventoryTransaction.sku)
            if sku_filter is not None:
                inventory_summary_query = inventory_summary_query.filter(sku_filter)
            inventory_summary_query = apply_datetime_range(inventory_summary_query, InventoryTransaction.created_at, start_date, end_date)
            inventory_map = {
                (warehouse_id, sku_id): qty
                for warehouse_id, sku_id, qty in inventory_summary_query.group_by(InventoryTransaction.warehouse_id, InventoryTransaction.sku_id).all()
            }
        else:
            inventory_map = get_inventory_by_warehouse()
        sku_query = SKU.query.filter(sku_filter) if sku_filter is not None else SKU.query
        for warehouse in Warehouse.query.order_by(Warehouse.name.asc()).all():
            for sku in sku_query.order_by(SKU.name.asc()).all():
                qty = inventory_map.get((warehouse.id, sku.id), 0)
                rows.append(
                    [warehouse.name, sku.sku_code, sku.barcode, sku.name, sku.color, sku.size, qty]
                    + ([float(sku.cost_price or 0), float(sku.sale_price or 0)] if show_prices else [])
                )
        return export_workbook(
            "库存报表.xlsx",
            "库存",
            ["仓库", "SKU编码", "条形码", "商品名称", "颜色", "尺码", "数量"]
            + (["成本价", "售价"] if show_prices else []),
            rows,
        )

    if report_type == "purchase":
        rows = []
        order_query = PurchaseOrder.query
        if sku_filter is not None:
            order_query = order_query.join(PurchaseOrder.items).join(PurchaseOrderItem.sku).filter(sku_filter).distinct()
        order_query = apply_datetime_range(order_query, PurchaseOrder.created_at, start_date, end_date)
        for order in order_query.order_by(PurchaseOrder.created_at.desc()).all():
            for item in order.items:
                if sku_filter is not None and not (
                    (item.sku.sku_code and sku_keyword.lower() in item.sku.sku_code.lower())
                    or (item.sku.name and sku_keyword.lower() in item.sku.name.lower())
                    or (item.sku.barcode and sku_keyword.lower() in item.sku.barcode.lower())
                ):
                    continue
                rows.append(
                    [order.order_no, order.supplier.name, order.warehouse.name, order.status_label, item.sku.sku_code, item.sku.name, item.quantity]
                    + ([float(item.unit_price or 0)] if show_prices else [])
                    + [order.created_at.strftime("%Y-%m-%d %H:%M")]
                )
        return export_workbook(
            "采购报表.xlsx",
            "采购",
            ["订单号", "供应商", "仓库", "状态", "SKU编码", "商品名称", "数量"]
            + (["单价"] if show_prices else [])
            + ["创建时间"],
            rows,
        )

    if report_type == "sales":
        rows = []
        order_query = SalesOrder.query
        if sku_filter is not None:
            order_query = order_query.join(SalesOrder.items).join(SalesOrderItem.sku).filter(sku_filter).distinct()
        order_query = apply_datetime_range(order_query, SalesOrder.created_at, start_date, end_date)
        for order in order_query.order_by(SalesOrder.created_at.desc()).all():
            for item in order.items:
                if sku_filter is not None and not (
                    (item.sku.sku_code and sku_keyword.lower() in item.sku.sku_code.lower())
                    or (item.sku.name and sku_keyword.lower() in item.sku.name.lower())
                    or (item.sku.barcode and sku_keyword.lower() in item.sku.barcode.lower())
                ):
                    continue
                rows.append(
                    [order.order_no, order.customer_name, order.warehouse.name, order.status_label, item.sku.sku_code, item.sku.name, item.quantity]
                    + ([float(item.unit_price or 0)] if show_prices else [])
                    + [order.created_at.strftime("%Y-%m-%d %H:%M")]
                )
        return export_workbook(
            "销售报表.xlsx",
            "销售",
            ["订单号", "客户", "仓库", "状态", "SKU编码", "商品名称", "数量"]
            + (["单价"] if show_prices else [])
            + ["创建时间"],
            rows,
        )

    flash("不支持的导出类型。", "error")
    return redirect(url_for("reports"))


def report_export_override(report_type):
    show_prices = can_view_prices()
    selected_warehouse_id = parse_int(request.args.get("warehouse_id"))
    sku_keyword, start_date, end_date = get_date_filters()
    sku_filter = build_sku_keyword_filter(sku_keyword)

    if report_type == "inventory":
        rows = []
        if start_date or end_date or sku_filter is not None:
            inventory_summary_query = db.session.query(
                InventoryTransaction.warehouse_id,
                InventoryTransaction.sku_id,
                func.coalesce(func.sum(InventoryTransaction.quantity), 0),
            ).join(InventoryTransaction.sku)
            if selected_warehouse_id:
                inventory_summary_query = inventory_summary_query.filter(InventoryTransaction.warehouse_id == selected_warehouse_id)
            if sku_filter is not None:
                inventory_summary_query = inventory_summary_query.filter(sku_filter)
            inventory_summary_query = apply_datetime_range(
                inventory_summary_query, InventoryTransaction.created_at, start_date, end_date
            )
            inventory_map = {
                (warehouse_id, sku_id): qty
                for warehouse_id, sku_id, qty in inventory_summary_query.group_by(
                    InventoryTransaction.warehouse_id, InventoryTransaction.sku_id
                ).all()
            }
        else:
            inventory_map = get_inventory_by_warehouse()

        defect_summary_query = DefectRepair.query
        if selected_warehouse_id:
            defect_summary_query = defect_summary_query.filter(DefectRepair.warehouse_id == selected_warehouse_id)
        if sku_filter is not None:
            defect_summary_query = defect_summary_query.join(DefectRepair.sku).filter(sku_filter)
        defect_summary_query = apply_datetime_range(
            defect_summary_query, DefectRepair.created_at, start_date, end_date
        )
        defect_status_map = {}
        for repair in defect_summary_query.all():
            if repair.remaining_quantity:
                status = "repairing" if repair.status == "repairing" else "pending"
                key = (repair.warehouse_id, repair.sku_id, status)
                defect_status_map[key] = defect_status_map.get(key, 0) + repair.remaining_quantity
            if repair.completed_quantity:
                key = (repair.warehouse_id, repair.sku_id, "done")
                defect_status_map[key] = defect_status_map.get(key, 0) + repair.completed_quantity
            if repair.scrapped_quantity:
                key = (repair.warehouse_id, repair.sku_id, "scrapped")
                defect_status_map[key] = defect_status_map.get(key, 0) + repair.scrapped_quantity

        sku_query = apply_sku_scope(SKU.query, SKU.id)
        if sku_filter is not None:
            sku_query = sku_query.filter(sku_filter)

        warehouse_query = Warehouse.query.order_by(Warehouse.name.asc())
        if selected_warehouse_id:
            warehouse_query = warehouse_query.filter(Warehouse.id == selected_warehouse_id)
        for warehouse in warehouse_query.all():
            for sku in sku_query.order_by(SKU.name.asc()).all():
                qty = inventory_map.get((warehouse.id, sku.id), 0)
                pending_qty = defect_status_map.get((warehouse.id, sku.id, "pending"), 0)
                repairing_qty = defect_status_map.get((warehouse.id, sku.id, "repairing"), 0)
                done_qty = defect_status_map.get((warehouse.id, sku.id, "done"), 0)
                scrapped_qty = defect_status_map.get((warehouse.id, sku.id, "scrapped"), 0)
                stock_value = qty * float(sku.cost_price or 0)
                rows.append(
                    [
                        warehouse.name,
                        sku.sku_code,
                        sku.barcode,
                        sku.name,
                        sku.color,
                        sku.size,
                        qty,
                        pending_qty,
                        repairing_qty,
                        done_qty,
                        scrapped_qty,
                        sku.safety_stock,
                    ]
                    + ([stock_value] if show_prices else [])
                    + ["库存偏低" if qty <= sku.safety_stock else "正常"]
                    + ([float(sku.cost_price or 0), float(sku.sale_price or 0)] if show_prices else [])
                )
        return export_workbook(
            "库存报表.xlsx",
            "库存",
            ["仓库", "SKU编码", "条形码", "商品名称", "颜色", "尺码", "当前库存", "待处理", "返修中", "已完成", "已报废", "安全库存"]
            + (["库存价值"] if show_prices else [])
            + ["状态"]
            + (["成本价", "售价"] if show_prices else []),
            rows,
        )

    return report_export(report_type)


app.view_functions["report_export"] = login_required(permission_required("report.view")(report_export_override))


@app.cli.command("init-db")
def init_db_command():
    db.create_all()
    ensure_barcode_optional_schema()
    ensure_defect_repair_schema()
    ensure_operator_name_schema()
    ensure_inventory_supplier_schema()
    ensure_login_security_schema()
    ensure_asset_schema()
    ensure_fba_inbound_schema()
    ensure_shipment_archive_schema()
    ensure_freight_compare_schema()
    seed_data()
    print("数据库初始化完成。")


with app.app_context():
    db.create_all()
    ensure_barcode_optional_schema()
    ensure_defect_repair_schema()
    ensure_operator_name_schema()
    ensure_inventory_supplier_schema()
    ensure_login_security_schema()
    ensure_asset_schema()
    ensure_fba_inbound_schema()
    ensure_shipment_archive_schema()
    ensure_freight_compare_schema()
    seed_data()


if __name__ == "__main__":
    app.run(host=os.getenv("ERP_HOST", "0.0.0.0"), port=parse_int(os.getenv("ERP_PORT"), 8000), debug=False)



